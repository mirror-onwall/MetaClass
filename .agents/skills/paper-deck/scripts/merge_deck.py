"""Merge Paper Deck raster and source-grounded slides into PPTX/PDF.

Without ``source-visual-manifest.json`` this preserves the original raster-only
behavior. With a manifest, every source-grounded page remains layered: generated
background, verified source pictures, and editable annotations are independent
PowerPoint objects.
"""
# ruff: noqa: I001

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import math
import os
import re
import sys
from pathlib import Path
from typing import Any


def _ensure_runtime_dependencies() -> None:
    if all(importlib.util.find_spec(name) is not None for name in ("PIL", "pptx")):
        return
    project_root = Path(__file__).resolve().parents[4]
    candidates = (
        project_root / "metaclass_env/bin/python",
        project_root / "apps/api/.venv/bin/python",
    )
    current = Path(sys.executable).resolve()
    for candidate in candidates:
        if candidate.is_file() and candidate.resolve() != current:
            os.execv(str(candidate), [str(candidate), str(Path(__file__).resolve()), *sys.argv[1:]])
    raise SystemExit(
        "paper-deck compositor requires Pillow and python-pptx; project runtime not found"
    )


_ensure_runtime_dependencies()

from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp"}
SLIDE_W_IN = 13.333333
SLIDE_H_IN = 7.5
MANIFEST_NAME = "source-visual-manifest.json"
RENDERED_DIR_NAME = "rendered"
SLIDE_W_PX = 1600
SLIDE_H_PX = 900
HEX_COLOR = re.compile(r"^#?([0-9a-fA-F]{6})$")


class MergeDeckError(ValueError):
    """Raised when a deck cannot be composed without losing source fidelity."""


def natural_key(path: Path) -> tuple[int, str]:
    match = re.match(r"^(\d+)", path.stem)
    return (int(match.group(1)) if match else 9999), path.name


def find_images(deck_dir: Path) -> list[Path]:
    image_dir = deck_dir / "images"
    if not image_dir.exists():
        raise MergeDeckError(f"Missing images directory: {image_dir}")
    images = [path for path in image_dir.iterdir() if path.suffix.lower() in IMAGE_EXTS]
    images.sort(key=natural_key)
    if not images:
        raise MergeDeckError(f"No slide images found in: {image_dir}")
    return images


def load_manifest(deck_dir: Path) -> dict[str, Any] | None:
    path = deck_dir / MANIFEST_NAME
    if not path.is_file():
        return None
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise MergeDeckError(f"Invalid {MANIFEST_NAME}") from exc
    if payload.get("schema_version") != "1.0" or not isinstance(payload.get("slides"), list):
        raise MergeDeckError(f"Unsupported or incomplete {MANIFEST_NAME}")
    slides = payload["slides"]
    if not slides:
        raise MergeDeckError("source visual manifest contains no slides")
    orders = [item.get("order") for item in slides if isinstance(item, dict)]
    if orders != list(range(1, len(slides) + 1)):
        raise MergeDeckError("manifest slide order must be continuous from 1")
    return payload


def _resolve_inside(root: Path, value: str, *, kind: str) -> Path:
    relative = Path(value)
    if relative.is_absolute() or ".." in relative.parts:
        raise MergeDeckError(f"unsafe {kind} path: {value}")
    resolved = (root / relative).resolve()
    try:
        resolved.relative_to(root.resolve())
    except ValueError as exc:
        raise MergeDeckError(f"{kind} path escapes workspace: {value}") from exc
    if not resolved.is_file():
        raise MergeDeckError(f"missing {kind}: {value}")
    return resolved


def _workspace_root(deck_dir: Path) -> Path:
    return deck_dir.parent if deck_dir.name == "provider_output" else deck_dir


def _background_path(deck_dir: Path, value: str) -> Path:
    return _resolve_inside(deck_dir, value, kind="background")


def _source_path(deck_dir: Path, value: str) -> Path:
    root = _workspace_root(deck_dir)
    path = _resolve_inside(root, value, kind="source asset")
    if deck_dir.name == "provider_output":
        expected = (root / "provider_input" / "assets").resolve()
        try:
            path.relative_to(expected)
        except ValueError as exc:
            raise MergeDeckError(
                "MetaClass source assets must stay inside provider_input/assets"
            ) from exc
    return path


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _verify_hash(path: Path, expected: str) -> None:
    normalized = expected.removeprefix("sha256:").casefold()
    if not re.fullmatch(r"[0-9a-f]{64}", normalized):
        raise MergeDeckError(f"invalid SHA-256 for source asset: {path.name}")
    if _sha256(path) != normalized:
        raise MergeDeckError(f"source asset SHA-256 mismatch: {path.name}")


def _normalized_box(item: dict[str, Any]) -> tuple[float, float, float, float]:
    try:
        x, y, w, h = (float(item[key]) for key in ("x", "y", "w", "h"))
    except (KeyError, TypeError, ValueError) as exc:
        raise MergeDeckError("layer requires numeric x, y, w, and h") from exc
    if x < 0 or y < 0 or w <= 0 or h <= 0 or x + w > 1 or y + h > 1:
        raise MergeDeckError("layer geometry must remain inside normalized slide bounds")
    return x, y, w, h


def _inches_box(item: dict[str, Any]) -> tuple[Any, Any, Any, Any]:
    x, y, w, h = _normalized_box(item)
    return (
        Inches(x * SLIDE_W_IN),
        Inches(y * SLIDE_H_IN),
        Inches(w * SLIDE_W_IN),
        Inches(h * SLIDE_H_IN),
    )


def _add_source_picture(slide: Any, path: Path, item: dict[str, Any]) -> Any:
    left, top, width, height = _inches_box(item)
    fit = item.get("fit", "contain")
    if item.get("crop", "full") != "full":
        raise MergeDeckError("manifest v1 only supports crop=full")
    with Image.open(path) as image:
        image_w, image_h = image.size
    if image_w < 1 or image_h < 1:
        raise MergeDeckError(f"source asset has invalid dimensions: {path.name}")
    frame_ratio = width / height
    image_ratio = image_w / image_h
    if fit == "contain":
        if image_ratio >= frame_ratio:
            rendered_w = width
            rendered_h = round(width / image_ratio)
            rendered_left = left
            rendered_top = top + round((height - rendered_h) / 2)
        else:
            rendered_h = height
            rendered_w = round(height * image_ratio)
            rendered_left = left + round((width - rendered_w) / 2)
            rendered_top = top
        picture = slide.shapes.add_picture(
            str(path), rendered_left, rendered_top, width=rendered_w, height=rendered_h
        )
    elif fit == "cover":
        picture = slide.shapes.add_picture(str(path), left, top, width=width, height=height)
        if image_ratio > frame_ratio:
            visible = frame_ratio / image_ratio
            picture.crop_left = picture.crop_right = (1 - visible) / 2
        elif image_ratio < frame_ratio:
            visible = image_ratio / frame_ratio
            picture.crop_top = picture.crop_bottom = (1 - visible) / 2
    else:
        raise MergeDeckError(f"unsupported image fit: {fit}")
    picture.name = f"source:{item['asset_id']}"
    return picture


def _rgb(value: str | None, default: str) -> RGBColor:
    matched = HEX_COLOR.match(value or default)
    if not matched:
        raise MergeDeckError(f"invalid RGB color: {value}")
    return RGBColor.from_string(matched.group(1).upper())


def _add_annotation(slide: Any, item: dict[str, Any], index: int) -> None:
    annotation_type = item.get("type", "text")
    left, top, width, height = _inches_box(item)
    line_color = _rgb(item.get("line_color"), "334155")
    if annotation_type == "text":
        text = item.get("text")
        if not isinstance(text, str) or not text.strip():
            raise MergeDeckError("text annotation requires non-empty text")
        shape = slide.shapes.add_textbox(left, top, width, height)
        shape.name = f"annotation:text:{index:02d}"
        frame = shape.text_frame
        frame.clear()
        frame.margin_left = frame.margin_right = Inches(0.04)
        frame.margin_top = frame.margin_bottom = Inches(0.02)
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = frame.paragraphs[0]
        paragraph.text = text
        paragraph.alignment = {
            "left": PP_ALIGN.LEFT,
            "center": PP_ALIGN.CENTER,
            "right": PP_ALIGN.RIGHT,
        }.get(item.get("align", "left"), PP_ALIGN.LEFT)
        run = paragraph.runs[0]
        run.font.name = item.get("font_name", "Aptos")
        run.font.size = Pt(float(item.get("font_size", 18)))
        run.font.bold = bool(item.get("bold", False))
        run.font.color.rgb = _rgb(item.get("color"), "0F172A")
        return
    if annotation_type == "rectangle":
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shape.name = f"annotation:rectangle:{index:02d}"
        shape.fill.background()
        shape.line.color.rgb = line_color
        shape.line.width = Pt(float(item.get("line_width", 1.5)))
        return
    if annotation_type == "arrow":
        shape = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, left, top, left + width, top + height
        )
        shape.name = f"annotation:arrow:{index:02d}"
        shape.line.color.rgb = line_color
        shape.line.width = Pt(float(item.get("line_width", 1.5)))
        from lxml import etree

        etree.SubElement(
            shape._element.spPr.ln,
            "{http://schemas.openxmlformats.org/drawingml/2006/main}headEnd",
            type="triangle",
        )
        return
    raise MergeDeckError(f"unsupported annotation type: {annotation_type}")


def make_pptx_from_images(images: list[Path], output: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_W_IN)
    presentation.slide_height = Inches(SLIDE_H_IN)
    blank = presentation.slide_layouts[6]
    for image in images:
        slide = presentation.slides.add_slide(blank)
        picture = slide.shapes.add_picture(
            str(image), 0, 0, width=presentation.slide_width, height=presentation.slide_height
        )
        picture.name = "paper-deck:native-raster"
    presentation.save(output)


def make_pptx_from_manifest(deck_dir: Path, manifest: dict[str, Any], output: Path) -> bool:
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_W_IN)
    presentation.slide_height = Inches(SLIDE_H_IN)
    blank = presentation.slide_layouts[6]
    contains_hybrid = False
    for entry in manifest["slides"]:
        mode = entry.get("render_mode")
        if mode not in {"native-raster", "source-grounded-hybrid"}:
            raise MergeDeckError(f"unsupported render mode: {mode}")
        background = _background_path(deck_dir, entry.get("background_path", ""))
        slide = presentation.slides.add_slide(blank)
        picture = slide.shapes.add_picture(
            str(background), 0, 0,
            width=presentation.slide_width, height=presentation.slide_height,
        )
        picture.name = f"background:{mode}"
        assets = entry.get("assets", [])
        annotations = entry.get("annotations", [])
        if not isinstance(assets, list) or not isinstance(annotations, list):
            raise MergeDeckError("manifest assets and annotations must be arrays")
        if mode == "native-raster":
            if assets or annotations:
                raise MergeDeckError("native-raster pages cannot contain hybrid layers")
            continue
        contains_hybrid = True
        if not assets:
            raise MergeDeckError("source-grounded-hybrid page requires a source asset")
        for asset in assets:
            source = _source_path(deck_dir, asset.get("source_path", ""))
            _verify_hash(source, asset.get("sha256", ""))
            _add_source_picture(slide, source, asset)
        for index, annotation in enumerate(annotations, start=1):
            _add_annotation(slide, annotation, index)
    presentation.save(output)
    return contains_hybrid


def make_raster_pdf(images: list[Path], output: Path) -> None:
    frames: list[Image.Image] = []
    for image_path in images:
        with Image.open(image_path) as image:
            frames.append(image.convert("RGB").copy())
    first, *rest = frames
    first.save(output, save_all=True, append_images=rest, resolution=150.0)


def update_composition_log(deck_dir: Path, manifest: dict[str, Any]) -> None:
    """Append a deterministic, validator-readable record without replacing model notes."""
    path = deck_dir / "generation-log.md"
    text = path.read_text(encoding="utf-8") if path.is_file() else "# Generation Log\n"
    start = "<!-- metaclass-composition-log:start -->"
    end = "<!-- metaclass-composition-log:end -->"
    if start in text and end in text:
        prefix, remainder = text.split(start, 1)
        _, suffix = remainder.split(end, 1)
        text = prefix.rstrip() + "\n" + suffix.lstrip()
    lines = [start, "## MetaClass deterministic composition"]
    for entry in manifest["slides"]:
        background = str(entry.get("background_path", ""))
        mode = str(entry.get("render_mode", ""))
        fields = [
            background,
            f"render_mode={mode}",
            "background_backend=imagegen",
        ]
        if mode == "source-grounded-hybrid":
            assets = entry.get("assets", [])
            fields.extend(
                [
                    "source_asset=" + ",".join(str(item.get("asset_id", "")) for item in assets),
                    "source_asset_sha256="
                    + ",".join(str(item.get("sha256", "")) for item in assets),
                    "composition_backend=metaclass",
                ]
            )
        lines.append(" ".join(fields))
    lines.append(end)
    path.write_text(text.rstrip() + "\n\n" + "\n".join(lines) + "\n", encoding="utf-8")


def _pixel_box(item: dict[str, Any]) -> tuple[int, int, int, int]:
    x, y, w, h = _normalized_box(item)
    return (
        round(x * SLIDE_W_PX),
        round(y * SLIDE_H_PX),
        round(w * SLIDE_W_PX),
        round(h * SLIDE_H_PX),
    )


def _font(size: int, *, bold: bool) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, size=size)
        except OSError:
            continue
    return ImageFont.load_default()


def _fit_source_raster(source: Image.Image, size: tuple[int, int], fit: str) -> Image.Image:
    if fit == "contain":
        return ImageOps.contain(source, size, Image.Resampling.LANCZOS)
    if fit == "cover":
        return ImageOps.fit(source, size, Image.Resampling.LANCZOS, centering=(0.5, 0.5))
    raise MergeDeckError(f"unsupported image fit: {fit}")


def _draw_text_annotation(draw: ImageDraw.ImageDraw, item: dict[str, Any]) -> None:
    left, top, width, height = _pixel_box(item)
    text = str(item.get("text", "")).strip()
    if not text:
        raise MergeDeckError("text annotation requires non-empty text")
    requested = max(8, round(float(item.get("font_size", 18)) * 5 / 3))
    font = _font(requested, bold=bool(item.get("bold", False)))
    spacing = max(2, requested // 5)
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=spacing)
    while requested > 8 and (bbox[2] > width - 10 or bbox[3] > height - 6):
        requested -= 1
        font = _font(requested, bold=bool(item.get("bold", False)))
        spacing = max(2, requested // 5)
        bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=spacing)
    align = str(item.get("align", "left"))
    text_width = bbox[2] - bbox[0]
    text_height = bbox[3] - bbox[1]
    x = left + 5
    if align == "center":
        x = left + (width - text_width) / 2
    elif align == "right":
        x = left + width - text_width - 5
    y = top + (height - text_height) / 2 - bbox[1]
    color = "#" + _rgb(item.get("color"), "0F172A").__str__()
    draw.multiline_text((x, y), text, font=font, fill=color, spacing=spacing, align=align)


def _draw_annotation(draw: ImageDraw.ImageDraw, item: dict[str, Any]) -> None:
    annotation_type = item.get("type", "text")
    if annotation_type == "text":
        _draw_text_annotation(draw, item)
        return
    left, top, width, height = _pixel_box(item)
    color = "#" + _rgb(item.get("line_color"), "334155").__str__()
    line_width = max(1, round(float(item.get("line_width", 1.5)) * 5 / 3))
    if annotation_type == "rectangle":
        draw.rectangle(
            (left, top, left + width, top + height), outline=color, width=line_width
        )
        return
    if annotation_type == "arrow":
        start = (left, top)
        end = (left + width, top + height)
        draw.line((start, end), fill=color, width=line_width)
        angle = math.atan2(end[1] - start[1], end[0] - start[0])
        length = max(10, line_width * 5)
        spread = 0.55
        points = [
            end,
            (
                end[0] - length * math.cos(angle - spread),
                end[1] - length * math.sin(angle - spread),
            ),
            (
                end[0] - length * math.cos(angle + spread),
                end[1] - length * math.sin(angle + spread),
            ),
        ]
        draw.polygon(points, fill=color)
        return
    raise MergeDeckError(f"unsupported annotation type: {annotation_type}")


def compose_rendered_pages(
    deck_dir: Path, manifest: dict[str, Any], output_dir: Path
) -> list[Path]:
    """Compose the same manifest layers into final raster pages without office software."""
    output_dir.mkdir(parents=True, exist_ok=True)
    rendered: list[Path] = []
    for index, entry in enumerate(manifest["slides"], start=1):
        background = _background_path(deck_dir, entry.get("background_path", ""))
        with Image.open(background) as opened:
            canvas = opened.convert("RGB").resize(
                (SLIDE_W_PX, SLIDE_H_PX), Image.Resampling.LANCZOS
            )
        for asset in entry.get("assets", []):
            source_path = _source_path(deck_dir, asset.get("source_path", ""))
            _verify_hash(source_path, asset.get("sha256", ""))
            if asset.get("crop", "full") != "full":
                raise MergeDeckError("manifest v1 only supports crop=full")
            left, top, width, height = _pixel_box(asset)
            with Image.open(source_path) as opened:
                source = opened.convert("RGB")
            fitted = _fit_source_raster(source, (width, height), asset.get("fit", "contain"))
            paste_left = left + (width - fitted.width) // 2
            paste_top = top + (height - fitted.height) // 2
            canvas.paste(fitted, (paste_left, paste_top))
        draw = ImageDraw.Draw(canvas)
        for annotation in entry.get("annotations", []):
            _draw_annotation(draw, annotation)
        destination = output_dir / f"{index:02d}-slide.png"
        canvas.save(destination)
        rendered.append(destination)
    expected = {path.name for path in rendered}
    for path in output_dir.iterdir():
        if (
            path.is_file()
            and path.suffix.lower() in IMAGE_EXTS
            and path.name not in expected
        ):
            path.unlink()
    return rendered


def render_native_pages(images: list[Path], output_dir: Path) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    rendered: list[Path] = []
    for index, image_path in enumerate(images, start=1):
        with Image.open(image_path) as opened:
            page = opened.convert("RGB").resize(
                (SLIDE_W_PX, SLIDE_H_PX), Image.Resampling.LANCZOS
            )
        destination = output_dir / f"{index:02d}-slide.png"
        page.save(destination)
        rendered.append(destination)
    return rendered


def merge_deck(
    deck_dir: Path,
    *,
    name: str | None = None,
) -> tuple[Path, Path, int]:
    deck_dir = deck_dir.expanduser().resolve()
    if not deck_dir.is_dir():
        raise MergeDeckError(f"Deck directory does not exist: {deck_dir}")
    base = name or deck_dir.name
    pptx_path = deck_dir / f"{base}.pptx"
    pdf_path = deck_dir / f"{base}.pdf"
    manifest = load_manifest(deck_dir)
    if manifest is None:
        images = find_images(deck_dir)
        make_pptx_from_images(images, pptx_path)
        rendered = render_native_pages(images, deck_dir / RENDERED_DIR_NAME)
        make_raster_pdf(rendered, pdf_path)
        return pptx_path, pdf_path, len(images)
    make_pptx_from_manifest(deck_dir, manifest, pptx_path)
    rendered = compose_rendered_pages(deck_dir, manifest, deck_dir / RENDERED_DIR_NAME)
    make_raster_pdf(rendered, pdf_path)
    update_composition_log(deck_dir, manifest)
    return pptx_path, pdf_path, len(manifest["slides"])


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("deck_dir", help="Deck directory containing images/ and optional manifest")
    parser.add_argument("--name", help="Output base name. Defaults to deck directory name.")
    args = parser.parse_args()
    try:
        pptx_path, pdf_path, slide_count = merge_deck(Path(args.deck_dir), name=args.name)
    except (MergeDeckError, OSError) as exc:
        raise SystemExit(str(exc)) from exc
    print(f"Merged {slide_count} slides")
    print(f"PPTX: {pptx_path}")
    print(f"PDF:  {pdf_path}")


if __name__ == "__main__":
    main()
