"""Run the real paper workflow against named PDFs and emit an auditable report.

This is intentionally not a pytest: it invokes authenticated Codex Skills, the
configured narration LLM, LibreOffice, and the real application repositories.
"""

from __future__ import annotations

import argparse
import asyncio
import json
import re
import shutil
import statistics
import subprocess
import sys
import time
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from fastapi import HTTPException, UploadFile
from pptx import Presentation

from metaclass.core.application import build_services
from metaclass.core.config import settings
from metaclass.core.llm_config import get_llm_runtime_config
from metaclass.modules.classroom.schemas import LearningMode
from metaclass.modules.paper_workflow.paper_classroom import (
    PaperNarrationValidator,
    PaperSlideEvidencePacket,
)
from metaclass.modules.paper_workflow.schemas import PaperDeckCourseResult, PaperWorkflowRequest

EXPECTED_SKILLS = {
    "paper-analyze",
    "extract-paper-images",
    "academic-pptx",
    "pptx",
}


def arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("samples", nargs="+", type=Path)
    parser.add_argument("--data-dir", type=Path, required=True)
    parser.add_argument("--report-dir", type=Path, required=True)
    parser.add_argument("--duration", type=int, default=30)
    parser.add_argument(
        "--resume-job",
        action="append",
        default=[],
        metavar="PDF_STEM=JOB_ID",
        help="Resume a durable job in --data-dir instead of uploading that sample again",
    )
    parser.add_argument(
        "--rebuild-course",
        action="store_true",
        help="Recompose the deterministic paper course even when it already exists",
    )
    parser.add_argument("--audience", default="具备基础专业背景的高校学生和研究生")
    return parser.parse_args()


def now() -> str:
    return datetime.now(UTC).isoformat()


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


async def upload_pdf(services: Any, path: Path) -> Any:
    with path.open("rb") as stream:
        material = await services.materials.create(UploadFile(stream, filename=path.name))
    pages = services.materials.parse(material.id)
    return services.materials.get(material.id), pages


def stage_reports(workspace: Path) -> list[dict[str, Any]]:
    reports: list[dict[str, Any]] = []
    for path in sorted(workspace.glob("stages/**/execution_report.json")):
        payload = read_json(path)
        payload["report_path"] = str(path.relative_to(workspace))
        reports.append(payload)
    return reports


def narration_metrics(plan: Any) -> dict[str, Any]:
    validator = PaperNarrationValidator()
    slides: list[dict[str, Any]] = []
    sources: dict[str, int] = {}
    for slide in plan.slides:
        source = slide.speaker_script_source or "unknown"
        sources[source] = sources.get(source, 0) + 1
        packet = PaperSlideEvidencePacket.model_validate(slide.paper_evidence_packet)
        unsupported = sorted(
            validator._numbers(slide.speaker_script) - validator._authorized_numbers(packet)
        )
        sentences = [item for item in re.split(r"[。！？!?]", slide.speaker_script) if item.strip()]
        slides.append(
            {
                "slide_id": slide.id,
                "title": slide.title,
                "source": source,
                "characters": len(slide.speaker_script),
                "sentence_count": len(sentences),
                "unsupported_numbers": unsupported,
                "original_context_count": len(packet.evidence_contexts),
                "resolved_block_count": sum(
                    1 for context in packet.evidence_contexts if context.block_id
                ),
            }
        )
    lengths = [item["characters"] for item in slides]
    llm_count = sources.get("paper_classroom_llm", 0)
    return {
        "source_counts": sources,
        "fallback_rate": round(1 - llm_count / len(slides), 4) if slides else 1,
        "average_characters": round(statistics.mean(lengths), 1) if lengths else 0,
        "min_characters": min(lengths, default=0),
        "max_characters": max(lengths, default=0),
        "unsupported_number_slides": [
            item["slide_id"] for item in slides if item["unsupported_numbers"]
        ],
        "slides": slides,
    }


def render_pptx(pptx: Path, report_dir: Path) -> dict[str, Any]:
    deck = Presentation(str(pptx))
    soffice = settings.libreoffice_bin or shutil.which("soffice")
    result: dict[str, Any] = {
        "python_pptx_opened": True,
        "slide_count": len(deck.slides),
        "libreoffice_binary": soffice,
        "libreoffice_rendered": False,
    }
    if not soffice:
        result["error"] = "LibreOffice executable not found"
        return result
    render_dir = report_dir / "libreoffice"
    render_dir.mkdir(parents=True, exist_ok=True)
    completed = subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(render_dir), str(pptx)],
        capture_output=True,
        text=True,
        timeout=180,
        check=False,
    )
    pdf = render_dir / f"{pptx.stem}.pdf"
    result.update(
        {
            "libreoffice_exit_code": completed.returncode,
            "libreoffice_rendered": completed.returncode == 0 and pdf.is_file(),
            "rendered_pdf": str(pdf) if pdf.is_file() else None,
            "stderr": completed.stderr[-1200:],
        }
    )
    return result


def acceptance_checks(sample: dict[str, Any], llm_provider: str) -> dict[str, bool]:
    reports = sample.get("stage_reports", [])
    skill_names = {item.get("skill_name") for item in reports}
    successful_stages = {item.get("stage") for item in reports if item.get("status") == "succeeded"}
    narration = sample.get("narration", {})
    source = sample.get("source_coverage", {})
    return {
        "workflow_succeeded": sample.get("workflow_status") == "succeeded",
        "required_workflow_stages_succeeded": {
            "analysis",
            "figures",
            "outline",
            "generation",
        }.issubset(successful_stages),
        "conditional_extraction_succeeded_if_invoked": (
            "extract-paper-images" not in skill_names
            or any(
                item.get("skill_name") == "extract-paper-images"
                and item.get("status") == "succeeded"
                for item in reports
            )
        ),
        "all_skill_reports_succeeded": bool(reports)
        and all(item.get("status") == "succeeded" for item in reports),
        "pptx_valid": bool(sample.get("rendering", {}).get("python_pptx_opened")),
        "libreoffice_rendered": bool(sample.get("rendering", {}).get("libreoffice_rendered")),
        "every_slide_has_source_block": source.get("block_coverage_rate") == 1,
        "real_llm_configured": llm_provider != "fake",
        "every_slide_uses_llm_narration": narration.get("fallback_rate") == 0,
        "narration_numbers_valid": not narration.get("unsupported_number_slides"),
        "narration_directly_teachable": (
            narration.get("min_characters", 0) >= 120
            and 180 <= narration.get("average_characters", 0) <= 900
        ),
        "classroom_created": bool(sample.get("classroom_session_id")),
    }


def run_sample(
    services: Any,
    args: argparse.Namespace,
    path: Path,
    resume_jobs: dict[str, str],
) -> dict[str, Any]:
    started = time.monotonic()
    sample_dir = args.report_dir / path.stem
    sample_dir.mkdir(parents=True, exist_ok=True)
    result: dict[str, Any] = {"sample": str(path), "started_at": now()}
    try:
        resume_job_id = resume_jobs.get(path.stem)
        if resume_job_id:
            job = services.paper_workflows.get(resume_job_id)
            material = services.materials.get(job.source_material_id)
            pages = services.materials.pages(material.id)
            if job.status.value == "failed":
                job = services.paper_workflows.resume(job.id)
            if job.status.value == "queued":
                print(f"[{path.stem}] workflow {job.id} resumed", flush=True)
                job = services.paper_workflows.run(job.id)
        else:
            material, pages = asyncio.run(upload_pdf(services, path))
            job = services.paper_workflows.create(
                PaperWorkflowRequest(
                    material_id=material.id,
                    strategy="composed_skills",
                    duration_minutes=args.duration,
                    audience=args.audience,
                    language="zh-CN",
                    depth="standard",
                )
            )
            print(f"[{path.stem}] workflow {job.id} started", flush=True)
            job = services.paper_workflows.run(job.id)
        result.update({"material_id": material.id, "paper_page_count": len(pages)})
        result.update(
            {
                "job_id": job.id,
                "workflow_status": job.status.value,
                "workflow_error": job.error,
                "fallback_reason": job.fallback_reason,
            }
        )
        workspace = args.data_dir / "runtime" / "paper_workflows" / job.id
        result["stage_reports"] = stage_reports(workspace)
        if job.status.value != "succeeded":
            return result

        outline = services.paper_workflows.outline(job.id)
        evidence = services.paper_workflows.slide_evidence(job.id)
        refs_by_slide = {
            item.slide_id: [ref for ref in item.source_refs if ref.block_id]
            for item in evidence.slides
        }
        covered = sum(bool(refs_by_slide.get(slide.id)) for slide in outline.slides)
        result["source_coverage"] = {
            "slide_count": len(outline.slides),
            "slides_with_block": covered,
            "block_coverage_rate": round(covered / len(outline.slides), 4),
            "missing_slide_ids": [
                slide.id for slide in outline.slides if not refs_by_slide.get(slide.id)
            ],
        }

        identity = job.id.removeprefix("paper_job_")
        try:
            if args.rebuild_course:
                raise LookupError("paper course rebuild requested")
            existing_plan = services.presentations.get_plan(f"plan_paper_{identity}")
            bundle = services.paper_workflows.result(job.id)
            course = PaperDeckCourseResult(
                paper_job_id=job.id,
                derived_material_id=bundle.derived_material_id,
                source_paper_material_id=job.source_material_id,
                artifact_bundle_id=bundle.id,
                content_id=existing_plan.content_id,
                presentation_plan_id=existing_plan.id,
            )
        except (HTTPException, LookupError):
            course = services.paper_workflows.create_paper_deck_course(job.id)
        plan = services.presentations.get_plan(course.presentation_plan_id)
        content = services.contents.get(course.content_id)
        result["paper_deck"] = course.model_dump(mode="json")
        result["knowledge_tree"] = {
            "node_count": len(content.knowledge_tree.nodes) if content.knowledge_tree else 0,
            "unit_count": len(content.knowledge_units),
            "warning_count": len(content.knowledge_tree.warnings) if content.knowledge_tree else 0,
        }
        result["narration"] = narration_metrics(plan)

        # A lecture acceptance verifies the classroom contract without invoking
        # the unrelated interactive Question Bank model fan-out.
        services.classrooms.question_banks = None
        classroom_plan = services.classrooms.create_plan(
            course.content_id, course.presentation_plan_id
        )
        session = services.classrooms.create_session(classroom_plan.id, LearningMode.LECTURE)
        result["classroom_session_id"] = session.id

        bundle = services.paper_workflows.result(job.id)
        root = args.data_dir / bundle.root_path
        result["rendering"] = render_pptx(root / "presentation.pptx", sample_dir)
    except Exception as exc:  # noqa: BLE001 - report must survive one bad sample
        result["fatal_error"] = f"{type(exc).__name__}: {exc}"
    finally:
        result["duration_seconds"] = round(time.monotonic() - started, 2)
        result["finished_at"] = now()
    return result


def markdown(report: dict[str, Any]) -> str:
    lines = ["# Paper workflow real acceptance", "", f"Generated: {report['generated_at']}", ""]
    lines += [
        f"- LLM provider: `{report['environment']['llm_provider']}`",
        f"- LLM model: `{report['environment']['llm_model']}`",
        f"- Overall: **{'PASS' if report['passed'] else 'FAIL'}**",
        "",
    ]
    for name, passed in report.get("suite_checks", {}).items():
        lines.append(f"- {'PASS' if passed else 'FAIL'} — {name}")
    lines.append("")
    for sample in report["samples"]:
        lines += [f"## {Path(sample['sample']).stem}", ""]
        if sample.get("fatal_error"):
            lines += [f"Fatal error: `{sample['fatal_error']}`", ""]
        for name, passed in sample["checks"].items():
            lines.append(f"- {'PASS' if passed else 'FAIL'} — {name}")
        narration = sample.get("narration", {})
        lines += [
            "",
            (
                f"Narration sources: `{narration.get('source_counts', {})}`; "
                f"average length: `{narration.get('average_characters', 0)}` characters."
            ),
            "",
        ]
    return "\n".join(lines)


def main() -> int:
    args = arguments()
    args.data_dir = args.data_dir.resolve()
    args.report_dir = args.report_dir.resolve()
    args.data_dir.mkdir(parents=True, exist_ok=True)
    args.report_dir.mkdir(parents=True, exist_ok=True)
    llm = get_llm_runtime_config()
    services = build_services(args.data_dir, force_fake_llm=False)
    try:
        resume_jobs = dict(item.split("=", 1) for item in args.resume_job)
        samples = [run_sample(services, args, path.resolve(), resume_jobs) for path in args.samples]
    finally:
        services.database.dispose()
    for sample in samples:
        sample["checks"] = acceptance_checks(sample, llm.provider)
        sample["passed"] = all(sample["checks"].values())
    observed_skills = {
        item.get("skill_name") for sample in samples for item in sample.get("stage_reports", [])
    }
    suite_checks = {
        "all_four_real_skills_observed_across_suite": EXPECTED_SKILLS.issubset(observed_skills),
        "both_samples_completed": len(samples) == 2
        and all(sample.get("workflow_status") == "succeeded" for sample in samples),
    }
    report = {
        "generated_at": now(),
        "environment": {
            "llm_provider": llm.provider,
            "llm_model": llm.model,
            "material_parser": settings.material_parser,
            "libreoffice": settings.libreoffice_bin or shutil.which("soffice"),
        },
        "samples": samples,
        "suite_checks": suite_checks,
        "passed": all(item["passed"] for item in samples) and all(suite_checks.values()),
    }
    (args.report_dir / "acceptance.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (args.report_dir / "acceptance.md").write_text(markdown(report), encoding="utf-8")
    print(
        json.dumps({"passed": report["passed"], "report": str(args.report_dir)}, ensure_ascii=False)
    )
    return 0 if report["passed"] else 1


if __name__ == "__main__":
    sys.exit(main())
