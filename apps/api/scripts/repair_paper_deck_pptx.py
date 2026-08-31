"""Apply narrow, fail-closed QA repairs to a generated paper-deck PPTX."""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def remove_title_accent_lines(source: Path, destination: Path) -> int:
    presentation = Presentation(source)
    removed = 0

    for slide_number, slide in enumerate(presentation.slides, start=1):
        if slide_number == 1:
            continue
        matches = [
            shape
            for shape in slide.shapes
            if shape.shape_type == MSO_SHAPE_TYPE.LINE
            and 600_000 <= shape.left <= 850_000
            and 1_100_000 <= shape.top <= 1_300_000
            and 700_000 <= shape.width <= 11_000_000
            and 0 <= shape.height <= 60_000
        ]
        if len(matches) != 1:
            raise RuntimeError(
                f"Expected one title accent line on slide {slide_number}, found {len(matches)}"
            )
        element = matches[0]._element
        element.getparent().remove(element)
        removed += 1

    expected = len(presentation.slides) - 1
    if removed != expected:
        raise RuntimeError(f"Expected to remove {expected} lines, removed {removed}")

    destination.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(destination)
    return removed


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()
    removed = remove_title_accent_lines(args.source.resolve(), args.destination.resolve())
    print(f"Removed {removed} title accent lines; wrote {args.destination.resolve()}")


if __name__ == "__main__":
    main()
