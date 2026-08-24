import hashlib
import json
import os
import shlex
import shutil
import subprocess
import tempfile
import time
from dataclasses import dataclass
from pathlib import Path
from threading import Event, Lock
from typing import Literal
from uuid import uuid4

import fitz
from fastapi import HTTPException, UploadFile
from PIL import Image, ImageDraw
from pptx import Presentation

from metaclass.core.schemas import utc_now
from metaclass.infrastructure.storage import safe_filename, save_upload
from metaclass.modules.materials.repository import MaterialRepository
from metaclass.modules.materials.schemas import (
    Material,
    MaterialCollection,
    MaterialProcessingJob,
    MaterialProcessingJobStatus,
    MaterialStatus,
    MaterialType,
    PageImage,
    PageMetadata,
    ProcessedMaterial,
    ProcessedMaterials,
    SourceRef,
)


class MaterialProcessingCancelled(RuntimeError):
    pass


@dataclass(frozen=True)
class MaterialRichParseResult:
    """A selected rich parser artifact; projection remains a caller concern."""

    format: Literal["content_list_v2", "content_list", "middle"]
    path: Path
    root: Path
    payload: list[object] | dict[str, object]


class MaterialService:
    def __init__(
        self,
        data_dir: Path,
        repository: MaterialRepository,
        *,
        parser_backend: str = "auto",
        mineru_command: str = "mineru",
        mineru_timeout_seconds: float = 180.0,
    ) -> None:
        self.data_dir = data_dir
        self.repository = repository
        self.parser_backend = parser_backend
        self.mineru_command = mineru_command
        self.mineru_timeout_seconds = mineru_timeout_seconds
        self._jobs: dict[str, MaterialProcessingJob] = {}
        self._cancel_events: dict[str, Event] = {}
        self._discard_on_cancel: set[str] = set()
        self._job_lock = Lock()
        self._jobs_dir = self.data_dir / "runtime" / "material_processing" / "jobs"
        self._restore_processing_jobs()

    def _restore_processing_jobs(self) -> None:
        if not self._jobs_dir.exists():
            return
        for path in self._jobs_dir.glob("*.json"):
            try:
                job = MaterialProcessingJob.model_validate_json(path.read_text(encoding="utf-8"))
            except (OSError, ValueError):
                continue
            if job.status == MaterialProcessingJobStatus.RUNNING:
                job.status = MaterialProcessingJobStatus.PAUSED
                job.step = "paused"
                job.message = "服务器中断，进度已保存，可以继续"
            self._jobs[job.id] = job
            self._cancel_events[job.id] = Event()

    async def create(self, upload: UploadFile) -> Material:
        filename = safe_filename(upload.filename)
        suffix = Path(filename).suffix.lower().lstrip(".")
        if suffix not in {"pdf", "pptx"}:
            raise HTTPException(415, "Only PDF and PPTX files are supported")
        material_id = f"mat_{uuid4().hex[:12]}"
        path = self.data_dir / "raw" / material_id / f"source.{suffix}"
        await save_upload(upload, path)
        file_hash = self._sha256(path)
        material = Material(
            id=material_id,
            filename=filename,
            file_type=suffix,
            file_hash=file_hash,
            storage_path=str(path),
        )
        self.repository.save_material(material)
        return material

    def list_materials(self) -> list[Material]:
        return self.repository.list_materials()

    def register_generated_pptx(
        self,
        source: Path,
        *,
        filename: str,
        derivation_key: str,
    ) -> Material:
        """Idempotently register a generated deck as a normal material.

        The workflow identity and content hash deliberately determine the material
        id. A retry after the deck was copied but before the workflow result was
        persisted therefore converges on the same material instead of creating an
        orphan duplicate.
        """
        if not source.is_file() or source.suffix.lower() != ".pptx":
            raise ValueError("generated material must be an existing PPTX")
        if not derivation_key.strip():
            raise ValueError("generated material derivation_key must not be empty")
        file_hash = self._sha256(source)
        identity = hashlib.sha256(
            f"generated-pptx\0{derivation_key}\0{file_hash}".encode("utf-8")
        ).hexdigest()
        material_id = f"mat_{identity[:24]}"
        destination = self.data_dir / "raw" / material_id / "source.pptx"
        with self._job_lock:
            existing = self.repository.get_material(material_id)
            if existing is not None:
                if existing.file_type != MaterialType.PPTX or existing.file_hash != file_hash:
                    raise ValueError("generated material identity collision")
                if Path(existing.storage_path).is_file():
                    return existing

            destination.parent.mkdir(parents=True, exist_ok=True)
            temporary = destination.with_name(f".{destination.name}.{uuid4().hex}.tmp")
            try:
                shutil.copyfile(source, temporary)
                if self._sha256(temporary) != file_hash:
                    raise OSError("generated PPTX changed while it was being registered")
                os.replace(temporary, destination)
            finally:
                temporary.unlink(missing_ok=True)

            material = Material(
                id=material_id,
                filename=safe_filename(filename),
                file_type=MaterialType.PPTX,
                file_hash=file_hash,
                storage_path=str(destination),
            )
            self.repository.save_material(material)
            return material

    def get(self, material_id: str) -> Material:
        material = self.repository.get_material(material_id)
        if not material:
            raise HTTPException(404, "Material not found")
        return material

    def create_collection(
        self,
        title: str,
        material_ids: list[str],
        primary_material_id: str | None = None,
    ) -> MaterialCollection:
        if not material_ids:
            raise HTTPException(400, "Material collection requires at least one material")
        for material_id in material_ids:
            self.get(material_id)
        primary = primary_material_id or material_ids[0]
        if primary not in material_ids:
            raise HTTPException(400, "Primary material must be included in material_ids")
        collection = MaterialCollection(
            id=f"col_{uuid4().hex[:12]}",
            title=title,
            material_ids=material_ids,
            primary_material_id=primary,
        )
        self.repository.save_collection(collection)
        return collection

    def get_collection(self, collection_id: str) -> MaterialCollection:
        collection = self.repository.get_collection(collection_id)
        if not collection:
            raise HTTPException(404, "Material collection not found")
        return collection

    def list_collections(self) -> list[MaterialCollection]:
        return self.repository.list_collections()

    async def create_processing_job(
        self,
        uploads: list[UploadFile],
        *,
        title: str = "上传课程资料集",
    ) -> MaterialProcessingJob:
        if not uploads:
            raise HTTPException(400, "Processing job requires at least one material")
        materials = [await self.create(upload) for upload in uploads]
        collection = self.create_collection(title, [material.id for material in materials])
        job = MaterialProcessingJob(
            id=f"mat_job_{uuid4().hex[:12]}",
            status=MaterialProcessingJobStatus.QUEUED,
            progress=0,
            step="queued",
            message="Waiting to process materials",
            material_ids=[material.id for material in materials],
            collection_id=collection.id,
        )
        self._save_job(job)
        with self._job_lock:
            self._cancel_events[job.id] = Event()
        return job

    def get_processing_job(self, job_id: str) -> MaterialProcessingJob:
        with self._job_lock:
            job = self._jobs.get(job_id)
            if not job:
                raise HTTPException(404, "Material processing job not found")
            return job.model_copy(deep=True)

    def list_processing_jobs(self) -> list[MaterialProcessingJob]:
        with self._job_lock:
            return sorted(
                (job.model_copy(deep=True) for job in self._jobs.values()),
                key=lambda job: job.updated_at,
                reverse=True,
            )

    def processing_job_result(self, job_id: str) -> ProcessedMaterials:
        job = self.get_processing_job(job_id)
        if job.status == MaterialProcessingJobStatus.FAILED:
            raise HTTPException(422, job.error or "Material processing job failed")
        if job.status == MaterialProcessingJobStatus.CANCELED:
            raise HTTPException(409, "Material processing job was canceled")
        if job.status != MaterialProcessingJobStatus.SUCCEEDED:
            raise HTTPException(409, "Material processing job is not finished")
        collection = self.get_collection(job.collection_id) if job.collection_id else None
        items = [
            ProcessedMaterial(material=self.get(material_id), pages=self.pages(material_id))
            for material_id in job.material_ids
        ]
        return ProcessedMaterials(items=items, collection=collection)

    def pause_processing_job(self, job_id: str) -> MaterialProcessingJob:
        with self._job_lock:
            job = self._jobs.get(job_id)
            if not job:
                raise HTTPException(404, "Material processing job not found")
            if job.status in {
                MaterialProcessingJobStatus.SUCCEEDED,
                MaterialProcessingJobStatus.FAILED,
            }:
                raise HTTPException(409, "Material processing job is already finished")
            event = self._cancel_events.setdefault(job_id, Event())
            event.set()
            job.status = MaterialProcessingJobStatus.PAUSED
            job.step = "paused"
            job.message = "处理已暂停，当前进度已保存"
            job.error = None
            job.updated_at = utc_now()
        self._save_job(job)
        return job

    def resume_processing_job(self, job_id: str) -> MaterialProcessingJob:
        job = self.get_processing_job(job_id)
        if job.status != MaterialProcessingJobStatus.PAUSED:
            raise HTTPException(409, "Only paused jobs can be resumed")
        with self._job_lock:
            self._cancel_events[job_id] = Event()
        job.status = MaterialProcessingJobStatus.QUEUED
        job.step = "queued"
        job.message = "等待从已保存进度继续"
        self._save_job(job)
        return job

    def discard_processing_job(self, job_id: str) -> None:
        job = self.get_processing_job(job_id)
        if job.status == MaterialProcessingJobStatus.RUNNING:
            raise HTTPException(409, "Pause the job before discarding it")
        self._discard_processing_job(job)
        with self._job_lock:
            self._jobs.pop(job_id, None)
            self._cancel_events.pop(job_id, None)
        (self._jobs_dir / f"{job_id}.json").unlink(missing_ok=True)

    def run_processing_job(self, job_id: str) -> None:
        job = self.get_processing_job(job_id)
        with self._job_lock:
            cancel_event = self._cancel_events.setdefault(job_id, Event())
        try:
            self._raise_if_cancelled(cancel_event)
            total = max(len(job.material_ids), 1)
            job.status = MaterialProcessingJobStatus.RUNNING
            job.step = "parsing"
            job.message = "Parsing uploaded materials"
            job.progress = 5
            job.updated_at = utc_now()
            self._save_job(job)
            for index, material_id in enumerate(job.material_ids, start=1):
                self._raise_if_cancelled(cancel_event)
                material = self.get(material_id)
                if material.status == MaterialStatus.PARSED and self.pages(material_id):
                    job.progress = min(95, 5 + int((index / total) * 90))
                    self._save_job(job)
                    continue
                job.step = f"parsing:{material.filename}"
                job.message = f"Parsing {material.filename}"
                job.progress = min(95, 5 + int(((index - 1) / total) * 90))
                job.updated_at = utc_now()
                self._save_job(job)
                self.parse(material_id, cancel_event=cancel_event)
                job.progress = min(95, 5 + int((index / total) * 90))
                job.updated_at = utc_now()
                self._save_job(job)
            job.status = MaterialProcessingJobStatus.SUCCEEDED
            job.step = "completed"
            job.message = "Material processing completed"
            job.progress = 100
            job.updated_at = utc_now()
            self._save_job(job)
        except MaterialProcessingCancelled:
            job.status = MaterialProcessingJobStatus.PAUSED
            job.step = "paused"
            job.message = "处理已暂停，当前进度已保存"
            job.error = None
            job.updated_at = utc_now()
            self._save_job(job)
        except Exception as exc:  # noqa: BLE001 - processing failures become durable job state
            job.status = MaterialProcessingJobStatus.FAILED
            job.step = "failed"
            job.message = "Material processing failed"
            job.progress = 100
            job.error = str(exc)
            job.updated_at = utc_now()
            self._save_job(job)

    def pages(self, material_id: str) -> list[PageMetadata]:
        self.get(material_id)
        return self.repository.list_pages(material_id)

    def rich_parse_result(self, material_id: str) -> MaterialRichParseResult | None:
        """Read the best available MinerU JSON without exposing selection internals."""
        self.get(material_id)
        root = (self.data_dir / "processed" / material_id / "mineru").resolve()
        if not root.is_dir():
            return None
        candidates = (
            ("content_list_v2", "content_list_v2.json"),
            ("content_list", "content_list.json"),
            ("middle", "middle.json"),
        )
        for format_name, filename in candidates:
            for path in sorted(root.rglob(filename)):
                resolved = path.resolve()
                resolved.relative_to(root)
                try:
                    payload = json.loads(resolved.read_text(encoding="utf-8"))
                except (OSError, json.JSONDecodeError):
                    continue
                if isinstance(payload, (list, dict)):
                    return MaterialRichParseResult(
                        format=format_name,
                        path=resolved,
                        root=resolved.parent,
                        payload=payload,
                    )
        return None

    def delete_project(self, material_id: str) -> None:
        self.get(material_id)
        try:
            cleanup = self.repository.delete_material_project(material_id)
        except ValueError as exc:
            raise HTTPException(409, str(exc)) from exc
        for path_value in cleanup.file_paths:
            path = Path(path_value)
            try:
                resolved = path.resolve()
                resolved.relative_to(self.data_dir.resolve())
            except (OSError, ValueError):
                continue
            if resolved.is_file():
                resolved.unlink(missing_ok=True)
        presentations_root = self.data_dir / "generated" / "presentations"
        for job_id in cleanup.presentation_job_ids:
            output_dir = presentations_root / job_id
            try:
                resolved_output_dir = output_dir.resolve()
                resolved_output_dir.relative_to(presentations_root.resolve())
            except (OSError, ValueError):
                continue
            if resolved_output_dir != presentations_root.resolve():
                shutil.rmtree(resolved_output_dir, ignore_errors=True)
        shutil.rmtree(self.data_dir / "raw" / material_id, ignore_errors=True)
        shutil.rmtree(self.data_dir / "processed" / material_id, ignore_errors=True)
        checkpoint = (
            self.data_dir
            / "runtime"
            / "content_generation"
            / "checkpoints"
            / f"source_deck_{material_id}.json"
        )
        checkpoint.unlink(missing_ok=True)
        narration_dir = self.data_dir / "runtime" / "presentation_narration"
        if narration_dir.exists():
            for narration_checkpoint in narration_dir.glob(f"*_{material_id}.json"):
                narration_checkpoint.unlink(missing_ok=True)

    def parse(self, material_id: str, *, cancel_event: Event | None = None) -> list[PageMetadata]:
        material = self.get(material_id)
        material.status = "parsing"
        material.updated_at = utc_now()
        self.repository.save_material(material)
        try:
            pages = (
                self._parse_pdf(material, cancel_event)
                if material.file_type == "pdf"
                else self._parse_pptx(material, cancel_event)
            )
            material.status = "parsed"
            material.page_count = len(pages)
            material.error = None
            material.updated_at = utc_now()
            self.repository.save_material(material)
            return pages
        except MaterialProcessingCancelled:
            material.status = "uploaded"
            material.error = None
            material.updated_at = utc_now()
            self.repository.save_material(material)
            raise
        except Exception as exc:
            material.status = "failed"
            material.error = str(exc)
            material.updated_at = utc_now()
            self.repository.save_material(material)
            raise HTTPException(422, f"Material parsing failed: {exc}") from exc

    @staticmethod
    def _raise_if_cancelled(cancel_event: Event | None) -> None:
        if cancel_event and cancel_event.is_set():
            raise MaterialProcessingCancelled("Material processing canceled")

    def _discard_processing_job(self, job: MaterialProcessingJob) -> None:
        if job.collection_id:
            self.repository.delete_collection(job.collection_id)
        for material_id in job.material_ids:
            self.repository.delete_material(material_id)
            shutil.rmtree(self.data_dir / "raw" / material_id, ignore_errors=True)
            shutil.rmtree(self.data_dir / "processed" / material_id, ignore_errors=True)

    def _save_job(self, job: MaterialProcessingJob) -> None:
        with self._job_lock:
            self._jobs[job.id] = job.model_copy(deep=True)
            self._jobs_dir.mkdir(parents=True, exist_ok=True)
            path = self._jobs_dir / f"{job.id}.json"
            temporary = path.with_suffix(".json.tmp")
            temporary.write_text(job.model_dump_json(indent=2), encoding="utf-8")
            temporary.replace(path)

    @staticmethod
    def _sha256(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as source:
            for chunk in iter(lambda: source.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()

    def _parse_pdf(
        self, material: Material, cancel_event: Event | None = None
    ) -> list[PageMetadata]:
        if self.parser_backend in {"auto", "mineru"}:
            try:
                return self._parse_pdf_with_mineru(material, cancel_event)
            except MaterialProcessingCancelled:
                raise
            except Exception:
                if self.parser_backend == "mineru":
                    raise

        return self._parse_pdf_locally(material, cancel_event)

    def _parse_pdf_with_mineru(
        self, material: Material, cancel_event: Event | None = None
    ) -> list[PageMetadata]:
        output_dir = self.data_dir / "processed" / material.id / "mineru"
        output_dir.mkdir(parents=True, exist_ok=True)
        command = self._mineru_command()
        self._run_cancelable_subprocess(
            [*command, "-p", material.storage_path, "-o", str(output_dir)],
            timeout=self.mineru_timeout_seconds,
            cancel_event=cancel_event,
        )
        self._raise_if_cancelled(cancel_event)
        pages = self._metadata_from_mineru_output(material, output_dir)
        if not pages:
            raise ValueError("MinerU did not produce page metadata")
        return self._save_pages(material.id, pages)

    def _parse_pdf_locally(
        self, material: Material, cancel_event: Event | None = None
    ) -> list[PageMetadata]:
        document = fitz.open(material.storage_path)
        images_dir = self.data_dir / "processed" / material.id / "pages"
        images_dir.mkdir(parents=True, exist_ok=True)
        embedded_images = self._extract_pdf_embedded_images(material)
        result = []
        for index, page in enumerate(document):
            self._raise_if_cancelled(cancel_event)
            number = index + 1
            text = page.get_text("text").strip()
            image_path = images_dir / f"page_{number:03d}.png"
            page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(image_path)
            result.append(
                self._metadata(
                    material.id,
                    number,
                    text,
                    image_path,
                    embedded_images=embedded_images.get(number, []),
                )
            )
        document.close()
        return self._save_pages(material.id, result)

    def _metadata_from_mineru_output(
        self, material: Material, output_dir: Path
    ) -> list[PageMetadata]:
        page_images = self._render_pdf_pages(material)
        embedded_images = self._extract_pdf_embedded_images(material)
        content_list = self._load_mineru_content_list(output_dir)
        if content_list:
            texts_by_page = self._mineru_text_by_page(content_list)
        else:
            texts_by_page = self._mineru_markdown_by_page(output_dir)

        pages = []
        max_page = max(len(page_images), max(texts_by_page.keys(), default=0))
        for number in range(1, max_page + 1):
            image_path = (
                page_images[number - 1]
                if number <= len(page_images)
                else self._placeholder(self.data_dir / "processed" / material.id / "pages", number)
            )
            pages.append(
                self._metadata(
                    material.id,
                    number,
                    texts_by_page.get(number, ""),
                    image_path,
                    embedded_images=embedded_images.get(number, []),
                )
            )
        return pages

    def _extract_pdf_embedded_images(self, material: Material) -> dict[int, list[PageImage]]:
        output_dir = self.data_dir / "processed" / material.id / "embedded_images"
        output_dir.mkdir(parents=True, exist_ok=True)
        result: dict[int, list[PageImage]] = {}
        document = fitz.open(material.storage_path)
        try:
            for page_index, page in enumerate(document):
                page_no = page_index + 1
                seen_xrefs: set[int] = set()
                for image_index, image_info in enumerate(page.get_images(full=True), start=1):
                    xref = int(image_info[0])
                    if xref in seen_xrefs:
                        continue
                    seen_xrefs.add(xref)
                    try:
                        extracted = document.extract_image(xref)
                    except (RuntimeError, ValueError):
                        continue
                    image_bytes = extracted.get("image")
                    if not image_bytes:
                        continue
                    width = int(extracted.get("width") or 0)
                    height = int(extracted.get("height") or 0)
                    if width <= 1 or height <= 1:
                        continue
                    extension = str(extracted.get("ext") or "png").lower()
                    if extension == "jpeg":
                        extension = "jpg"
                    image_id = f"{material.id}_page_{page_no:03d}_image_{image_index:02d}"
                    image_path = (
                        output_dir / f"page_{page_no:03d}_image_{image_index:02d}.{extension}"
                    )
                    image_path.write_bytes(image_bytes)
                    result.setdefault(page_no, []).append(
                        PageImage(
                            id=image_id,
                            material_id=material.id,
                            page_id=f"{material.id}_page_{page_no:03d}",
                            page_no=page_no,
                            image_path=str(image_path),
                            width=width,
                            height=height,
                            description=f"Embedded image {image_index} from page {page_no}.",
                        )
                    )
        finally:
            document.close()
        return result

    def _render_pdf_pages(self, material: Material) -> list[Path]:
        document = fitz.open(material.storage_path)
        images_dir = self.data_dir / "processed" / material.id / "pages"
        images_dir.mkdir(parents=True, exist_ok=True)
        paths = []
        for index, page in enumerate(document):
            image_path = images_dir / f"page_{index + 1:03d}.png"
            page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(image_path)
            paths.append(image_path)
        document.close()
        return paths

    def _load_mineru_content_list(self, output_dir: Path) -> list[dict] | None:
        for path in sorted(output_dir.rglob("*content_list*.json")):
            payload = json.loads(path.read_text(encoding="utf-8"))
            if isinstance(payload, list):
                return [item for item in payload if isinstance(item, dict)]
        return None

    def _mineru_text_by_page(self, content_list: list[dict]) -> dict[int, str]:
        pages: dict[int, list[str]] = {}
        for item in content_list:
            page_no = self._mineru_page_no(item)
            text = self._mineru_item_text(item)
            if text:
                pages.setdefault(page_no, []).append(text)
        return {page_no: "\n".join(parts).strip() for page_no, parts in pages.items()}

    @staticmethod
    def _mineru_page_no(item: dict) -> int:
        value = item.get("page_no", item.get("page", item.get("page_idx", 0)))
        try:
            number = int(value)
        except (TypeError, ValueError):
            return 1
        return number + 1 if "page_idx" in item and "page_no" not in item else max(number, 1)

    @staticmethod
    def _mineru_item_text(item: dict) -> str:
        candidates = [
            item.get("text"),
            item.get("html"),
            item.get("table_body"),
            item.get("latex"),
            item.get("caption"),
        ]
        return "\n".join(str(value).strip() for value in candidates if value).strip()

    def _mineru_markdown_by_page(self, output_dir: Path) -> dict[int, str]:
        markdown_files = sorted(output_dir.rglob("*.md"))
        if not markdown_files:
            return {}
        text = markdown_files[0].read_text(encoding="utf-8").strip()
        if not text:
            return {}
        parts = [part.strip() for part in text.split("\f")]
        return {index + 1: part for index, part in enumerate(parts) if part}

    def _mineru_command(self) -> list[str]:
        return shlex.split(self.mineru_command, posix=os.name != "nt") or ["mineru"]

    def _parse_pptx(
        self, material: Material, cancel_event: Event | None = None
    ) -> list[PageMetadata]:
        presentation = Presentation(material.storage_path)
        texts = []
        for slide in presentation.slides:
            self._raise_if_cancelled(cancel_event)
            texts.append("\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text")))
        images = self._render_pptx(material, len(texts), cancel_event)
        pages = [
            self._metadata(material.id, index + 1, text, images[index])
            for index, text in enumerate(texts)
        ]
        return self._save_pages(material.id, pages)

    def _render_pptx(
        self,
        material: Material,
        page_count: int,
        cancel_event: Event | None = None,
    ) -> list[Path]:
        images_dir = self.data_dir / "processed" / material.id / "pages"
        images_dir.mkdir(parents=True, exist_ok=True)
        try:
            with tempfile.TemporaryDirectory() as temp_dir:
                self._run_cancelable_subprocess(
                    [
                        "soffice",
                        "--headless",
                        "--convert-to",
                        "pdf",
                        "--outdir",
                        temp_dir,
                        material.storage_path,
                    ],
                    timeout=60,
                    cancel_event=cancel_event,
                )
                pdf_path = Path(temp_dir) / f"{Path(material.storage_path).stem}.pdf"
                document = fitz.open(pdf_path)
                paths = []
                for index, page in enumerate(document):
                    self._raise_if_cancelled(cancel_event)
                    path = images_dir / f"page_{index + 1:03d}.png"
                    page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(path)
                    paths.append(path)
                document.close()
                if len(paths) == page_count:
                    return paths
        except MaterialProcessingCancelled:
            raise
        except (FileNotFoundError, subprocess.SubprocessError, fitz.FileDataError):
            pass
        return [self._placeholder(images_dir, index + 1) for index in range(page_count)]

    def _run_cancelable_subprocess(
        self,
        command: list[str],
        *,
        timeout: float,
        cancel_event: Event | None,
    ) -> None:
        process = subprocess.Popen(
            command,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )
        started_at = time.monotonic()
        try:
            while process.poll() is None:
                if cancel_event and cancel_event.wait(0.1):
                    process.terminate()
                    try:
                        process.wait(timeout=3)
                    except subprocess.TimeoutExpired:
                        process.kill()
                    raise MaterialProcessingCancelled("Material processing canceled")
                if time.monotonic() - started_at > timeout:
                    process.kill()
                    raise subprocess.TimeoutExpired(command, timeout)
                if not cancel_event:
                    time.sleep(0.1)
            stdout, stderr = process.communicate()
            if process.returncode:
                raise subprocess.CalledProcessError(
                    process.returncode,
                    command,
                    output=stdout,
                    stderr=stderr,
                )
        finally:
            if process.poll() is None:
                process.kill()

    @staticmethod
    def _placeholder(directory: Path, number: int) -> Path:
        path = directory / f"page_{number:03d}.png"
        image = Image.new("RGB", (1280, 720), "white")
        ImageDraw.Draw(image).text((60, 60), f"Slide {number}", fill="black")
        image.save(path)
        return path

    def _metadata(
        self,
        material_id: str,
        number: int,
        text: str,
        image_path: Path,
        *,
        embedded_images: list[PageImage] | None = None,
    ) -> PageMetadata:
        title = next(
            (line.strip() for line in text.splitlines() if line.strip()), f"第 {number} 页"
        )
        ref = SourceRef(
            material_id=material_id,
            page_id=f"{material_id}_page_{number:03d}",
            page_no=number,
            text_span=text[:120] or None,
            image_path=str(image_path),
        )
        return PageMetadata(
            id=f"{material_id}_page_{number:03d}",
            material_id=material_id,
            page_no=number,
            title=title[:100],
            raw_text=text,
            image_path=str(image_path),
            embedded_images=embedded_images or [],
            source_refs=[ref],
        )

    def _save_pages(self, material_id: str, pages: list[PageMetadata]) -> list[PageMetadata]:
        self.repository.replace_pages(material_id, pages)
        return pages
