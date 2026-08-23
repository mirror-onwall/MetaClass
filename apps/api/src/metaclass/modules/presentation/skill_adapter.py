import json
import os
import platform
import re
import shutil
import subprocess
import tempfile
import textwrap
from pathlib import Path
from uuid import uuid4

import fitz
from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from metaclass.modules.presentation.schemas import PPTArtifact, PPTSlideImage, PresentationPlan
from metaclass.modules.presentation.brand_palette import BRAND_PALETTE
from metaclass.modules.presentation.themes import (
    PresentationTheme,
    apply_presentation_theme,
    get_presentation_theme,
)


class PPTSkillAdapter:
    """Boundary for the PPT generation skill.

    This first implementation turns PresentationPlan into a real, simple PPTX
    using the project's existing python-pptx dependency. The structured request
    file is kept so a richer external skill can replace this renderer later.
    """

    def __init__(self, *, libreoffice_bin: str | Path | None = None) -> None:
        configured = (
            str(libreoffice_bin).strip()
            if libreoffice_bin is not None
            else os.getenv("METACLASS_LIBREOFFICE_BIN", "").strip()
        )
        self.libreoffice_bin = configured or None

    def prepare_request(
        self,
        *,
        plan: PresentationPlan,
        job_id: str,
        output_dir: Path,
        theme: PresentationTheme | None = None,
    ) -> PPTArtifact:
        selected_theme = theme or get_presentation_theme()
        themed_plan = apply_presentation_theme(plan, selected_theme)
        output_dir.mkdir(parents=True, exist_ok=True)
        pptx_path = output_dir / "deck.pptx"
        speaker_scripts_path = output_dir / "speaker_scripts.json"
        self._render_basic_pptx(themed_plan, pptx_path)
        slide_images = self._render_slide_images(
            themed_plan,
            pptx_path,
            output_dir / "slides",
        )
        self._write_speaker_scripts(plan, speaker_scripts_path)
        request_path = output_dir / "skill_request.json"
        request_path.write_text(
            json.dumps(
                {
                    "presentation_plan": plan.model_dump(mode="json"),
                    "theme": selected_theme.prompt_payload(),
                    "planner_skill": str(Path(__file__).with_name("pptx_skill") / "SKILL.md"),
                    "expected_output": str(pptx_path),
                    "speaker_scripts_output": str(speaker_scripts_path),
                    "instructions": [
                        "Generate a polished PPTX from presentation_plan.slides.",
                        "Use each slide.speaker_script as speaker notes.",
                        "Render each slide.layout and slide.visual_payload as the visible structure.",
                    ],
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        return PPTArtifact(
            id=f"ppt_artifact_{uuid4().hex[:12]}",
            job_id=job_id,
            presentation_plan_id=plan.id,
            pptx_path=str(pptx_path),
            skill_request_path=str(request_path),
            slide_images=slide_images,
        )

    def render_declarative_pptx(self, plan: PresentationPlan, destination: Path) -> None:
        """Compile an already validated declarative scene into an editable PPTX."""
        missing = [slide.id for slide in plan.slides if not slide.elements]
        if missing:
            raise ValueError(
                "Declarative PPTX rendering requires elements on every slide: " + ", ".join(missing)
            )
        destination.parent.mkdir(parents=True, exist_ok=True)
        self._render_basic_pptx(plan, destination)

    @staticmethod
    def _write_speaker_scripts(plan: PresentationPlan, destination: Path) -> None:
        destination.write_text(
            json.dumps(
                {
                    "presentation_plan_id": plan.id,
                    "slides": [
                        {
                            "slide_id": slide.id,
                            "order": slide.order,
                            "title": slide.title,
                            "speaker_script": slide.speaker_script,
                        }
                        for slide in plan.slides
                    ],
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )

    def prepare_external_pptx(
        self,
        *,
        plan: PresentationPlan,
        job_id: str,
        output_dir: Path,
        pptx_path: Path,
        provider_name: str,
        provider_metadata: dict,
        external_slide_images: list[PPTSlideImage] | None = None,
        preview_plan: PresentationPlan | None = None,
    ) -> PPTArtifact:
        """Package and preview a PPTX produced by an external presentation service."""
        output_dir.mkdir(parents=True, exist_ok=True)
        speaker_scripts_path = output_dir / "speaker_scripts.json"
        self._write_speaker_scripts(plan, speaker_scripts_path)
        request_path = output_dir / "skill_request.json"
        request_path.write_text(
            json.dumps(
                {
                    "presentation_plan": plan.model_dump(mode="json"),
                    "provider": provider_name,
                    "provider_metadata": provider_metadata,
                    "speaker_scripts_output": str(speaker_scripts_path),
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )
        slide_images = external_slide_images
        if slide_images is None:
            rendering_plan = preview_plan or plan
            slide_images = self._render_slide_images(
                rendering_plan,
                pptx_path,
                output_dir / "slides",
                allow_placeholder=False,
                allow_declarative_fallback=preview_plan is not None,
            )
        return PPTArtifact(
            id=f"ppt_artifact_{uuid4().hex[:12]}",
            job_id=job_id,
            presentation_plan_id=plan.id,
            pptx_path=str(pptx_path),
            skill_request_path=str(request_path),
            slide_images=slide_images,
        )

    def _render_slide_images(
        self,
        plan: PresentationPlan,
        pptx_path: Path,
        output_dir: Path,
        *,
        allow_placeholder: bool = True,
        allow_declarative_fallback: bool = True,
    ) -> list[PPTSlideImage]:
        output_dir.mkdir(parents=True, exist_ok=True)
        operating_system = platform.system().lower()
        has_declarative_scenes = all(slide.elements for slide in plan.slides)

        # Codex decks already have an authoritative, validated SlideElement scene.
        # On macOS render that scene immediately: it avoids Keynote automation
        # permissions/timeouts and keeps the browser preview deterministic. External
        # decks without a scene still use the native PPTX renderer chain below.
        if operating_system == "darwin" and allow_declarative_fallback and has_declarative_scenes:
            self._clear_rendered_slide_images(output_dir)
            (output_dir.parent / "render_error.txt").unlink(missing_ok=True)
            return self._render_declarative_preview_images(plan, output_dir)

        renderers = []
        if operating_system == "windows":
            renderers.append(("PowerPoint", self._render_with_powerpoint))
        if operating_system == "darwin":
            # Keynote uses macOS-native fonts and produces one PDF page per slide.
            renderers.append(("Keynote", self._render_with_keynote))
        renderers.append(("LibreOffice", self._render_with_libreoffice))

        renderer_errors: list[tuple[str, Exception]] = []
        for renderer_name, renderer in renderers:
            self._clear_rendered_slide_images(output_dir)
            try:
                slide_images = renderer(plan, pptx_path, output_dir)
                (output_dir.parent / "render_error.txt").unlink(missing_ok=True)
                return slide_images
            except (
                FileNotFoundError,
                OSError,
                subprocess.SubprocessError,
                fitz.FileDataError,
                RuntimeError,
            ) as exc:
                renderer_errors.append((renderer_name, exc))

        self._clear_rendered_slide_images(output_dir)
        error = RuntimeError(
            "; ".join(f"{name}: {exc}" for name, exc in renderer_errors)
            or "No PPTX preview renderer is available"
        )
        (output_dir.parent / "render_error.txt").write_text(
            "\n\n".join(
                f"[{name}] {PPTSkillAdapter._render_error_message(exc)}"
                for name, exc in renderer_errors
            )
            or PPTSkillAdapter._render_error_message(error),
            encoding="utf-8",
        )

        if allow_declarative_fallback and has_declarative_scenes:
            return self._render_declarative_preview_images(plan, output_dir)
        if allow_placeholder:
            return self._render_placeholder_images(plan, output_dir)
        raise RuntimeError(
            "Real PPTX preview rendering failed; placeholder previews are disabled"
        ) from error

    @staticmethod
    def _clear_rendered_slide_images(output_dir: Path) -> None:
        for image_path in output_dir.glob("slide_*.png"):
            image_path.unlink(missing_ok=True)

    def _resolve_libreoffice_binary(self) -> str:
        candidates = [
            self.libreoffice_bin,
            shutil.which("soffice"),
            shutil.which("libreoffice"),
            "/Applications/LibreOffice.app/Contents/MacOS/soffice",
            "/opt/homebrew/bin/soffice",
            "/usr/local/bin/soffice",
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ]
        for candidate in candidates:
            if candidate and Path(candidate).is_file():
                return str(Path(candidate))
        raise FileNotFoundError(
            "LibreOffice was not found; set METACLASS_LIBREOFFICE_BIN to its soffice binary"
        )

    def _render_with_libreoffice(
        self,
        plan: PresentationPlan,
        pptx_path: Path,
        output_dir: Path,
    ) -> list[PPTSlideImage]:
        libreoffice = self._resolve_libreoffice_binary()
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            libreoffice_profile = temp_path / "lo_profile"
            libreoffice_profile.mkdir(parents=True, exist_ok=True)
            completed = subprocess.run(
                [
                    libreoffice,
                    "--headless",
                    "--nologo",
                    "--nodefault",
                    "--nofirststartwizard",
                    "--norestore",
                    f"-env:UserInstallation={libreoffice_profile.as_uri()}",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(temp_path),
                    str(pptx_path.resolve()),
                ],
                check=True,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=90,
            )
            pdf_path = temp_path / f"{pptx_path.stem}.pdf"
            if not pdf_path.is_file():
                details = "\n".join(
                    part.strip()
                    for part in (completed.stdout, completed.stderr)
                    if part and part.strip()
                )
                raise FileNotFoundError(
                    f"LibreOffice converted PDF was not found: {pdf_path}"
                    + (f"\n{details}" if details else "")
                )
            return PPTSkillAdapter._render_pdf_pages(plan, pdf_path, output_dir)

    @staticmethod
    def _render_with_keynote(
        plan: PresentationPlan,
        pptx_path: Path,
        output_dir: Path,
    ) -> list[PPTSlideImage]:
        osascript = Path(shutil.which("osascript") or "/usr/bin/osascript")
        keynote_app = Path("/Applications/Keynote.app")
        if not osascript.is_file():
            raise FileNotFoundError("macOS osascript was not found")
        if not keynote_app.exists():
            raise FileNotFoundError("Keynote was not found in /Applications")

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            script_path = temp_path / "render_keynote.applescript"
            pdf_path = temp_path / "deck.pdf"
            script_path.write_text(
                textwrap.dedent(
                    """
                    on run argv
                        set inputPath to item 1 of argv
                        set outputPath to item 2 of argv
                        tell application id "com.apple.iWork.Keynote"
                            set openedDocument to open POSIX file inputPath
                            try
                                export openedDocument to POSIX file outputPath as PDF
                            on error errorMessage number errorNumber
                                close openedDocument saving no
                                error errorMessage number errorNumber
                            end try
                            close openedDocument saving no
                        end tell
                    end run
                    """
                ).strip(),
                encoding="utf-8",
            )
            subprocess.run(
                [
                    str(osascript),
                    str(script_path),
                    str(pptx_path.resolve()),
                    str(pdf_path),
                ],
                check=True,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=120,
            )
            if not pdf_path.is_file():
                raise FileNotFoundError(f"Keynote exported PDF was not found: {pdf_path}")
            return PPTSkillAdapter._render_pdf_pages(plan, pdf_path, output_dir)

    @staticmethod
    def _render_with_powerpoint(
        plan: PresentationPlan,
        pptx_path: Path,
        output_dir: Path,
    ) -> list[PPTSlideImage]:
        powershell = Path(r"C:\Windows\System32\WindowsPowerShell\v1.0\powershell.exe")
        if not powershell.exists():
            raise FileNotFoundError("Windows PowerShell was not found")

        with tempfile.TemporaryDirectory() as temp_dir:
            script_path = Path(temp_dir) / "render_powerpoint.ps1"
            script_path.write_text(
                textwrap.dedent(
                    """
                    param(
                        [Parameter(Mandatory=$true)][string]$InputPath,
                        [Parameter(Mandatory=$true)][string]$OutputDir
                    )
                    $ErrorActionPreference = 'Stop'
                    $powerpoint = New-Object -ComObject PowerPoint.Application
                    $presentation = $null
                    try {
                        $presentation = $powerpoint.Presentations.Open(
                            $InputPath, $true, $false, $false
                        )
                        for ($index = 1; $index -le $presentation.Slides.Count; $index++) {
                            $outputPath = Join-Path $OutputDir (
                                'slide_{0:D3}.png' -f $index
                            )
                            $presentation.Slides.Item($index).Export(
                                $outputPath, 'PNG', 1600, 900
                            )
                        }
                    }
                    finally {
                        if ($null -ne $presentation) {
                            $presentation.Close()
                            [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject(
                                $presentation
                            )
                        }
                        $powerpoint.Quit()
                        [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject(
                            $powerpoint
                        )
                    }
                    """
                ).strip(),
                encoding="utf-8-sig",
            )
            completed = subprocess.run(
                [
                    str(powershell),
                    "-NoProfile",
                    "-NonInteractive",
                    "-ExecutionPolicy",
                    "Bypass",
                    "-File",
                    str(script_path),
                    "-InputPath",
                    str(pptx_path.resolve()),
                    "-OutputDir",
                    str(output_dir.resolve()),
                ],
                check=True,
                capture_output=True,
                timeout=90,
            )
            if completed.stderr:
                stderr = completed.stderr.decode("utf-8", errors="replace").strip()
                if stderr:
                    raise RuntimeError(stderr)

        slide_images = []
        for index, slide_plan in enumerate(plan.slides, start=1):
            image_path = output_dir / f"slide_{index:03d}.png"
            if not image_path.exists():
                raise RuntimeError(f"PowerPoint preview page-count mismatch: missing page {index}")
            with Image.open(image_path) as image:
                image.verify()
                width, height = image.size
            slide_images.append(
                PPTSlideImage(
                    slide_id=slide_plan.id,
                    slide_no=index,
                    image_path=str(image_path),
                    width=width,
                    height=height,
                )
            )
        return slide_images

    @staticmethod
    def _render_error_message(exc: Exception) -> str:
        message = f"PPTX to slide image rendering failed: {exc}"
        if isinstance(exc, subprocess.CalledProcessError):

            def output_text(value) -> str:
                if not value:
                    return ""
                if isinstance(value, bytes):
                    return value.decode("utf-8", errors="replace")
                return str(value)

            stderr = output_text(exc.stderr)
            stdout = output_text(exc.stdout)
            details = "\n".join(part for part in [stdout.strip(), stderr.strip()] if part)
            if details:
                message = f"{message}\n{details}"
        return message

    @staticmethod
    def _render_pdf_pages(
        plan: PresentationPlan,
        pdf_path: Path,
        output_dir: Path,
    ) -> list[PPTSlideImage]:
        document = fitz.open(pdf_path)
        slide_images = []
        try:
            if document.page_count != len(plan.slides):
                raise fitz.FileDataError("Rendered slide count does not match PresentationPlan")
            for index, page in enumerate(document):
                slide_plan = plan.slides[index]
                image_path = output_dir / f"slide_{index + 1:03d}.png"
                pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                pixmap.save(image_path)
                slide_images.append(
                    PPTSlideImage(
                        slide_id=slide_plan.id,
                        slide_no=index + 1,
                        image_path=str(image_path),
                        width=pixmap.width,
                        height=pixmap.height,
                    )
                )
        finally:
            document.close()
        return slide_images

    @staticmethod
    def _load_preview_font(
        size: int,
        font_role: str = "sans",
        *,
        bold: bool = False,
    ) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
        role_paths = {
            "sans": {
                False: [
                    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
                    "/usr/share/fonts/opentype/noto/NotoSansCJKsc-Regular.otf",
                    "C:/Windows/Fonts/msyh.ttc",
                    "/System/Library/Fonts/PingFang.ttc",
                ],
                True: [
                    "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc",
                    "/usr/share/fonts/opentype/noto/NotoSansCJKsc-Bold.otf",
                    "C:/Windows/Fonts/msyhbd.ttc",
                    "/System/Library/Fonts/PingFang.ttc",
                ],
            },
            "display": {
                False: [
                    "/usr/share/fonts/opentype/noto/NotoSansCJK-Black.ttc",
                    "/usr/share/fonts/opentype/noto/NotoSansCJKsc-Black.otf",
                    "C:/Windows/Fonts/simhei.ttf",
                    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
                    "/usr/share/fonts/opentype/noto/NotoSansCJKsc-Regular.otf",
                    "C:/Windows/Fonts/msyh.ttc",
                    "/System/Library/Fonts/PingFang.ttc",
                ],
                True: [
                    "/usr/share/fonts/opentype/noto/NotoSansCJK-Black.ttc",
                    "/usr/share/fonts/opentype/noto/NotoSansCJKsc-Black.otf",
                    "C:/Windows/Fonts/simhei.ttf",
                    "C:/Windows/Fonts/msyhbd.ttc",
                    "/System/Library/Fonts/PingFang.ttc",
                ],
            },
            "serif": {
                False: [
                    "/usr/share/fonts/opentype/noto/NotoSerifCJK-Regular.ttc",
                    "/usr/share/fonts/opentype/noto/NotoSerifCJKsc-Regular.otf",
                    "C:/Windows/Fonts/simsun.ttc",
                    "/System/Library/Fonts/Supplemental/Songti.ttc",
                ],
                True: [
                    "/usr/share/fonts/opentype/noto/NotoSerifCJK-Bold.ttc",
                    "/usr/share/fonts/opentype/noto/NotoSerifCJKsc-Bold.otf",
                    "C:/Windows/Fonts/simsun.ttc",
                    "/System/Library/Fonts/Supplemental/Songti.ttc",
                ],
            },
            "handwritten": {
                False: [
                    "C:/Windows/Fonts/simkai.ttf",
                    "/System/Library/Fonts/Kaiti.ttc",
                    "/System/Library/Fonts/Supplemental/Kaiti.ttc",
                    "/usr/share/fonts/opentype/noto/NotoSerifCJK-Regular.ttc",
                ],
                True: [
                    "C:/Windows/Fonts/simkai.ttf",
                    "/System/Library/Fonts/Kaiti.ttc",
                    "/System/Library/Fonts/Supplemental/Kaiti.ttc",
                    "/usr/share/fonts/opentype/noto/NotoSerifCJK-Bold.ttc",
                ],
            },
            # A CJK-capable sans face is preferred to a latin-only monospace
            # face because a single PIL font cannot perform glyph fallback.
            "mono": {
                False: [
                    "/usr/share/fonts/opentype/noto/NotoSansMonoCJK-Regular.ttc",
                    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
                    "C:/Windows/Fonts/msyh.ttc",
                    "/System/Library/Fonts/PingFang.ttc",
                ],
                True: [
                    "/usr/share/fonts/opentype/noto/NotoSansMonoCJK-Bold.ttc",
                    "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc",
                    "C:/Windows/Fonts/msyhbd.ttc",
                    "/System/Library/Fonts/PingFang.ttc",
                ],
            },
        }
        for path in [
            *role_paths.get(font_role, role_paths["sans"])[bold],
            "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
            "/usr/share/fonts/opentype/noto/NotoSansCJKsc-Regular.otf",
            "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
            "/usr/share/fonts/truetype/noto/NotoSansSC-Regular.otf",
            "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
            "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
            "C:/Windows/Fonts/msyh.ttc",
            "C:/Windows/Fonts/simhei.ttf",
            "C:/Windows/Fonts/arial.ttf",
            "/System/Library/Fonts/PingFang.ttc",
            "/System/Library/Fonts/STHeiti Medium.ttc",
            "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
            "/System/Library/Fonts/Supplemental/Songti.ttc",
            "/System/Library/Fonts/Helvetica.ttc",
        ]:
            try:
                return ImageFont.truetype(path, size=size)
            except OSError:
                continue
        return ImageFont.load_default()

    @staticmethod
    def _preview_text_width(
        draw: ImageDraw.ImageDraw,
        text: str,
        *,
        font: ImageFont.FreeTypeFont | ImageFont.ImageFont,
        stroke_width: int,
    ) -> float:
        if not text:
            return 0
        left, _, right, _ = draw.textbbox(
            (0, 0),
            text,
            font=font,
            stroke_width=stroke_width,
        )
        return right - left

    @staticmethod
    def _wrap_preview_text(
        draw: ImageDraw.ImageDraw,
        text: str,
        *,
        font: ImageFont.FreeTypeFont | ImageFont.ImageFont,
        max_width: int,
        stroke_width: int,
    ) -> str:
        """Wrap exact copy by measured glyph width without shrinking or truncating it."""

        def fits(value: str) -> bool:
            return (
                PPTSkillAdapter._preview_text_width(
                    draw,
                    value,
                    font=font,
                    stroke_width=stroke_width,
                )
                <= max_width
            )

        def split_oversized(value: str) -> list[str]:
            pieces: list[str] = []
            current = ""
            for character in value:
                candidate = f"{current}{character}"
                if current and not fits(candidate):
                    pieces.append(current)
                    current = character
                else:
                    current = candidate
            if current or not pieces:
                pieces.append(current)
            return pieces

        wrapped: list[str] = []
        # Preserve author-provided line breaks. Within each paragraph, keep
        # western words together when possible and allow CJK text to break at
        # character boundaries only when its measured width requires it.
        for paragraph in text.split("\n"):
            if not paragraph:
                wrapped.append("")
                continue
            tokens = re.findall(r"\s+|[^\s]+", paragraph)
            current = ""
            for token in tokens:
                candidate = f"{current}{token}"
                if fits(candidate):
                    current = candidate
                    continue
                if current.strip():
                    wrapped.append(current.rstrip())
                    current = ""
                token = token.lstrip()
                if not token:
                    continue
                if fits(token):
                    current = token
                    continue
                pieces = split_oversized(token)
                wrapped.extend(pieces[:-1])
                current = pieces[-1]
            wrapped.append(current.rstrip())
        return "\n".join(wrapped)

    @staticmethod
    def _font_faces(font_role: str = "sans") -> tuple[str, str]:
        system = platform.system().lower()
        if system == "darwin":
            # LibreOffice on macOS may classify Chinese glyphs as latin text
            # when importing python-pptx output. Using the CJK-capable face for
            # both font slots prevents it from substituting empty Arial glyphs.
            return {
                "display": ("Avenir Next", "Heiti SC"),
                "serif": ("Songti SC", "Songti SC"),
                "handwritten": ("Kaiti SC", "Kaiti SC"),
                "mono": ("Menlo", "PingFang SC"),
            }.get(font_role, ("PingFang SC", "PingFang SC"))
        if system == "windows":
            return {
                "display": ("Arial Black", "SimHei"),
                "serif": ("SimSun", "SimSun"),
                "handwritten": ("KaiTi", "KaiTi"),
                "mono": ("Consolas", "Microsoft YaHei"),
            }.get(font_role, ("Microsoft YaHei", "Microsoft YaHei"))
        return {
            "display": ("Noto Sans CJK SC", "Noto Sans CJK SC"),
            "serif": ("Noto Serif CJK SC", "Noto Serif CJK SC"),
            "handwritten": ("Noto Serif CJK SC", "Noto Serif CJK SC"),
            "mono": ("Noto Sans Mono CJK SC", "Noto Sans CJK SC"),
        }.get(font_role, ("Noto Sans CJK SC", "Noto Sans CJK SC"))

    @staticmethod
    def _apply_font(
        paragraph,
        *,
        size: Pt,
        color: RGBColor,
        bold: bool = False,
        opacity: int = 100,
        font_role: str = "sans",
    ) -> None:
        latin_font, cjk_font = PPTSkillAdapter._font_faces(font_role)
        paragraph.font.size = size
        paragraph.font.bold = bold
        paragraph.font.name = latin_font
        paragraph.font.color.rgb = color
        PPTSkillAdapter._apply_color_opacity(paragraph.font.color, opacity)

        run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
        run.font.name = latin_font
        run.font.size = size
        run.font.bold = bold
        run.font.color.rgb = color
        PPTSkillAdapter._apply_color_opacity(run.font.color, opacity)
        rpr = run._r.get_or_add_rPr()
        for tag, font_name in {
            "a:latin": latin_font,
            "a:ea": cjk_font,
            "a:cs": latin_font,
        }.items():
            element = rpr.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                rpr.append(element)
            element.set("typeface", font_name)

    @staticmethod
    def _apply_color_opacity(color_format, opacity: int) -> None:
        """Apply DrawingML alpha because python-pptx has no public transparency API."""

        normalized = max(0, min(100, int(opacity)))
        try:
            color_element = color_format._color._xClr
        except AttributeError:
            return
        for alpha in list(color_element.findall(qn("a:alpha"))):
            color_element.remove(alpha)
        if normalized < 100:
            alpha = OxmlElement("a:alpha")
            alpha.set("val", str(normalized * 1000))
            color_element.append(alpha)

    @staticmethod
    def _fit_text(text: str, *, max_chars: int, max_lines: int) -> str:
        clean = " ".join(text.split())
        if not clean:
            return ""
        lines = textwrap.wrap(
            clean,
            width=max_chars,
            max_lines=max_lines,
            placeholder="...",
            break_long_words=True,
            break_on_hyphens=False,
        )
        return "\n".join(lines) if lines else clean[:max_chars]

    @staticmethod
    def _render_declarative_preview_images(
        plan: PresentationPlan,
        output_dir: Path,
    ) -> list[PPTSlideImage]:
        """Render the exact validated SlideElement scenes used to build the PPTX."""

        missing = [slide.id for slide in plan.slides if not slide.elements]
        if missing:
            raise ValueError(
                "Declarative preview requires elements on every slide: " + ", ".join(missing)
            )

        slide_images = []
        for index, slide_plan in enumerate(plan.slides, start=1):
            image_path = output_dir / f"slide_{index:03d}.png"
            image = PPTSkillAdapter._render_scene_preview(slide_plan)
            image.save(image_path)
            slide_images.append(
                PPTSlideImage(
                    slide_id=slide_plan.id,
                    slide_no=index,
                    image_path=str(image_path),
                    width=image.width,
                    height=image.height,
                )
            )
        return slide_images

    @staticmethod
    def _render_placeholder_images(
        plan: PresentationPlan,
        output_dir: Path,
    ) -> list[PPTSlideImage]:
        slide_images = []
        for index, slide_plan in enumerate(plan.slides, start=1):
            image_path = output_dir / f"slide_{index:03d}.png"
            if slide_plan.elements:
                image = PPTSkillAdapter._render_scene_preview(slide_plan)
                image.save(image_path)
                slide_images.append(
                    PPTSlideImage(
                        slide_id=slide_plan.id,
                        slide_no=index,
                        image_path=str(image_path),
                        width=image.width,
                        height=image.height,
                    )
                )
                continue
            image = Image.new("RGB", (1280, 720), "white")
            draw = ImageDraw.Draw(image)
            title_font = PPTSkillAdapter._load_preview_font(34)
            body_font = PPTSkillAdapter._load_preview_font(24)
            caption_font = PPTSkillAdapter._load_preview_font(20)
            draw.rectangle((0, 0, 28, 720), fill="#D1495B")
            title = PPTSkillAdapter._fit_text(slide_plan.title, max_chars=42, max_lines=2)
            draw.multiline_text(
                (70, 54),
                title,
                fill="#1F2937",
                font=title_font,
                spacing=8,
            )
            y = 150
            for point in slide_plan.key_points[:5]:
                text = PPTSkillAdapter._fit_text(point, max_chars=46, max_lines=2)
                bullet = f"- {text}"
                draw.multiline_text(
                    (90, y),
                    bullet,
                    fill="#374151",
                    font=body_font,
                    spacing=7,
                )
                y += 38 * (bullet.count("\n") + 1) + 14
                if y > 480:
                    break
            draw.rectangle((720, 150, 1190, 500), outline="#D1495B", width=3)
            draw.text((760, 190), "Visual direction", fill="#D1495B", font=caption_font)
            visual = PPTSkillAdapter._fit_text(
                slide_plan.suggested_visual,
                max_chars=34,
                max_lines=8,
            )
            draw.multiline_text(
                (760, 230),
                visual,
                fill="#374151",
                font=caption_font,
                spacing=7,
            )
            image.save(image_path)
            slide_images.append(
                PPTSlideImage(
                    slide_id=slide_plan.id,
                    slide_no=index,
                    image_path=str(image_path),
                    width=image.width,
                    height=image.height,
                )
            )
        return slide_images

    @staticmethod
    def _render_scene_preview(slide_plan) -> Image.Image:
        """Render the declarative scene for the browser when LibreOffice is absent."""
        width, height = 1280, 720
        image = Image.new("RGB", (width, height), f"#{slide_plan.background}")
        draw = ImageDraw.Draw(image)
        for element in sorted(slide_plan.elements, key=lambda item: item.z):
            x0, y0 = round(element.x * width), round(element.y * height)
            x1 = round((element.x + element.w) * width)
            y1 = round((element.y + element.h) * height)
            style = element.style
            fill = f"#{style.fill}" if style.fill else None
            outline = f"#{style.line_color}" if style.line_color else None

            if element.type == "shape":
                if element.shape == "oval":
                    draw.ellipse(
                        (x0, y0, x1, y1),
                        fill=fill,
                        outline=outline,
                        width=max(1, round(style.line_width)),
                    )
                elif element.shape == "chevron":
                    middle_y = (y0 + y1) // 2
                    shoulder_x = x0 + round((x1 - x0) * 0.72)
                    notch_x = x0 + round((x1 - x0) * 0.28)
                    points = [
                        (x0, y0),
                        (shoulder_x, y0),
                        (x1, middle_y),
                        (shoulder_x, y1),
                        (x0, y1),
                        (notch_x, middle_y),
                    ]
                    draw.polygon(points, fill=fill)
                    if outline:
                        draw.line(
                            [*points, points[0]],
                            fill=outline,
                            width=max(1, round(style.line_width)),
                            joint="curve",
                        )
                else:
                    radius = 16 if element.shape == "rounded_rectangle" else 0
                    draw.rounded_rectangle(
                        (x0, y0, x1, y1),
                        radius=radius,
                        fill=fill,
                        outline=outline,
                        width=max(1, round(style.line_width)),
                    )
                continue
            if element.type == "line":
                draw.line(
                    (x0, y0, x1, y1),
                    fill=outline or f"#{style.color}",
                    width=max(1, round(style.line_width * 2)),
                )
                continue
            if element.type == "image" and element.image_path:
                try:
                    with Image.open(element.image_path) as opened:
                        source = ImageOps.exif_transpose(opened).convert("RGBA")
                    target_size = (max(1, x1 - x0), max(1, y1 - y0))
                    if element.image_fit == "contain":
                        rendered = ImageOps.contain(source, target_size)
                        paste_x = x0 + (target_size[0] - rendered.width) // 2
                        paste_y = y0 + (target_size[1] - rendered.height) // 2
                    else:
                        rendered = ImageOps.fit(source, target_size)
                        paste_x, paste_y = x0, y0
                    image.paste(rendered, (paste_x, paste_y), rendered)
                except (FileNotFoundError, OSError):
                    draw.rectangle((x0, y0, x1, y1), fill="#E5E7EB", outline="#94A3B8", width=2)
                continue
            if element.type == "text":
                if fill:
                    draw.rounded_rectangle(
                        (x0, y0, x1, y1),
                        radius=0,
                        fill=fill,
                        outline=(outline if style.line_color and style.line_width > 0 else None),
                        width=max(1, round(style.line_width)),
                    )
                text = element.text or "\n".join(element.items)
                # The preview canvas is 96 px/in while PowerPoint font sizes
                # are points (72/in). Preserve the Skill-selected size instead
                # of running a second auto-fit pass that changes its design.
                font_size = max(1, round(style.font_size * 96 / 72))
                font = PPTSkillAdapter._load_preview_font(
                    font_size,
                    style.font_role,
                    bold=style.bold,
                )
                font_path = str(getattr(font, "path", "")).lower()
                has_bold_face = any(
                    marker in Path(font_path).stem
                    for marker in ("bold", "black", "heavy", "bd")
                )
                stroke_width = 1 if style.bold and not has_bold_face else 0
                margin_x = round(style.text_margin_x * 96)
                margin_y = round(style.text_margin_y * 96)
                inner_x0, inner_y0 = x0 + margin_x, y0 + margin_y
                inner_x1, inner_y1 = x1 - margin_x, y1 - margin_y
                fitted = PPTSkillAdapter._wrap_preview_text(
                    draw,
                    text,
                    font=font,
                    max_width=max(1, inner_x1 - inner_x0),
                    stroke_width=stroke_width,
                )
                spacing = max(1, round(font_size * 0.18))
                bbox = draw.multiline_textbbox(
                    (0, 0),
                    fitted,
                    font=font,
                    spacing=spacing,
                    align=style.align,
                    stroke_width=stroke_width,
                )
                text_width = bbox[2] - bbox[0]
                text_height = bbox[3] - bbox[1]
                if style.align == "center":
                    target_x = inner_x0 + (inner_x1 - inner_x0 - text_width) / 2
                elif style.align == "right":
                    target_x = inner_x1 - text_width
                else:
                    target_x = inner_x0
                if style.valign == "middle":
                    target_y = inner_y0 + (inner_y1 - inner_y0 - text_height) / 2
                elif style.valign == "bottom":
                    target_y = inner_y1 - text_height
                else:
                    target_y = inner_y0
                draw.multiline_text(
                    (target_x - bbox[0], target_y - bbox[1]),
                    fitted,
                    fill=f"#{style.color}",
                    font=font,
                    spacing=spacing,
                    align=style.align,
                    stroke_width=stroke_width,
                    stroke_fill=f"#{style.color}",
                )
                continue
            if element.type in {"table", "chart"}:
                draw.rounded_rectangle(
                    (x0, y0, x1, y1),
                    radius=10,
                    fill=fill or "#F8FAFC",
                    outline=outline or "#94A3B8",
                    width=2,
                )
                label = element.text or ("表格" if element.type == "table" else "图表")
                draw.text(
                    (x0 + 12, y0 + 10),
                    label,
                    fill=f"#{style.color}",
                    font=PPTSkillAdapter._load_preview_font(max(13, round(style.font_size))),
                )
        return image

    def _render_basic_pptx(self, plan: PresentationPlan, output: Path) -> None:
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        presentation.core_properties.title = plan.title
        presentation.core_properties.author = "MetaClass"

        for index, slide_plan in enumerate(plan.slides):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            if slide_plan.elements:
                self._render_scene(slide, slide_plan)
                continue
            self._add_background(slide, index)
            if slide_plan.layout == "hero":
                self._add_hero(
                    slide,
                    slide_plan.title,
                    slide_plan.key_points,
                    slide_plan.visual_payload,
                )
            else:
                self._add_title(slide, slide_plan.title, slide_plan.order)
            if slide_plan.layout == "cards":
                self._add_cards(slide, slide_plan.visual_payload or slide_plan.key_points)
            elif slide_plan.layout == "process":
                self._add_process(slide, slide_plan.visual_payload or slide_plan.key_points)
            elif slide_plan.layout == "comparison":
                self._add_comparison(slide, slide_plan.visual_payload or slide_plan.key_points)
            elif slide_plan.layout == "timeline":
                self._add_timeline(slide, slide_plan.visual_payload or slide_plan.key_points)
            elif slide_plan.layout == "pyramid":
                self._add_pyramid(slide, slide_plan.visual_payload or slide_plan.key_points)
            elif slide_plan.layout == "spotlight":
                self._add_spotlight(slide, slide_plan.visual_payload or slide_plan.key_points)
            elif slide_plan.layout == "two_column":
                self._add_key_points(slide, slide_plan.key_points)
                self._add_visual_panel(
                    slide, slide_plan.suggested_visual, slide_plan.visual_payload
                )
            self._add_script_summary(slide, slide_plan.speaker_script)

        presentation.save(output)

    @staticmethod
    def _render_scene(slide, slide_plan) -> None:
        """Compile an LLM-authored declarative scene into editable PPTX objects."""
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor.from_string(slide_plan.background)
        for element in sorted(slide_plan.elements, key=lambda item: item.z):
            PPTSkillAdapter._add_scene_element(slide, element)

    @staticmethod
    def _add_scene_element(slide, element) -> None:
        sw, sh = 13.333, 7.5
        x, y = Inches(element.x * sw), Inches(element.y * sh)
        w, h = Inches(element.w * sw), Inches(element.h * sh)
        style = element.style
        color = RGBColor.from_string(style.color)
        alignments = {
            "left": PP_ALIGN.LEFT,
            "center": PP_ALIGN.CENTER,
            "right": PP_ALIGN.RIGHT,
        }
        anchors = {
            "top": MSO_ANCHOR.TOP,
            "middle": MSO_ANCHOR.MIDDLE,
            "bottom": MSO_ANCHOR.BOTTOM,
        }

        if element.type == "text":
            box = slide.shapes.add_textbox(x, y, w, h)
            if element.object_id:
                box.name = f"MetaClass {element.contract_role} {element.object_id}"
            elif element.contract_role == "visual_placeholder":
                box.name = "MetaClass Visual Placeholder"
            if style.fill:
                box.fill.solid()
                box.fill.fore_color.rgb = RGBColor.from_string(style.fill)
                PPTSkillAdapter._apply_color_opacity(
                    box.fill.fore_color,
                    style.opacity,
                )
            else:
                box.fill.background()
            if style.line_color and style.line_width > 0:
                box.line.color.rgb = RGBColor.from_string(style.line_color)
                box.line.width = Pt(style.line_width)
                PPTSkillAdapter._apply_color_opacity(
                    box.line.color,
                    style.opacity,
                )
            else:
                box.line.fill.background()
            frame = box.text_frame
            frame.clear()
            frame.word_wrap = True
            # Geometry and font fitting are completed before rendering. Letting
            # PowerPoint auto-fit here introduces a second, platform-dependent
            # layout engine and makes previews differ from the exported deck.
            frame.auto_size = None
            frame.margin_left = Inches(style.text_margin_x)
            frame.margin_right = Inches(style.text_margin_x)
            frame.margin_top = Inches(style.text_margin_y)
            frame.margin_bottom = Inches(style.text_margin_y)
            frame.vertical_anchor = anchors[style.valign]
            text = element.text or "\n".join(element.items)
            # Keep one immutable Plan string in one paragraph. python-pptx serializes
            # embedded newlines as soft line breaks, which preserves the text object
            # while allowing the contract validator to reconstruct the original copy.
            paragraph = frame.paragraphs[0]
            paragraph.text = text
            paragraph.alignment = alignments[style.align]
            paragraph.line_spacing = Pt(style.font_size * 1.18)
            paragraph.space_before = Pt(0)
            paragraph.space_after = Pt(0)
            PPTSkillAdapter._apply_font(
                paragraph,
                size=Pt(style.font_size),
                bold=style.bold,
                color=color,
                opacity=style.opacity,
                font_role=style.font_role,
            )
            return

        if element.type == "shape":
            shape_types = {
                "rectangle": MSO_SHAPE.RECTANGLE,
                "rounded_rectangle": MSO_SHAPE.ROUNDED_RECTANGLE,
                "oval": MSO_SHAPE.OVAL,
                "chevron": MSO_SHAPE.CHEVRON,
            }
            shape = slide.shapes.add_shape(shape_types[element.shape], x, y, w, h)
            if element.object_id:
                shape.name = f"MetaClass {element.contract_role} {element.object_id}"
            if style.fill:
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(style.fill)
                PPTSkillAdapter._apply_color_opacity(
                    shape.fill.fore_color,
                    style.opacity,
                )
            else:
                shape.fill.background()
            if style.line_color and style.line_width > 0:
                shape.line.color.rgb = RGBColor.from_string(style.line_color)
                shape.line.width = Pt(style.line_width)
                PPTSkillAdapter._apply_color_opacity(
                    shape.line.color,
                    style.opacity,
                )
            else:
                shape.line.fill.background()
            return

        if element.type == "line":
            connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x, y, x + w, y + h)
            if element.object_id:
                connector.name = f"MetaClass {element.contract_role} {element.object_id}"
            connector.line.color.rgb = RGBColor.from_string(style.line_color or style.color)
            connector.line.width = Pt(max(style.line_width, 1))
            PPTSkillAdapter._apply_color_opacity(
                connector.line.color,
                style.opacity,
            )
            return

        if element.type == "image" and element.image_path:
            image_path = Path(element.image_path)
            if image_path.is_file():
                try:
                    with Image.open(image_path) as source:
                        source_width, source_height = source.size
                    source_ratio = source_width / source_height
                    frame_ratio = int(w) / int(h)
                    if element.image_fit == "contain":
                        if source_ratio > frame_ratio:
                            picture_width = int(w)
                            picture_height = round(int(w) / source_ratio)
                            picture_x = int(x)
                            picture_y = int(y) + (int(h) - picture_height) // 2
                        else:
                            picture_height = int(h)
                            picture_width = round(int(h) * source_ratio)
                            picture_x = int(x) + (int(w) - picture_width) // 2
                            picture_y = int(y)
                        picture = slide.shapes.add_picture(
                            str(image_path),
                            picture_x,
                            picture_y,
                            picture_width,
                            picture_height,
                        )
                    else:
                        picture = slide.shapes.add_picture(str(image_path), x, y, w, h)
                        if source_ratio > frame_ratio:
                            crop = (1 - frame_ratio / source_ratio) / 2
                            picture.crop_left = crop
                            picture.crop_right = crop
                        elif source_ratio < frame_ratio:
                            crop = (1 - source_ratio / frame_ratio) / 2
                            picture.crop_top = crop
                            picture.crop_bottom = crop
                    if element.object_id:
                        picture.name = (
                            f"MetaClass {element.contract_role} {element.object_id}"
                        )
                except (OSError, ZeroDivisionError):
                    picture = slide.shapes.add_picture(str(image_path), x, y, w, h)
                    if element.object_id:
                        picture.name = (
                            f"MetaClass {element.contract_role} {element.object_id}"
                        )
            return

        if element.type == "table" and element.table_rows:
            rows = element.table_rows
            column_count = max(len(row) for row in rows)
            table = slide.shapes.add_table(len(rows), column_count, x, y, w, h).table
            for row_index, row in enumerate(rows):
                for column_index in range(column_count):
                    cell = table.cell(row_index, column_index)
                    cell.text = row[column_index] if column_index < len(row) else ""
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string(
                        style.fill
                        or ("334E68" if row_index == 0 else "EFF6FF" if row_index % 2 else "F8FAFC")
                    )
                    paragraph = cell.text_frame.paragraphs[0]
                    PPTSkillAdapter._apply_font(
                        paragraph,
                        size=Pt(style.font_size),
                        bold=style.bold or row_index == 0,
                        color=RGBColor(255, 255, 255) if row_index == 0 else color,
                    )
            return

        if element.type == "chart" and element.chart_categories and element.chart_series:
            valid_series = [
                values
                for values in element.chart_series
                if len(values) == len(element.chart_categories)
            ]
            if not valid_series:
                return
            data = ChartData()
            data.categories = element.chart_categories
            for index, values in enumerate(valid_series):
                name = (
                    element.chart_series_names[index]
                    if index < len(element.chart_series_names)
                    else f"Series {index + 1}"
                )
                data.add_series(name, values)
            chart_types = {
                "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
                "line": XL_CHART_TYPE.LINE,
                "pie": XL_CHART_TYPE.PIE,
                "doughnut": XL_CHART_TYPE.DOUGHNUT,
            }
            chart = slide.shapes.add_chart(chart_types[element.chart_type], x, y, w, h, data).chart
            chart.has_legend = len(valid_series) > 1 or element.chart_type in {"pie", "doughnut"}
            chart.has_title = bool(element.text)
            if element.text:
                chart.chart_title.text_frame.text = element.text

    @staticmethod
    def _add_background(slide, index: int) -> None:
        if index == 0:
            background, primary, accent = (
                BRAND_PALETTE.board,
                BRAND_PALETTE.chalk,
                BRAND_PALETTE.amber,
            )
        else:
            background, primary, accent = (
                BRAND_PALETTE.paper,
                BRAND_PALETTE.ink,
                BRAND_PALETTE.amber,
            )
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor.from_string(background)
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0),
            Inches(0),
            Inches(0.18),
            Inches(7.5),
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor.from_string(accent)
        bar.line.fill.background()
        slide._metaclass_primary = primary
        slide._metaclass_accent = accent

    @staticmethod
    def _add_title(slide, title: str, order: int) -> None:
        title_box = slide.shapes.add_textbox(Inches(0.65), Inches(0.45), Inches(9), Inches(0.8))
        frame = title_box.text_frame
        frame.clear()
        paragraph = frame.paragraphs[0]
        paragraph.text = PPTSkillAdapter._fit_text(title, max_chars=42, max_lines=2)
        PPTSkillAdapter._apply_font(
            paragraph,
            size=Pt(30),
            bold=True,
            color=RGBColor.from_string(slide._metaclass_primary),
        )

        badge = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(11.6),
            Inches(0.45),
            Inches(0.78),
            Inches(0.78),
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = RGBColor.from_string(slide._metaclass_primary)
        badge.line.fill.background()
        number = badge.text_frame.paragraphs[0]
        number.text = f"{order:02d}"
        number.alignment = PP_ALIGN.CENTER
        PPTSkillAdapter._apply_font(
            number,
            size=Pt(18),
            bold=True,
            color=RGBColor(255, 255, 255),
        )

    @staticmethod
    def _add_key_points(slide, key_points: list[str]) -> None:
        points = key_points[:5] or ["核心概念", "关键例子", "课堂小结"]
        box = slide.shapes.add_textbox(Inches(0.75), Inches(1.55), Inches(5.95), Inches(4.35))
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = True
        for index, point in enumerate(points):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = PPTSkillAdapter._fit_text(point, max_chars=50, max_lines=2)
            paragraph.level = 0
            PPTSkillAdapter._apply_font(
                paragraph,
                size=Pt(18 if index == 0 else 15),
                bold=index == 0,
                color=RGBColor.from_string(slide._metaclass_primary),
            )
            paragraph.space_after = Pt(8)

    @staticmethod
    def _add_visual_panel(
        slide, suggested_visual: str, visual_payload: list[str] | None = None
    ) -> None:
        panel = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(7.35),
            Inches(1.55),
            Inches(5.2),
            Inches(3.9),
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = RGBColor(255, 255, 255)
        panel.line.color.rgb = RGBColor.from_string(slide._metaclass_accent)
        panel.line.width = Pt(2)

        label = slide.shapes.add_textbox(Inches(7.75), Inches(1.95), Inches(4.4), Inches(0.45))
        label_frame = label.text_frame
        label_frame.clear()
        label_text = label_frame.paragraphs[0]
        label_text.text = "Visual direction"
        PPTSkillAdapter._apply_font(
            label_text,
            size=Pt(14),
            bold=True,
            color=RGBColor.from_string(slide._metaclass_accent),
        )

        body = slide.shapes.add_textbox(Inches(7.75), Inches(2.55), Inches(4.35), Inches(2.35))
        body_frame = body.text_frame
        body_frame.clear()
        body_frame.word_wrap = True
        paragraph = body_frame.paragraphs[0]
        visible_content = "\n".join(f"• {item}" for item in (visual_payload or []))
        paragraph.text = PPTSkillAdapter._fit_text(
            visible_content or suggested_visual,
            max_chars=36,
            max_lines=5,
        )
        PPTSkillAdapter._apply_font(
            paragraph,
            size=Pt(14),
            color=RGBColor.from_string(slide._metaclass_primary),
        )

    @staticmethod
    def _add_cards(slide, items: list[str]) -> None:
        values = (items or ["核心概念", "关键方法", "应用提示"])[:6]
        columns = 3 if len(values) in {3, 6} else 2
        card_width = 3.7 if columns == 3 else 5.55
        for index, item in enumerate(values):
            row, column = divmod(index, columns)
            x = 0.72 + column * (card_width + 0.38)
            y = 1.55 + row * 2.12
            card = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(x),
                Inches(y),
                Inches(card_width),
                Inches(1.68),
            )
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor(255, 255, 255)
            card.line.color.rgb = RGBColor.from_string(slide._metaclass_accent)
            frame = card.text_frame
            frame.clear()
            frame.margin_left = frame.margin_right = Inches(0.22)
            paragraph = frame.paragraphs[0]
            paragraph.text = PPTSkillAdapter._fit_text(item, max_chars=24, max_lines=3)
            PPTSkillAdapter._apply_font(
                paragraph,
                size=Pt(16),
                bold=index == 0,
                color=RGBColor.from_string(slide._metaclass_primary),
            )

    @staticmethod
    def _add_hero(slide, title: str, key_points: list[str], items: list[str]) -> None:
        title_box = slide.shapes.add_textbox(Inches(1.05), Inches(1.25), Inches(11.2), Inches(1.45))
        paragraph = title_box.text_frame.paragraphs[0]
        paragraph.text = PPTSkillAdapter._fit_text(title, max_chars=30, max_lines=2)
        paragraph.alignment = PP_ALIGN.CENTER
        PPTSkillAdapter._apply_font(
            paragraph,
            size=Pt(34),
            bold=True,
            color=RGBColor.from_string(slide._metaclass_primary),
        )
        subtitle = slide.shapes.add_textbox(Inches(2.15), Inches(2.9), Inches(9.0), Inches(0.85))
        sub = subtitle.text_frame.paragraphs[0]
        sub.text = PPTSkillAdapter._fit_text(
            (key_points or ["建立整体认识"])[0], max_chars=52, max_lines=2
        )
        sub.alignment = PP_ALIGN.CENTER
        PPTSkillAdapter._apply_font(
            sub, size=Pt(18), color=RGBColor.from_string(slide._metaclass_primary)
        )
        labels = (items or key_points or ["概念", "方法", "应用"])[:4]
        width = min(2.45, 9.8 / max(len(labels), 1))
        total = len(labels) * width + max(0, len(labels) - 1) * 0.25
        start_x = (13.333 - total) / 2
        for index, item in enumerate(labels):
            pill = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(start_x + index * (width + 0.25)),
                Inches(4.2),
                Inches(width),
                Inches(0.72),
            )
            pill.fill.solid()
            pill.fill.fore_color.rgb = RGBColor.from_string(slide._metaclass_primary)
            pill.line.fill.background()
            text = pill.text_frame.paragraphs[0]
            text.text = PPTSkillAdapter._fit_text(item, max_chars=14, max_lines=1)
            text.alignment = PP_ALIGN.CENTER
            PPTSkillAdapter._apply_font(text, size=Pt(13), bold=True, color=RGBColor(255, 255, 255))

    @staticmethod
    def _add_comparison(slide, items: list[str]) -> None:
        values = (items or ["方案 A", "方案 B", "适用条件 A", "适用条件 B"])[:6]
        midpoint = max(1, (len(values) + 1) // 2)
        groups = [values[:midpoint], values[midpoint:] or ["另一种视角"]]
        for column, group in enumerate(groups):
            x = 0.75 + column * 6.05
            panel = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(x),
                Inches(1.55),
                Inches(5.55),
                Inches(4.25),
            )
            panel.fill.solid()
            panel.fill.fore_color.rgb = RGBColor(255, 255, 255)
            panel.line.color.rgb = RGBColor.from_string(
                slide._metaclass_accent if column else slide._metaclass_primary
            )
            heading = panel.text_frame.paragraphs[0]
            heading.text = "视角 A" if column == 0 else "视角 B"
            PPTSkillAdapter._apply_font(
                heading,
                size=Pt(18),
                bold=True,
                color=RGBColor.from_string(slide._metaclass_primary),
            )
            for item in group:
                paragraph = panel.text_frame.add_paragraph()
                paragraph.text = PPTSkillAdapter._fit_text(item, max_chars=30, max_lines=2)
                paragraph.space_before = Pt(10)
                PPTSkillAdapter._apply_font(
                    paragraph,
                    size=Pt(15),
                    color=RGBColor.from_string(slide._metaclass_primary),
                )

    @staticmethod
    def _add_timeline(slide, items: list[str]) -> None:
        values = (items or ["起点", "发展", "转折", "结果"])[:5]
        count = len(values)
        start_x, end_x, y = 1.15, 12.15, 3.2
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(start_x),
            Inches(y),
            Inches(end_x - start_x),
            Inches(0.06),
        )
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor.from_string(slide._metaclass_accent)
        line.line.fill.background()
        for index, item in enumerate(values):
            x = start_x + index * ((end_x - start_x) / max(count - 1, 1))
            dot = slide.shapes.add_shape(
                MSO_SHAPE.OVAL, Inches(x - 0.18), Inches(y - 0.15), Inches(0.36), Inches(0.36)
            )
            dot.fill.solid()
            dot.fill.fore_color.rgb = RGBColor.from_string(slide._metaclass_primary)
            dot.line.fill.background()
            box_y = 1.65 if index % 2 == 0 else 3.75
            box = slide.shapes.add_textbox(Inches(x - 1.0), Inches(box_y), Inches(2.0), Inches(1.0))
            paragraph = box.text_frame.paragraphs[0]
            paragraph.text = PPTSkillAdapter._fit_text(item, max_chars=16, max_lines=3)
            paragraph.alignment = PP_ALIGN.CENTER
            PPTSkillAdapter._apply_font(
                paragraph,
                size=Pt(14),
                bold=index in {0, count - 1},
                color=RGBColor.from_string(slide._metaclass_primary),
            )

    @staticmethod
    def _add_pyramid(slide, items: list[str]) -> None:
        values = (items or ["基础", "方法", "能力", "目标"])[:5]
        # The prompt contract supplies items bottom-to-top; draw top-to-bottom.
        for index, item in enumerate(reversed(values)):
            level = len(values) - index
            width = 3.0 + index * 1.25
            x = (13.333 - width) / 2
            y = 1.45 + index * 0.88
            layer = slide.shapes.add_shape(
                MSO_SHAPE.TRAPEZOID,
                Inches(x),
                Inches(y),
                Inches(width),
                Inches(0.72),
            )
            layer.fill.solid()
            layer.fill.fore_color.rgb = RGBColor.from_string(
                slide._metaclass_accent if index % 2 == 0 else slide._metaclass_primary
            )
            layer.line.fill.background()
            paragraph = layer.text_frame.paragraphs[0]
            paragraph.text = (
                f"{level}  {PPTSkillAdapter._fit_text(item, max_chars=24, max_lines=1)}"
            )
            paragraph.alignment = PP_ALIGN.CENTER
            PPTSkillAdapter._apply_font(
                paragraph, size=Pt(14), bold=True, color=RGBColor(255, 255, 255)
            )

    @staticmethod
    def _add_spotlight(slide, items: list[str]) -> None:
        values = items or ["本页最重要的结论", "为什么重要", "如何应用"]
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(4.55), Inches(1.45), Inches(4.2), Inches(3.55)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = RGBColor.from_string(slide._metaclass_primary)
        circle.line.color.rgb = RGBColor.from_string(slide._metaclass_accent)
        circle.line.width = Pt(4)
        main = circle.text_frame.paragraphs[0]
        main.text = PPTSkillAdapter._fit_text(values[0], max_chars=24, max_lines=4)
        main.alignment = PP_ALIGN.CENTER
        PPTSkillAdapter._apply_font(main, size=Pt(22), bold=True, color=RGBColor(255, 255, 255))
        for index, item in enumerate(values[1:4]):
            x = 0.8 if index % 2 == 0 else 9.25
            y = 1.8 + index * 1.35
            box = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(x),
                Inches(y),
                Inches(3.25),
                Inches(1.0),
            )
            box.fill.solid()
            box.fill.fore_color.rgb = RGBColor(255, 255, 255)
            box.line.color.rgb = RGBColor.from_string(slide._metaclass_accent)
            paragraph = box.text_frame.paragraphs[0]
            paragraph.text = PPTSkillAdapter._fit_text(item, max_chars=20, max_lines=2)
            paragraph.alignment = PP_ALIGN.CENTER
            PPTSkillAdapter._apply_font(
                paragraph,
                size=Pt(14),
                color=RGBColor.from_string(slide._metaclass_primary),
            )

    @staticmethod
    def _add_process(slide, items: list[str]) -> None:
        values = (items or ["理解问题", "分析方法", "形成结论"])[:5]
        count = len(values)
        width = min(2.15, 10.8 / max(count, 1))
        gap = 0.28
        total = count * width + max(0, count - 1) * gap
        start_x = (13.333 - total) / 2
        for index, item in enumerate(values):
            x = start_x + index * (width + gap)
            circle = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                Inches(x + width / 2 - 0.28),
                Inches(1.62),
                Inches(0.56),
                Inches(0.56),
            )
            circle.fill.solid()
            circle.fill.fore_color.rgb = RGBColor.from_string(slide._metaclass_accent)
            circle.line.fill.background()
            number = circle.text_frame.paragraphs[0]
            number.text = str(index + 1)
            number.alignment = PP_ALIGN.CENTER
            PPTSkillAdapter._apply_font(
                number, size=Pt(13), bold=True, color=RGBColor(255, 255, 255)
            )
            box = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(x),
                Inches(2.52),
                Inches(width),
                Inches(2.25),
            )
            box.fill.solid()
            box.fill.fore_color.rgb = RGBColor(255, 255, 255)
            box.line.color.rgb = RGBColor.from_string(slide._metaclass_primary)
            paragraph = box.text_frame.paragraphs[0]
            paragraph.text = PPTSkillAdapter._fit_text(item, max_chars=18, max_lines=4)
            paragraph.alignment = PP_ALIGN.CENTER
            PPTSkillAdapter._apply_font(
                paragraph,
                size=Pt(15),
                bold=index == 0,
                color=RGBColor.from_string(slide._metaclass_primary),
            )

    @staticmethod
    def _add_script_summary(slide, speaker_script: str) -> None:
        footer = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.65),
            Inches(6.5),
            Inches(11.8),
            Inches(0.4),
        )
        footer.fill.solid()
        footer.fill.fore_color.rgb = RGBColor(255, 255, 255)
        footer.line.fill.background()
        frame = footer.text_frame
        frame.clear()
        paragraph = frame.paragraphs[0]
        paragraph.text = PPTSkillAdapter._fit_text(
            f"讲稿摘要：{speaker_script}",
            max_chars=92,
            max_lines=1,
        )
        PPTSkillAdapter._apply_font(
            paragraph,
            size=Pt(11),
            color=RGBColor(75, 85, 99),
        )
