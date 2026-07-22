import json
import logging
import math
import re
import shutil
import ssl
import unicodedata
from collections import Counter
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Protocol
from urllib.error import HTTPError, URLError
from urllib.parse import quote, urljoin, urlsplit, urlunsplit
from urllib.request import Request, urlopen
from zipfile import BadZipFile, ZipFile

import certifi
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Pt

from metaclass.modules.presentation.schemas import (
    PPTArtifact,
    PPTSlideImage,
    PresentationPlan,
)
from metaclass.modules.presentation.skill_adapter import PPTSkillAdapter
from metaclass.modules.presentation.themes import (
    PresentationTheme,
    get_presentation_theme,
)


logger = logging.getLogger(__name__)


@dataclass
class _TextStyle:
    alignment: Any = None
    font_name: str | None = None
    font_size: Any = None
    bold: bool | None = None
    italic: bool | None = None
    color: Any = None


@dataclass
class _TextTarget:
    frame: Any
    text: str
    left: int
    top: int
    width: int
    height: int
    style: _TextStyle
    z_index: int = 0
    occluded: bool = False

    @property
    def area(self) -> int:
        return max(self.width, 1) * max(self.height, 1)


class PPTProvider(Protocol):
    def prepare_request(
        self,
        *,
        plan: PresentationPlan,
        job_id: str,
        output_dir: Path,
        theme: PresentationTheme | None = None,
    ) -> PPTArtifact: ...


class PresentonPPTProvider:
    """Generate a deck with Presenton while keeping PresentationPlan authoritative."""

    GENERATE_PATH = "/api/v1/ppt/presentation/generate"
    EXPORT_PATH = "/api/v1/ppt/presentation/export"

    def __init__(
        self,
        *,
        base_url: str,
        api_key: str,
        adapter: PPTSkillAdapter,
        timeout_seconds: float = 300.0,
        template: str = "general",
    ) -> None:
        self.base_url = base_url.rstrip("/")
        self.api_key = api_key
        self.adapter = adapter
        self.timeout_seconds = timeout_seconds
        self.template = template
        self._ssl_context = ssl.create_default_context(cafile=certifi.where())

    def prepare_request(
        self,
        *,
        plan: PresentationPlan,
        job_id: str,
        output_dir: Path,
        theme: PresentationTheme | None = None,
    ) -> PPTArtifact:
        selected_theme = theme or get_presentation_theme()
        output_dir.mkdir(parents=True, exist_ok=True)
        response = self._post_json(
            self.GENERATE_PATH,
            self.build_payload(plan, selected_theme),
        )
        download_path = response.get("path")
        if not isinstance(download_path, str) or not download_path.strip():
            raise RuntimeError("Presenton response did not include a PPTX path")

        pptx_path = output_dir / "deck.pptx"
        self._download_file(download_path, pptx_path)
        content_lock = self.lock_deck_content(plan, pptx_path)
        validation = self.validate_deck(plan, pptx_path)
        presentation_id = response.get("presentation_id")

        return self.adapter.prepare_external_pptx(
            plan=plan,
            job_id=job_id,
            output_dir=output_dir,
            pptx_path=pptx_path,
            provider_name="presenton",
            provider_metadata={
                "presentation_id": presentation_id,
                "edit_path": response.get("edit_path"),
                "credits_consumed": response.get("credits_consumed"),
                "template": self.template,
                "theme": selected_theme.prompt_payload(),
                "content_generation": "preserve",
                "content_lock": content_lock,
                "validation": validation,
                # Presenton's remote PNG still contains its rewritten text. Render the
                # locally locked PPTX so browser previews match the downloadable deck.
                "preview_source": "local-renderer-content-locked",
            },
            external_slide_images=None,
        )

    def build_payload(
        self,
        plan: PresentationPlan,
        theme: PresentationTheme | None = None,
    ) -> dict:
        selected_theme = theme or get_presentation_theme()
        slide_markdown = []
        visual_directions = []
        for index, slide in enumerate(plan.slides, start=1):
            lines = [f"# {slide.title}"]
            lines.extend(f"- {point}" for point in slide.key_points if point.strip())
            slide_markdown.append("\n".join(lines))
            visual_directions.append(f"第 {index} 页（{slide.title}）：{slide.suggested_visual}")

        instructions = "\n".join(
            [
                "严格按 slides_markdown 的数组顺序逐页制作，一项只能生成一页。",
                "不得新增、删除、合并、拆分或重排页面；不得改写、压缩或扩写任何标题与要点。",
                "不得添加额外正文、解释句、摘要、页脚或结论；版式中的多余文本字段必须留空。",
                "标题和要点必须作为可编辑文本完整出现在对应页面，不要放入演讲者备注。",
                "只负责版式、配色、图形与图片增强，不要把下面的视觉说明显示成正文。",
                *visual_directions,
            ]
        )
        instructions += (
            "\nSelected visual theme (do not change visible text):\n"
            + json.dumps(selected_theme.prompt_payload(), ensure_ascii=False)
        )
        return {
            "slides_markdown": slide_markdown,
            "n_slides": len(plan.slides),
            "instructions": instructions,
            "tone": "educational",
            "verbosity": "concise",
            "content_generation": "preserve",
            "markdown_emphasis": True,
            "web_search": False,
            "image_type": "stock",
            "language": "Simplified Chinese",
            "template": self.template,
            "include_table_of_contents": False,
            "include_title_slide": False,
            "allow_access_to_user_info": False,
            "export_as": "pptx",
        }

    def lock_deck_content(self, plan: PresentationPlan, pptx_path: Path) -> dict:
        """Make the PresentationPlan the only visible text source in a deck.

        Presenton may split or paraphrase a key point to fit a template. We retain
        the selected template, shapes, charts, and imagery, but deterministically
        replace all visible slide text with the exact title and key points from the
        corresponding plan page. This turns prompt preservation into an enforced
        postcondition rather than a best-effort instruction.
        """
        if not pptx_path.exists() or pptx_path.stat().st_size == 0:
            raise RuntimeError("Presenton returned an empty PPTX file")
        try:
            deck = Presentation(str(pptx_path))
        except Exception as exc:
            raise RuntimeError("Presenton returned an invalid PPTX file") from exc

        expected_count = len(plan.slides)
        actual_count = len(deck.slides)
        if actual_count != expected_count:
            raise RuntimeError(
                f"Presenton page-count mismatch: expected {expected_count}, got {actual_count}"
            )

        repaired_pages = 0
        removed_generated_text_items = 0
        locked_text_items = 0
        for planned_slide, generated_slide in zip(plan.slides, deck.slides, strict=True):
            targets = self._collect_text_targets(generated_slide)
            expected_values = [
                value for value in [planned_slide.title, *planned_slide.key_points] if value.strip()
            ]
            original_values = [
                target.text for target in targets if target.text.strip() and not target.occluded
            ]
            if Counter(map(self._normalize_text, original_values)) != Counter(
                map(self._normalize_text, expected_values)
            ):
                repaired_pages += 1

            title_target = self._select_title_target(
                targets,
                planned_slide.title,
                deck.slide_width,
                deck.slide_height,
            )
            body_targets = [
                target
                for target in targets
                if target is not title_target
                and not target.occluded
                and target.width >= int(deck.slide_width * 0.12)
            ]
            assignments = self._match_key_points(
                body_targets,
                [point for point in planned_slide.key_points if point.strip()],
            )

            assigned_targets = {
                target_index
                for _, _, target_indexes in assignments
                for target_index in target_indexes
            }
            removed_generated_text_items += sum(
                1
                for index, target in enumerate(body_targets)
                if target.text.strip() and index not in assigned_targets
            )

            for target in targets:
                target.frame.clear()

            if planned_slide.title.strip():
                if title_target is None:
                    title_target = self._add_text_target(
                        generated_slide,
                        left=int(deck.slide_width * 0.06),
                        top=int(deck.slide_height * 0.05),
                        width=int(deck.slide_width * 0.88),
                        height=int(deck.slide_height * 0.14),
                        font_size=Pt(28),
                        bold=True,
                    )
                self._set_target_text(title_target, planned_slide.title)
                locked_text_items += 1

            for point_index, (point, anchor, _) in enumerate(assignments):
                if anchor is None:
                    anchor = self._add_body_text_target(
                        generated_slide,
                        point_index,
                        len(assignments),
                        deck.slide_width,
                        deck.slide_height,
                    )
                self._set_target_text(anchor, point)
                locked_text_items += 1

        locked_path = pptx_path.with_name(f"{pptx_path.stem}.content-locked.pptx")
        deck.save(locked_path)
        locked_path.replace(pptx_path)
        return {
            "status": "locked",
            "slide_count": expected_count,
            "locked_text_items": locked_text_items,
            "repaired_pages": repaired_pages,
            "removed_generated_text_items": removed_generated_text_items,
        }

    def validate_deck(self, plan: PresentationPlan, pptx_path: Path) -> dict:
        return validate_deck_against_plan(plan, pptx_path, provider_name="Presenton")

    def _post_json(self, path: str, payload: dict) -> dict:
        request = Request(
            urljoin(f"{self.base_url}/", path.lstrip("/")),
            data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
            headers={
                "Authorization": f"Bearer {self.api_key}",
                "Content-Type": "application/json",
                "Accept": "application/json",
            },
            method="POST",
        )
        try:
            with urlopen(
                request,
                timeout=self.timeout_seconds,
                context=self._ssl_context,
            ) as response:
                body = response.read().decode("utf-8")
        except HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="replace")[:500]
            raise RuntimeError(
                f"Presenton generation failed with HTTP {exc.code}: {detail}"
            ) from exc
        except (URLError, TimeoutError) as exc:
            raise RuntimeError("Presenton generation request failed") from exc

        try:
            result = json.loads(body)
        except json.JSONDecodeError as exc:
            raise RuntimeError("Presenton returned a non-JSON response") from exc
        if not isinstance(result, dict):
            raise RuntimeError("Presenton returned an unexpected response")
        return result

    def _download_file(self, remote_path: str, destination: Path) -> None:
        url = urljoin(f"{self.base_url}/", remote_path)
        parts = urlsplit(url)
        url = urlunsplit(
            (
                parts.scheme,
                parts.netloc,
                quote(parts.path, safe="/%"),
                parts.query,
                parts.fragment,
            )
        )
        headers = {}
        if parts.netloc.casefold() == urlsplit(self.base_url).netloc.casefold():
            headers["Authorization"] = f"Bearer {self.api_key}"
        request = Request(
            url,
            headers=headers,
            method="GET",
        )
        try:
            with urlopen(
                request,
                timeout=self.timeout_seconds,
                context=self._ssl_context,
            ) as response:
                content = response.read()
        except HTTPError as exc:
            raise RuntimeError(f"Presenton PPTX download failed with HTTP {exc.code}") from exc
        except (URLError, TimeoutError) as exc:
            raise RuntimeError("Presenton PPTX download failed") from exc
        if not content.startswith(b"PK"):
            raise RuntimeError("Presenton download was not a valid Office or ZIP archive")
        destination.write_bytes(content)

    def _export_slide_images(
        self,
        plan: PresentationPlan,
        presentation_id: str,
        output_dir: Path,
    ) -> list[PPTSlideImage]:
        output_dir.mkdir(parents=True, exist_ok=True)
        response = self._post_json(
            self.EXPORT_PATH,
            {"id": presentation_id, "export_as": "png"},
        )
        remote_path = response.get("path")
        if not isinstance(remote_path, str) or not remote_path.strip():
            raise RuntimeError("Presenton PNG export did not include a download path")

        archive_path = output_dir / "presenton-previews.zip"
        self._download_file(remote_path, archive_path)
        slides_dir = output_dir / "slides"
        slides_dir.mkdir(parents=True, exist_ok=True)
        try:
            with ZipFile(archive_path) as archive:
                members = [
                    member
                    for member in archive.infolist()
                    if not member.is_dir() and Path(member.filename).suffix.lower() == ".png"
                ]
                members.sort(key=self._image_member_sort_key)
                if len(members) != len(plan.slides):
                    raise RuntimeError(
                        "Presenton preview page-count mismatch: "
                        f"expected {len(plan.slides)}, got {len(members)}"
                    )

                slide_images = []
                for index, (planned_slide, member) in enumerate(
                    zip(plan.slides, members, strict=True), start=1
                ):
                    image_path = slides_dir / f"slide_{index:03d}.png"
                    with archive.open(member) as source, image_path.open("wb") as target:
                        shutil.copyfileobj(source, target)
                    with Image.open(image_path) as image:
                        image.verify()
                        width, height = image.size
                    slide_images.append(
                        PPTSlideImage(
                            slide_id=planned_slide.id,
                            slide_no=index,
                            image_path=str(image_path),
                            width=width,
                            height=height,
                        )
                    )
        except BadZipFile as exc:
            raise RuntimeError("Presenton returned an invalid PNG preview archive") from exc
        return slide_images

    @staticmethod
    def _image_member_sort_key(member) -> tuple[int, str]:
        match = re.search(r"(\d+)(?!.*\d)", Path(member.filename).stem)
        return (int(match.group(1)) if match else 0, member.filename)

    @classmethod
    def _extract_slide_text(cls, slide) -> str:
        return "\n".join(cls._extract_slide_text_items(slide))

    @classmethod
    def _extract_slide_text_items(cls, slide) -> list[str]:
        values: list[str] = []

        def collect(shape) -> None:
            if getattr(shape, "has_text_frame", False):
                values.append(shape.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        values.append(cell.text)
            for child in getattr(shape, "shapes", []):
                collect(child)

        for shape in slide.shapes:
            collect(shape)
        return values

    @classmethod
    def _collect_text_targets(cls, slide) -> list[_TextTarget]:
        targets: list[_TextTarget] = []
        top_level_shapes = list(slide.shapes)

        def collect(shape, z_index: int) -> None:
            if getattr(shape, "has_text_frame", False):
                targets.append(
                    _TextTarget(
                        frame=shape.text_frame,
                        text=shape.text,
                        left=int(getattr(shape, "left", 0)),
                        top=int(getattr(shape, "top", 0)),
                        width=int(getattr(shape, "width", 0)),
                        height=int(getattr(shape, "height", 0)),
                        style=cls._capture_text_style(shape.text_frame),
                        z_index=z_index,
                    )
                )
            if getattr(shape, "has_table", False):
                table = shape.table
                row_top = int(getattr(shape, "top", 0))
                for row in table.rows:
                    column_left = int(getattr(shape, "left", 0))
                    for column_index, cell in enumerate(row.cells):
                        column_width = int(table.columns[column_index].width)
                        targets.append(
                            _TextTarget(
                                frame=cell.text_frame,
                                text=cell.text,
                                left=column_left,
                                top=row_top,
                                width=column_width,
                                height=int(row.height),
                                style=cls._capture_text_style(cell.text_frame),
                                z_index=z_index,
                            )
                        )
                        column_left += column_width
                    row_top += int(row.height)
            for child in getattr(shape, "shapes", []):
                collect(child, z_index)

        for z_index, shape in enumerate(top_level_shapes):
            collect(shape, z_index)
        for target in targets:
            target.occluded = any(
                shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                and cls._coverage_ratio(target, shape) >= 0.85
                for shape in top_level_shapes[target.z_index + 1 :]
            )
        return targets

    @staticmethod
    def _coverage_ratio(target: _TextTarget, covering_shape) -> float:
        target_right = target.left + target.width
        target_bottom = target.top + target.height
        cover_left = int(getattr(covering_shape, "left", 0))
        cover_top = int(getattr(covering_shape, "top", 0))
        cover_right = cover_left + int(getattr(covering_shape, "width", 0))
        cover_bottom = cover_top + int(getattr(covering_shape, "height", 0))
        overlap_width = max(0, min(target_right, cover_right) - max(target.left, cover_left))
        overlap_height = max(0, min(target_bottom, cover_bottom) - max(target.top, cover_top))
        return (overlap_width * overlap_height) / max(target.area, 1)

    @staticmethod
    def _capture_text_style(frame) -> _TextStyle:
        paragraph = next(
            (item for item in frame.paragraphs if item.text.strip()),
            frame.paragraphs[0] if frame.paragraphs else None,
        )
        if paragraph is None:
            return _TextStyle()
        run = next((item for item in paragraph.runs if item.text.strip()), None)
        if run is None:
            return _TextStyle(alignment=paragraph.alignment)
        color = None
        try:
            color = run.font.color.rgb
        except (AttributeError, ValueError):
            pass
        return _TextStyle(
            alignment=paragraph.alignment,
            font_name=run.font.name,
            font_size=run.font.size,
            bold=run.font.bold,
            italic=run.font.italic,
            color=color,
        )

    @classmethod
    def _select_title_target(
        cls,
        targets: list[_TextTarget],
        title: str,
        slide_width: int,
        slide_height: int,
    ) -> _TextTarget | None:
        visible_targets = [
            target for target in targets if target.text.strip() and not target.occluded
        ]
        if not visible_targets:
            return None
        title_region_targets = [
            target
            for target in visible_targets
            if target.top <= int(slide_height * 0.25) and target.width >= int(slide_width * 0.25)
        ]
        if title_region_targets:
            visible_targets = title_region_targets
        expected = cls._normalize_text(title)

        def score(target: _TextTarget) -> float:
            actual = cls._normalize_text(target.text)
            similarity = SequenceMatcher(None, expected, actual).ratio()
            exact_bonus = 2.0 if expected and expected == actual else 0.0
            containment_bonus = 0.7 if expected and expected in actual else 0.0
            top_bonus = max(0.0, 1.0 - target.top / max(slide_height, 1)) * 0.35
            width_bonus = target.width / max(slide_width, 1) * 0.15
            font_bonus = 0.0
            if target.style.font_size is not None:
                font_bonus = min(float(target.style.font_size.pt) / 100.0, 0.5)
            return (
                similarity + exact_bonus + containment_bonus + top_bonus + width_bonus + font_bonus
            )

        return max(visible_targets, key=score)

    @classmethod
    def _match_key_points(
        cls,
        targets: list[_TextTarget],
        key_points: list[str],
    ) -> list[tuple[str, _TextTarget | None, tuple[int, ...]]]:
        assignments: list[tuple[str, _TextTarget | None, tuple[int, ...]]] = []
        used: set[int] = set()
        target_count = len(targets)
        point_count = max(len(key_points), 1)

        for point_index, point in enumerate(key_points):
            expected = cls._normalize_text(point)
            best: tuple[float, tuple[int, ...]] | None = None
            for start in range(target_count):
                for group_length in range(1, min(3, target_count - start) + 1):
                    indexes = tuple(range(start, start + group_length))
                    if used.intersection(indexes):
                        continue
                    actual = "".join(cls._normalize_text(targets[index].text) for index in indexes)
                    if not actual:
                        continue
                    similarity = SequenceMatcher(None, expected, actual).ratio()
                    containment_bonus = 0.0
                    if expected in actual or actual in expected:
                        containment_bonus = 0.35
                    expected_position = point_index / point_count
                    actual_position = start / max(target_count, 1)
                    order_bonus = max(0.0, 0.08 - abs(expected_position - actual_position) * 0.08)
                    length_penalty = (group_length - 1) * 0.01
                    score = similarity + containment_bonus + order_bonus - length_penalty
                    if best is None or score > best[0]:
                        best = (score, indexes)

            indexes: tuple[int, ...] = ()
            anchor = None
            if best is not None and best[0] >= 0.3:
                indexes = best[1]
                used.update(indexes)
                anchor = max((targets[index] for index in indexes), key=lambda item: item.area)
            else:
                remaining = [
                    (index, target) for index, target in enumerate(targets) if index not in used
                ]
                if remaining:
                    index, anchor = max(remaining, key=lambda item: item[1].area)
                    indexes = (index,)
                    used.add(index)
            assignments.append((point, anchor, indexes))
        return assignments

    @staticmethod
    def _set_target_text(target: _TextTarget, value: str) -> None:
        frame = target.frame
        frame.clear()
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        paragraph = frame.paragraphs[0]
        paragraph.alignment = target.style.alignment
        run = paragraph.add_run()
        run.text = value
        run.font.name = target.style.font_name or "Aptos"
        run.font.size = PresentonPPTProvider._fit_font_size(target, value)
        if target.style.bold is not None:
            run.font.bold = target.style.bold
        if target.style.italic is not None:
            run.font.italic = target.style.italic
        if target.style.color is not None:
            run.font.color.rgb = target.style.color
        else:
            # Some Presenton templates inherit text styling from HTML and leave
            # the PPTX run style empty. PowerPoint can then render repaired text
            # as white-on-white. Use the template's standard dark foreground.
            run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)

    @staticmethod
    def _fit_font_size(target: _TextTarget, value: str):
        preferred = target.style.font_size or Pt(18)
        preferred_points = max(10, int(preferred.pt))
        horizontal_margins = int(target.frame.margin_left or 0) + int(
            target.frame.margin_right or 0
        )
        vertical_margins = int(target.frame.margin_top or 0) + int(target.frame.margin_bottom or 0)
        usable_width_points = max((target.width - horizontal_margins) / 12700, 10)
        usable_height_points = max((target.height - vertical_margins) / 12700, 10)
        character_units = sum(
            1.0 if unicodedata.east_asian_width(character) in {"W", "F"} else 0.55
            for character in value
        )
        for font_points in range(preferred_points, 9, -1):
            units_per_line = max(usable_width_points / font_points, 1)
            required_lines = max(1, math.ceil(character_units / units_per_line))
            if required_lines * font_points * 1.15 <= usable_height_points:
                return Pt(font_points)
        return Pt(10)

    @classmethod
    def _add_text_target(
        cls,
        slide,
        *,
        left: int,
        top: int,
        width: int,
        height: int,
        font_size,
        bold: bool,
    ) -> _TextTarget:
        shape = slide.shapes.add_textbox(left, top, width, height)
        return _TextTarget(
            frame=shape.text_frame,
            text="",
            left=left,
            top=top,
            width=width,
            height=height,
            style=_TextStyle(font_size=font_size, bold=bold),
        )

    @classmethod
    def _add_body_text_target(
        cls,
        slide,
        point_index: int,
        point_count: int,
        slide_width: int,
        slide_height: int,
    ) -> _TextTarget:
        content_top = int(slide_height * 0.25)
        content_height = int(slide_height * 0.62)
        row_height = max(int(content_height / max(point_count, 1)), 1)
        return cls._add_text_target(
            slide,
            left=int(slide_width * 0.1),
            top=content_top + point_index * row_height,
            width=int(slide_width * 0.8),
            height=int(row_height * 0.78),
            font_size=Pt(20),
            bold=False,
        )

    @staticmethod
    def _normalize_text(value: str) -> str:
        normalized = unicodedata.normalize("NFKC", value).casefold()
        return re.sub(r"[^\w]+", "", normalized, flags=re.UNICODE)


def validate_deck_against_plan(
    plan: PresentationPlan,
    pptx_path: Path,
    *,
    provider_name: str,
) -> dict:
    """Enforce the immutable visible-text contract shared by every provider.

    Comparisons trim only the outer whitespace of a text object. Unicode code
    points, punctuation, case, internal whitespace, wording, page count, and
    page assignment therefore cannot be changed.
    """
    if not pptx_path.exists() or pptx_path.stat().st_size == 0:
        raise RuntimeError(f"{provider_name} returned an empty PPTX file")
    try:
        deck = Presentation(str(pptx_path))
    except Exception as exc:
        raise RuntimeError(f"{provider_name} returned an invalid PPTX file") from exc

    expected_count = len(plan.slides)
    actual_count = len(deck.slides)
    if actual_count != expected_count:
        raise RuntimeError(
            f"{provider_name} page-count mismatch: expected {expected_count}, got {actual_count}"
        )

    checked_items = 0
    for index, (planned_slide, generated_slide) in enumerate(
        zip(plan.slides, deck.slides, strict=True), start=1
    ):
        actual_items: list[str] = []
        for target in PresentonPPTProvider._collect_text_targets(generated_slide):
            if target.occluded:
                continue
            paragraphs = [
                paragraph.text for paragraph in target.frame.paragraphs if paragraph.text.strip()
            ]
            if paragraphs:
                actual_items.extend(paragraphs)
            elif target.text.strip():
                actual_items.append(target.text)

        required_items = [
            item for item in [planned_slide.title, *planned_slide.key_points] if item.strip()
        ]
        actual_counter = Counter(_contract_text(item) for item in actual_items)
        required_counter = Counter(_contract_text(item) for item in required_items)
        missing_counter = required_counter - actual_counter
        extra_counter = actual_counter - required_counter
        if missing_counter or extra_counter:
            missing = list(missing_counter.elements())
            extra = list(extra_counter.elements())
            details = []
            if missing:
                details.append(f"missing: {missing!r}")
            if extra:
                details.append(f"extra or rewritten: {extra!r}")
            raise RuntimeError(
                f"{provider_name} content mismatch on page {index}; " + "; ".join(details)
            )
        checked_items += len(required_items)

    return {
        "expected_slide_count": expected_count,
        "actual_slide_count": actual_count,
        "checked_text_items": checked_items,
        "status": "passed",
    }


def _contract_text(value: str) -> str:
    return value.strip()


class UnavailablePPTProvider:
    def __init__(self, message: str) -> None:
        self.message = message

    def prepare_request(
        self,
        *,
        plan: PresentationPlan,
        job_id: str,
        output_dir: Path,
        theme: PresentationTheme | None = None,
    ) -> PPTArtifact:
        raise RuntimeError(self.message)


class FallbackPPTProvider:
    def __init__(
        self,
        *,
        primary: PPTProvider,
        fallback: PPTProvider,
        primary_name: str,
        fallback_name: str = "fallback",
        fallback_exceptions: tuple[type[Exception], ...] = (Exception,),
    ) -> None:
        self.primary = primary
        self.fallback = fallback
        self.primary_name = primary_name
        self.fallback_name = fallback_name
        self.fallback_exceptions = fallback_exceptions

    def prepare_request(
        self,
        *,
        plan: PresentationPlan,
        job_id: str,
        output_dir: Path,
        theme: PresentationTheme | None = None,
    ) -> PPTArtifact:
        provider_kwargs = {
            "plan": plan,
            "job_id": job_id,
            "output_dir": output_dir,
        }
        if theme is not None:
            provider_kwargs["theme"] = theme
        try:
            return self.primary.prepare_request(**provider_kwargs)
        except self.fallback_exceptions as primary_error:
            logger.warning(
                "PPT provider %s failed; falling back to %s: %s",
                self.primary_name,
                self.fallback_name,
                primary_error,
            )
            try:
                artifact = self.fallback.prepare_request(**provider_kwargs)
            except Exception as fallback_error:
                raise RuntimeError(
                    f"{self.primary_name} generation failed: {primary_error}; "
                    f"{self.fallback_name} fallback failed: {fallback_error}"
                ) from fallback_error
            self._record_fallback(artifact, primary_error)
            return artifact

    def _record_fallback(self, artifact: PPTArtifact, primary_error: Exception) -> None:
        request_path_value = getattr(artifact, "skill_request_path", None)
        if not isinstance(request_path_value, (str, Path)):
            return
        request_path = Path(request_path_value)
        if not request_path.exists():
            return
        try:
            payload = json.loads(request_path.read_text(encoding="utf-8"))
            if not isinstance(payload, dict):
                return
            payload["fallback"] = {
                "from": self.primary_name,
                "to": self.fallback_name,
                "reason": str(primary_error)[:1000],
            }
            request_path.write_text(
                json.dumps(payload, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
        except (OSError, json.JSONDecodeError):
            logger.warning("Could not record PPT fallback metadata in %s", request_path)
