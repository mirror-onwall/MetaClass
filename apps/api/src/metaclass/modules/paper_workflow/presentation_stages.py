import hashlib
import json
import math
import re
import shutil
import subprocess
import tempfile
from difflib import SequenceMatcher
from pathlib import Path
from typing import Literal

import fitz
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from metaclass.modules.paper_workflow.runtime import CodexSkillInvocation, SkillRuntime
from metaclass.modules.paper_workflow.schemas import (
    ComposedStage,
    FigureCatalog,
    PaperAnalysis,
    PresentationOutline,
    SlideEvidence,
    StageExecutionReport,
)

_NUMBER = re.compile(r"(?<![A-Za-z])\d+(?:\.\d+)?%?")
_PLACEHOLDER = re.compile(
    r"\b(?:lorem ipsum|placeholder|todo|tbd|replace me|xxxx)\b", re.IGNORECASE
)
_RESULT = re.compile(r"result|finding|experiment|结果|发现|实验|性能|提升|降低", re.IGNORECASE)


class OutlineValidationError(ValueError):
    pass


class OutlineStructureError(OutlineValidationError):
    pass


class OutlineEvidenceError(OutlineValidationError):
    pass


class OutlineNarrativeError(OutlineValidationError):
    pass


class RepairableGenerationError(ValueError):
    pass


class FatalGenerationError(ValueError):
    pass


def validate_outline_contract(
    outline: PresentationOutline,
    evidence: SlideEvidence,
    *,
    analysis: PaperAnalysis,
    figures: FigureCatalog,
    duration_minutes: int,
) -> None:
    slide_ids = [slide.id for slide in outline.slides]
    if [item.slide_id for item in evidence.slides] != slide_ids:
        raise OutlineStructureError("slide evidence must cover slides in exact order")
    orders = [slide.order for slide in outline.slides]
    if orders != list(range(1, len(orders) + 1)):
        raise OutlineStructureError("slide order must be continuous")
    positions = {slide.id: index for index, slide in enumerate(outline.slides)}
    for section in outline.sections:
        indexes = [positions[item] for item in section.slide_ids]
        if indexes != list(range(min(indexes), max(indexes) + 1)):
            raise OutlineStructureError(f"section {section.id} is not contiguous")
    claim_ids = {claim.id for claim in analysis.claims}
    core_claims = {claim.id for claim in analysis.claims if claim.importance == "core"}
    asset_ids = {figure.id for figure in figures.figures}
    covered_claims = {claim for entry in evidence.slides for claim in entry.claim_ids}
    if not core_claims.issubset(covered_claims):
        raise OutlineEvidenceError(
            f"core claims lack slide coverage: {sorted(core_claims - covered_claims)}"
        )
    source_numbers = set(_NUMBER.findall(analysis.model_dump_json()))
    evidence_by_slide = {item.slide_id: item for item in evidence.slides}
    for slide in outline.slides:
        if len(slide.key_points) > 4:
            raise OutlineNarrativeError(f"slide {slide.id} contains multiple major arguments")
        entry = evidence_by_slide[slide.id]
        unknown_claims = set(entry.claim_ids) - claim_ids
        unknown_assets = {*slide.asset_ids, *entry.asset_ids} - asset_ids
        if unknown_claims:
            raise OutlineEvidenceError(f"slide {slide.id} references unknown claims")
        if unknown_assets:
            raise OutlineEvidenceError(f"slide {slide.id} references unknown assets")
        is_result = bool(_RESULT.search(f"{slide.title} {slide.purpose} {slide.layout_intent}"))
        if is_result and not (entry.claim_ids or entry.source_refs):
            raise OutlineEvidenceError(f"result slide {slide.id} lacks evidence")
        if is_result:
            slide_numbers = set(
                _NUMBER.findall(" ".join([slide.title, slide.purpose, *slide.key_points]))
            )
            unsupported = slide_numbers - source_numbers
            if unsupported:
                raise OutlineNarrativeError(
                    f"slide {slide.id} introduces unsupported quantitative values: "
                    f"{sorted(unsupported)}"
                )
    minimum = max(4, math.floor(duration_minutes * 0.6))
    maximum = max(minimum, math.ceil(duration_minutes * 0.9) + 1)
    if not minimum <= len(outline.slides) <= maximum:
        raise OutlineNarrativeError(
            f"outline slide count {len(outline.slides)} is outside {minimum}-{maximum}"
        )


class AcademicOutlineAdapter:
    skill_name = "academic-pptx"
    prompt_version = "academic-pptx-v2"

    def __init__(self, runtime: SkillRuntime, skill_directory: Path) -> None:
        self.runtime = runtime
        self.skill_directory = skill_directory.resolve()

    @staticmethod
    def write_contract_schemas(stage_root: Path) -> tuple[Path, Path, Path]:
        """Materialize the exact Stage 3 contracts so the Skill never guesses JSON shapes."""
        stage_root.mkdir(parents=True, exist_ok=True)
        outline_schema = stage_root / "presentation_outline_schema.json"
        evidence_schema = stage_root / "slide_evidence_schema.json"
        runtime_schema = stage_root / "runtime_result_schema.json"
        schemas = {
            outline_schema: PresentationOutline.model_json_schema(),
            evidence_schema: SlideEvidence.model_json_schema(),
            runtime_schema: {
                "type": "object",
                "additionalProperties": False,
                "required": ["status", "summary", "warnings"],
                "properties": {
                    "status": {"type": "string", "enum": ["completed", "blocked"]},
                    "summary": {"type": "string"},
                    "warnings": {"type": "array", "items": {"type": "string"}},
                },
            },
        }
        for path, schema in schemas.items():
            temporary = path.with_name(f".{path.name}.tmp")
            temporary.write_text(
                json.dumps(schema, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
                encoding="utf-8",
            )
            temporary.replace(path)
        return outline_schema, evidence_schema, runtime_schema

    def run(
        self,
        *,
        workspace: Path,
        output: Path,
        input_hash: str,
        skill_version: str,
        attempt: int,
        outline_schema: Path,
        evidence_schema: Path,
        runtime_schema: Path,
        mode: Literal["initial", "structure_repair", "evidence_repair", "rerun"] = "initial",
    ) -> StageExecutionReport:
        prompt_names = {
            "initial": ("academic_outline_contract.md", "prompt.md"),
            "structure_repair": (
                "academic_outline_structure_repair.md",
                "structure_repair_prompt.md",
            ),
            "evidence_repair": ("academic_outline_evidence_repair.md", "outline_repair_prompt.md"),
            "rerun": ("academic_outline_rerun.md", "rerun_prompt.md"),
        }
        source_prompt, stage_prompt = prompt_names[mode]
        prompt = output.parent / stage_prompt
        prompt.write_text(
            (Path(__file__).parent / "prompts" / source_prompt).read_text(encoding="utf-8"),
            encoding="utf-8",
        )
        input_paths = [
            workspace / "stages/01_analysis/output/paper_analysis.json",
            workspace / "stages/02_figures/output/figures.json",
            workspace / "resolved_request.json",
            workspace / "source/paper_source.json",
            outline_schema,
            evidence_schema,
        ]
        if mode in {"structure_repair", "evidence_repair"}:
            input_paths.extend(
                [
                    output / "presentation_outline.json",
                    output / "slide_evidence.json",
                ]
            )
        if mode != "initial":
            input_paths.append(output.parent / "validation_errors.json")
        return self.runtime.run(
            CodexSkillInvocation(
                stage=ComposedStage.OUTLINE,
                skill_name=self.skill_name,
                skill_directory=self.skill_directory,
                skill_version=skill_version,
                prompt_version=self.prompt_version + (f"-{mode}" if mode != "initial" else ""),
                input_hash=input_hash,
                workspace=workspace,
                prompt_path=prompt,
                input_paths=tuple(input_paths),
                output_directory=output,
                expected_outputs=(
                    "outline.md",
                    "presentation_outline.json",
                    "slide_evidence.json",
                ),
                output_schema_path=runtime_schema,
                attempt=attempt,
                network_enabled=False,
            )
        )


class PptxGenerateAdapter:
    skill_name = "pptx"
    prompt_version = "pptx-academic-contract-v1"

    def __init__(self, runtime: SkillRuntime, skill_directory: Path) -> None:
        self.runtime = runtime
        self.skill_directory = skill_directory.resolve()

    def run(
        self,
        *,
        workspace: Path,
        output: Path,
        input_hash: str,
        skill_version: str,
        attempt: int,
        repair: bool = False,
    ) -> StageExecutionReport:
        prompt_name = "pptx_generate_repair.md" if repair else "pptx_generate_contract.md"
        prompt = output.parent / ("repair_prompt.md" if repair else "prompt.md")
        prompt.write_text(
            (Path(__file__).parent / "prompts" / prompt_name).read_text(encoding="utf-8"),
            encoding="utf-8",
        )
        inputs = [
            workspace / "stages/03_outline/output/outline.md",
            workspace / "stages/03_outline/output/presentation_outline.json",
            workspace / "stages/03_outline/output/slide_evidence.json",
            workspace / "stages/02_figures/output/figures.json",
        ]
        inputs.extend(
            sorted(
                item
                for item in (workspace / "stages/02_figures/output/assets").rglob("*")
                if item.is_file()
            )
        )
        return self.runtime.run(
            CodexSkillInvocation(
                stage=ComposedStage.GENERATION,
                skill_name=self.skill_name,
                skill_directory=self.skill_directory,
                skill_version=skill_version,
                prompt_version=self.prompt_version + ("-repair" if repair else ""),
                input_hash=input_hash,
                workspace=workspace,
                prompt_path=prompt,
                input_paths=tuple(inputs),
                output_directory=output,
                expected_outputs=(
                    "presentation.pptx",
                    "slide_plan.json",
                    "speaker_notes.json",
                    "qa_report.json",
                ),
                attempt=attempt,
                timeout_seconds=1200,
                network_enabled=False,
            )
        )


def validate_and_render_pptx(
    output: Path,
    *,
    outline: PresentationOutline,
    evidence: SlideEvidence,
) -> SlideEvidence:
    pptx_path = output / "presentation.pptx"
    try:
        presentation = Presentation(pptx_path)
    except Exception as exc:
        raise FatalGenerationError("presentation.pptx is damaged or unreadable") from exc
    actual = len(presentation.slides)
    expected = len(outline.slides)
    if actual != expected:
        raise FatalGenerationError(f"PPTX slide count {actual} does not match outline {expected}")
    expected_assets = {asset for slide in outline.slides for asset in slide.asset_ids}
    expected_assets.update(asset for entry in evidence.slides for asset in entry.asset_ids)
    image_count = 0
    title_scores: list[float] = []
    for index, (slide, contract) in enumerate(zip(presentation.slides, outline.slides), start=1):
        texts = [
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text_frame") and shape.text.strip()
        ]
        combined = "\n".join(texts)
        if _PLACEHOLDER.search(combined):
            raise FatalGenerationError(f"slide {index} contains placeholder text")
        # Academic templates often place a short section eyebrow before the
        # assertion-style headline and may not use PowerPoint title
        # placeholders at all. Validate the frozen title against every visible
        # text shape and retain the best match; requiring the first text shape
        # incorrectly treats labels such as "Evidence" as the slide title.
        title_scores.append(
            max(
                (
                    SequenceMatcher(None, candidate, contract.title).ratio()
                    for candidate in texts
                ),
                default=0.0,
            )
        )
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image_count += 1
            if shape.left < 0 or shape.top < 0:
                raise RepairableGenerationError(
                    f"slide {index} has a shape outside the page"
                )
            if (
                shape.left + shape.width > presentation.slide_width
                or shape.top + shape.height > presentation.slide_height
            ):
                raise RepairableGenerationError(
                    f"slide {index} has a shape outside the page"
                )
            if hasattr(shape, "text_frame") and shape.text.strip():
                chars = len(shape.text.strip())
                area = max((shape.width / 914400) * (shape.height / 914400), 0.1)
                if chars / area > 90:
                    raise FatalGenerationError(f"slide {index} has obvious text overflow risk")
    if expected_assets and image_count < len(expected_assets):
        raise FatalGenerationError("critical figures are not embedded in the PPTX")
    if any(score < 0.55 for score in title_scores):
        raise FatalGenerationError("PPTX titles diverge severely from the frozen outline")
    if any(score < 0.78 for score in title_scores):
        raise RepairableGenerationError("PPTX titles need a light mapping repair")
    documents: dict[str, object] = {}
    for name in ("slide_plan.json", "speaker_notes.json", "qa_report.json"):
        try:
            documents[name] = json.loads((output / name).read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise RepairableGenerationError(f"{name} is invalid") from exc
    slide_ids = {slide.id for slide in outline.slides}
    for name in ("slide_plan.json", "speaker_notes.json"):
        document = documents[name]
        entries = document.get("slides", []) if isinstance(document, dict) else document
        if not isinstance(entries, list):
            entries = []
        covered_ids = {
            entry.get("slide_id")
            for entry in entries
            if isinstance(entry, dict) and isinstance(entry.get("slide_id"), str)
        }
        if covered_ids != slide_ids:
            raise RepairableGenerationError(f"{name} does not cover every frozen slide")
    _render_with_libreoffice(pptx_path, output / "rendered", expected)
    rebound = SlideEvidence.model_validate(
        {
            **evidence.model_dump(mode="python"),
            "slides": [
                {**entry.model_dump(mode="python"), "page_no": index}
                for index, entry in enumerate(evidence.slides, start=1)
            ],
        }
    )
    (output / "slide_evidence.json").write_text(rebound.model_dump_json(indent=2), encoding="utf-8")
    return rebound


def _render_with_libreoffice(pptx_path: Path, rendered: Path, expected: int) -> None:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    if not Path(soffice).is_file():
        raise FatalGenerationError("LibreOffice is required for Stage 4 acceptance")
    rendered.mkdir(parents=True, exist_ok=True)
    for path in rendered.glob("slide-*.png"):
        path.unlink()
    with tempfile.TemporaryDirectory(prefix="paper-ppt-render-") as directory:
        result = subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", directory, str(pptx_path)],
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
        )
        pdf_path = Path(directory) / f"{pptx_path.stem}.pdf"
        if result.returncode != 0 or not pdf_path.is_file():
            raise FatalGenerationError("LibreOffice could not render presentation.pptx")
        document = fitz.open(pdf_path)
        if document.page_count != expected:
            document.close()
            raise FatalGenerationError("LibreOffice rendered the wrong number of pages")
        dimensions: set[tuple[int, int]] = set()
        for index, page in enumerate(document, start=1):
            pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
            dimensions.add((pixmap.width, pixmap.height))
            pixmap.save(rendered / f"slide-{index:03d}.png")
        document.close()
        if len(dimensions) != 1 or len(list(rendered.glob("slide-*.png"))) != expected:
            raise FatalGenerationError("rendered slide image dimensions or count are inconsistent")


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()
