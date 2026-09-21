from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path
from typing import Any

import fitz
from PIL import Image, UnidentifiedImageError
from pptx import Presentation

from metaclass.modules.paper_workflow.schemas import PaperPresentationArtifact


class SourceVisualValidationError(ValueError):
    """Raised when a Paper Deck artifact cannot prove source-visual fidelity."""


class SourceVisualValidator:
    """Validate Paper Deck source-grounded output without trusting model attestations."""

    _heading = re.compile(r"^##\s+(\d+)[.)]?\s*(.+?)\s*$", re.MULTILINE)
    _field = re.compile(r"^-\s*([^:：]+)[:：]\s*(.*?)\s*$", re.MULTILINE)
    _asset_id = re.compile(r"^\s*-?\s*asset_id\s*[:：]\s*([^\s#]+)\s*$", re.MULTILINE)
    _source_visual_terms = re.compile(
        r"\b(?:figure|fig\.?|table|plot|chart|screenshot)\b|图\s*\d|表\s*\d|曲线|截图",
        re.IGNORECASE,
    )
    _empty_source_visual = re.compile(
        r"^(?:none|no|n/?a|generated|无|无。|不使用|不直接嵌图)", re.IGNORECASE
    )
    _forbidden_prompt_phrases = (
        "use the supplied",
        "use source image",
        "insert the paper figure",
        "embed the paper figure",
    )
    _figure_label = re.compile(
        r"^\s*((?:fig(?:ure)?\.?|table)\s*[A-Za-z0-9._-]+|(?:图|表)\s*[A-Za-z0-9._-]+)",
        re.IGNORECASE,
    )

    def validate(self, artifact: PaperPresentationArtifact, *, workspace: Path) -> None:
        root = workspace.resolve()
        manifest_path = self._resolve(root, artifact.source_visual_manifest_path)
        outline_path = self._resolve(root, artifact.outline_path)
        prompts_dir = self._resolve(root, artifact.prompts_dir, directory=True)
        generation_log = self._resolve(root, artifact.generation_log_path)
        pdf_path = self._resolve(root, artifact.presentation_pdf_path)
        pptx_value = artifact.debug_pptx_path or str(
            Path(artifact.presentation_pdf_path).with_suffix(".pptx")
        )
        pptx_path = self._resolve(root, pptx_value)
        manifest = self._load_json(manifest_path, "source visual manifest")
        slides = manifest.get("slides")
        if manifest.get("schema_version") != "1.0" or not isinstance(slides, list):
            raise SourceVisualValidationError("source visual manifest schema is invalid")
        outline_slides = self._parse_outline(outline_path)
        prompt_paths = self._numbered_files(prompts_dir, {".md"})
        presentation = Presentation(pptx_path)
        with fitz.open(pdf_path) as document:
            pdf_pages = document.page_count
        counts = {
            "manifest": len(slides),
            "outline": len(outline_slides),
            "prompts": len(prompt_paths),
            "pptx": len(presentation.slides),
            "pdf": pdf_pages,
        }
        if len(set(counts.values())) != 1 or not slides:
            raise SourceVisualValidationError(
                "PDF, PPTX, manifest, outline, and prompt page counts must match: "
                + ", ".join(f"{key}={value}" for key, value in counts.items())
            )
        source_assets = self._source_assets(root)
        selected_assets: set[str] = set()
        log_text = generation_log.read_text(encoding="utf-8")
        for index, (entry, outline, prompt_path, ppt_slide) in enumerate(
            zip(slides, outline_slides, prompt_paths, presentation.slides, strict=True),
            start=1,
        ):
            if not isinstance(entry, dict) or entry.get("order") != index:
                raise SourceVisualValidationError("manifest slide order must be continuous from 1")
            mode = entry.get("render_mode")
            if mode not in {"native-raster", "source-grounded-hybrid"}:
                raise SourceVisualValidationError(f"slide {index} has an invalid render_mode")
            if outline["render_mode"] != mode:
                raise SourceVisualValidationError(
                    f"slide {index} render_mode differs between outline and manifest"
                )
            assets = entry.get("assets")
            if not isinstance(assets, list):
                raise SourceVisualValidationError(f"slide {index} assets must be an array")
            outline_asset_ids = set(outline["asset_ids"])
            manifest_asset_ids = {
                str(item.get("asset_id")) for item in assets if isinstance(item, dict)
            }
            if outline_asset_ids != manifest_asset_ids:
                raise SourceVisualValidationError(
                    f"slide {index} asset_ids differ between outline and manifest"
                )
            requires_source = bool(outline_asset_ids) or self._requires_source_visual(outline)
            if requires_source and mode != "source-grounded-hybrid":
                raise SourceVisualValidationError(
                    f"slide {index} uses Figure/Table evidence but is not source-grounded-hybrid"
                )
            if mode == "native-raster":
                if assets:
                    raise SourceVisualValidationError(
                        f"native-raster slide {index} cannot contain source assets"
                    )
                if len(ppt_slide.shapes) != 1 or not ppt_slide.shapes[0].name.startswith(
                    "background:native-raster"
                ):
                    raise SourceVisualValidationError(
                        f"native-raster slide {index} must remain one full-slide raster object"
                    )
            else:
                if not assets:
                    raise SourceVisualValidationError(
                        f"source-grounded-hybrid slide {index} has no source asset"
                    )
                self._validate_prompt(prompt_path, assets, slide_no=index)
            self._validate_generation_log(log_text, entry, slide_no=index)
            for asset in assets:
                asset_id = str(asset.get("asset_id", ""))
                if asset_id not in source_assets:
                    raise SourceVisualValidationError(
                        f"slide {index} references unknown asset_id: {asset_id}"
                    )
                source_record = source_assets[asset_id]
                source_path = self._validated_source_path(root, asset)
                self._validate_source_identity(
                    root,
                    source_path,
                    asset,
                    source_record,
                    slide_no=index,
                )
                self._validate_hash(source_path, str(asset.get("sha256", "")))
                self._validate_picture_object(
                    ppt_slide,
                    source_path,
                    asset,
                    slide_width=presentation.slide_width,
                    slide_height=presentation.slide_height,
                    slide_no=index,
                )
                selected_assets.add(asset_id)
        self._validate_core_figure_selection(root, source_assets, selected_assets)

    @classmethod
    def _parse_outline(cls, path: Path) -> list[dict[str, Any]]:
        text = path.read_text(encoding="utf-8")
        headings = list(cls._heading.finditer(text))
        if not headings:
            raise SourceVisualValidationError("outline contains no numbered slides")
        slides: list[dict[str, Any]] = []
        for index, heading in enumerate(headings):
            end = headings[index + 1].start() if index + 1 < len(headings) else len(text)
            body = text[heading.end() : end]
            fields = {
                match.group(1).strip().casefold(): match.group(2).strip()
                for match in cls._field.finditer(body)
            }
            render_mode = fields.get("render mode")
            if render_mode not in {"native-raster", "source-grounded-hybrid"}:
                raise SourceVisualValidationError(
                    f"outline slide {index + 1} is missing a valid Render mode"
                )
            slides.append(
                {
                    "order": int(heading.group(1)),
                    "role": fields.get("role", "").casefold(),
                    "source_visual": fields.get("source visual", ""),
                    "render_mode": render_mode,
                    "asset_ids": cls._asset_id.findall(body),
                }
            )
        if [slide["order"] for slide in slides] != list(range(1, len(slides) + 1)):
            raise SourceVisualValidationError("outline slide order must be continuous from 1")
        return slides

    @classmethod
    def _requires_source_visual(cls, outline: dict[str, Any]) -> bool:
        value = str(outline.get("source_visual", "")).strip()
        if not value or cls._empty_source_visual.match(value):
            return False
        return bool(cls._source_visual_terms.search(value)) and outline.get("role") in {
            "evidence",
            "result",
            "comparison",
            "method",
        }

    @classmethod
    def _source_assets(cls, root: Path) -> dict[str, dict[str, Any]]:
        path = root / "provider_input/paper_source.json"
        if not path.is_file():
            raise SourceVisualValidationError("provider_input/paper_source.json is missing")
        payload = cls._load_json(path, "paper source")
        assets = payload.get("assets", [])
        if not isinstance(assets, list):
            raise SourceVisualValidationError("paper source assets must be an array")
        return {
            str(item["id"]): item
            for item in assets
            if isinstance(item, dict) and isinstance(item.get("id"), str)
        }

    @staticmethod
    def _validated_source_path(root: Path, asset: dict[str, Any]) -> Path:
        value = str(asset.get("source_path", ""))
        relative = Path(value)
        if relative.is_absolute() or ".." in relative.parts:
            raise SourceVisualValidationError(f"unsafe source asset path: {value}")
        allowed = (root / "provider_input/assets").resolve()
        resolved = (root / relative).resolve()
        try:
            resolved.relative_to(allowed)
        except ValueError as exc:
            raise SourceVisualValidationError(
                f"source asset path must stay inside provider_input/assets: {value}"
            ) from exc
        if not resolved.is_file():
            raise SourceVisualValidationError(f"source asset file is missing: {value}")
        return resolved

    @staticmethod
    def _validate_hash(path: Path, expected: str) -> None:
        normalized = expected.removeprefix("sha256:").casefold()
        if not re.fullmatch(r"[0-9a-f]{64}", normalized):
            raise SourceVisualValidationError(f"invalid source asset SHA-256: {path.name}")
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        if digest != normalized:
            raise SourceVisualValidationError(f"source asset SHA-256 mismatch: {path.name}")

    @classmethod
    def _validate_source_identity(
        cls,
        root: Path,
        source_path: Path,
        manifest_asset: dict[str, Any],
        source_asset: dict[str, Any],
        *,
        slide_no: int,
    ) -> None:
        raw_path = Path(str(source_asset.get("path", "")))
        if raw_path.is_absolute() or ".." in raw_path.parts or not raw_path.parts:
            raise SourceVisualValidationError(
                f"paper_source.json has an unsafe path for slide {slide_no}"
            )
        relative = (
            Path(*raw_path.parts[1:])
            if raw_path.parts[0] in {"assets", "existing_assets"}
            else raw_path
        )
        expected_path = (root / "provider_input/assets" / relative).resolve()
        if source_path != expected_path:
            raise SourceVisualValidationError(
                f"slide {slide_no} source_path does not match paper_source.json"
            )
        if manifest_asset.get("source_page") != source_asset.get("page_no"):
            raise SourceVisualValidationError(
                f"slide {slide_no} source_page does not match paper_source.json"
            )
        caption = str(source_asset.get("caption") or "")
        matched = cls._figure_label.match(caption)
        if not matched:
            raise SourceVisualValidationError(
                f"paper_source.json lacks a canonical Figure/Table label for slide {slide_no}"
            )
        expected_label = re.sub(r"\s+", " ", matched.group(1)).casefold()
        actual_label = re.sub(
            r"\s+", " ", str(manifest_asset.get("figure_label", "")).strip()
        ).casefold()
        if actual_label != expected_label:
            raise SourceVisualValidationError(
                f"slide {slide_no} figure_label does not match paper_source.json"
            )
        source_digest = hashlib.sha256(expected_path.read_bytes()).hexdigest()
        manifest_digest = str(manifest_asset.get("sha256", "")).removeprefix("sha256:").casefold()
        if manifest_digest != source_digest:
            raise SourceVisualValidationError(
                f"slide {slide_no} sha256 does not match paper_source.json asset"
            )

    @classmethod
    def _validate_picture_object(
        cls,
        slide: Any,
        source_path: Path,
        asset: dict[str, Any],
        *,
        slide_width: int,
        slide_height: int,
        slide_no: int,
    ) -> None:
        expected_name = f"source:{asset['asset_id']}"
        pictures = [shape for shape in slide.shapes if shape.name == expected_name]
        if len(pictures) != 1:
            raise SourceVisualValidationError(
                f"slide {slide_no} does not contain exactly one independent {expected_name} picture"
            )
        picture = pictures[0]
        embedded_digest = hashlib.sha256(picture.image.blob).hexdigest()
        expected_digest = str(asset.get("sha256", "")).removeprefix("sha256:").casefold()
        if embedded_digest != expected_digest:
            raise SourceVisualValidationError(
                f"slide {slide_no} embedded picture bytes do not match source asset"
            )
        try:
            with Image.open(source_path) as image:
                source_ratio = image.width / image.height
        except (OSError, UnidentifiedImageError, ZeroDivisionError) as exc:
            raise SourceVisualValidationError(
                f"source asset is not a readable raster image: {source_path.name}"
            ) from exc
        shape_ratio = picture.width / picture.height
        fit = asset.get("fit", "contain")
        if fit == "contain" and abs(shape_ratio / source_ratio - 1) > 0.02:
            raise SourceVisualValidationError(
                f"slide {slide_no} source picture aspect ratio changed unexpectedly"
            )
        crop_values = (
            picture.crop_left,
            picture.crop_right,
            picture.crop_top,
            picture.crop_bottom,
        )
        if fit == "contain" and any(abs(value) > 1e-6 for value in crop_values):
            raise SourceVisualValidationError(
                f"slide {slide_no} contain picture must not be cropped"
            )
        if fit == "cover" and any(value < 0 or value >= 1 for value in crop_values):
            raise SourceVisualValidationError(
                f"slide {slide_no} cover picture has invalid crop values"
            )
        if fit == "cover" and not any(
            value > 0
            for value in crop_values
        ) and abs(shape_ratio / source_ratio - 1) > 0.02:
            raise SourceVisualValidationError(
                f"slide {slide_no} cover picture is stretched instead of cropped"
            )
        cls._validate_picture_geometry(
            picture,
            slide_width=slide_width,
            slide_height=slide_height,
            source_ratio=source_ratio,
            asset=asset,
            slide_no=slide_no,
        )

    @staticmethod
    def _validate_picture_geometry(
        picture: Any,
        *,
        slide_width: int,
        slide_height: int,
        source_ratio: float,
        asset: dict[str, Any],
        slide_no: int,
    ) -> None:
        x, y, w, h = (float(asset[key]) for key in ("x", "y", "w", "h"))
        left = x * slide_width
        top = y * slide_height
        width = w * slide_width
        height = h * slide_height
        frame_ratio = width / height
        if asset.get("fit", "contain") == "contain":
            if source_ratio >= frame_ratio:
                expected = (left, top + (height - width / source_ratio) / 2, width, width / source_ratio)
            else:
                expected = (left + (width - height * source_ratio) / 2, top, height * source_ratio, height)
        else:
            expected = (left, top, width, height)
        actual = (picture.left, picture.top, picture.width, picture.height)
        if any(
            abs(actual_value - expected_value) > 3
            for actual_value, expected_value in zip(actual, expected, strict=True)
        ):
            raise SourceVisualValidationError(
                f"slide {slide_no} source picture geometry differs from manifest placement"
            )

    @classmethod
    def _validate_prompt(
        cls, path: Path, assets: list[dict[str, Any]], *, slide_no: int
    ) -> None:
        text = path.read_text(encoding="utf-8").casefold()
        forbidden = list(cls._forbidden_prompt_phrases)
        for asset in assets:
            source_path = str(asset.get("source_path", "")).casefold()
            if source_path:
                forbidden.extend([source_path, Path(source_path).name])
        matched = next((token for token in forbidden if token and token in text), None)
        if matched:
            raise SourceVisualValidationError(
                f"slide {slide_no} sends source evidence to image generation: {matched}"
            )
        reserves_empty_bay = bool(
            re.search(r"reserve\s+(?:an?|one|two|multiple)?\s*empty evidence bays?", text)
        )
        defers_insertion = bool(
            re.search(
                r"(?:inserted|added).{0,60}(?:deterministically|after generation|afterward|later)",
                text,
            )
        )
        prohibits_redrawing = "do not draw" in text and "insert" in text
        if not (reserves_empty_bay and defers_insertion and prohibits_redrawing):
            raise SourceVisualValidationError(
                f"slide {slide_no} hybrid prompt does not reserve a deterministic evidence bay"
            )

    @staticmethod
    def _validate_generation_log(text: str, entry: dict[str, Any], *, slide_no: int) -> None:
        mode = re.escape(str(entry["render_mode"]))
        background = Path(str(entry.get("background_path", ""))).name
        lines = [line for line in text.splitlines() if background in line]
        if not lines or not any(
            re.search(rf"render_mode\s*[:=]\s*{mode}\b", line, re.IGNORECASE)
            for line in lines
        ):
            raise SourceVisualValidationError(
                f"generation-log.md does not record render_mode for slide {slide_no}"
            )
        if entry["render_mode"] == "source-grounded-hybrid":
            joined = " ".join(lines).casefold()
            required = ("background_backend", "source_asset", "source_asset_sha256", "composition_backend")
            if not all(token in joined for token in required):
                raise SourceVisualValidationError(
                    f"generation-log.md lacks source-grounded provenance for slide {slide_no}"
                )

    @classmethod
    def _validate_core_figure_selection(
        cls,
        root: Path,
        source_assets: dict[str, dict[str, Any]],
        selected_assets: set[str],
    ) -> None:
        analysis_path = root / "provider_input/paper_analysis.json"
        if not analysis_path.is_file():
            return
        analysis = cls._load_json(analysis_path, "paper analysis")
        core_claims = {
            str(item.get("id"))
            for item in analysis.get("claims", [])
            if isinstance(item, dict) and item.get("importance") == "core"
        }
        core_assets: set[str] = set()
        for candidate in analysis.get("figure_candidates", []):
            if not isinstance(candidate, dict) or not core_claims.intersection(
                map(str, candidate.get("claim_ids", []))
            ):
                continue
            for ref in candidate.get("source_refs", []):
                if isinstance(ref, dict) and ref.get("asset_id") in source_assets:
                    core_assets.add(str(ref["asset_id"]))
        if core_assets and not core_assets.intersection(selected_assets):
            raise SourceVisualValidationError(
                "paper contains usable core Figure/Table assets but none were selected"
            )

    @classmethod
    def _numbered_files(cls, directory: Path, suffixes: set[str]) -> list[Path]:
        numbered: list[tuple[int, Path]] = []
        for path in directory.iterdir():
            matched = re.match(r"^(\d+)[-_].+", path.name)
            if matched and path.is_file() and path.suffix.casefold() in suffixes:
                numbered.append((int(matched.group(1)), path))
        numbered.sort(key=lambda item: item[0])
        if [number for number, _ in numbered] != list(range(1, len(numbered) + 1)):
            raise SourceVisualValidationError("prompt numbering must be continuous from 1")
        return [path for _, path in numbered]

    @staticmethod
    def _load_json(path: Path, label: str) -> dict[str, Any]:
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise SourceVisualValidationError(f"{label} is not valid JSON") from exc
        if not isinstance(payload, dict):
            raise SourceVisualValidationError(f"{label} must be a JSON object")
        return payload

    @staticmethod
    def _resolve(root: Path, value: str, *, directory: bool = False) -> Path:
        relative = Path(value)
        if relative.is_absolute() or ".." in relative.parts:
            raise SourceVisualValidationError("artifact path is unsafe")
        resolved = (root / relative).resolve()
        try:
            resolved.relative_to(root)
        except ValueError as exc:
            raise SourceVisualValidationError("artifact path escapes workspace") from exc
        exists = resolved.is_dir() if directory else resolved.is_file()
        if not exists:
            raise SourceVisualValidationError(f"artifact is missing: {value}")
        return resolved
