from collections.abc import Iterator
from io import BytesIO
from pathlib import Path

import fitz
import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation

from metaclass.main import create_app


@pytest.fixture
def client(tmp_path: Path) -> Iterator[TestClient]:
    with TestClient(create_app(tmp_path)) as test_client:
        yield test_client


def make_pdf() -> bytes:
    document = fitz.open()
    page = document.new_page(width=800, height=450)
    page.insert_text((72, 90), "Matrix Multiplication", fontsize=28)
    page.insert_text((72, 140), "Rows are multiplied by columns.", fontsize=16)
    payload = document.tobytes()
    document.close()
    return payload


def make_pdf_with_embedded_image() -> bytes:
    image_stream = BytesIO()
    Image.new("RGB", (80, 40), "#336699").save(image_stream, format="PNG")
    document = fitz.open()
    page = document.new_page(width=800, height=450)
    page.insert_text((72, 90), "Image Example", fontsize=28)
    page.insert_image(fitz.Rect(72, 140, 232, 220), stream=image_stream.getvalue())
    payload = document.tobytes()
    document.close()
    return payload


def make_pptx() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Classroom Plan"
    slide.placeholders[1].text = "A plan contains typed teaching actions."
    stream = BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def test_pptx_upload_and_parse(client: TestClient) -> None:
    upload = client.post(
        "/api/v1/materials",
        files={
            "file": (
                "lesson.pptx",
                make_pptx(),
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )
        },
    )
    parsed = client.post(f"/api/v1/materials/{upload.json()['id']}/parse")
    assert parsed.status_code == 200, parsed.text
    assert parsed.json()[0]["title"] == "Classroom Plan"
    assert Path(parsed.json()[0]["image_path"]).exists()


def test_combined_upload_and_parse(client: TestClient) -> None:
    response = client.post(
        "/api/v1/materials/process",
        files={"file": ("lesson.pdf", make_pdf(), "application/pdf")},
    )
    assert response.status_code == 201
    assert response.json()["material"]["status"] == "parsed"
    assert response.json()["pages"][0]["source_refs"][0]["page_no"] == 1


def test_pdf_parse_extracts_embedded_images(client: TestClient) -> None:
    response = client.post(
        "/api/v1/materials/process",
        files={"file": ("lesson.pdf", make_pdf_with_embedded_image(), "application/pdf")},
    )

    assert response.status_code == 201, response.text
    images = response.json()["pages"][0]["embedded_images"]
    assert len(images) == 1
    assert images[0]["page_no"] == 1
    assert images[0]["width"] == 80
    assert images[0]["height"] == 40
    assert Path(images[0]["image_path"]).exists()


def test_batch_upload_and_parse_accepts_pdf_and_pptx(client: TestClient) -> None:
    response = client.post(
        "/api/v1/materials/batch-process",
        files=[
            ("files", ("lesson.pdf", make_pdf(), "application/pdf")),
            (
                "files",
                (
                    "lesson.pptx",
                    make_pptx(),
                    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                ),
            ),
        ],
    )

    assert response.status_code == 201, response.text
    payload = response.json()
    items = payload["items"]
    assert len(items) == 2
    assert {item["material"]["file_type"] for item in items} == {"pdf", "pptx"}
    assert all(item["material"]["file_hash"] for item in items)
    assert all(item["material"]["status"] == "parsed" for item in items)
    assert all(item["pages"] for item in items)
    assert payload["collection"]["material_ids"] == [item["material"]["id"] for item in items]

    listed = client.get("/api/v1/materials")
    assert listed.status_code == 200
    assert len(listed.json()) == 2

    collections = client.get("/api/v1/materials/collections")
    assert collections.status_code == 200
    assert collections.json()[0]["id"] == payload["collection"]["id"]


def test_material_processing_job_can_be_polled_for_result(client: TestClient) -> None:
    created = client.post(
        "/api/v1/materials/processing-jobs",
        files=[
            ("files", ("lesson.pdf", make_pdf(), "application/pdf")),
            (
                "files",
                (
                    "lesson.pptx",
                    make_pptx(),
                    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                ),
            ),
        ],
    )
    assert created.status_code == 202, created.text
    job_id = created.json()["id"]

    job = client.get(f"/api/v1/materials/processing-jobs/{job_id}")
    assert job.status_code == 200
    assert job.json()["status"] == "succeeded"
    assert job.json()["progress"] == 100

    result = client.get(f"/api/v1/materials/processing-jobs/{job_id}/result")
    assert result.status_code == 200
    assert len(result.json()["items"]) == 2
    assert result.json()["collection"]["material_ids"] == job.json()["material_ids"]


def test_learning_content_job_can_be_polled_for_result(client: TestClient) -> None:
    processed = client.post(
        "/api/v1/materials/process",
        files={"file": ("lesson.pdf", make_pdf(), "application/pdf")},
    )
    material_id = processed.json()["material"]["id"]

    created = client.post(f"/api/v1/materials/{material_id}/learning-content-jobs")
    assert created.status_code == 202, created.text
    job_id = created.json()["id"]

    job = client.get(f"/api/v1/learning-content-jobs/{job_id}")
    assert job.status_code == 200
    assert job.json()["status"] == "succeeded"
    assert job.json()["progress"] == 100
    assert job.json()["content_id"]

    result = client.get(f"/api/v1/learning-content-jobs/{job_id}/result")
    assert result.status_code == 200
    payload = result.json()
    assert payload["id"] == job.json()["content_id"]
    assert payload["sections"]
    assert payload["knowledge_tree"]["root_node_ids"]
    assert payload["sections"][0]["tree_node_ids"]
    assert payload["quality"]["coverage_score"] == 1.0


def test_collection_learning_content_job_uses_all_materials(client: TestClient) -> None:
    processed = client.post(
        "/api/v1/materials/batch-process",
        files=[
            ("files", ("lesson.pdf", make_pdf(), "application/pdf")),
            (
                "files",
                (
                    "lesson.pptx",
                    make_pptx(),
                    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                ),
            ),
        ],
    )
    assert processed.status_code == 201, processed.text
    collection = processed.json()["collection"]

    created = client.post(f"/api/v1/material-collections/{collection['id']}/learning-content-jobs")
    assert created.status_code == 202, created.text
    job_id = created.json()["id"]

    job = client.get(f"/api/v1/learning-content-jobs/{job_id}")
    assert job.status_code == 200
    assert job.json()["status"] == "succeeded"

    result = client.get(f"/api/v1/learning-content-jobs/{job_id}/result")
    assert result.status_code == 200
    payload = result.json()
    assert payload["collection_id"] == collection["id"]
    assert payload["material_ids"] == collection["material_ids"]
    assert payload["knowledge_units"]
    assert payload["knowledge_tree"]["nodes"]
    assert payload["knowledge_units"][0]["source_excerpts"]
    assert payload["sections"]
    assert payload["sections"][0]["page_refs"]
    assert payload["sections"][0]["tree_node_ids"]
    assert len(payload["sections"]) <= len(payload["knowledge_units"])
    assert any(section["source_excerpts"] for section in payload["sections"])
    assert payload["quality"]["material_coverage"] == 1.0
    assert payload["quality"]["knowledge_coverage"] == 1.0

    diagnostics = client.get(f"/api/v1/learning-contents/{payload['id']}/diagnostics")
    assert diagnostics.status_code == 200
    assert diagnostics.json()["knowledge_tree"]["id"] == payload["knowledge_tree"]["id"]
    assert diagnostics.json()["quality"]["coverage_score"] == 1.0


def test_complete_mvp_flow(client: TestClient) -> None:

    upload = client.post(
        "/api/v1/materials",
        files={"file": ("lesson.pdf", make_pdf(), "application/pdf")},
    )
    assert upload.status_code == 201
    material_id = upload.json()["id"]

    parsed = client.post(f"/api/v1/materials/{material_id}/parse")
    assert parsed.status_code == 200
    assert parsed.json()[0]["source_refs"][0]["page_no"] == 1
    assert "summary" not in parsed.json()[0]
    assert "knowledge_points" not in parsed.json()[0]
    assert Path(parsed.json()[0]["image_path"]).exists()

    content = client.post(f"/api/v1/materials/{material_id}/learning-content")
    assert content.status_code == 201
    content_id = content.json()["id"]
    assert content.json()["sections"][0]["source_refs"]
    assert content.json()["sections"][0]["quiz_items"]
    assert content.json()["sections"][0]["quiz_items"][0]["question"] != "本页主要讲解的内容是？"
    understandings = client.get(f"/api/v1/materials/{material_id}/understandings")
    assert understandings.status_code == 200
    assert understandings.json()[0]["provider"] == "fake"
    assert understandings.json()[0]["summary"]

    plan = client.post(f"/api/v1/learning-contents/{content_id}/classroom-plans")
    assert plan.status_code == 201
    assert [item["type"] for item in plan.json()["scenes"][0]["actions"]] == [
        "SHOW_PAGE",
        "EXPLAIN",
        "PROBE",
        "ASK_QUIZ",
        "GIVE_FEEDBACK",
        "END",
    ]
    plan_meta = client.get(f"/api/v1/classroom-plans/{plan.json()['id']}/generation-meta")
    assert plan_meta.status_code == 200
    assert plan_meta.json()["plan_id"] == plan.json()["id"]
    assert plan_meta.json()["source"] == "llm"
    assert plan_meta.json()["provider"] == "fake"

    session = client.post(
        f"/api/v1/classroom-plans/{plan.json()['id']}/sessions",
        json={"mode": "interactive"},
    )
    session_id = session.json()["id"]
    assert session.json()["mode"] == "interactive"
    assert len(session.json()["student_states"]) == 4
    state = client.get(f"/api/v1/classroom-sessions/{session_id}/state")
    assert state.status_code == 200
    assert state.json()["current_action_type"] == "SHOW_PAGE"
    assert [student["display_name"] for student in state.json()["students"]] == [
        "课堂气氛调节者",
        "深度思考者",
        "课堂笔记员",
        "研究型同学",
    ]
    assert (
        client.post(f"/api/v1/classroom-sessions/{session_id}/next").json()["action"]["type"]
        == "SHOW_PAGE"
    )
    assert (
        client.post(f"/api/v1/classroom-sessions/{session_id}/next").json()["action"]["type"]
        == "EXPLAIN"
    )
    probe = client.post(f"/api/v1/classroom-sessions/{session_id}/next").json()
    assert probe["action"]["type"] == "PROBE"
    quiz = client.post(f"/api/v1/classroom-sessions/{session_id}/next").json()
    assert quiz["action"]["type"] == "ASK_QUIZ"
    directed = client.post(f"/api/v1/classroom-sessions/{session_id}/agent-turns/next")
    assert directed.status_code == 200
    assert directed.json()["decision"]["next_role"] == "student"
    assert directed.json()["turns"][0]["role"] == "student"
    recorded_session = client.get(f"/api/v1/classroom-sessions/{session_id}")
    assert recorded_session.json()["events"][-1]["type"] == "AGENT_TURN"
    assert (
        recorded_session.json()["events"][-1]["payload"]["turn"]["agent_id"]
        == directed.json()["turns"][0]["agent_id"]
    )

    answer = client.post(
        f"/api/v1/classroom-sessions/{session_id}/answers", json={"selected_index": 0}
    )
    assert answer.json()["correct"] is True
    assert answer.json()["session"]["mastery"][0]["value"] == 1.0

    question = client.post(
        f"/api/v1/classroom-sessions/{session_id}/questions",
        json={"question": "How does it work?"},
    )
    assert "How does it work?" in question.json()["feedback"]
    assert question.json()["source_refs"][0]["page_no"] == 1

    services = client.app.state.services
    services.classrooms.repository.save_session(
        services.classrooms.get_session(session_id).model_copy(update={"status": "completed"})
    )
    replay = client.post(f"/api/v1/classroom-sessions/{session_id}/replay")
    assert replay.status_code == 201
    assert replay.json()["id"] != session_id
    assert replay.json()["plan_id"] == plan.json()["id"]
    assert replay.json()["status"] == "running"
    assert replay.json()["scene_index"] == 0
    assert replay.json()["action_index"] == 0
    assert replay.json()["events"] == []

    lecture_plan = client.post(
        f"/api/v1/classroom-plans/{plan.json()['id']}/lecture-variant"
    )
    assert lecture_plan.status_code == 201
    assert lecture_plan.json()["id"] != plan.json()["id"]
    lecture_action_types = {
        action["type"]
        for scene in lecture_plan.json()["scenes"]
        for action in scene["actions"]
    }
    assert lecture_action_types <= {"SHOW_PAGE", "EXPLAIN", "END"}
    assert {"SHOW_PAGE", "EXPLAIN"} <= lecture_action_types

    after_class_question = client.post(
        f"/api/v1/classroom-sessions/{session_id}/questions",
        json={"question": "Can I still ask after class?"},
    )
    assert after_class_question.status_code == 200
    assert after_class_question.json()["status"] == "answered"
    assert after_class_question.json()["session"]["status"] == "completed"

    teacher_turn = client.post(
        f"/api/v1/classroom-sessions/{session_id}/teacher-turn",
        json={"prompt": "请追问一下这个知识点"},
    )
    assert teacher_turn.status_code == 200
    assert teacher_turn.json()["role"] == "teacher"
    assert teacher_turn.json()["speech"]

    student_turns = client.post(
        f"/api/v1/classroom-sessions/{session_id}/student-turns",
        json={"prompt": "老师刚刚讲完，请同学们反馈"},
    )
    assert student_turns.status_code == 200
    assert len(student_turns.json()) == 4
    assert {turn["role"] for turn in student_turns.json()} == {"student"}

    presentation_plan = client.post(
        f"/api/v1/learning-contents/{content_id}/presentation-plans"
    )
    assert presentation_plan.status_code == 201
    ppt_job = client.post(
        f"/api/v1/presentation-plans/{presentation_plan.json()['id']}/ppt-jobs"
    )
    assert ppt_job.status_code == 202
    finished_ppt_job = client.get(f"/api/v1/ppt-jobs/{ppt_job.json()['id']}")
    assert finished_ppt_job.json()["status"] == "finished", finished_ppt_job.json().get(
        "error"
    )
    ppt_artifact = client.get(
        f"/api/v1/ppt-jobs/{finished_ppt_job.json()['id']}/artifact"
    )
    assert ppt_artifact.status_code == 200

    job = client.post(
        f"/api/v1/learning-contents/{content_id}/videos",
        params={"presentation_artifact_id": ppt_artifact.json()["id"]},
    )
    assert job.status_code == 202
    assert job.json()["status"] == "pending"
    finished_job = client.get(f"/api/v1/video-jobs/{job.json()['id']}")
    assert finished_job.json()["status"] == "finished", finished_job.json().get("error")
    assert finished_job.json()["progress"] == 1.0
    result = client.get(f"/api/v1/video-jobs/{finished_job.json()['id']}/result")
    assert result.status_code == 200
    assert result.json()["job_id"] == job.json()["id"]
    assert Path(result.json()["video_path"]).stat().st_size > 0
    assert Path(result.json()["subtitles_path"]).read_text(encoding="utf-8")
    assert result.json()["duration_seconds"] > 0


def test_classroom_plan_job_generates_plan_without_replacing_sync_api(client: TestClient) -> None:
    processed = client.post(
        "/api/v1/materials/process",
        files={"file": ("lesson.pdf", make_pdf(), "application/pdf")},
    )
    material_id = processed.json()["material"]["id"]
    content = client.post(f"/api/v1/materials/{material_id}/learning-content")
    content_id = content.json()["id"]

    created = client.post(f"/api/v1/learning-contents/{content_id}/classroom-plan-jobs")

    assert created.status_code == 202
    job_id = created.json()["id"]
    assert created.json()["status"] in {"queued", "succeeded"}

    job = client.get(f"/api/v1/classroom-plan-jobs/{job_id}")
    assert job.status_code == 200
    assert job.json()["status"] == "succeeded"
    assert job.json()["progress"] == 100
    assert job.json()["plan_id"]

    plan = client.get(f"/api/v1/classroom-plans/{job.json()['plan_id']}")
    assert plan.status_code == 200
    assert plan.json()["content_id"] == content_id
    assert plan.json()["scenes"]


def test_classroom_session_uses_selected_student_agent_types(client: TestClient) -> None:
    processed = client.post(
        "/api/v1/materials/process",
        files={"file": ("lesson.pdf", make_pdf(), "application/pdf")},
    )
    material_id = processed.json()["material"]["id"]
    content = client.post(f"/api/v1/materials/{material_id}/learning-content")
    plan = client.post(f"/api/v1/learning-contents/{content.json()['id']}/classroom-plans")

    session = client.post(
        f"/api/v1/classroom-plans/{plan.json()['id']}/sessions",
        json={
            "mode": "interactive",
            "student_agent_types": ["foundation_weak", "concept_confused"],
        },
    )

    assert session.status_code == 201
    assert [student["agent_type"] for student in session.json()["student_states"]] == [
        "foundation_weak",
        "concept_confused",
    ]
    session_id = session.json()["id"]

    client.post(f"/api/v1/classroom-sessions/{session_id}/auto-step")
    client.post(f"/api/v1/classroom-sessions/{session_id}/auto-step")
    client.post(f"/api/v1/classroom-sessions/{session_id}/auto-step")
    dialog = client.post(f"/api/v1/classroom-sessions/{session_id}/auto-step")

    assert dialog.status_code == 200
    assert dialog.json()["status"] == "agent_turn"
    assert dialog.json()["directed_turn"]["turns"][0]["agent_id"] in {
        "student_agent_001",
        "student_agent_002",
    }
    assert dialog.json()["session"]["student_states"][1]["agent_type"] == "concept_confused"


def test_presentation_plan_and_ppt_skill_request_flow(client: TestClient) -> None:
    processed = client.post(
        "/api/v1/materials/process",
        files={"file": ("lesson.pdf", make_pdf(), "application/pdf")},
    )
    material_id = processed.json()["material"]["id"]
    content = client.post(f"/api/v1/materials/{material_id}/learning-content")
    content_id = content.json()["id"]

    plan = client.post(f"/api/v1/learning-contents/{content_id}/presentation-plans")

    assert plan.status_code == 201
    assert plan.json()["content_id"] == content_id
    assert plan.json()["slides"][0]["source_section_ids"]
    assert plan.json()["slides"][0]["speaker_script"]

    question_bank = client.get(
        f"/api/v1/presentation-plans/{plan.json()['id']}/question-bank"
    )
    assert question_bank.status_code == 200
    assert question_bank.json()["items"]
    prepared_question = question_bank.json()["items"][0]
    assert prepared_question["slide_id"] == plan.json()["slides"][0]["id"]
    assert prepared_question["agent_type"]
    assert prepared_question["canonical_answer"]
    assert prepared_question["teacher_answer"]
    source_slide = next(
        slide
        for slide in plan.json()["slides"]
        if slide["id"] == prepared_question["slide_id"]
    )
    assert prepared_question["teacher_answer"] != source_slide["speaker_script"]
    assert prepared_question["canonical_question"] in prepared_question["teacher_answer"]
    assert prepared_question["moment"] == "after_explanation"

    search = client.get(
        f"/api/v1/presentation-plans/{plan.json()['id']}/question-bank/search",
        params={"q": prepared_question["canonical_question"]},
    )
    assert search.status_code == 200
    assert search.json()[0]["item"]["id"] == prepared_question["id"]

    latest = client.get(f"/api/v1/learning-contents/{content_id}/presentation-plan")
    assert latest.status_code == 200
    assert latest.json()["id"] == plan.json()["id"]

    themes = client.get("/api/v1/ppt-themes")
    assert themes.status_code == 200
    assert {theme["id"] for theme in themes.json()} >= {
        "academic_blue",
        "scholar_green",
        "deep_technology",
    }

    invalid_theme_job = client.post(
        f"/api/v1/presentation-plans/{plan.json()['id']}/ppt-jobs",
        json={"theme_id": "unknown_theme"},
    )
    assert invalid_theme_job.status_code == 422

    job = client.post(
        f"/api/v1/presentation-plans/{plan.json()['id']}/ppt-jobs",
        json={"theme_id": "scholar_green"},
    )
    assert job.status_code == 202
    assert job.json()["theme_id"] == "scholar_green"
    assert job.json()["status"] in {"queued", "running", "finished"}

    job = client.get(f"/api/v1/ppt-jobs/{job.json()['id']}")
    assert job.status_code == 200
    assert job.json()["status"] == "finished", job.json().get("error")
    assert job.json()["artifact_id"]

    artifact = client.get(f"/api/v1/ppt-jobs/{job.json()['id']}/artifact")
    assert artifact.status_code == 200
    assert artifact.json()["pptx_path"]
    assert Path(artifact.json()["pptx_path"]).exists()
    assert artifact.json()["slide_images"]
    assert Path(artifact.json()["slide_images"][0]["image_path"]).exists()
    skill_request_path = Path(artifact.json()["skill_request_path"])
    assert skill_request_path.exists()
    assert "presentation_plan" in skill_request_path.read_text(encoding="utf-8")

    download = client.get(f"/api/v1/ppt-artifacts/{artifact.json()['id']}/download")
    assert download.status_code == 200
    assert download.content[:2] == b"PK"

    slides = client.get(f"/api/v1/ppt-artifacts/{artifact.json()['id']}/slides")
    assert slides.status_code == 200
    assert slides.json()[0]["slide_no"] == 1

    slide_image = client.get(f"/api/v1/ppt-artifacts/{artifact.json()['id']}/slides/1/image")
    assert slide_image.status_code == 200
    assert slide_image.content[:8] == b"\x89PNG\r\n\x1a\n"


def test_lecture_presentation_job_skips_question_bank(client: TestClient) -> None:
    processed = client.post(
        "/api/v1/materials/process",
        files={"file": ("lecture.pdf", make_pdf(), "application/pdf")},
    )
    material_id = processed.json()["material"]["id"]
    content = client.post(f"/api/v1/materials/{material_id}/learning-content").json()

    job = client.post(
        f"/api/v1/learning-contents/{content['id']}/presentation-plan-jobs",
        params={"prepare_question_bank": False},
    ).json()
    finished = client.get(f"/api/v1/presentation-plan-jobs/{job['id']}").json()

    assert job["prepare_question_bank"] is False
    assert finished["status"] == "succeeded"
    plan = client.get(
        f"/api/v1/presentation-plan-jobs/{job['id']}/result"
    ).json()
    question_bank = client.get(
        f"/api/v1/presentation-plans/{plan['id']}/question-bank"
    ).json()
    assert question_bank["items"] == []


def test_prepared_question_bank_runs_as_classroom_script(client: TestClient) -> None:
    processed = client.post(
        "/api/v1/materials/process",
        files={"file": ("lesson.pdf", make_pdf(), "application/pdf")},
    )
    material_id = processed.json()["material"]["id"]
    content = client.post(f"/api/v1/materials/{material_id}/learning-content").json()
    presentation = client.post(
        f"/api/v1/learning-contents/{content['id']}/presentation-plans"
    ).json()
    question_bank = client.get(
        f"/api/v1/presentation-plans/{presentation['id']}/question-bank"
    ).json()
    prepared = question_bank["items"][0]

    missing_artifact = client.post(
        f"/api/v1/learning-contents/{content['id']}/classroom-plans",
        params={"presentation_plan_id": presentation["id"]},
    )
    assert missing_artifact.status_code == 409
    ppt_job = client.post(f"/api/v1/presentation-plans/{presentation['id']}/ppt-jobs")
    assert ppt_job.status_code == 202
    finished_ppt_job = client.get(f"/api/v1/ppt-jobs/{ppt_job.json()['id']}")
    assert finished_ppt_job.json()["status"] == "finished"

    classroom_plan_response = client.post(
        f"/api/v1/learning-contents/{content['id']}/classroom-plans",
        params={"presentation_plan_id": presentation["id"]},
    )
    assert classroom_plan_response.status_code == 201, classroom_plan_response.text
    classroom_plan = classroom_plan_response.json()
    action_types = [
        action["type"] for action in classroom_plan["scenes"][0]["actions"]
    ]
    end_action = next(
        action
        for action in classroom_plan["scenes"][0]["actions"]
        if action["type"] == "END"
    )
    assert not end_action["payload"]["summary"].startswith("本页要点：")
    student_index = action_types.index("STUDENT_QUESTION")
    assert action_types[student_index + 1] == "TEACHER_QA_RESPONSE"
    assert (
        classroom_plan["scenes"][0]["actions"][student_index]["payload"]["qa_id"]
        == prepared["id"]
    )

    session = client.post(
        f"/api/v1/classroom-plans/{classroom_plan['id']}/sessions",
        json={"mode": "interactive", "student_agent_types": ["researcher"]},
    ).json()
    session_id = session["id"]

    first = client.post(f"/api/v1/classroom-sessions/{session_id}/auto-step")
    second = client.post(f"/api/v1/classroom-sessions/{session_id}/auto-step")
    user_question = client.post(
        f"/api/v1/classroom-sessions/{session_id}/questions",
        json={"question": prepared["canonical_question"]},
    )
    assert user_question.status_code == 200
    assert user_question.json()["status"] == "answered"
    assert user_question.json()["source_refs"] == prepared["source_refs"]
    student_step = client.post(f"/api/v1/classroom-sessions/{session_id}/auto-step")
    blocked_question = client.post(
        f"/api/v1/classroom-sessions/{session_id}/questions",
        json={"question": "这里可以再解释一下吗？"},
    )
    assert blocked_question.status_code == 409
    teacher_step = client.post(f"/api/v1/classroom-sessions/{session_id}/auto-step")

    assert first.json()["action"]["type"] == "SHOW_SLIDE"
    assert second.json()["action"]["type"] == "EXPLAIN"
    assert student_step.json()["status"] == "agent_turn"
    assert student_step.json()["directed_turn"]["turns"][0]["role"] == "student"
    assert student_step.json()["directed_turn"]["turns"][0]["speech"] == prepared[
        "student_question"
    ]
    assert student_step.json()["directed_turn"]["turns"][0]["agent_id"] == (
        "student_agent_001"
    )
    assert teacher_step.json()["directed_turn"]["turns"][0]["role"] == "teacher"
    assert teacher_step.json()["directed_turn"]["turns"][0]["speech"] == prepared[
        "teacher_answer"
    ]

    events = teacher_step.json()["session"]["events"]
    qa_event = next(event for event in events if event["type"] == "QA_INTERACTION_EXECUTED")
    assert qa_event["payload"]["qa_id"] == prepared["id"]
    assert qa_event["payload"]["student_agent_id"] == "student_agent_001"

    continued = client.post(f"/api/v1/classroom-sessions/{session_id}/auto-step")
    assert continued.status_code == 200
    assert continued.json()["action"]["type"] not in {
        "STUDENT_QUESTION",
        "TEACHER_QA_RESPONSE",
    }


def test_tts_artifact_flow(client: TestClient) -> None:
    artifact = client.post(
        "/api/v1/tts-artifacts",
        json={
            "text": "欢迎来到互动课堂。",
            "scope": "slide_script",
            "ref_id": "slide_001",
            "voice": "teacher",
        },
    )

    assert artifact.status_code == 201
    payload = artifact.json()
    assert payload["scope"] == "slide_script"
    assert payload["ref_id"] == "slide_001"
    assert payload["voice"] == "teacher"
    assert payload["audio_url"] == f"/api/v1/tts-artifacts/{payload['id']}/audio"
    assert payload["duration_ms"] > 0
    assert payload["duration_seconds"] > 0
    assert Path(payload["audio_path"]).exists()

    cached = client.post(
        "/api/v1/tts-artifacts",
        json={
            "text": "欢迎来到互动课堂。",
            "scope": "slide_script",
            "ref_id": "slide_001",
            "voice": "teacher",
        },
    )
    assert cached.status_code == 201
    assert cached.json()["id"] == payload["id"]

    fetched = client.get(f"/api/v1/tts-artifacts/{payload['id']}")
    assert fetched.status_code == 200
    assert fetched.json()["id"] == payload["id"]

    audio = client.get(payload["audio_url"])
    assert audio.status_code == 200
    assert audio.content[:4] == b"RIFF"


def test_rejects_unsupported_file_type(client: TestClient) -> None:
    response = client.post(
        "/api/v1/materials",
        files={"file": ("notes.txt", b"plain text", "text/plain")},
    )
    assert response.status_code == 415


def test_cannot_build_content_before_parse(client: TestClient) -> None:
    upload = client.post(
        "/api/v1/materials",
        files={"file": ("lesson.pdf", make_pdf(), "application/pdf")},
    )

    response = client.post(f"/api/v1/materials/{upload.json()['id']}/learning-content")
    assert response.status_code == 409


def test_repeated_parse_replaces_pages(client: TestClient) -> None:
    upload = client.post(
        "/api/v1/materials",
        files={"file": ("lesson.pdf", make_pdf(), "application/pdf")},
    )
    material_id = upload.json()["id"]

    first = client.post(f"/api/v1/materials/{material_id}/parse")
    second = client.post(f"/api/v1/materials/{material_id}/parse")
    listed = client.get(f"/api/v1/materials/{material_id}/pages")

    assert first.status_code == second.status_code == listed.status_code == 200
    assert len(first.json()) == len(second.json()) == len(listed.json()) == 1
    assert first.json()[0]["id"] == second.json()[0]["id"]


def create_classroom_session(client: TestClient) -> str:
    processed = client.post(
        "/api/v1/materials/process",
        files={"file": ("lesson.pdf", make_pdf(), "application/pdf")},
    )
    material_id = processed.json()["material"]["id"]
    content = client.post(f"/api/v1/materials/{material_id}/learning-content")
    plan = client.post(f"/api/v1/learning-contents/{content.json()['id']}/classroom-plans")
    session = client.post(f"/api/v1/classroom-plans/{plan.json()['id']}/sessions")
    return session.json()["id"]


def test_lecture_mode_agent_turn_keeps_teacher_in_control(client: TestClient) -> None:
    session_id = create_classroom_session(client)
    for _ in range(3):
        client.post(f"/api/v1/classroom-sessions/{session_id}/next")

    directed = client.post(f"/api/v1/classroom-sessions/{session_id}/agent-turns/next")

    assert directed.status_code == 200
    assert directed.json()["decision"]["next_role"] == "teacher"
    assert directed.json()["turns"][0]["role"] == "teacher"


def test_auto_classroom_waits_for_user_quiz_after_agent_turn(client: TestClient) -> None:
    processed = client.post(
        "/api/v1/materials/process",
        files={"file": ("lesson.pdf", make_pdf(), "application/pdf")},
    )
    material_id = processed.json()["material"]["id"]
    content = client.post(f"/api/v1/materials/{material_id}/learning-content")
    plan = client.post(f"/api/v1/learning-contents/{content.json()['id']}/classroom-plans")
    session = client.post(
        f"/api/v1/classroom-plans/{plan.json()['id']}/sessions",
        json={"mode": "interactive"},
    )
    session_id = session.json()["id"]

    first = client.post(f"/api/v1/classroom-sessions/{session_id}/auto-step")
    second = client.post(f"/api/v1/classroom-sessions/{session_id}/auto-step")
    third = client.post(f"/api/v1/classroom-sessions/{session_id}/auto-step")
    fourth = client.post(f"/api/v1/classroom-sessions/{session_id}/auto-step")
    fifth = client.post(f"/api/v1/classroom-sessions/{session_id}/auto-step")
    sixth = client.post(f"/api/v1/classroom-sessions/{session_id}/auto-step")
    seventh = client.post(f"/api/v1/classroom-sessions/{session_id}/auto-step")

    assert first.json()["action"]["type"] == "SHOW_PAGE"
    assert second.json()["action"]["type"] == "EXPLAIN"
    assert third.json()["action"]["type"] == "PROBE"
    assert fourth.json()["status"] == "agent_turn"
    assert fourth.json()["directed_turn"]["turns"][0]["role"] == "student"
    assert fifth.json()["directed_turn"]["turns"][0]["role"] == "teacher"
    assert sixth.json()["action"]["type"] == "ASK_QUIZ"
    assert seventh.json()["status"] == "waiting"
    assert seventh.json()["session"]["waiting_for"] == "quiz_answer"
    assert seventh.json()["session"]["mastery"] == []

    answer = client.post(
        f"/api/v1/classroom-sessions/{session_id}/answers", json={"selected_index": 0}
    )
    assert answer.json()["correct"] is True
    assert answer.json()["feedback"].startswith("很棒！回答正确了！")
    assert answer.json()["session"]["waiting_for"] is None
    assert answer.json()["session"]["mastery"][0]["value"] == 1.0


def test_cannot_answer_before_quiz(client: TestClient) -> None:
    session_id = create_classroom_session(client)

    response = client.post(
        f"/api/v1/classroom-sessions/{session_id}/answers",
        json={"selected_index": 0},
    )
    assert response.status_code == 409


def test_rejects_answer_index_outside_quiz_options(client: TestClient) -> None:
    session_id = create_classroom_session(client)
    for _ in range(4):
        client.post(f"/api/v1/classroom-sessions/{session_id}/next")

    response = client.post(
        f"/api/v1/classroom-sessions/{session_id}/answers",
        json={"selected_index": 999},
    )
    assert response.status_code == 422
