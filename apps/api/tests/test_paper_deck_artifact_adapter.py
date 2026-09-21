import json
from pathlib import Path

import fitz
import pytest
from PIL import Image
from pptx import Presentation

from metaclass.modules.paper_workflow.paper_deck_artifact_adapter import (
    PaperDeckArtifactAdapter,
    PaperDeckArtifactError,
)
from metaclass.modules.paper_workflow.schemas import PaperPresentationArtifact


def _artifact_root(
    tmp_path: Path, *, backend: str = "imagegen", slide_count: int = 2
) -> Path:
    output = tmp_path / "provider_output"
    prompts = output / "prompts"
    images = output / "images"
    rendered = output / "rendered"
    prompts.mkdir(parents=True)
    images.mkdir()
    rendered.mkdir()
    provider_input = tmp_path / "provider_input"
    provider_input.mkdir()
    (provider_input / "paper_source.json").write_text(
        json.dumps({"assets": []}), encoding="utf-8"
    )
    (provider_input / "paper_analysis.json").write_text("{}", encoding="utf-8")
    (output / "analysis.md").write_text("# Analysis\n", encoding="utf-8")
    (output / "deck-brief.md").write_text(
        "# Deck Brief\n\n- style_preset: `journal-minimal`\n- language: zh-CN\n",
        encoding="utf-8",
    )
    outline = ["# Outline\n"]
    log = []
    for index in range(1, slide_count + 1):
        prompt = prompts / f"{index:02d}-slide-topic.md"
        image = images / f"{index:02d}-slide-topic.png"
        prompt.write_text(f"Create slide {index}.\n", encoding="utf-8")
        canvas = Image.new("RGB", (1600, 900), (index * 17 % 256, 40, 80))
        canvas.putpixel((index, index), (255, index % 256, 0))
        canvas.save(image)
        canvas.save(rendered / f"{index:02d}-slide.png")
        outline.append(
            f"""## {index:02d}. Slide title {index}
- Role: evidence
- Message: Explain result {index}.
- Render mode: native-raster
- Visual: Figure {index} with one callout.
- Text: Result {index}; Meaning {index}
- Evidence: Figure {index}, page {index + 2}
- Source visual: None
"""
        )
        log.append(
            f"images/{image.name}: backend={backend} render_mode=native-raster"
        )
    (output / "outline.md").write_text("\n".join(outline), encoding="utf-8")
    (output / "generation-log.md").write_text("\n".join(log), encoding="utf-8")
    (output / "source-visual-manifest.json").write_text(
        json.dumps(
            {
                "schema_version": "1.0",
                "slides": [
                    {
                        "slide_id": f"slide_{index:03d}",
                        "order": index,
                        "render_mode": "native-raster",
                        "background_path": f"images/{index:02d}-slide-topic.png",
                        "assets": [],
                        "annotations": [],
                    }
                    for index in range(1, slide_count + 1)
                ],
            }
        ),
        encoding="utf-8",
    )
    document = fitz.open()
    for index in range(1, slide_count + 1):
        document.new_page(width=1600, height=900).insert_text(
            (72, 72), f"Slide {index}"
        )
    document.save(output / "presentation.pdf")
    document.close()
    presentation = Presentation()
    for index in range(1, slide_count + 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        picture = slide.shapes.add_picture(
            str(images / f"{index:02d}-slide-topic.png"),
            0,
            0,
            width=presentation.slide_width,
            height=presentation.slide_height,
        )
        picture.name = "background:native-raster"
    presentation.save(output / "presentation.pptx")
    return output


def _artifact() -> PaperPresentationArtifact:
    return PaperPresentationArtifact(
        provider="native_paper_deck",
        presentation_pdf_path="provider_output/presentation.pdf",
        source_images_dir="provider_output/rendered",
        analysis_path="provider_output/analysis.md",
        deck_brief_path="provider_output/deck-brief.md",
        outline_path="provider_output/outline.md",
        prompts_dir="provider_output/prompts",
        generation_log_path="provider_output/generation-log.md",
        source_visual_manifest_path="provider_output/source-visual-manifest.json",
        debug_pptx_path="provider_output/presentation.pptx",
    )


def test_adapter_builds_deterministic_manifest_from_native_outputs(tmp_path: Path) -> None:
    _artifact_root(tmp_path)

    manifest = PaperDeckArtifactAdapter().adapt(_artifact(), workspace=tmp_path)

    assert manifest.provider == "native_paper_deck"
    assert manifest.style_preset == "journal-minimal"
    assert manifest.language == "zh-CN"
    assert manifest.slide_count == 2
    assert [slide.id for slide in manifest.slides] == [
        "paper_deck_slide_001",
        "paper_deck_slide_002",
    ]
    assert [slide.order for slide in manifest.slides] == [1, 2]
    assert [slide.pdf_page_no for slide in manifest.slides] == [1, 2]
    assert manifest.slides[0].title_hint == "Slide title 1"
    assert manifest.slides[0].planned_text == ["Result 1", "Meaning 1"]
    assert manifest.slides[0].source_visual_hint == "None"
    assert len({slide.image_hash for slide in manifest.slides}) == 2


def test_adapter_accepts_markdown_formatted_deck_brief_keys(tmp_path: Path) -> None:
    output = _artifact_root(tmp_path)
    (output / "deck-brief.md").write_text(
        "# Deck Brief\n\n- `style_preset`: `journal-minimal`\n- **language**: zh-CN\n",
        encoding="utf-8",
    )

    manifest = PaperDeckArtifactAdapter().adapt(_artifact(), workspace=tmp_path)

    assert manifest.style_preset == "journal-minimal"
    assert manifest.language == "zh-CN"


def test_adapter_rejects_duplicate_slide_images(tmp_path: Path) -> None:
    output = _artifact_root(tmp_path)
    duplicate = (output / "rendered/01-slide.png").read_bytes()
    (output / "rendered/02-slide.png").write_bytes(duplicate)

    with pytest.raises(PaperDeckArtifactError, match="image hashes must be unique"):
        PaperDeckArtifactAdapter().adapt(_artifact(), workspace=tmp_path)


def test_adapter_rejects_inconsistent_pdf_page_dimensions(tmp_path: Path) -> None:
    output = _artifact_root(tmp_path)
    document = fitz.open()
    document.new_page(width=1600, height=900)
    document.new_page(width=1200, height=900)
    replacement = output / "replacement.pdf"
    document.save(replacement)
    document.close()
    (output / "presentation.pdf").write_bytes(replacement.read_bytes())

    with pytest.raises(PaperDeckArtifactError, match="same dimensions"):
        PaperDeckArtifactAdapter().adapt(_artifact(), workspace=tmp_path)


def test_adapter_rejects_non_raster_generation_log_backend(tmp_path: Path) -> None:
    _artifact_root(tmp_path, backend="Pillow template renderer")

    with pytest.raises(PaperDeckArtifactError, match="non-raster backend"):
        PaperDeckArtifactAdapter().adapt(_artifact(), workspace=tmp_path)


def test_source_grounded_artifact_requires_layered_pptx(tmp_path: Path) -> None:
    _artifact_root(tmp_path)

    manifest = PaperDeckArtifactAdapter().adapt(_artifact(), workspace=tmp_path)

    assert _artifact().debug_pptx_path == "provider_output/presentation.pptx"
    assert manifest.slide_count == 2


def test_adapter_accepts_builtin_image_gen_backend_name(tmp_path: Path) -> None:
    _artifact_root(tmp_path, backend="Codex built-in image_gen")

    manifest = PaperDeckArtifactAdapter().adapt(_artifact(), workspace=tmp_path)

    assert manifest.slide_count == 2


def test_adapter_aligns_thirteen_native_pages_by_file_number(tmp_path: Path) -> None:
    _artifact_root(tmp_path, slide_count=13)

    manifest = PaperDeckArtifactAdapter().adapt(_artifact(), workspace=tmp_path)

    assert manifest.slide_count == 13
    assert [item.order for item in manifest.slides] == list(range(1, 14))
    assert [item.pdf_page_no for item in manifest.slides] == list(range(1, 14))
    assert manifest.slides[-1].prompt_path.startswith("provider_output/prompts/13-")
    assert manifest.slides[-1].image_path.startswith("provider_output/rendered/13-")


@pytest.mark.parametrize("kind", ["prompt", "image", "pdf_page", "extra_image"])
def test_adapter_rejects_missing_duplicate_or_extra_pages(tmp_path: Path, kind: str) -> None:
    output = _artifact_root(tmp_path, slide_count=13)
    if kind == "prompt":
        (output / "prompts/07-slide-topic.md").unlink()
    elif kind == "image":
        (output / "images/07-slide-topic.png").unlink()
    elif kind == "extra_image":
        Image.new("RGB", (1600, 900), (1, 2, 3)).save(output / "images/14-extra.png")
    else:
        document = fitz.open(output / "presentation.pdf")
        document.delete_page(12)
        replacement = output / "short.pdf"
        document.save(replacement)
        document.close()
        (output / "presentation.pdf").write_bytes(replacement.read_bytes())

    with pytest.raises(PaperDeckArtifactError):
        PaperDeckArtifactAdapter().adapt(_artifact(), workspace=tmp_path)
