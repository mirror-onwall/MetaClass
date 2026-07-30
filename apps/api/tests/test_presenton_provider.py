from pathlib import Path
from unittest.mock import Mock
from zipfile import ZipFile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from metaclass.modules.presentation.providers import (
    FallbackPPTProvider,
    PresentonPPTProvider,
)
from metaclass.modules.presentation.schemas import PresentationPlan, SlideElement, SlidePlan
from metaclass.modules.presentation.skill_adapter import PPTSkillAdapter


def make_plan() -> PresentationPlan:
    return PresentationPlan(
        id="plan_001",
        content_id="content_001",
        title="光合作用",
        slides=[
            SlidePlan(
                id="slide_001",
                order=1,
                source_section_ids=["section_001"],
                title="光合作用概览",
                key_points=["植物利用光能合成有机物", "同时释放氧气"],
                speaker_script="这一页先解释光合作用的整体过程。",
                suggested_visual="叶片、阳光和气体交换的流程图",
            ),
            SlidePlan(
                id="slide_002",
                order=2,
                source_section_ids=["section_001"],
                title="反应条件",
                key_points=["光照", "叶绿体", "水和二氧化碳"],
                speaker_script="这一页逐项说明反应条件。",
                suggested_visual="四项反应条件的图标矩阵",
            ),
        ],
    )


def make_provider(adapter: Mock | None = None) -> PresentonPPTProvider:
    return PresentonPPTProvider(
        base_url="https://api.presenton.test",
        api_key="test-key",
        adapter=adapter or Mock(),
    )


def make_content_lock_plan() -> PresentationPlan:
    return PresentationPlan(
        id="plan_content_lock",
        content_id="content_content_lock",
        title="Spatial analysis",
        slides=[
            SlidePlan(
                id="slide_content_lock",
                order=1,
                source_section_ids=["section_001"],
                title="Spatial autocorrelation",
                key_points=[
                    "Traditional correlation analysis applies to non-spatial observations.",
                    "Spatial autocorrelation considers geographic position.",
                ],
                speaker_script="Explain the distinction without changing slide text.",
                suggested_visual="A spatial comparison diagram.",
            )
        ],
    )


def write_pptx(path: Path, slide_texts: list[list[str]]) -> None:
    deck = Presentation()
    for texts in slide_texts:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        for index, value in enumerate(texts):
            textbox = slide.shapes.add_textbox(
                Inches(1), Inches(1 + index), Inches(10), Inches(0.6)
            )
            textbox.text_frame.text = value
    deck.save(path)


def test_presenton_payload_preserves_plan_pages_and_excludes_speaker_scripts() -> None:
    plan = make_plan()

    payload = make_provider().build_payload(plan)

    assert payload["slides_markdown"] == [
        "# 光合作用概览\n- 植物利用光能合成有机物\n- 同时释放氧气",
        "# 反应条件\n- 光照\n- 叶绿体\n- 水和二氧化碳",
    ]
    assert payload["n_slides"] == len(plan.slides)
    assert payload["content_generation"] == "preserve"
    assert payload["verbosity"] == "concise"
    assert payload["include_title_slide"] is False
    assert payload["include_table_of_contents"] is False
    assert "content" not in payload
    assert all(slide.speaker_script not in str(payload) for slide in plan.slides)


def test_presenton_validation_checks_every_page_title_and_key_point(tmp_path: Path) -> None:
    plan = make_plan()
    path = tmp_path / "deck.pptx"
    write_pptx(
        path,
        [
            ["光合作用概览", "植物利用光能合成有机物", "同时释放氧气"],
            ["反应条件", "光照", "叶绿体", "水和二氧化碳"],
        ],
    )

    result = make_provider().validate_deck(plan, path)

    assert result == {
        "expected_slide_count": 2,
        "actual_slide_count": 2,
        "checked_text_items": 7,
        "status": "passed",
    }


def test_presenton_validation_rejects_merged_or_rewritten_pages(tmp_path: Path) -> None:
    plan = make_plan()
    merged_path = tmp_path / "merged.pptx"
    write_pptx(merged_path, [["光合作用概览", "反应条件"]])

    with pytest.raises(RuntimeError, match="page-count mismatch"):
        make_provider().validate_deck(plan, merged_path)

    rewritten_path = tmp_path / "rewritten.pptx"
    write_pptx(
        rewritten_path,
        [
            ["光合作用概览", "植物吸收阳光", "同时释放氧气"],
            ["反应条件", "光照", "叶绿体", "水和二氧化碳"],
        ],
    )
    with pytest.raises(RuntimeError, match="content mismatch on page 1"):
        make_provider().validate_deck(plan, rewritten_path)


def test_presenton_content_lock_restores_exact_plan_text_and_removes_extras(
    tmp_path: Path,
) -> None:
    plan = make_content_lock_plan()
    path = tmp_path / "rewritten-by-presenton.pptx"
    write_pptx(
        path,
        [
            [
                "Spatial autocorrelation",
                "Traditional correlation",
                "applies to non-spatial observations.",
                "Spatial autocorrelation",
                "considers geographic position.",
                "Generated conclusion that is not in the plan.",
            ]
        ],
    )
    provider = make_provider()

    lock = provider.lock_deck_content(plan, path)
    validation = provider.validate_deck(plan, path)

    assert lock["status"] == "locked"
    assert lock["slide_count"] == 1
    assert lock["locked_text_items"] == 3
    assert lock["repaired_pages"] == 1
    assert lock["removed_generated_text_items"] >= 1
    assert validation["status"] == "passed"
    deck = Presentation(path)
    actual = [item for item in provider._extract_slide_text_items(deck.slides[0]) if item.strip()]
    assert actual == [plan.slides[0].title, *plan.slides[0].key_points]
    point_shape = next(
        shape
        for shape in deck.slides[0].shapes
        if getattr(shape, "has_text_frame", False) and shape.text == plan.slides[0].key_points[0]
    )
    point_font = point_shape.text_frame.paragraphs[0].runs[0].font
    assert point_font.size.pt == 18
    assert str(point_font.color.rgb) == "111827"


def test_presenton_validation_rejects_extra_generated_text(tmp_path: Path) -> None:
    plan = make_content_lock_plan()
    path = tmp_path / "extra-text.pptx"
    write_pptx(
        path,
        [
            [
                plan.slides[0].title,
                *plan.slides[0].key_points,
                "Generated conclusion that is not in the plan.",
            ]
        ],
    )

    with pytest.raises(RuntimeError, match="extra or rewritten"):
        make_provider().validate_deck(plan, path)


def test_content_lock_moves_plan_text_out_from_behind_a_picture(tmp_path: Path) -> None:
    point = "The plan text must remain visible."
    plan = PresentationPlan(
        id="plan_occlusion",
        content_id="content_occlusion",
        title="Occlusion",
        slides=[
            SlidePlan(
                id="slide_occlusion",
                order=1,
                source_section_ids=["section_001"],
                title="Visible title",
                key_points=[point],
                speaker_script="Explain the visible point.",
                suggested_visual="A picture beside a text panel.",
            )
        ],
    )
    image_path = tmp_path / "cover.png"
    Image.new("RGB", (600, 400), "navy").save(image_path)
    path = tmp_path / "occluded.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(0.8), Inches(0.4), Inches(11), Inches(0.7)
    ).text_frame.text = plan.slides[0].title
    hidden = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.8),
        Inches(1.6),
        Inches(5.2),
        Inches(4.2),
    )
    hidden.text_frame.text = point
    slide.shapes.add_picture(
        str(image_path),
        Inches(0.8),
        Inches(1.6),
        Inches(5.2),
        Inches(4.2),
    )
    slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(6.6),
        Inches(1.6),
        Inches(5.6),
        Inches(2.2),
    )
    deck.save(path)

    provider = make_provider()
    provider.lock_deck_content(plan, path)
    provider.validate_deck(plan, path)

    locked = Presentation(path)
    assert locked.slides[0].shapes[1].text == ""
    assert locked.slides[0].shapes[3].text == point


def test_content_lock_keeps_title_in_the_visual_title_region(tmp_path: Path) -> None:
    plan = PresentationPlan(
        id="plan_visual_roles",
        content_id="content_visual_roles",
        title="Visual roles",
        slides=[
            SlidePlan(
                id="slide_visual_roles",
                order=1,
                source_section_ids=["section_001"],
                title="The authoritative page title that must fit its title box",
                key_points=["A short decorative label"],
                speaker_script="Explain the page.",
                suggested_visual="A title above a content panel.",
            )
        ],
    )
    path = tmp_path / "visual-roles.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    top = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(5), Inches(0.8))
    top.text_frame.text = plan.slides[0].key_points[0]
    top.text_frame.paragraphs[0].runs[0].font.size = Pt(45)
    body = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(10.5), Inches(2.0))
    body.text_frame.text = plan.slides[0].title
    deck.save(path)

    provider = make_provider()
    provider.lock_deck_content(plan, path)
    provider.validate_deck(plan, path)

    locked = Presentation(path)
    assert locked.slides[0].shapes[0].text == plan.slides[0].title
    assert locked.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.size.pt < 45
    assert locked.slides[0].shapes[1].text == plan.slides[0].key_points[0]


def test_fallback_provider_keeps_existing_local_renderer(tmp_path: Path) -> None:
    plan = make_plan()
    primary = Mock()
    primary.prepare_request.side_effect = RuntimeError("provider unavailable")
    fallback = Mock()
    fallback.prepare_request.return_value = Mock(id="local_artifact")
    provider = FallbackPPTProvider(
        primary=primary,
        fallback=fallback,
        primary_name="presenton",
    )

    artifact = provider.prepare_request(
        plan=plan,
        job_id="job_001",
        output_dir=tmp_path,
    )

    assert artifact.id == "local_artifact"
    fallback.prepare_request.assert_called_once_with(
        plan=plan,
        job_id="job_001",
        output_dir=tmp_path,
    )


def test_presenton_runtime_is_fail_closed(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from metaclass.core.application import build_services
    from metaclass.core.config import settings

    monkeypatch.setattr(settings, "ppt_provider", "presenton")
    monkeypatch.setattr(settings, "presenton_api_key", "test-key")
    services = build_services(tmp_path, force_fake_llm=False)
    try:
        assert isinstance(services.presentations.ppt_adapter, PresentonPPTProvider)
        assert not isinstance(services.presentations.ppt_adapter, FallbackPPTProvider)
    finally:
        services.database.dispose()


def test_windows_preview_renders_the_final_pptx_with_powerpoint(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    plan = make_content_lock_plan()
    deck_path = tmp_path / "deck.pptx"
    write_pptx(
        deck_path,
        [[plan.slides[0].title, *plan.slides[0].key_points]],
    )

    def fake_run(command, **kwargs):
        output_dir = Path(command[command.index("-OutputDir") + 1])
        Image.new("RGB", (1600, 900), "white").save(output_dir / "slide_001.png")
        return Mock(stderr=b"")

    monkeypatch.setattr(
        "metaclass.modules.presentation.skill_adapter.platform.system",
        lambda: "Windows",
    )
    original_exists = Path.exists
    monkeypatch.setattr(
        Path,
        "exists",
        lambda path: True
        if str(path) == r"C:\Windows\System32\WindowsPowerShell\v1.0\powershell.exe"
        else original_exists(path),
    )
    monkeypatch.setattr(
        "metaclass.modules.presentation.skill_adapter.subprocess.run",
        fake_run,
    )

    images = PPTSkillAdapter()._render_slide_images(
        plan,
        deck_path,
        tmp_path / "slides",
    )

    assert len(images) == 1
    assert images[0].slide_id == plan.slides[0].id
    assert (images[0].width, images[0].height) == (1600, 900)


def test_external_preview_never_falls_back_to_placeholder(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    plan = make_content_lock_plan()
    deck_path = tmp_path / "deck.pptx"
    write_pptx(
        deck_path,
        [[plan.slides[0].title, *plan.slides[0].key_points]],
    )

    def unavailable(*args, **kwargs):
        raise FileNotFoundError("renderer unavailable")

    monkeypatch.setattr(
        "metaclass.modules.presentation.skill_adapter.platform.system",
        lambda: "Windows",
    )
    monkeypatch.setattr(PPTSkillAdapter, "_render_with_powerpoint", unavailable)
    monkeypatch.setattr(
        "metaclass.modules.presentation.skill_adapter.subprocess.run",
        unavailable,
    )

    output_dir = tmp_path / "slides"
    with pytest.raises(RuntimeError, match="placeholder previews are disabled"):
        PPTSkillAdapter()._render_slide_images(
            plan,
            deck_path,
            output_dir,
            allow_placeholder=False,
        )

    assert not list(output_dir.glob("*.png"))


def test_macos_external_deck_uses_declarative_scene_preview(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    plan = make_content_lock_plan()
    source_slide = plan.slides[0]
    slide = source_slide.model_copy(
        update={
            "elements": [
                SlideElement(
                    type="text",
                    x=0.08,
                    y=0.08,
                    w=0.84,
                    h=0.16,
                    text=source_slide.title,
                ),
                SlideElement(
                    type="shape",
                    x=0.08,
                    y=0.3,
                    w=0.84,
                    h=0.42,
                ),
            ]
        }
    )
    plan = plan.model_copy(update={"slides": [slide]})
    deck_path = tmp_path / "deck.pptx"
    write_pptx(deck_path, [[slide.title, *slide.key_points]])
    monkeypatch.setattr(
        "metaclass.modules.presentation.skill_adapter.platform.system",
        lambda: "Darwin",
    )

    images = PPTSkillAdapter()._render_slide_images(
        plan,
        deck_path,
        tmp_path / "slides",
        allow_placeholder=False,
    )

    assert len(images) == 1
    assert Path(images[0].image_path).is_file()


def test_presenton_download_does_not_forward_api_key_to_external_storage(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    captured_requests = []

    class FakeResponse:
        def __enter__(self):
            return self

        def __exit__(self, *args) -> None:
            return None

        def read(self) -> bytes:
            return b"PK fake-pptx"

    def fake_urlopen(request, **kwargs):
        captured_requests.append(request)
        return FakeResponse()

    monkeypatch.setattr(
        "metaclass.modules.presentation.providers.urlopen",
        fake_urlopen,
    )
    provider = make_provider()

    provider._download_file(
        "https://storage.example.com/exports/中文课件.pptx",
        tmp_path / "external.pptx",
    )
    provider._download_file(
        "/static/user_data/deck.pptx",
        tmp_path / "presenton.pptx",
    )

    assert captured_requests[0].get_header("Authorization") is None
    assert captured_requests[1].get_header("Authorization") == "Bearer test-key"
    assert "%E4%B8%AD%E6%96%87" in captured_requests[0].full_url


def test_presenton_png_export_is_mapped_to_plan_slide_ids(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    archive_path = tmp_path / "source.zip"
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    for page_no, color in [(2, "blue"), (1, "red")]:
        Image.new("RGB", (1280, 720), color).save(source_dir / f"page_{page_no}.png")
    with ZipFile(archive_path, "w") as archive:
        archive.write(source_dir / "page_2.png", "page_2.png")
        archive.write(source_dir / "page_1.png", "page_1.png")

    provider = make_provider()
    monkeypatch.setattr(
        provider,
        "_post_json",
        lambda path, payload: {"path": "https://storage.example.com/slides.zip"},
    )
    monkeypatch.setattr(
        provider,
        "_download_file",
        lambda remote_path, destination: destination.write_bytes(archive_path.read_bytes()),
    )

    images = provider._export_slide_images(
        make_plan(),
        "presentation_001",
        tmp_path / "output",
    )

    assert [image.slide_id for image in images] == ["slide_001", "slide_002"]
    assert [image.slide_no for image in images] == [1, 2]
    assert all((image.width, image.height) == (1280, 720) for image in images)
    with Image.open(images[0].image_path) as first_image:
        assert first_image.getpixel((0, 0)) == (255, 0, 0)
