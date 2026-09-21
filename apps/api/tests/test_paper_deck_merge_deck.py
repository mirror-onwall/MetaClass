import hashlib
import importlib.util
import json
from pathlib import Path

import fitz
import pytest
from PIL import Image
from pptx import Presentation

SCRIPT = (
    Path(__file__).parents[3]
    / ".agents/skills/paper-deck/scripts/merge_deck.py"
)
SPEC = importlib.util.spec_from_file_location("paper_deck_merge_deck", SCRIPT)
assert SPEC and SPEC.loader
merge_deck = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(merge_deck)


def _image(path: Path, size: tuple[int, int], color: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", size, color).save(path)


def test_native_page_remains_full_raster(tmp_path: Path) -> None:
    deck = tmp_path / "deck"
    _image(deck / "images/01-slide.png", (1600, 900), "navy")
    _image(deck / "images/02-slide.png", (1600, 900), "orange")

    pptx_path, pdf_path, count = merge_deck.merge_deck(deck, name="presentation")

    assert count == 2
    presentation = Presentation(pptx_path)
    assert len(presentation.slides) == 2
    assert [len(slide.shapes) for slide in presentation.slides] == [1, 1]
    with fitz.open(pdf_path) as document:
        assert document.page_count == 2
    rendered = sorted((deck / "rendered").glob("*.png"))
    assert [path.name for path in rendered] == ["01-slide.png", "02-slide.png"]
    assert all(Image.open(path).size == (1600, 900) for path in rendered)


def test_hybrid_slide_contains_independent_picture(tmp_path: Path) -> None:
    workspace = tmp_path / "job"
    deck = workspace / "provider_output"
    source = workspace / "provider_input/assets/figure-2.png"
    _image(deck / "images/01-slide.png", (1600, 900), "white")
    _image(deck / "images/02-slide-background.png", (1600, 900), "white")
    _image(source, (800, 400), "teal")
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    manifest = {
        "schema_version": "1.0",
        "slides": [
            {
                "slide_id": "slide_001",
                "order": 1,
                "render_mode": "native-raster",
                "background_path": "images/01-slide.png",
                "assets": [],
                "annotations": [],
            },
            {
                "slide_id": "slide_002",
                "order": 2,
                "render_mode": "source-grounded-hybrid",
                "background_path": "images/02-slide-background.png",
                "assets": [
                    {
                        "asset_id": "asset_figure_2",
                        "source_path": "provider_input/assets/figure-2.png",
                        "source_page": 5,
                        "figure_label": "Figure 2",
                        "x": 0.07,
                        "y": 0.23,
                        "w": 0.62,
                        "h": 0.66,
                        "fit": "contain",
                        "crop": "full",
                        "preserve": ["axes", "legend", "labels", "values"],
                        "sha256": source_hash,
                    }
                ],
                "annotations": [
                    {
                        "type": "text",
                        "text": "高质量子集持续提升",
                        "x": 0.73,
                        "y": 0.38,
                        "w": 0.21,
                        "h": 0.12,
                    },
                    {
                        "type": "rectangle",
                        "x": 0.065,
                        "y": 0.225,
                        "w": 0.63,
                        "h": 0.67,
                    },
                    {
                        "type": "arrow",
                        "x": 0.7,
                        "y": 0.45,
                        "w": 0.04,
                        "h": 0.02,
                    },
                ],
            },
        ],
    }
    (deck / "source-visual-manifest.json").write_text(
        json.dumps(manifest), encoding="utf-8"
    )
    output = deck / "presentation.pptx"

    contains_hybrid = merge_deck.make_pptx_from_manifest(deck, manifest, output)

    assert contains_hybrid is True
    presentation = Presentation(output)
    assert len(presentation.slides) == 2
    names = [shape.name for shape in presentation.slides[1].shapes]
    assert names == [
        "background:source-grounded-hybrid",
        "source:asset_figure_2",
        "annotation:text:01",
        "annotation:rectangle:02",
        "annotation:arrow:03",
    ]
    source_shape = presentation.slides[1].shapes[1]
    assert source_shape.width / source_shape.height == pytest.approx(2.0, rel=0.01)


def test_hybrid_merge_uses_no_office_runtime_and_records_provenance(tmp_path: Path) -> None:
    workspace = tmp_path / "job"
    deck = workspace / "provider_output"
    source = workspace / "provider_input/assets/figure.png"
    _image(deck / "images/01-background.png", (1600, 900), "white")
    _image(source, (800, 400), "teal")
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    manifest = {
        "schema_version": "1.0",
        "slides": [
            {
                "slide_id": "slide_001",
                "order": 1,
                "render_mode": "source-grounded-hybrid",
                "background_path": "images/01-background.png",
                "assets": [
                    {
                        "asset_id": "figure",
                        "source_path": "provider_input/assets/figure.png",
                        "source_page": 1,
                        "figure_label": "Figure 1",
                        "x": 0.1,
                        "y": 0.2,
                        "w": 0.6,
                        "h": 0.6,
                        "fit": "contain",
                        "crop": "full",
                        "preserve": ["labels"],
                        "sha256": source_hash,
                    }
                ],
                "annotations": [],
            }
        ],
    }
    (deck / "source-visual-manifest.json").write_text(json.dumps(manifest), encoding="utf-8")

    pptx_path, pdf_path, count = merge_deck.merge_deck(deck, name="presentation")

    assert count == 1
    assert pptx_path.is_file() and pdf_path.is_file()
    assert (deck / "rendered/01-slide.png").is_file()
    log = (deck / "generation-log.md").read_text(encoding="utf-8")
    assert "render_mode=source-grounded-hybrid" in log
    assert "source_asset=figure" in log
    assert f"source_asset_sha256={source_hash}" in log


def test_source_asset_hash_is_preserved(tmp_path: Path) -> None:
    workspace = tmp_path / "job"
    deck = workspace / "provider_output"
    source = workspace / "provider_input/assets/figure.png"
    _image(deck / "images/01-background.png", (1600, 900), "white")
    _image(source, (800, 400), "teal")
    manifest = {
        "schema_version": "1.0",
        "slides": [
            {
                "slide_id": "slide_001",
                "order": 1,
                "render_mode": "source-grounded-hybrid",
                "background_path": "images/01-background.png",
                "assets": [
                    {
                        "asset_id": "figure",
                        "source_path": "provider_input/assets/figure.png",
                        "source_page": 1,
                        "figure_label": "Figure 1",
                        "x": 0.1,
                        "y": 0.2,
                        "w": 0.6,
                        "h": 0.6,
                        "fit": "contain",
                        "crop": "full",
                        "preserve": ["labels"],
                        "sha256": "0" * 64,
                    }
                ],
                "annotations": [],
            }
        ],
    }

    with pytest.raises(merge_deck.MergeDeckError, match="SHA-256 mismatch"):
        merge_deck.make_pptx_from_manifest(
            deck, manifest, deck / "presentation.pptx"
        )
