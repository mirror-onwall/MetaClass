import json
from collections.abc import Callable
from pathlib import Path

import fitz
import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation

from metaclass.core.schemas import utc_now
from metaclass.main import create_app
from metaclass.modules.paper_workflow import presentation_stages
from metaclass.modules.paper_workflow.orchestrator import PaperWorkflowOrchestrator
from metaclass.modules.paper_workflow.providers.composed_skills import ComposedSkillsProvider
from metaclass.modules.paper_workflow.runtime import (
    CodexSkillInvocation,
    FakeCodexSkillRuntime,
)
from metaclass.modules.paper_workflow.schemas import (
    PresentationOutline,
    SlideEvidence,
    StageExecutionReport,
    StageStatus,
)


@pytest.fixture(autouse=True)
def _stub_libreoffice_render(monkeypatch: pytest.MonkeyPatch) -> None:
    def render(_: Path, rendered: Path, expected: int) -> None:
        rendered.mkdir(parents=True, exist_ok=True)
        for index in range(1, expected + 1):
            Image.new("RGB", (1600, 900), "white").save(rendered / f"slide-{index:03d}.png")

    monkeypatch.setattr(presentation_stages, "_render_with_libreoffice", render)


def _pdf_bytes() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text(
        (72, 72),
        "Evidence-Aware Paper\nOur method reaches 91.2 accuracy.\nThe evaluation is limited.",
    )
    payload = document.tobytes()
    document.close()
    return payload


def _material(client: TestClient) -> str:
    response = client.post(
        "/api/v1/materials",
        files={"file": ("paper.pdf", _pdf_bytes(), "application/pdf")},
    )
    assert response.status_code == 201
    return response.json()["id"]


def _skill(tmp_path: Path, name: str = "paper-analyze") -> Path:
    directory = tmp_path / "skills" / name
    directory.mkdir(parents=True)
    (directory / "SKILL.md").write_text(
        f"# {name}\nRun the requested paper stage.\n",
        encoding="utf-8",
    )
    return directory


def test_resolve_academic_outline_skill_accepts_distribution_directory(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    installed = _skill(tmp_path, "academic-pptx-skill")
    monkeypatch.chdir(tmp_path)
    monkeypatch.delenv("METACLASS_ACADEMIC_PPTX_SKILL_DIR", raising=False)

    resolved = ComposedSkillsProvider._resolve_skill_directory_for(
        "academic-pptx", "METACLASS_ACADEMIC_PPTX_SKILL_DIR"
    )

    assert resolved == installed.resolve()


def _valid_analysis(invocation: CodexSkillInvocation) -> dict:
    source = json.loads(
        (invocation.workspace / "source" / "paper_source.json").read_text(encoding="utf-8")
    )
    block = source["blocks"][0]
    reference = {"page_no": block["page_no"], "block_id": block["id"]}
    return {
        "schema_version": "1.0",
        "paper": {"title": "Evidence-Aware Paper"},
        "paper_type": "methods",
        "central_question": "How can the method improve accuracy?",
        "knowledge_gap": "Prior methods leave an accuracy gap.",
        "main_claim": "The proposed method improves the reported result.",
        "method_summary": {"overview": "A test method."},
        "claims": [
            {
                "id": "claim_core_01",
                "statement": "The paper proposes an evidence-aware method.",
                "importance": "core",
                "confidence": 0.9,
                "source_refs": [reference],
            }
        ],
        "experiments": [],
        "quantitative_results": [
            {
                "id": "result_01",
                "statement": "The method reports 91.2 accuracy.",
                "metric": "accuracy",
                "value": 91.2,
                "source_refs": [reference],
            }
        ],
        "limitations": [
            {
                "id": "limitation_01",
                "statement": "The evaluation is limited.",
                "source_refs": [reference],
            }
        ],
        "figure_candidates": [],
        "terminology": [],
        "critical_assessment": {},
        "warnings": [],
    }


def _success_report(invocation: CodexSkillInvocation) -> StageExecutionReport:
    now = utc_now()
    return StageExecutionReport(
        stage=invocation.stage,
        skill_name=invocation.skill_name,
        skill_version=invocation.skill_version,
        prompt_version=invocation.prompt_version,
        runtime_version="fake-paper-runtime-v1",
        status=StageStatus.SUCCEEDED,
        attempt=invocation.attempt,
        started_at=now,
        finished_at=now,
        input_hash=invocation.input_hash,
        exit_code=0,
        outputs=[
            str(path.relative_to(invocation.workspace))
            for path in invocation.output_directory.rglob("*")
            if path.is_file()
        ],
        validation_passed=True,
    )


def _runtime(
    writer: Callable[[CodexSkillInvocation, int], None],
) -> FakeCodexSkillRuntime:
    calls = 0

    def handler(invocation: CodexSkillInvocation) -> StageExecutionReport:
        nonlocal calls
        assert invocation.output_directory is not None
        invocation.output_directory.mkdir(parents=True, exist_ok=True)
        if invocation.stage.value == "figures":
            _write_valid_figure(invocation)
        elif invocation.stage.value == "outline":
            _write_valid_outline(invocation)
        elif invocation.stage.value == "generation":
            _write_valid_generation(invocation)
        else:
            calls += 1
            writer(invocation, calls)
        return _success_report(invocation)

    return FakeCodexSkillRuntime(handler)


def _runtime_with_outline_writer(
    writer: Callable[[CodexSkillInvocation, int], None],
) -> FakeCodexSkillRuntime:
    outline_calls = 0

    def handler(invocation: CodexSkillInvocation) -> StageExecutionReport:
        nonlocal outline_calls
        assert invocation.output_directory is not None
        invocation.output_directory.mkdir(parents=True, exist_ok=True)
        if invocation.stage.value == "analysis":
            _write_valid(invocation, 1)
        elif invocation.stage.value == "figures":
            _write_valid_figure(invocation)
        elif invocation.stage.value == "outline":
            outline_calls += 1
            writer(invocation, outline_calls)
        else:
            _write_valid_generation(invocation)
        return _success_report(invocation)

    return FakeCodexSkillRuntime(handler)


def _write_valid(invocation: CodexSkillInvocation, _: int) -> None:
    assert invocation.output_directory is not None
    (invocation.output_directory / "paper_analysis.md").write_text(
        "# Analysis\n\nCore claim [page=1].\n", encoding="utf-8"
    )
    (invocation.output_directory / "paper_analysis.json").write_text(
        json.dumps(_valid_analysis(invocation), ensure_ascii=False), encoding="utf-8"
    )


def _write_valid_figure(invocation: CodexSkillInvocation) -> None:
    assert invocation.output_directory is not None
    assets = invocation.output_directory / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (800, 500), "white").save(assets / "figure_1.png")
    payload = {
        "schema_version": "1.0",
        "figures": [
            {
                "id": "fig_extracted_01",
                "path": "assets/figure_1.png",
                "original_figure": "Figure 1",
                "page_no": 1,
                "caption": "Core method and evaluation result.",
                "source_method": "pdf_crop",
                "supports_claim_ids": ["claim_core_01"],
                "quality": {"width": 800, "height": 500, "readable": True},
                "crop_notes": "axes, legend, caption, and panel labels preserved",
            }
        ],
    }
    (invocation.output_directory / "figures.json").write_text(json.dumps(payload), encoding="utf-8")


def _write_valid_outline(invocation: CodexSkillInvocation) -> None:
    assert invocation.output_directory is not None
    slide_ids = [f"slide_{index:02d}" for index in range(1, 10)]
    outline = {
        "title": "Evidence-aware paper presentation",
        "paper_type": "methods",
        "narrative_arc": "problem-to-solution",
        "structure_summary": "Question, method, evidence, and conclusion.",
        "sections": [
            {
                "id": "section_01",
                "title": "Main argument",
                "role": "argument",
                "content_goal": "Explain the paper.",
                "slide_ids": slide_ids,
            }
        ],
        "slides": [
            {
                "id": slide_id,
                "order": index,
                "title": f"The paper advances its argument in step {chr(64 + index)}",
                "purpose": "Explain one part of the argument.",
                "key_points": ["One supported point."],
                "asset_ids": [],
                "speaker_note": "Explain the evidence.",
                "layout_intent": "academic_content",
            }
            for index, slide_id in enumerate(slide_ids, start=1)
        ],
    }
    evidence = {
        "slides": [
            {
                "slide_id": slide_id,
                "claim_ids": ["claim_core_01"],
                "source_refs": [{"page_no": 1}],
                "asset_ids": [],
                "evidence_strength": "direct",
            }
            for slide_id in slide_ids
        ]
    }
    (invocation.output_directory / "outline.md").write_text(
        "# Evidence-aware paper presentation\n", encoding="utf-8"
    )
    (invocation.output_directory / "presentation_outline.json").write_text(
        json.dumps(outline), encoding="utf-8"
    )
    (invocation.output_directory / "slide_evidence.json").write_text(
        json.dumps(evidence), encoding="utf-8"
    )


def _write_valid_generation(invocation: CodexSkillInvocation) -> None:
    assert invocation.output_directory is not None
    outline = json.loads(
        (invocation.workspace / "stages/03_outline/output/presentation_outline.json").read_text(
            encoding="utf-8"
        )
    )
    presentation = Presentation()
    presentation.slide_width = 12192000
    presentation.slide_height = 6858000
    for contract in outline["slides"]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = contract["title"]
    presentation.save(invocation.output_directory / "presentation.pptx")
    (invocation.output_directory / "slide_plan.json").write_text(
        json.dumps(
            [
                {
                    "slide_id": item["id"],
                    "order": item["order"],
                    "title": item["title"],
                    "asset_ids": item["asset_ids"],
                }
                for item in outline["slides"]
            ]
        ),
        encoding="utf-8",
    )
    (invocation.output_directory / "speaker_notes.json").write_text(
        json.dumps(
            [{"slide_id": item["id"], "note": item["speaker_note"]} for item in outline["slides"]]
        ),
        encoding="utf-8",
    )
    (invocation.output_directory / "qa_report.json").write_text(
        json.dumps({"status": "passed"}), encoding="utf-8"
    )


def _provider(tmp_path: Path, runtime: FakeCodexSkillRuntime) -> ComposedSkillsProvider:
    return ComposedSkillsProvider(
        runtime,
        skill_directory=_skill(tmp_path),
        figure_skill_directory=_skill(tmp_path, "extract-paper-images"),
        outline_skill_directory=_skill(tmp_path, "academic-pptx"),
        generation_skill_directory=_skill(tmp_path, "academic-pptx-generate"),
    )


def test_full_composed_workflow_is_available_via_api(tmp_path: Path) -> None:
    app = create_app(tmp_path)
    runtime = _runtime(_write_valid)
    provider = _provider(tmp_path, runtime)
    registered_paths: list[Path] = []

    def register(path: Path, filename: str, derivation_key: str) -> str:
        assert filename == "paper-presentation.pptx"
        assert derivation_key.startswith("paper_job_")
        assert path.is_file()
        registered_paths.append(path)
        return "mat_normalized_presentation"

    provider.register_presentation = register
    app.state.services.paper_workflows.orchestrator = PaperWorkflowOrchestrator(
        tmp_path, [provider]
    )

    with TestClient(app) as client:
        material_id = _material(client)
        job_id = client.post(
            "/api/v1/paper-workflows",
            json={"material_id": material_id, "strategy": "composed_skills"},
        ).json()["id"]
        completed = client.post(f"/api/v1/paper-workflows/{job_id}/run")
        repeated = client.post(f"/api/v1/paper-workflows/{job_id}/run")
        analysis = client.get(f"/api/v1/paper-workflows/{job_id}/analysis")
        figures = client.get(f"/api/v1/paper-workflows/{job_id}/figures")
        outline = client.get(f"/api/v1/paper-workflows/{job_id}/outline")
        evidence = client.get(f"/api/v1/paper-workflows/{job_id}/slide-evidence")
        bundle = client.get(f"/api/v1/paper-workflows/{job_id}/result")

    assert completed.json()["status"] == "succeeded"
    assert completed.json()["stage"] == "completed"
    assert repeated.json() == completed.json()
    assert analysis.status_code == 200
    assert analysis.json()["claims"][0]["importance"] == "core"
    assert figures.status_code == 200
    assert figures.json()["figures"][0]["source_method"] == "pdf_crop"
    assert outline.status_code == 200
    assert len(outline.json()["slides"]) == 9
    assert evidence.status_code == 200
    assert evidence.json()["slides"][0]["slide_id"] == "slide_01"
    assert bundle.status_code == 200
    assert bundle.json()["root_path"].endswith(f"paper_workflows/{job_id}/final")
    bundle_paths = {item["path"] for item in bundle.json()["files"]}
    assert bundle_paths >= {
        "presentation.pptx",
        "paper_analysis.json",
        "presentation_outline.json",
        "slide_evidence.json",
        "asset_manifest.json",
        "speaker_notes.json",
        "generation_report.json",
        "qa_report.json",
    }
    assert len([path for path in bundle_paths if path.startswith("assets/")]) == 1
    assert len(runtime.invocations) == 4
    assert [item.attempt for item in runtime.invocations] == [1, 1, 1, 1]
    outline_invocation = next(item for item in runtime.invocations if item.stage.value == "outline")
    assert outline_invocation.prompt_version == "academic-pptx-v2"
    assert {path.name for path in outline_invocation.input_paths} >= {
        "presentation_outline_schema.json",
        "slide_evidence_schema.json",
    }
    assert outline_invocation.output_schema_path is not None
    assert outline_invocation.output_schema_path.name == "runtime_result_schema.json"
    workspace = tmp_path / "runtime" / "paper_workflows" / job_id
    resolved = json.loads((workspace / "resolved_request.json").read_text(encoding="utf-8"))
    assert resolved["duration_minutes"] == 15
    assert resolved["audience"] == "具备基础专业背景的高校学生和研究生"
    output = workspace / "stages" / "01_analysis" / "output"
    assert {path.name for path in output.iterdir()} == {
        "paper_analysis.md",
        "paper_analysis.json",
        "execution_report.json",
    }
    report = StageExecutionReport.model_validate_json(
        (output / "execution_report.json").read_text(encoding="utf-8")
    )
    assert report.validation_passed is True
    final = workspace / "final"
    assert registered_paths == [final / "presentation.pptx"]
    assert {path.name for path in final.iterdir()} == {
        "presentation.pptx",
        "paper_analysis.json",
        "presentation_outline.json",
        "slide_evidence.json",
        "asset_manifest.json",
        "speaker_notes.json",
        "generation_report.json",
        "qa_report.json",
        "assets",
    }


def test_final_validation_failure_does_not_register_a_material(tmp_path: Path) -> None:
    app = create_app(tmp_path)
    runtime = _runtime(_write_valid)
    provider = _provider(tmp_path, runtime)
    registered_paths: list[Path] = []
    provider.register_presentation = lambda path, _filename, _key: (
        registered_paths.append(path) or "mat_must_not_exist"
    )
    provider.validator.validate_final_directory = lambda _root, mode: [  # type: ignore[method-assign]
        f"forced final validation failure in {mode} mode"
    ]
    app.state.services.paper_workflows.orchestrator = PaperWorkflowOrchestrator(
        tmp_path, [provider]
    )

    with TestClient(app) as client:
        material_id = _material(client)
        job_id = client.post(
            "/api/v1/paper-workflows",
            json={"material_id": material_id, "strategy": "composed_skills"},
        ).json()["id"]
        result = client.post(f"/api/v1/paper-workflows/{job_id}/run")

    assert result.json()["status"] == "failed"
    assert "forced final validation failure" in result.json()["error"]
    assert registered_paths == []
    assert not (tmp_path / "runtime" / "paper_workflows" / job_id / "final").exists()


def test_generation_validator_ignores_empty_text_boxes_before_fallback_title(
    tmp_path: Path,
) -> None:
    output = tmp_path / "output"
    output.mkdir()
    outline = PresentationOutline.model_validate(
        {
            "title": "Test deck",
            "paper_type": "methods",
            "narrative_arc": "problem-to-solution",
            "objectives": ["Explain the result"],
            "structure_summary": "One-slide test deck.",
            "sections": [
                {
                    "id": "section_01",
                    "title": "Result",
                    "role": "result",
                    "content_goal": "Explain the result.",
                    "slide_ids": ["slide_01"],
                }
            ],
            "slides": [
                {
                    "id": "slide_01",
                    "order": 1,
                    "title": "The visible title matches the frozen outline",
                    "purpose": "Explain the result.",
                    "key_points": ["Supported point"],
                    "asset_ids": [],
                    "speaker_note": "Explain the result.",
                    "layout_intent": "academic_content",
                }
            ],
        }
    )
    evidence = SlideEvidence.model_validate(
        {
            "slides": [
                {
                    "slide_id": "slide_01",
                    "claim_ids": ["claim_01"],
                    "source_refs": [{"page_no": 1}],
                    "asset_ids": [],
                    "evidence_strength": "direct",
                }
            ]
        }
    )
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, 100, 100).text = ""
    slide.shapes.add_textbox(100, 50, 1000000, 300000).text = "RESULTS"
    slide.shapes.add_textbox(100, 100, 5000000, 500000).text = outline.slides[0].title
    presentation.save(output / "presentation.pptx")
    (output / "slide_plan.json").write_text(
        json.dumps([{"slide_id": "slide_01"}]), encoding="utf-8"
    )
    (output / "speaker_notes.json").write_text(
        json.dumps([{"slide_id": "slide_01", "note": "Explain."}]), encoding="utf-8"
    )
    (output / "qa_report.json").write_text("{}", encoding="utf-8")

    presentation_stages.validate_and_render_pptx(
        output,
        outline=outline,
        evidence=evidence,
    )


def test_stage3_contract_schemas_are_valid_and_deterministic(tmp_path: Path) -> None:
    stage_root = tmp_path / "runtime" / "paper_workflows" / "job" / "stages/03_outline"
    outline_path, evidence_path, runtime_path = (
        presentation_stages.AcademicOutlineAdapter.write_contract_schemas(stage_root)
    )
    first = {path.name: path.read_bytes() for path in (outline_path, evidence_path, runtime_path)}

    presentation_stages.AcademicOutlineAdapter.write_contract_schemas(stage_root)

    assert outline_path.read_bytes() == first[outline_path.name]
    assert evidence_path.read_bytes() == first[evidence_path.name]
    assert runtime_path.read_bytes() == first[runtime_path.name]
    outline_schema = json.loads(outline_path.read_text(encoding="utf-8"))
    evidence_schema = json.loads(evidence_path.read_text(encoding="utf-8"))
    runtime_schema = json.loads(runtime_path.read_text(encoding="utf-8"))
    assert set(outline_schema["required"]) >= {"title", "sections", "slides"}
    assert set(evidence_schema["required"]) >= {"slides"}
    assert runtime_schema["properties"]["status"]["enum"] == ["completed", "blocked"]


def test_stage3_repairs_json_structure_once(tmp_path: Path) -> None:
    def writer(invocation: CodexSkillInvocation, call: int) -> None:
        _write_valid_outline(invocation)
        if call == 1:
            assert invocation.output_directory is not None
            (invocation.output_directory / "presentation_outline.json").write_text(
                "{invalid", encoding="utf-8"
            )
        else:
            assert invocation.prompt_path.name == "structure_repair_prompt.md"

    app = create_app(tmp_path)
    runtime = _runtime_with_outline_writer(writer)
    app.state.services.paper_workflows.orchestrator = PaperWorkflowOrchestrator(
        tmp_path, [_provider(tmp_path, runtime)]
    )
    with TestClient(app) as client:
        material_id = _material(client)
        job_id = client.post("/api/v1/paper-workflows", json={"material_id": material_id}).json()[
            "id"
        ]
        result = client.post(f"/api/v1/paper-workflows/{job_id}/run")

    outline_invocations = [item for item in runtime.invocations if item.stage.value == "outline"]
    assert result.json()["status"] == "succeeded"
    assert [item.attempt for item in outline_invocations] == [1, 2]
    assert outline_invocations[1].prompt_version.endswith("-structure_repair")
    assert "validation_errors.json" in {path.name for path in outline_invocations[1].input_paths}


def test_stage3_repairs_small_evidence_errors_once(tmp_path: Path) -> None:
    def writer(invocation: CodexSkillInvocation, call: int) -> None:
        _write_valid_outline(invocation)
        if call == 1:
            assert invocation.output_directory is not None
            path = invocation.output_directory / "slide_evidence.json"
            payload = json.loads(path.read_text(encoding="utf-8"))
            payload["slides"][0]["claim_ids"] = ["claim_does_not_exist"]
            path.write_text(json.dumps(payload), encoding="utf-8")
        else:
            assert invocation.prompt_path.name == "outline_repair_prompt.md"

    app = create_app(tmp_path)
    runtime = _runtime_with_outline_writer(writer)
    app.state.services.paper_workflows.orchestrator = PaperWorkflowOrchestrator(
        tmp_path, [_provider(tmp_path, runtime)]
    )
    with TestClient(app) as client:
        material_id = _material(client)
        job_id = client.post("/api/v1/paper-workflows", json={"material_id": material_id}).json()[
            "id"
        ]
        result = client.post(f"/api/v1/paper-workflows/{job_id}/run")

    outline_invocations = [item for item in runtime.invocations if item.stage.value == "outline"]
    assert result.json()["status"] == "succeeded"
    assert [item.attempt for item in outline_invocations] == [1, 2]
    assert outline_invocations[1].prompt_version.endswith("-evidence_repair")
    assert "validation_errors.json" in {path.name for path in outline_invocations[1].input_paths}


def test_stage3_fully_reruns_for_narrative_errors(tmp_path: Path) -> None:
    def writer(invocation: CodexSkillInvocation, call: int) -> None:
        _write_valid_outline(invocation)
        if call == 1:
            assert invocation.output_directory is not None
            path = invocation.output_directory / "presentation_outline.json"
            payload = json.loads(path.read_text(encoding="utf-8"))
            payload["slides"][0]["key_points"] = [f"Point {index}" for index in range(5)]
            path.write_text(json.dumps(payload), encoding="utf-8")
        else:
            assert invocation.prompt_path.name == "rerun_prompt.md"

    app = create_app(tmp_path)
    runtime = _runtime_with_outline_writer(writer)
    app.state.services.paper_workflows.orchestrator = PaperWorkflowOrchestrator(
        tmp_path, [_provider(tmp_path, runtime)]
    )
    with TestClient(app) as client:
        material_id = _material(client)
        job_id = client.post("/api/v1/paper-workflows", json={"material_id": material_id}).json()[
            "id"
        ]
        result = client.post(f"/api/v1/paper-workflows/{job_id}/run")

    outline_invocations = [item for item in runtime.invocations if item.stage.value == "outline"]
    assert result.json()["status"] == "succeeded"
    assert [item.attempt for item in outline_invocations] == [1, 2]
    assert outline_invocations[1].prompt_version.endswith("-rerun")


@pytest.mark.parametrize(
    ("failure_kind", "expected_prompt"),
    [
        ("structure", "structure_repair_prompt.md"),
        ("evidence", "outline_repair_prompt.md"),
        ("narrative", "rerun_prompt.md"),
    ],
)
def test_stage3_repair_and_rerun_are_bounded(
    tmp_path: Path,
    failure_kind: str,
    expected_prompt: str,
) -> None:
    def writer(invocation: CodexSkillInvocation, call: int) -> None:
        _write_valid_outline(invocation)
        assert invocation.output_directory is not None
        if call == 2:
            assert invocation.prompt_path.name == expected_prompt
        if failure_kind == "structure":
            (invocation.output_directory / "presentation_outline.json").write_text(
                "{invalid", encoding="utf-8"
            )
        elif failure_kind == "evidence":
            path = invocation.output_directory / "slide_evidence.json"
            payload = json.loads(path.read_text(encoding="utf-8"))
            payload["slides"][0]["claim_ids"] = ["claim_does_not_exist"]
            path.write_text(json.dumps(payload), encoding="utf-8")
        else:
            path = invocation.output_directory / "presentation_outline.json"
            payload = json.loads(path.read_text(encoding="utf-8"))
            payload["slides"][0]["key_points"] = [f"Point {index}" for index in range(5)]
            path.write_text(json.dumps(payload), encoding="utf-8")

    app = create_app(tmp_path)
    runtime = _runtime_with_outline_writer(writer)
    app.state.services.paper_workflows.orchestrator = PaperWorkflowOrchestrator(
        tmp_path, [_provider(tmp_path, runtime)]
    )
    with TestClient(app) as client:
        material_id = _material(client)
        job_id = client.post("/api/v1/paper-workflows", json={"material_id": material_id}).json()[
            "id"
        ]
        result = client.post(f"/api/v1/paper-workflows/{job_id}/run")

    outline_invocations = [item for item in runtime.invocations if item.stage.value == "outline"]
    assert result.json()["status"] == "failed"
    assert len(outline_invocations) == 2


def test_stage1_uses_one_structure_repair_for_invalid_json(tmp_path: Path) -> None:
    def writer(invocation: CodexSkillInvocation, call: int) -> None:
        assert invocation.output_directory is not None
        (invocation.output_directory / "paper_analysis.md").write_text(
            "# Analysis", encoding="utf-8"
        )
        if call == 1:
            (invocation.output_directory / "paper_analysis.json").write_text(
                "{invalid", encoding="utf-8"
            )
        else:
            assert invocation.prompt_path.name == "repair_prompt.md"
            (invocation.output_directory / "paper_analysis.json").write_text(
                json.dumps(_valid_analysis(invocation)), encoding="utf-8"
            )

    app = create_app(tmp_path)
    runtime = _runtime(writer)
    app.state.services.paper_workflows.orchestrator = PaperWorkflowOrchestrator(
        tmp_path,
        [_provider(tmp_path, runtime)],
    )
    with TestClient(app) as client:
        material_id = _material(client)
        job_id = client.post("/api/v1/paper-workflows", json={"material_id": material_id}).json()[
            "id"
        ]
        result = client.post(f"/api/v1/paper-workflows/{job_id}/run")

    assert result.json()["status"] == "succeeded"
    assert len(runtime.invocations) == 5
    assert [item.attempt for item in runtime.invocations] == [1, 2, 1, 1, 1]


def test_stage1_falls_back_to_markdown_when_pdf_preflight_fails(
    tmp_path: Path, monkeypatch
) -> None:
    app = create_app(tmp_path)
    runtime = _runtime(_write_valid)
    provider = _provider(tmp_path, runtime)
    monkeypatch.setattr(provider, "_pdf_readable", lambda _: False)
    app.state.services.paper_workflows.orchestrator = PaperWorkflowOrchestrator(
        tmp_path, [provider]
    )

    with TestClient(app) as client:
        material_id = _material(client)
        job_id = client.post("/api/v1/paper-workflows", json={"material_id": material_id}).json()[
            "id"
        ]
        result = client.post(f"/api/v1/paper-workflows/{job_id}/run")
        analysis = client.get(f"/api/v1/paper-workflows/{job_id}/analysis")

    assert result.json()["status"] == "succeeded"
    assert "source_pdf_unreadable_used_paper_content_markdown" in analysis.json()["warnings"]
    assert all(path.name != "paper.pdf" for path in runtime.invocations[0].input_paths)


def test_stage1_fully_reruns_when_evidence_is_missing(tmp_path: Path) -> None:
    def writer(invocation: CodexSkillInvocation, call: int) -> None:
        assert invocation.output_directory is not None
        payload = _valid_analysis(invocation)
        if call == 1:
            payload["claims"][0]["source_refs"] = []
        else:
            assert invocation.prompt_path.name == "rerun_prompt.md"
        (invocation.output_directory / "paper_analysis.md").write_text(
            "# Analysis", encoding="utf-8"
        )
        (invocation.output_directory / "paper_analysis.json").write_text(
            json.dumps(payload), encoding="utf-8"
        )

    app = create_app(tmp_path)
    runtime = _runtime(writer)
    app.state.services.paper_workflows.orchestrator = PaperWorkflowOrchestrator(
        tmp_path,
        [_provider(tmp_path, runtime)],
    )
    with TestClient(app) as client:
        material_id = _material(client)
        job_id = client.post("/api/v1/paper-workflows", json={"material_id": material_id}).json()[
            "id"
        ]
        result = client.post(f"/api/v1/paper-workflows/{job_id}/run")

    assert result.json()["status"] == "succeeded"
    assert len(runtime.invocations) == 5


def test_failed_stage2_resumes_without_rerunning_stage1(tmp_path: Path) -> None:
    figure_calls = 0

    def handler(invocation: CodexSkillInvocation) -> StageExecutionReport:
        nonlocal figure_calls
        assert invocation.output_directory is not None
        invocation.output_directory.mkdir(parents=True, exist_ok=True)
        if invocation.stage.value == "analysis":
            _write_valid(invocation, 1)
            return _success_report(invocation)
        if invocation.stage.value == "figures":
            figure_calls += 1
        if invocation.stage.value == "figures" and figure_calls == 1:
            now = utc_now()
            return StageExecutionReport(
                stage=invocation.stage,
                skill_name=invocation.skill_name,
                skill_version=invocation.skill_version,
                prompt_version=invocation.prompt_version,
                runtime_version="fake-paper-runtime-v1",
                status=StageStatus.FAILED,
                attempt=invocation.attempt,
                started_at=now,
                finished_at=now,
                input_hash=invocation.input_hash,
                exit_code=1,
                validation_passed=False,
            )
        if invocation.stage.value == "figures":
            _write_valid_figure(invocation)
        elif invocation.stage.value == "outline":
            _write_valid_outline(invocation)
        else:
            _write_valid_generation(invocation)
        return _success_report(invocation)

    app = create_app(tmp_path)
    runtime = FakeCodexSkillRuntime(handler)
    app.state.services.paper_workflows.orchestrator = PaperWorkflowOrchestrator(
        tmp_path, [_provider(tmp_path, runtime)]
    )
    with TestClient(app) as client:
        material_id = _material(client)
        job_id = client.post("/api/v1/paper-workflows", json={"material_id": material_id}).json()[
            "id"
        ]
        failed = client.post(f"/api/v1/paper-workflows/{job_id}/run")
        resumed = client.post(f"/api/v1/paper-workflows/{job_id}/resume")
        completed = client.post(f"/api/v1/paper-workflows/{job_id}/run")

    assert failed.json()["status"] == "failed"
    assert resumed.json()["status"] == "queued"
    assert completed.json()["status"] == "succeeded"
    assert [item.stage.value for item in runtime.invocations] == [
        "analysis",
        "figures",
        "figures",
        "outline",
        "generation",
    ]
    assert [item.attempt for item in runtime.invocations] == [1, 1, 2, 1, 1]


def test_failed_generation_resumes_without_rerunning_prior_stages(tmp_path: Path) -> None:
    generation_calls = 0

    def handler(invocation: CodexSkillInvocation) -> StageExecutionReport:
        nonlocal generation_calls
        assert invocation.output_directory is not None
        invocation.output_directory.mkdir(parents=True, exist_ok=True)
        if invocation.stage.value == "analysis":
            _write_valid(invocation, 1)
        elif invocation.stage.value == "figures":
            _write_valid_figure(invocation)
        elif invocation.stage.value == "outline":
            _write_valid_outline(invocation)
        else:
            generation_calls += 1
            if generation_calls == 1:
                now = utc_now()
                return StageExecutionReport(
                    stage=invocation.stage,
                    skill_name=invocation.skill_name,
                    skill_version=invocation.skill_version,
                    prompt_version=invocation.prompt_version,
                    runtime_version="fake-paper-runtime-v1",
                    status=StageStatus.FAILED,
                    attempt=invocation.attempt,
                    started_at=now,
                    finished_at=now,
                    input_hash=invocation.input_hash,
                    exit_code=1,
                    validation_passed=False,
                )
            _write_valid_generation(invocation)
        return _success_report(invocation)

    app = create_app(tmp_path)
    runtime = FakeCodexSkillRuntime(handler)
    app.state.services.paper_workflows.orchestrator = PaperWorkflowOrchestrator(
        tmp_path, [_provider(tmp_path, runtime)]
    )
    with TestClient(app) as client:
        material_id = _material(client)
        job_id = client.post("/api/v1/paper-workflows", json={"material_id": material_id}).json()[
            "id"
        ]
        failed = client.post(f"/api/v1/paper-workflows/{job_id}/run")
        outline = client.get(f"/api/v1/paper-workflows/{job_id}/outline")
        resumed = client.post(f"/api/v1/paper-workflows/{job_id}/resume")
        completed = client.post(f"/api/v1/paper-workflows/{job_id}/run")

    assert failed.json()["status"] == "failed"
    assert outline.status_code == 200
    assert resumed.json()["status"] == "queued"
    assert completed.json()["status"] == "succeeded"
    assert [item.stage.value for item in runtime.invocations] == [
        "analysis",
        "figures",
        "outline",
        "generation",
        "generation",
    ]
    assert [item.attempt for item in runtime.invocations] == [1, 1, 1, 1, 2]
