import hashlib
import json
from pathlib import Path

import fitz
import pytest
from PIL import Image
from pptx import Presentation

from metaclass.core.schemas import utc_now
from metaclass.modules.paper_workflow.providers.base import PaperProviderContext
from metaclass.modules.paper_workflow.providers.native_paper_deck import (
    NativePaperDeckError,
    NativePaperDeckProvider,
)
from metaclass.modules.paper_workflow.runtime import (
    CodexSkillInvocation,
    FakeCodexSkillRuntime,
)
from metaclass.modules.paper_workflow.schemas import (
    PaperSourceBundle,
    PaperWorkflowCheckpoint,
    PaperWorkflowRequest,
    StageExecutionReport,
    StageStatus,
)


def _skill(tmp_path: Path) -> Path:
    root = tmp_path / "skills/paper-deck"
    (root / "references").mkdir(parents=True)
    (root / "scripts").mkdir()
    (root / "SKILL.md").write_text("# paper-deck\nRaster-first workflow.\n", encoding="utf-8")
    (root / "references/quality-gate.md").write_text("# QA\n", encoding="utf-8")
    (root / "scripts/merge_deck.py").write_text("# merge\n", encoding="utf-8")
    return root


def _workspace(tmp_path: Path) -> tuple[Path, PaperSourceBundle]:
    workspace = tmp_path / "job"
    source = workspace / "source"
    assets = source / "existing_assets"
    analysis = workspace / "stages/01_analysis/output"
    assets.mkdir(parents=True)
    analysis.mkdir(parents=True)

    document = fitz.open()
    document.new_page().insert_text((72, 72), "Paper source")
    document.save(source / "paper.pdf")
    document.close()
    (source / "paper_content.md").write_text("# Paper\n", encoding="utf-8")
    (source / "paper_source.json").write_text('{"blocks": []}', encoding="utf-8")
    Image.new("RGB", (640, 360), "white").save(assets / "figure_1.png")
    (analysis / "paper_analysis.json").write_text(
        json.dumps({"main_claim": "A supported claim."}), encoding="utf-8"
    )
    (workspace / "resolved_request.json").write_text(
        json.dumps({"language": "zh-CN", "duration_minutes": 30}), encoding="utf-8"
    )
    bundle = PaperSourceBundle(
        material_id="material_paper",
        file_hash="a" * 64,
        page_count=1,
        pdf_path="paper.pdf",
        paper_source_path="paper_source.json",
        paper_content_path="paper_content.md",
        asset_directory="existing_assets",
    )
    return workspace, bundle


def _write_native_outputs(invocation: CodexSkillInvocation, *, valid: bool = True) -> None:
    assert invocation.output_directory is not None
    output = invocation.output_directory
    prompts = output / "prompts"
    images = output / "images"
    rendered = output / "rendered"
    prompts.mkdir(parents=True, exist_ok=True)
    images.mkdir(parents=True, exist_ok=True)
    rendered.mkdir(parents=True, exist_ok=True)
    (output / "analysis.md").write_text("# Analysis\n", encoding="utf-8")
    (output / "deck-brief.md").write_text(
        "# Deck Brief\n\n- style_preset: `journal-minimal`\n- language: zh-CN\n",
        encoding="utf-8",
    )
    outline_parts = ["# Outline\n"]
    slide_count = 2
    for index in range(1, slide_count + 1):
        (prompts / f"{index:02d}-slide.md").write_text("Generate a slide.\n", encoding="utf-8")
        outline_parts.append(
            f"""## {index:02d}. Slide {index}
- Role: method
- Message: Explain supported point {index}.
- Render mode: native-raster
- Visual: A factual method diagram.
- Text: Point {index}; Evidence {index}
- Evidence: Paper page {index}
- Source visual: None
"""
        )
        if valid or index == 1:
            Image.new("RGB", (1600, 900), (index * 30, index * 40, index * 50)).save(
                images / f"{index:02d}-slide.png"
            )
            Image.new("RGB", (1600, 900), (index * 30, index * 40, index * 50)).save(
                rendered / f"{index:02d}-slide.png"
            )
    (output / "outline.md").write_text("\n".join(outline_parts), encoding="utf-8")
    (output / "generation-log.md").write_text(
        "\n".join(
            f"images/{index:02d}-slide.png: backend=imagegen render_mode=native-raster"
            for index in range(1, 3)
        ),
        encoding="utf-8",
    )
    (output / "source-visual-manifest.json").write_text(
        json.dumps(
            {
                "schema_version": "1.0",
                "slides": [
                    {
                        "slide_id": f"slide_{index:03d}",
                        "order": index,
                        "render_mode": "native-raster",
                        "background_path": f"images/{index:02d}-slide.png",
                        "assets": [],
                        "annotations": [],
                    }
                    for index in range(1, slide_count + 1)
                ],
            }
        ),
        encoding="utf-8",
    )

    pdf = fitz.open()
    for _ in range(slide_count):
        pdf.new_page(width=1600, height=900)
    pdf.save(output / "presentation.pdf")
    pdf.close()

    pptx = Presentation()
    for index in range(1, slide_count + 1):
        image_path = images / f"{index:02d}-slide.png"
        if not image_path.is_file():
            continue
        slide = pptx.slides.add_slide(pptx.slide_layouts[6])
        picture = slide.shapes.add_picture(
            str(image_path),
            0,
            0,
            width=pptx.slide_width,
            height=pptx.slide_height,
        )
        picture.name = "background:native-raster"
    pptx.save(output / "presentation.pptx")


def _success_report(invocation: CodexSkillInvocation) -> StageExecutionReport:
    now = utc_now()
    return StageExecutionReport(
        stage=invocation.stage,
        skill_name=invocation.skill_name,
        skill_version=invocation.skill_version,
        prompt_version=invocation.prompt_version,
        runtime_version="fake-runtime-v1",
        status=StageStatus.SUCCEEDED,
        attempt=invocation.attempt,
        started_at=now,
        finished_at=now,
        input_hash=invocation.input_hash,
        exit_code=0,
        outputs=[
            str(path.relative_to(invocation.workspace))
            for path in invocation.output_directory.rglob("*")
            if path.is_file()
        ],
        validation_passed=True,
    )


def _context(tmp_path: Path) -> PaperProviderContext:
    workspace, bundle = _workspace(tmp_path)
    checkpoint = PaperWorkflowCheckpoint(
        job_id="paper_job_native",
        request_hash="b" * 64,
    )
    return PaperProviderContext(
        job_id="paper_job_native",
        workspace=workspace,
        request=PaperWorkflowRequest(material_id="material_paper", duration_minutes=30),
        source_bundle=bundle,
        checkpoint=checkpoint,
        report_progress=lambda _progress, _stage: None,
        is_pause_requested=lambda: False,
        persist_checkpoint=lambda _checkpoint: None,
    )


def test_native_provider_runs_complete_skill_without_presentation_plan(tmp_path: Path) -> None:
    def handler(invocation: CodexSkillInvocation) -> StageExecutionReport:
        _write_native_outputs(invocation)
        return _success_report(invocation)

    runtime = FakeCodexSkillRuntime(handler)
    context = _context(tmp_path)
    provider = NativePaperDeckProvider(runtime, skill_directory=_skill(tmp_path))

    artifacts = provider.run(context)

    assert artifacts.format == "pdf"
    assert artifacts.presentation_pdf_path == "provider_output/presentation.pdf"
    assert artifacts.source_visual_manifest_path == (
        "provider_output/source-visual-manifest.json"
    )
    assert artifacts.debug_pptx_path == "provider_output/presentation.pptx"
    assert len(runtime.invocations) == 1
    invocation = runtime.invocations[0]
    assert invocation.skill_name == "paper-deck"
    assert invocation.output_directory == context.workspace / "provider_output"
    assert invocation.timeout_seconds == 3600
    prompt = invocation.prompt_path.read_text(encoding="utf-8")
    assert "最多" in prompt and "重试 2 次" in prompt
    assert "预留 10 分钟" in prompt
    assert "不得尝试 LibreOffice/PowerPoint/Keynote/WPS" in prompt
    assert invocation.network_enabled is False
    assert "source-visual-manifest.json" in invocation.expected_outputs
    assert {path.name for path in invocation.input_paths} >= {
        "paper.pdf",
        "paper_content.md",
        "paper_source.json",
        "paper_analysis.json",
        "resolved_request.json",
        "figure_1.png",
    }
    prompt = invocation.prompt_path.read_text(encoding="utf-8")
    assert "完整保留 paper-deck 的分析、叙事" in prompt
    assert "source-grounded-hybrid" in prompt
    assert "不要创建或消费 MetaClass PresentationPlan" in prompt
    assert "不要调用 CodexPPTProvider" in prompt
    assert "保持为独立 PowerPoint 对象" in prompt


def test_native_provider_reuses_valid_checkpoint(tmp_path: Path) -> None:
    def handler(invocation: CodexSkillInvocation) -> StageExecutionReport:
        _write_native_outputs(invocation)
        return _success_report(invocation)

    runtime = FakeCodexSkillRuntime(handler)
    context = _context(tmp_path)
    provider = NativePaperDeckProvider(runtime, skill_directory=_skill(tmp_path))

    first = provider.run(context)
    second = provider.run(context)

    assert first.presentation_pdf_path == second.presentation_pdf_path
    assert len(runtime.invocations) == 1


def test_native_provider_repairs_only_named_pages_and_revalidates_pdf(tmp_path: Path) -> None:
    def handler(invocation: CodexSkillInvocation) -> StageExecutionReport:
        if "targeted-repair" in invocation.prompt_version:
            prompt = invocation.prompt_path.read_text(encoding="utf-8")
            assert '"slide_id": "paper_deck_slide_002"' in prompt
            (invocation.output_directory / "prompts/02-slide.md").write_text(
                "Remove the unsupported number and regenerate this page only.\n",
                encoding="utf-8",
            )
        else:
            _write_native_outputs(invocation)
        return _success_report(invocation)

    runtime = FakeCodexSkillRuntime(handler)
    context = _context(tmp_path)
    provider = NativePaperDeckProvider(runtime, skill_directory=_skill(tmp_path))
    provider.run(context)
    analysis_before = hashlib.sha256(
        (context.workspace / "provider_output/analysis.md").read_bytes()
    ).hexdigest()
    outline_before = hashlib.sha256(
        (context.workspace / "provider_output/outline.md").read_bytes()
    ).hexdigest()

    repaired = provider.repair_slides(
        context,
        directives=[
            {
                "slide_id": "paper_deck_slide_002",
                "prompt_path": "provider_output/prompts/02-slide.md",
                "unverified_mentions": ["99.9%"],
                "instruction": "Remove unsupported 99.9%.",
                "scope": "single_slide",
            }
        ],
    )

    assert repaired.presentation_pdf_path == "provider_output/presentation.pdf"
    assert len(runtime.invocations) == 2
    assert runtime.invocations[-1].prompt_version.endswith("targeted-repair-v1")
    assert runtime.invocations[-1].attempt == 2
    assert hashlib.sha256(
        (context.workspace / "provider_output/analysis.md").read_bytes()
    ).hexdigest() == analysis_before
    assert hashlib.sha256(
        (context.workspace / "provider_output/outline.md").read_bytes()
    ).hexdigest() == outline_before


def test_native_analysis_and_outline_are_not_replaced(tmp_path: Path) -> None:
    test_native_provider_repairs_only_named_pages_and_revalidates_pdf(tmp_path)


def test_native_provider_rejects_prompt_image_page_mismatch(tmp_path: Path) -> None:
    def handler(invocation: CodexSkillInvocation) -> StageExecutionReport:
        _write_native_outputs(invocation, valid=False)
        return _success_report(invocation)

    context = _context(tmp_path)
    provider = NativePaperDeckProvider(
        FakeCodexSkillRuntime(handler),
        skill_directory=_skill(tmp_path),
    )

    with pytest.raises(ValueError, match="page counts must match"):
        provider.run(context)


def test_native_provider_requires_precomputed_paper_analysis(tmp_path: Path) -> None:
    context = _context(tmp_path)
    (context.workspace / "stages/01_analysis/output/paper_analysis.json").unlink()
    provider = NativePaperDeckProvider(
        FakeCodexSkillRuntime(lambda invocation: _success_report(invocation)),
        skill_directory=_skill(tmp_path),
    )

    with pytest.raises(NativePaperDeckError, match="paper_analysis.json is required"):
        provider.run(context)
