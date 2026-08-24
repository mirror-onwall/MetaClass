import json
import os
from pathlib import Path
from threading import Event, Thread

import fitz
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation

import metaclass.modules.paper_workflow.stage_cache as stage_cache_module
from metaclass.core.schemas import utc_now
from metaclass.main import create_app
from metaclass.modules.paper_workflow.orchestrator import PaperWorkflowOrchestrator
from metaclass.modules.paper_workflow.providers.composed_skills import ComposedSkillsProvider
from metaclass.modules.paper_workflow.runtime import FakeCodexSkillRuntime
from metaclass.modules.paper_workflow.schemas import (
    ArtifactFile,
    ComposedStage,
    PaperArtifactBundle,
    PaperWorkflowCheckpoint,
    PaperWorkflowJob,
    PaperWorkflowStatus,
    StageExecutionReport,
    StageStatus,
)
from metaclass.modules.paper_workflow.stage_cache import (
    PaperWorkflowCheckpointStore,
    StageCache,
)

HASH = "a" * 64
REQUIRED_ROLES = [
    "presentation",
    "analysis",
    "outline",
    "slide_evidence",
    "asset_manifest",
    "speaker_notes",
    "generation_report",
    "qa_report",
]


class SuccessfulPaperProvider:
    name = "composed_skills"

    def __init__(self) -> None:
        self.executions = {stage: 0 for stage in ComposedStage}

    @staticmethod
    def _write_output(context, relative_path: str, value: str) -> str:
        path = context.workspace / relative_path
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(value, encoding="utf-8")
        return value

    @staticmethod
    def _restore_output(context, outputs: dict[str, str]) -> str:
        return (context.workspace / outputs["marker"]).read_text(encoding="utf-8")

    def run(self, context) -> PaperArtifactBundle:
        for stage, skill_name in [
            (ComposedStage.ANALYSIS, "paper-analyze"),
            (ComposedStage.FIGURES, "extract-paper-images"),
            (ComposedStage.OUTLINE, "academic-pptx"),
            (ComposedStage.GENERATION, "academic-pptx-generate"),
        ]:

            def execute(current_stage=stage):
                self.executions[current_stage] += 1
                relative_path = f"stages/{current_stage.value}/completed.json"
                value = self._write_output(context, relative_path, current_stage.value)
                return value, {"marker": relative_path}

            restored = context.run_stage(
                stage=stage,
                input_hash="sha256:"
                + {
                    ComposedStage.ANALYSIS: "c",
                    ComposedStage.FIGURES: "d",
                    ComposedStage.OUTLINE: "e",
                    ComposedStage.GENERATION: "f",
                }[stage]
                * 64,
                skill_name=skill_name,
                skill_version="test-v1",
                prompt_version="v1",
                execute=execute,
                validate_cached=StageCache(context.workspace).outputs_exist,
                restore=lambda outputs: self._restore_output(context, outputs),
            )
            assert restored == stage.value
        return self._bundle(context)

    def _bundle(self, context) -> PaperArtifactBundle:
        return PaperArtifactBundle(
            id=f"paper_bundle_{context.job_id}",
            job_id=context.job_id,
            source_material_id=context.request.material_id,
            provider=self.name,
            root_path=str(context.workspace / "final"),
            files=[
                ArtifactFile(
                    role=role,
                    path="presentation.pptx" if role == "presentation" else f"{role}.json",
                    sha256=HASH,
                    media_type=(
                        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
                        if role == "presentation"
                        else "application/json"
                    ),
                )
                for role in REQUIRED_ROLES
            ],
            validation_status="passed",
            derived_material_id="mat_derived_test",
        )


class BlockingPaperProvider(SuccessfulPaperProvider):
    def __init__(self) -> None:
        super().__init__()
        self.started = Event()
        self.release = Event()

    def run(self, context) -> PaperArtifactBundle:
        def execute():
            self.executions[ComposedStage.ANALYSIS] += 1
            self.started.set()
            assert self.release.wait(timeout=5)
            relative_path = "stages/analysis/completed.json"
            value = self._write_output(context, relative_path, "analysis")
            return value, {"marker": relative_path}

        context.run_stage(
            stage=ComposedStage.ANALYSIS,
            input_hash="sha256:" + "b" * 64,
            skill_name="paper-analyze",
            skill_version="test-v1",
            prompt_version="v1",
            execute=execute,
            validate_cached=StageCache(context.workspace).outputs_exist,
            restore=lambda outputs: self._restore_output(context, outputs),
        )
        return super().run(context)


class FailOncePaperProvider(SuccessfulPaperProvider):
    def __init__(self) -> None:
        super().__init__()
        self.failed_once = False

    def run(self, context) -> PaperArtifactBundle:
        def analyze():
            self.executions[ComposedStage.ANALYSIS] += 1
            relative_path = "stages/analysis/paper_analysis.json"
            value = self._write_output(context, relative_path, "analysis")
            return value, {"analysis": relative_path}

        context.run_stage(
            stage=ComposedStage.ANALYSIS,
            input_hash="sha256:" + "c" * 64,
            skill_name="paper-analyze",
            skill_version="test-v1",
            prompt_version="v1",
            execute=analyze,
            validate_cached=StageCache(context.workspace).outputs_exist,
            restore=lambda outputs: (context.workspace / outputs["analysis"]).read_text(
                encoding="utf-8"
            ),
        )

        def prepare_figures():
            self.executions[ComposedStage.FIGURES] += 1
            if not self.failed_once:
                self.failed_once = True
                raise RuntimeError("injected stage failure")
            relative_path = "stages/figures/figures.json"
            value = self._write_output(context, relative_path, "figures")
            return value, {"figures": relative_path}

        context.run_stage(
            stage=ComposedStage.FIGURES,
            input_hash="sha256:" + "d" * 64,
            skill_name="extract-paper-images",
            skill_version="test-v1",
            prompt_version="v1",
            execute=prepare_figures,
            validate_cached=StageCache(context.workspace).outputs_exist,
            restore=lambda outputs: (context.workspace / outputs["figures"]).read_text(
                encoding="utf-8"
            ),
        )
        for stage, skill_name in [
            (ComposedStage.OUTLINE, "academic-pptx"),
            (ComposedStage.GENERATION, "academic-pptx-generate"),
        ]:

            def execute(current_stage=stage):
                self.executions[current_stage] += 1
                relative_path = f"stages/{current_stage.value}/completed.json"
                value = self._write_output(context, relative_path, current_stage.value)
                return value, {"marker": relative_path}

            context.run_stage(
                stage=stage,
                input_hash="sha256:" + ("e" if stage == ComposedStage.OUTLINE else "f") * 64,
                skill_name=skill_name,
                skill_version="test-v1",
                prompt_version="v1",
                execute=execute,
                validate_cached=StageCache(context.workspace).outputs_exist,
                restore=lambda outputs: self._restore_output(context, outputs),
            )
        return self._bundle(context)


def _pdf_bytes() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Paper workflow API fixture")
    payload = document.tobytes()
    document.close()
    return payload


def _create_pdf_material(client: TestClient) -> str:
    response = client.post(
        "/api/v1/materials",
        files={"file": ("paper.pdf", _pdf_bytes(), "application/pdf")},
    )
    assert response.status_code == 201
    return response.json()["id"]


def test_paper_workflow_api_lifecycle_and_result(tmp_path: Path) -> None:
    app = create_app(tmp_path)
    service = app.state.services.paper_workflows
    service.orchestrator = PaperWorkflowOrchestrator(tmp_path, [SuccessfulPaperProvider()])

    with TestClient(app) as client:
        material_id = _create_pdf_material(client)
        created = client.post(
            "/api/v1/paper-workflows",
            json={"material_id": material_id, "strategy": "composed_skills"},
        )
        assert created.status_code == 201
        job_id = created.json()["id"]
        assert created.json()["status"] == "queued"

        fetched = client.get(f"/api/v1/paper-workflows/{job_id}")
        assert fetched.status_code == 200
        assert fetched.json()["source_material_id"] == material_id

        paused = client.post(f"/api/v1/paper-workflows/{job_id}/pause")
        assert paused.json()["status"] == "paused"
        resumed = client.post(f"/api/v1/paper-workflows/{job_id}/resume")
        assert resumed.json()["status"] == "queued"

        completed = client.post(f"/api/v1/paper-workflows/{job_id}/run")
        assert completed.status_code == 200
        assert completed.json()["status"] == "succeeded"
        assert completed.json()["progress"] == 1

        result = client.get(f"/api/v1/paper-workflows/{job_id}/result")
        assert result.status_code == 200
        assert result.json()["job_id"] == job_id
        assert result.json()["validation_status"] == "passed"

        checkpoint_path = tmp_path / "runtime" / "paper_workflows" / job_id / "checkpoint.json"
        checkpoint = json.loads(checkpoint_path.read_text(encoding="utf-8"))
        analysis = checkpoint["stages"]["analysis"]
        assert analysis["status"] == "succeeded"
        assert analysis["skill_name"] == "paper-analyze"
        assert analysis["skill_version"] == "test-v1"
        assert analysis["prompt_version"] == "v1"
        assert analysis["started_at"] and analysis["finished_at"]
        assert analysis["outputs"]["marker"].endswith("completed.json")
        assert analysis["validation"]["passed"] is True
        assert not list(checkpoint_path.parent.glob(".checkpoint-*.tmp"))

        repeated = client.post(f"/api/v1/paper-workflows/{job_id}/run")
        assert repeated.json()["status"] == "succeeded"
        assert service.orchestrator.providers["composed_skills"].executions == {
            stage: 1 for stage in ComposedStage
        }


def test_succeeded_workflow_creates_reconciled_paper_deck_course(
    tmp_path: Path, monkeypatch
) -> None:
    app = create_app(tmp_path)
    service = app.state.services.paper_workflows
    with TestClient(app) as client:
        source_material_id = _create_pdf_material(client)
        created = client.post(
            "/api/v1/paper-workflows",
            json={"material_id": source_material_id, "audience": "研究生"},
        ).json()
        job_id = created["id"]
        workspace = tmp_path / "runtime" / "paper_workflows" / job_id
        service.source_bundles.build(
            material_id=source_material_id,
            workspace=workspace,
        )
        final = workspace / "final"
        final.mkdir(parents=True)

        deck_path = final / "presentation.pptx"
        presentation = Presentation()
        for title in ("Research question", "Evidence supports the claim"):
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            slide.shapes.title.text = title
        presentation.save(deck_path)
        deck = service.materials.register_generated_pptx(
            deck_path,
            filename="paper-presentation.pptx",
            derivation_key=job_id,
        )

        outline = {
            "title": "Paper class",
            "paper_type": "methods",
            "narrative_arc": "problem-to-solution",
            "objectives": ["Explain the evidence"],
            "structure_summary": "Question then evidence.",
            "sections": [
                {
                    "id": "section_main",
                    "title": "Main argument",
                    "role": "method",
                    "content_goal": "Explain the paper.",
                    "slide_ids": ["slide_01", "slide_02"],
                }
            ],
            "slides": [
                {
                    "id": "slide_01",
                    "order": 1,
                    "title": "Research question",
                    "purpose": "Introduce the question.",
                    "key_points": ["Question"],
                    "speaker_note": "Explain the question.",
                    "layout_intent": "title",
                },
                {
                    "id": "slide_02",
                    "order": 2,
                    "title": "Evidence supports the claim",
                    "purpose": "Interpret the evidence.",
                    "key_points": ["Evidence"],
                    "speaker_note": "Explain the evidence.",
                    "layout_intent": "result",
                },
            ],
        }
        evidence = {
            "slides": [
                {
                    "slide_id": "slide_01",
                    "claim_ids": ["claim_01"],
                    "source_refs": [{"page_no": 1}],
                },
                {
                    "slide_id": "slide_02",
                    "claim_ids": ["claim_02"],
                    "source_refs": [{"page_no": 1}],
                },
            ]
        }
        (final / "presentation_outline.json").write_text(json.dumps(outline), encoding="utf-8")
        (final / "paper_analysis.json").write_text(
            json.dumps(
                {
                    "paper_type": "methods",
                    "central_question": "What problem does the method solve?",
                    "knowledge_gap": "Prior work lacks grounded evidence.",
                    "main_claim": "The method solves the target problem.",
                    "method_summary": {"approach": "Evidence-aware processing."},
                    "claims": [
                        {
                            "id": "claim_01",
                            "statement": "The method addresses the research question.",
                            "importance": "core",
                            "confidence": 0.9,
                            "source_refs": [{"page_no": 1}],
                        },
                        {
                            "id": "claim_02",
                            "statement": "The evidence supports the method.",
                            "importance": "supporting",
                            "confidence": 0.8,
                            "source_refs": [{"page_no": 1}],
                        },
                    ],
                }
            ),
            encoding="utf-8",
        )
        (final / "asset_manifest.json").write_text(json.dumps({"figures": []}), encoding="utf-8")
        (final / "slide_evidence.json").write_text(json.dumps(evidence), encoding="utf-8")
        (final / "speaker_notes.json").write_text(
            json.dumps(
                [
                    {"slide_id": "slide_01", "note": "Authoring note one."},
                    {"slide_id": "slide_02", "note": "Authoring note two."},
                ]
            ),
            encoding="utf-8",
        )
        bundle_id = f"paper_bundle_{job_id}"
        bundle = PaperArtifactBundle(
            id=bundle_id,
            job_id=job_id,
            source_material_id=source_material_id,
            provider="composed_skills",
            root_path=str(final.relative_to(tmp_path)),
            files=[
                ArtifactFile(
                    role=role,
                    path="presentation.pptx" if role == "presentation" else f"{role}.json",
                    sha256=HASH,
                    media_type=(
                        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
                        if role == "presentation"
                        else "application/json"
                    ),
                )
                for role in REQUIRED_ROLES
            ],
            validation_status="passed",
            derived_material_id=deck.id,
        )
        service.repository.save_bundle(bundle)
        request = service.repository.get_request(job_id)
        checkpoint = service.checkpoints.load(job_id)
        assert request is not None and checkpoint is not None
        succeeded = PaperWorkflowJob.model_validate(
            {
                **created,
                "status": "succeeded",
                "stage": "completed",
                "progress": 1,
                "artifact_bundle_id": bundle.id,
                "derived_material_id": deck.id,
            }
        )
        service.repository.save_job(succeeded, request, checkpoint)

        def render(material, page_count, cancel_event=None):
            del cancel_event
            root = tmp_path / "processed" / material.id / "pages"
            root.mkdir(parents=True, exist_ok=True)
            paths = []
            for number in range(1, page_count + 1):
                path = root / f"page_{number:03d}.png"
                Image.new("RGB", (1280, 720), "white").save(path)
                paths.append(path)
            return paths

        monkeypatch.setattr(service.materials, "_render_pptx", render)
        result = client.post(f"/api/v1/paper-workflows/{job_id}/create-paper-deck-course")

        assert result.status_code == 201
        payload = result.json()
        assert payload["derived_material_id"] == deck.id
        content = client.get(f"/api/v1/learning-contents/{payload['content_id']}").json()
        assert content["organization_mode"] == "paper_deck"
        plan = client.get(f"/api/v1/presentation-plans/{payload['presentation_plan_id']}").json()
        assert plan["mode"] == "paper_deck"
        assert plan["source_paper_material_id"] == source_material_id
        assert [slide["source_page_no"] for slide in plan["slides"]] == [1, 2]
        assert plan["slides"][1]["paper_claim_ids"] == ["claim_02"]
        resource = client.get(
            f"/api/v1/presentation-plans/{payload['presentation_plan_id']}/resource"
        ).json()
        assert resource["kind"] == "paper_deck"
        assert [slide["source_page_no"] for slide in resource["slides"]] == [1, 2]
        classroom = client.post(
            f"/api/v1/learning-contents/{payload['content_id']}/classroom-plans",
            params={"presentation_plan_id": payload["presentation_plan_id"]},
        )
        assert classroom.status_code == 201
        session = client.post(f"/api/v1/classroom-plans/{classroom.json()['id']}/sessions")
        assert session.status_code == 201
        updated = client.patch(
            f"/api/v1/presentation-plans/{payload['presentation_plan_id']}"
            "/slides/slide_01/speaker-script",
            json={"speaker_script": "教师人工修订后的第一页讲稿。"},
        )
        assert updated.status_code == 200
        assert updated.json()["slides"][0]["speaker_script_source"] == "teacher_override"
        repeated = client.post(f"/api/v1/paper-workflows/{job_id}/create-paper-deck-course")
        assert repeated.status_code == 201
        assert repeated.json()["content_id"] == payload["content_id"]
        assert repeated.json()["presentation_plan_id"] == payload["presentation_plan_id"]
        preserved = client.get(
            f"/api/v1/presentation-plans/{payload['presentation_plan_id']}"
        ).json()
        assert preserved["slides"][0]["speaker_script"] == "教师人工修订后的第一页讲稿。"
        assert preserved["slides"][0]["speaker_script_source"] == "teacher_override"


def test_default_provider_reports_real_runtime_failure(tmp_path: Path) -> None:
    app = create_app(tmp_path)
    skill_directory = tmp_path / "skills" / "paper-analyze"
    skill_directory.mkdir(parents=True)
    (skill_directory / "SKILL.md").write_text("# paper-analyze\n", encoding="utf-8")

    def fail_runtime(invocation):
        now = utc_now()
        return StageExecutionReport(
            stage=invocation.stage,
            skill_name=invocation.skill_name,
            skill_version=invocation.skill_version,
            prompt_version=invocation.prompt_version,
            runtime_version="fake-failing-runtime-v1",
            status=StageStatus.FAILED,
            attempt=invocation.attempt,
            started_at=now,
            finished_at=now,
            input_hash=invocation.input_hash,
            exit_code=1,
            validation_passed=False,
        )

    runtime = FakeCodexSkillRuntime(fail_runtime)
    app.state.services.paper_workflows.orchestrator = PaperWorkflowOrchestrator(
        tmp_path,
        [
            ComposedSkillsProvider(
                runtime,
                skill_directory=skill_directory,
            )
        ],
    )
    with TestClient(app) as client:
        material_id = _create_pdf_material(client)
        created = client.post(
            "/api/v1/paper-workflows",
            json={"material_id": material_id, "strategy": "composed_skills"},
        )
        job_id = created.json()["id"]
        run = client.post(f"/api/v1/paper-workflows/{job_id}/run")
        assert run.status_code == 200
        assert run.json()["status"] == "failed"
        assert "paper-analyze runtime execution failed" in run.json()["error"]
        assert client.get(f"/api/v1/paper-workflows/{job_id}/result").status_code == 422


def test_paper_workflow_rejects_non_pdf_material(tmp_path: Path) -> None:
    app = create_app(tmp_path)
    with TestClient(app) as client:
        uploaded = client.post(
            "/api/v1/materials",
            files={"file": ("slides.pptx", b"fixture", "application/octet-stream")},
        )
        assert uploaded.status_code == 201
        workflow = client.post(
            "/api/v1/paper-workflows",
            json={"material_id": uploaded.json()["id"]},
        )
        assert workflow.status_code == 422


def test_running_pause_resume_and_duplicate_calls_are_idempotent(tmp_path: Path) -> None:
    app = create_app(tmp_path)
    provider = BlockingPaperProvider()
    service = app.state.services.paper_workflows
    service.orchestrator = PaperWorkflowOrchestrator(tmp_path, [provider])
    with TestClient(app) as client:
        material_id = _create_pdf_material(client)
        job_id = client.post("/api/v1/paper-workflows", json={"material_id": material_id}).json()[
            "id"
        ]

        runner = Thread(target=service.run, args=(job_id,))
        runner.start()
        assert provider.started.wait(timeout=5)
        assert service.run(job_id).status == PaperWorkflowStatus.RUNNING
        assert provider.executions[ComposedStage.ANALYSIS] == 1

        assert service.pause(job_id).status == PaperWorkflowStatus.PAUSED
        provider.release.set()
        runner.join(timeout=5)
        assert not runner.is_alive()
        assert service.get(job_id).status == PaperWorkflowStatus.PAUSED

        assert service.resume(job_id).status == PaperWorkflowStatus.QUEUED
        assert service.resume(job_id).status == PaperWorkflowStatus.QUEUED
        resumed_provider = SuccessfulPaperProvider()
        service.orchestrator = PaperWorkflowOrchestrator(tmp_path, [resumed_provider])
        assert service.run(job_id).status == PaperWorkflowStatus.SUCCEEDED


def test_successful_checkpoint_stage_is_reused_after_retry(tmp_path: Path) -> None:
    app = create_app(tmp_path)
    provider = FailOncePaperProvider()
    service = app.state.services.paper_workflows
    service.orchestrator = PaperWorkflowOrchestrator(tmp_path, [provider])
    with TestClient(app) as client:
        material_id = _create_pdf_material(client)
        job_id = client.post("/api/v1/paper-workflows", json={"material_id": material_id}).json()[
            "id"
        ]
        assert service.run(job_id).status == PaperWorkflowStatus.FAILED
        failed_checkpoint = service.checkpoints.load(job_id)
        assert failed_checkpoint is not None
        assert failed_checkpoint.stages[ComposedStage.FIGURES].status.value == "failed"
        assert failed_checkpoint.stages[ComposedStage.FIGURES].validation["passed"] is False
        assert service.resume(job_id).status == PaperWorkflowStatus.QUEUED
        assert service.resume(job_id).status == PaperWorkflowStatus.QUEUED
        assert service.run(job_id).status == PaperWorkflowStatus.SUCCEEDED
        assert provider.executions[ComposedStage.ANALYSIS] == 1
        assert provider.executions[ComposedStage.FIGURES] == 2


def test_missing_cached_output_reruns_successful_stage(tmp_path: Path) -> None:
    app = create_app(tmp_path)
    provider = FailOncePaperProvider()
    service = app.state.services.paper_workflows
    service.orchestrator = PaperWorkflowOrchestrator(tmp_path, [provider])
    with TestClient(app) as client:
        material_id = _create_pdf_material(client)
        job_id = client.post("/api/v1/paper-workflows", json={"material_id": material_id}).json()[
            "id"
        ]
        assert service.run(job_id).status == PaperWorkflowStatus.FAILED
        analysis_path = (
            tmp_path
            / "runtime"
            / "paper_workflows"
            / job_id
            / "stages"
            / "analysis"
            / "paper_analysis.json"
        )
        analysis_path.unlink()

        assert service.resume(job_id).status == PaperWorkflowStatus.QUEUED
        assert service.run(job_id).status == PaperWorkflowStatus.SUCCEEDED
        assert provider.executions[ComposedStage.ANALYSIS] == 2


def test_nature_workflow_waits_for_missing_presentation_context(tmp_path: Path) -> None:
    app = create_app(tmp_path)
    with TestClient(app) as client:
        material_id = _create_pdf_material(client)
        created = client.post(
            "/api/v1/paper-workflows",
            json={"material_id": material_id, "strategy": "nature_paper2ppt"},
        )

        assert created.status_code == 201
        payload = created.json()
        assert payload["status"] == "waiting_for_input"
        assert payload["required_input"] == {
            "reason": "nature_paper2ppt_requires_presentation_context",
            "fields": ["duration_minutes", "audience"],
        }
        initial_checkpoint = app.state.services.paper_workflows.checkpoints.load(payload["id"])
        assert initial_checkpoint is not None
        run = client.post(f"/api/v1/paper-workflows/{payload['id']}/run")
        assert run.status_code == 409

        supplied = client.patch(
            f"/api/v1/paper-workflows/{payload['id']}/request",
            json={"duration_minutes": 20, "audience": "计算机专业研究生"},
        )
        assert supplied.status_code == 200
        assert supplied.json()["status"] == "queued"
        assert supplied.json()["required_input"] is None
        request = app.state.services.paper_workflows.repository.get_request(payload["id"])
        assert request is not None
        assert request.duration_minutes == 20
        assert request.audience == "计算机专业研究生"
        updated_checkpoint = app.state.services.paper_workflows.checkpoints.load(payload["id"])
        assert updated_checkpoint is not None
        assert updated_checkpoint.request_hash != initial_checkpoint.request_hash
        assert updated_checkpoint.version == initial_checkpoint.version + 1

        repeated = client.patch(
            f"/api/v1/paper-workflows/{payload['id']}/request",
            json={"duration_minutes": 20, "audience": "计算机专业研究生"},
        )
        assert repeated.status_code == 200
        assert repeated.json()["status"] == "queued"

        conflicting = client.patch(
            f"/api/v1/paper-workflows/{payload['id']}/request",
            json={"duration_minutes": 30, "audience": "高校教师"},
        )
        assert conflicting.status_code == 409


def test_service_restart_recovers_running_job_as_paused(tmp_path: Path) -> None:
    first_app = create_app(tmp_path)
    with TestClient(first_app) as client:
        material_id = _create_pdf_material(client)
        job_id = client.post("/api/v1/paper-workflows", json={"material_id": material_id}).json()[
            "id"
        ]
        service = first_app.state.services.paper_workflows
        job = service.get(job_id)
        request = service.repository.get_request(job_id)
        checkpoint = service.checkpoints.load(job_id)
        assert request is not None and checkpoint is not None
        job.status = PaperWorkflowStatus.RUNNING
        service.repository.save_job(job, request, checkpoint)

    second_app = create_app(tmp_path)
    with TestClient(second_app) as client:
        restored = client.get(f"/api/v1/paper-workflows/{job_id}")
        assert restored.status_code == 200
        assert restored.json()["status"] == "paused"


def test_checkpoint_store_commits_with_atomic_replace(tmp_path: Path, monkeypatch) -> None:
    calls: list[tuple[Path, Path]] = []
    real_replace = os.replace

    def record_replace(source, destination) -> None:
        calls.append((Path(source), Path(destination)))
        real_replace(source, destination)

    monkeypatch.setattr(stage_cache_module.os, "replace", record_replace)
    store = PaperWorkflowCheckpointStore(tmp_path)
    checkpoint = PaperWorkflowCheckpoint(job_id="paper_job_atomic", request_hash=HASH)
    path = store.save(checkpoint)

    assert len(calls) == 1
    assert calls[0][1] == path
    assert calls[0][0].parent == path.parent
    assert calls[0][0].suffix == ".tmp"
    assert store.load(checkpoint.job_id) == checkpoint
    assert not calls[0][0].exists()
