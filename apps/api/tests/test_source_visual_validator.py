import hashlib
import importlib.util
import json
from pathlib import Path

import fitz
import pytest
from PIL import Image
from pptx import Presentation

from metaclass.modules.paper_workflow.schemas import PaperPresentationArtifact
from metaclass.modules.paper_workflow.source_visual_validator import (
    SourceVisualValidationError,
    SourceVisualValidator,
)

SCRIPT = Path(__file__).parents[3] / ".agents/skills/paper-deck/scripts/merge_deck.py"
SPEC = importlib.util.spec_from_file_location("validator_merge_deck", SCRIPT)
assert SPEC and SPEC.loader
merge_deck = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(merge_deck)


def _write_json(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload), encoding="utf-8")


def _fixture(tmp_path: Path) -> tuple[Path, PaperPresentationArtifact]:
    root = tmp_path / "job"
    provider_input = root / "provider_input"
    output = root / "provider_output"
    assets_dir = provider_input / "assets"
    images = output / "images"
    prompts = output / "prompts"
    assets_dir.mkdir(parents=True)
    images.mkdir(parents=True)
    prompts.mkdir()
    source = assets_dir / "figure-1.png"
    Image.new("RGB", (800, 400), "teal").save(source)
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    _write_json(
        provider_input / "paper_source.json",
        {
            "assets": [
                {
                    "id": "asset_figure_1",
                    "type": "figure",
                    "page_no": 5,
                    "path": "existing_assets/figure-1.png",
                    "caption": "Figure 1: Main result",
                }
            ]
        },
    )
    _write_json(provider_input / "paper_analysis.json", {"claims": [], "figure_candidates": []})
    Image.new("RGB", (1600, 900), "navy").save(images / "01-slide.png")
    Image.new("RGB", (1600, 900), "white").save(images / "02-background.png")
    (prompts / "01-slide.md").write_text("Create a native raster slide.", encoding="utf-8")
    (prompts / "02-slide.md").write_text(
        "Create the journal-minimal background. Reserve an empty evidence bay. "
        "Do not draw or reconstruct source evidence. "
        "The verified original source asset will be inserted deterministically after generation.",
        encoding="utf-8",
    )
    (output / "analysis.md").write_text("# Native paper-deck analysis\n", encoding="utf-8")
    (output / "deck-brief.md").write_text(
        "# Deck Brief\n- style_preset: `journal-minimal`\n- language: zh-CN\n",
        encoding="utf-8",
    )
    (output / "outline.md").write_text(
        """# Outline

## 01. Context
- Role: context
- Message: Establish context.
- Render mode: native-raster
- Visual: Generated context composition.
- Text: Context
- Evidence: Introduction
- Source visual: None

## 02. Evidence
- Role: evidence
- Message: Show the paper result.
- Render mode: source-grounded-hybrid
- Visual: Figure 1 with a concise callout.
- Text: Main result
- Evidence: Figure 1, page 5
- Source visual: Figure 1 from PDF page 5
- Source assets:
  - asset_id: asset_figure_1
    placement: [0.07, 0.23, 0.62, 0.66]
    fit: contain
    crop: full
    preserve: [axes, legend, labels, values]
- Generated layer: white background with empty evidence bay
- Overlay annotations:
  - Main result
""",
        encoding="utf-8",
    )
    manifest = {
        "schema_version": "1.0",
        "slides": [
            {
                "slide_id": "slide_001",
                "order": 1,
                "render_mode": "native-raster",
                "background_path": "images/01-slide.png",
                "assets": [],
                "annotations": [],
            },
            {
                "slide_id": "slide_002",
                "order": 2,
                "render_mode": "source-grounded-hybrid",
                "background_path": "images/02-background.png",
                "assets": [
                    {
                        "asset_id": "asset_figure_1",
                        "source_path": "provider_input/assets/figure-1.png",
                        "source_page": 5,
                        "figure_label": "Figure 1",
                        "x": 0.07,
                        "y": 0.23,
                        "w": 0.62,
                        "h": 0.66,
                        "fit": "contain",
                        "crop": "full",
                        "preserve": ["axes", "legend", "labels", "values"],
                        "sha256": source_hash,
                    }
                ],
                "annotations": [
                    {
                        "text": "Main result",
                        "x": 0.73,
                        "y": 0.38,
                        "w": 0.21,
                        "h": 0.12,
                    }
                ],
            },
        ],
    }
    _write_json(output / "source-visual-manifest.json", manifest)
    (output / "generation-log.md").write_text(
        "images/01-slide.png: background_backend=imagegen render_mode=native-raster\n"
        "images/02-background.png: background_backend=imagegen "
        "render_mode=source-grounded-hybrid source_asset=asset_figure_1 "
        f"source_asset_sha256={source_hash} composition_backend=metaclass\n",
        encoding="utf-8",
    )
    merge_deck.make_pptx_from_manifest(output, manifest, output / "presentation.pptx")
    pdf = fitz.open()
    pdf.new_page(width=1600, height=900)
    pdf.new_page(width=1600, height=900)
    pdf.save(output / "presentation.pdf")
    pdf.close()
    artifact = PaperPresentationArtifact(
        provider="native_paper_deck",
        presentation_pdf_path="provider_output/presentation.pdf",
        source_images_dir="provider_output/rendered",
        analysis_path="provider_output/analysis.md",
        deck_brief_path="provider_output/deck-brief.md",
        outline_path="provider_output/outline.md",
        prompts_dir="provider_output/prompts",
        generation_log_path="provider_output/generation-log.md",
        source_visual_manifest_path="provider_output/source-visual-manifest.json",
        debug_pptx_path="provider_output/presentation.pptx",
    )
    return root, artifact


def _manifest(root: Path) -> dict:
    return json.loads((root / "provider_output/source-visual-manifest.json").read_text())


def test_evidence_page_requires_source_grounded_mode(tmp_path: Path) -> None:
    root, artifact = _fixture(tmp_path)
    manifest = _manifest(root)
    manifest["slides"][1].update(
        {"render_mode": "native-raster", "assets": [], "annotations": []}
    )
    _write_json(root / artifact.source_visual_manifest_path, manifest)
    outline = (root / artifact.outline_path).read_text().replace(
        "Render mode: source-grounded-hybrid", "Render mode: native-raster"
    ).replace("  - asset_id: asset_figure_1", "  - removed_asset: asset_figure_1")
    (root / artifact.outline_path).write_text(outline, encoding="utf-8")
    with pytest.raises(SourceVisualValidationError, match="Figure/Table evidence"):
        SourceVisualValidator().validate(artifact, workspace=root)


def test_manifest_rejects_unknown_asset(tmp_path: Path) -> None:
    root, artifact = _fixture(tmp_path)
    manifest = _manifest(root)
    manifest["slides"][1]["assets"][0]["asset_id"] = "asset_unknown"
    _write_json(root / artifact.source_visual_manifest_path, manifest)
    outline = (root / artifact.outline_path).read_text().replace(
        "asset_id: asset_figure_1", "asset_id: asset_unknown"
    )
    (root / artifact.outline_path).write_text(outline, encoding="utf-8")
    with pytest.raises(SourceVisualValidationError, match="unknown asset_id"):
        SourceVisualValidator().validate(artifact, workspace=root)


def test_manifest_rejects_path_traversal(tmp_path: Path) -> None:
    root, artifact = _fixture(tmp_path)
    manifest = _manifest(root)
    manifest["slides"][1]["assets"][0]["source_path"] = "provider_input/assets/../figure.png"
    _write_json(root / artifact.source_visual_manifest_path, manifest)
    with pytest.raises(SourceVisualValidationError, match="unsafe source asset path"):
        SourceVisualValidator().validate(artifact, workspace=root)


def test_figure_aspect_ratio_is_preserved(tmp_path: Path) -> None:
    root, artifact = _fixture(tmp_path)
    path = root / artifact.debug_pptx_path
    presentation = Presentation(path)
    picture = next(
        shape for shape in presentation.slides[1].shapes if shape.name == "source:asset_figure_1"
    )
    picture.width = picture.height
    presentation.save(path)
    with pytest.raises(SourceVisualValidationError, match="aspect ratio changed"):
        SourceVisualValidator().validate(artifact, workspace=root)


def test_embedded_picture_bytes_must_match_source_asset(tmp_path: Path) -> None:
    root, artifact = _fixture(tmp_path)
    path = root / artifact.debug_pptx_path
    replacement = root / "provider_input/assets/replacement.png"
    Image.new("RGB", (800, 400), "purple").save(replacement)
    presentation = Presentation(path)
    slide = presentation.slides[1]
    source = next(shape for shape in slide.shapes if shape.name == "source:asset_figure_1")
    geometry = (source.left, source.top, source.width, source.height)
    slide.shapes._spTree.remove(source._element)
    wrong = slide.shapes.add_picture(str(replacement), *geometry)
    wrong.name = "source:asset_figure_1"
    presentation.save(path)
    with pytest.raises(SourceVisualValidationError, match="embedded picture bytes"):
        SourceVisualValidator().validate(artifact, workspace=root)


def test_picture_geometry_must_match_manifest_placement(tmp_path: Path) -> None:
    root, artifact = _fixture(tmp_path)
    path = root / artifact.debug_pptx_path
    presentation = Presentation(path)
    picture = next(
        shape for shape in presentation.slides[1].shapes if shape.name == "source:asset_figure_1"
    )
    picture.left += 1000
    presentation.save(path)
    with pytest.raises(SourceVisualValidationError, match="geometry differs"):
        SourceVisualValidator().validate(artifact, workspace=root)


def test_contain_picture_must_not_be_cropped(tmp_path: Path) -> None:
    root, artifact = _fixture(tmp_path)
    path = root / artifact.debug_pptx_path
    presentation = Presentation(path)
    picture = next(
        shape for shape in presentation.slides[1].shapes if shape.name == "source:asset_figure_1"
    )
    picture.crop_left = 0.1
    presentation.save(path)
    with pytest.raises(SourceVisualValidationError, match="must not be cropped"):
        SourceVisualValidator().validate(artifact, workspace=root)


@pytest.mark.parametrize(
    ("field", "value", "message"),
    [
        ("source_path", "provider_input/assets/replacement.png", "source_path does not match"),
        ("source_page", 6, "source_page does not match"),
        ("figure_label", "Figure 2", "figure_label does not match"),
    ],
)
def test_manifest_asset_identity_must_match_paper_source(
    tmp_path: Path, field: str, value: object, message: str
) -> None:
    root, artifact = _fixture(tmp_path)
    if field == "source_path":
        Image.new("RGB", (800, 400), "purple").save(
            root / "provider_input/assets/replacement.png"
        )
    manifest = _manifest(root)
    manifest["slides"][1]["assets"][0][field] = value
    _write_json(root / artifact.source_visual_manifest_path, manifest)
    with pytest.raises(SourceVisualValidationError, match=message):
        SourceVisualValidator().validate(artifact, workspace=root)


def test_no_source_figure_is_sent_to_imagegen(tmp_path: Path) -> None:
    root, artifact = _fixture(tmp_path)
    (root / "provider_output/prompts/02-slide.md").write_text(
        "Use the supplied Figure 1 and preserve it exactly.", encoding="utf-8"
    )
    with pytest.raises(SourceVisualValidationError, match="sends source evidence"):
        SourceVisualValidator().validate(artifact, workspace=root)


def test_hybrid_prompt_accepts_equivalent_deferred_insertion_wording(tmp_path: Path) -> None:
    root, artifact = _fixture(tmp_path)
    (root / "provider_output/prompts/02-slide.md").write_text(
        "Reserve an empty evidence bay. Do not draw or insert the paper plot. "
        "The verified source will be inserted deterministically; annotations are added later.",
        encoding="utf-8",
    )
    SourceVisualValidator().validate(artifact, workspace=root)


def test_deck_fails_when_selected_figure_is_not_composited(tmp_path: Path) -> None:
    root, artifact = _fixture(tmp_path)
    path = root / artifact.debug_pptx_path
    presentation = Presentation(path)
    source = next(
        shape for shape in presentation.slides[1].shapes if shape.name == "source:asset_figure_1"
    )
    source.name = "missing-source-object"
    presentation.save(path)
    with pytest.raises(SourceVisualValidationError, match="independent source:asset_figure_1"):
        SourceVisualValidator().validate(artifact, workspace=root)


def test_pdf_pptx_outline_and_prompt_counts_must_match(tmp_path: Path) -> None:
    root, artifact = _fixture(tmp_path)
    (root / "provider_output/prompts/02-slide.md").unlink()
    with pytest.raises(SourceVisualValidationError, match="page counts must match"):
        SourceVisualValidator().validate(artifact, workspace=root)


def test_generation_log_records_each_render_mode(tmp_path: Path) -> None:
    root, artifact = _fixture(tmp_path)
    log_path = root / artifact.generation_log_path
    log_path.write_text(
        log_path.read_text(encoding="utf-8").replace(
            "render_mode=source-grounded-hybrid", ""
        ),
        encoding="utf-8",
    )
    with pytest.raises(SourceVisualValidationError, match="does not record render_mode"):
        SourceVisualValidator().validate(artifact, workspace=root)


def test_core_figure_cannot_be_silently_omitted(tmp_path: Path) -> None:
    root, artifact = _fixture(tmp_path)
    _write_json(
        root / "provider_input/paper_analysis.json",
        {
            "claims": [{"id": "claim_core", "importance": "core"}],
            "figure_candidates": [
                {
                    "id": "candidate_1",
                    "claim_ids": ["claim_core"],
                    "source_refs": [{"asset_id": "asset_figure_1", "page_no": 5}],
                }
            ],
        },
    )
    manifest = _manifest(root)
    manifest["slides"][1].update(
        {"render_mode": "native-raster", "assets": [], "annotations": []}
    )
    _write_json(root / artifact.source_visual_manifest_path, manifest)
    outline = (root / artifact.outline_path).read_text().replace(
        "Render mode: source-grounded-hybrid", "Render mode: native-raster"
    ).replace("Figure 1 from PDF page 5", "None").replace(
        "  - asset_id: asset_figure_1", "  - removed_asset: asset_figure_1"
    )
    (root / artifact.outline_path).write_text(outline, encoding="utf-8")
    log_path = root / artifact.generation_log_path
    log_path.write_text(
        log_path.read_text(encoding="utf-8").replace(
            "render_mode=source-grounded-hybrid", "render_mode=native-raster"
        ),
        encoding="utf-8",
    )
    presentation = Presentation(root / artifact.debug_pptx_path)
    for shape in list(presentation.slides[1].shapes)[1:]:
        presentation.slides[1].shapes._spTree.remove(shape._element)
    presentation.slides[1].shapes[0].name = "background:native-raster"
    presentation.save(root / artifact.debug_pptx_path)
    with pytest.raises(SourceVisualValidationError, match="core Figure/Table"):
        SourceVisualValidator().validate(artifact, workspace=root)
