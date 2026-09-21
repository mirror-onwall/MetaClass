import json
from pathlib import Path

import fitz
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation

from metaclass.main import create_app
from metaclass.modules.paper_workflow.final_paper_deck_page_analyzer import TesseractOCRProvider
from metaclass.modules.paper_workflow.native_runner import NativePaperDeckWorkflowRunner
from metaclass.modules.paper_workflow.orchestrator import PaperWorkflowOrchestrator
from metaclass.modules.paper_workflow.providers.base import PaperProviderContext
from metaclass.modules.paper_workflow.schemas import (
    PaperAnalysis,
    PaperClaim,
    PaperPresentationArtifact,
    PaperSourceBundle,
    SourceReference,
)


class RunnerLLM:
    name = "runner-test"
    model = "runner-test"

    def __init__(self) -> None:
        self.needs_repair = True
        self.vision_calls = 0
        self.mapping_calls = 0
        self.narration_calls = 0

    def complete_image_json(self, prompt, image_path, *, temperature=0.2):
        self.vision_calls += 1
        order = int(Path(image_path).name[:2])
        return json.dumps(
            {
                "visible_title": "Method overview" if order == 1 else "Supported conclusion",
                "visible_text": (
                    ["Unsupported draft result 99.9%"]
                    if order == 2 and self.needs_repair
                    else ["Supported method"]
                ),
                "visual_summary": "A method diagram grounded in the paper.",
                "figure_labels": [],
                "quantitative_mentions": [],
                "formula_mentions": [],
                "page_type": "method" if order == 1 else "conclusion",
                "detected_warnings": [],
            }
        )

    def complete_json(self, messages, *, temperature=0.2):
        marker = messages[0].content
        payload = json.loads(messages[1].content)
        if "GROUNDED_FINAL_SLIDE_ALIGNMENT_V1" in marker:
            self.mapping_calls += 1
            return json.dumps(
                {
                    "slides": [
                        {
                            "slide_id": item["slide_id"],
                            "claim_ids": ["claim_method"],
                            "result_ids": [],
                            "asset_ids": [],
                            "evidence_strength": "direct",
                        }
                        for item in payload["slides"]
                    ]
                }
            )
        if "GROUNDED_PAPER_CLASSROOM_NARRATION_V2" in marker:
            self.narration_calls += 1
            slides = []
            for index, item in enumerate(payload["slides"]):
                slide = item["slide"]
                slides.append(
                    {
                        "slide_id": slide["slide_id"],
                        "speaker_script": (
                            "这一页围绕论文明确提出的方法主张展开。页面负责呈现作者的论证意图，"
                            "而讲解补充原文证据与适用边界，使观众能够理解这项方法为什么被提出、"
                            "它回应了什么问题，以及结论应当限制在论文实际支持的范围之内。"
                        ),
                        "transition": "下一页继续沿着论文的证据顺序推进。" if index == 0 else "",
                        "used_claim_ids": ["claim_method"],
                        "used_source_refs": slide["source_refs"],
                        "used_asset_ids": [],
                        "validation_status": "pending",
                    }
                )
            return json.dumps({"slides": slides}, ensure_ascii=False)
        if "SHARED_INTERACTION_QUESTION_GENERATOR_V1" in marker:
            return json.dumps(
                {
                    "items": [
                        {
                            "id": f"llm_interaction_{index}",
                            "blueprint_id": node["blueprint"]["id"],
                            "slide_id": node["blueprint"]["slide_id"],
                            "question": f"如何依据当前证据理解机制与边界？问题批次{chr(65 + index)}",
                            "answer": "当前证据支持页面所述关系，但解释必须保留论文给出的条件与边界。",
                            "source_refs": node["blueprint"]["source_refs"],
                            "claim_ids": node["blueprint"]["claim_ids"],
                            "result_ids": node["blueprint"]["result_ids"],
                            "asset_ids": node["blueprint"]["asset_ids"],
                        }
                        for index, node in enumerate(payload["nodes"])
                    ]
                },
                ensure_ascii=False,
            )
        raise AssertionError(marker)


class AnalysisProvider:
    def run_analysis(self, context):
        source = json.loads(
            (context.workspace / "source/paper_source.json").read_text(encoding="utf-8")
        )
        block = source["blocks"][0]
        analysis = PaperAnalysis(
            paper={"title": "Native Runner Paper"},
            paper_type="methods",
            central_question="How does the method address the problem?",
            knowledge_gap="The existing approach leaves a documented gap.",
            main_claim="The paper proposes a supported method.",
            method_summary={"overview": "A supported method."},
            claims=[
                PaperClaim(
                    id="claim_method",
                    statement="The paper proposes a supported method.",
                    importance="core",
                    confidence=0.9,
                    source_refs=[
                        SourceReference(
                            page_no=block["page_no"],
                            block_id=block["id"],
                        )
                    ],
                )
            ],
        )
        output = context.workspace / "stages/01_analysis/output"
        output.mkdir(parents=True, exist_ok=True)
        (output / "paper_analysis.json").write_text(
            analysis.model_dump_json(indent=2), encoding="utf-8"
        )
        return analysis


class DeckProvider:
    skill_directory = Path(".")

    def __init__(self, llm: RunnerLLM) -> None:
        self.llm = llm
        self.repair_calls = 0

    @staticmethod
    def _directory_hash(_path):
        return "deck-skill-test"

    @staticmethod
    def _resolve_skill_directory():
        return Path(".")

    def run(self, context):
        output = context.workspace / "provider_output"
        if (output / "presentation.pdf").is_file():
            return self._artifact()
        provider_input = context.workspace / "provider_input"
        provider_input.mkdir(parents=True, exist_ok=True)
        (provider_input / "paper_source.json").write_text(
            (context.workspace / "source" / context.source_bundle.paper_source_path).read_text(
                encoding="utf-8"
            ),
            encoding="utf-8",
        )
        analysis_source = context.workspace / "stages/01_analysis/output/paper_analysis.json"
        (provider_input / "paper_analysis.json").write_text(
            analysis_source.read_text(encoding="utf-8"), encoding="utf-8"
        )
        prompts = output / "prompts"
        images = output / "images"
        rendered = output / "rendered"
        prompts.mkdir(parents=True, exist_ok=True)
        images.mkdir(parents=True, exist_ok=True)
        rendered.mkdir(parents=True, exist_ok=True)
        (output / "analysis.md").write_text("# Native analysis\n", encoding="utf-8")
        (output / "deck-brief.md").write_text(
            "# Deck Brief\n\n- style_preset: `journal-minimal`\n- language: zh-CN\n",
            encoding="utf-8",
        )
        outline = ["# Outline\n"]
        for order in range(1, 3):
            name = f"{order:02d}-slide"
            (prompts / f"{name}.md").write_text("Use paper evidence.\n", encoding="utf-8")
            image = images / f"{name}.png"
            Image.new("RGB", (1600, 900), (40 * order, 80, 120)).save(image)
            Image.new("RGB", (1600, 900), (40 * order, 80, 120)).save(
                rendered / f"{name}.png"
            )
            outline.append(
                f"""## {order:02d}. Slide {order}
- Role: method
- Message: Explain the supported method.
- Render mode: native-raster
- Visual: Method diagram.
- Text: Supported method
- Evidence: Paper page 1
- Source visual: None
"""
            )
        (output / "outline.md").write_text("\n".join(outline), encoding="utf-8")
        (output / "generation-log.md").write_text(
            "images/01-slide.png: backend=imagegen render_mode=native-raster\n"
            "images/02-slide.png: backend=imagegen render_mode=native-raster\n",
            encoding="utf-8",
        )
        (output / "source-visual-manifest.json").write_text(
            json.dumps(
                {
                    "schema_version": "1.0",
                    "slides": [
                        {
                            "slide_id": f"slide_{index:03d}",
                            "order": index,
                            "render_mode": "native-raster",
                            "background_path": f"images/{index:02d}-slide.png",
                            "assets": [],
                            "annotations": [],
                        }
                        for index in (1, 2)
                    ],
                }
            ),
            encoding="utf-8",
        )
        self._compose_pdf(output)
        return self._artifact()

    @staticmethod
    def _artifact() -> PaperPresentationArtifact:
        return PaperPresentationArtifact(
            provider="native_paper_deck",
            presentation_pdf_path="provider_output/presentation.pdf",
            source_images_dir="provider_output/rendered",
            analysis_path="provider_output/analysis.md",
            deck_brief_path="provider_output/deck-brief.md",
            outline_path="provider_output/outline.md",
            prompts_dir="provider_output/prompts",
            generation_log_path="provider_output/generation-log.md",
            source_visual_manifest_path="provider_output/source-visual-manifest.json",
            debug_pptx_path="provider_output/presentation.pptx",
        )

    def repair_slides(self, context, *, directives):
        assert [item["slide_id"] for item in directives] == ["paper_deck_slide_002"]
        self.repair_calls += 1
        output = context.workspace / "provider_output"
        Image.new("RGB", (1600, 900), (20, 160, 80)).save(output / "images/02-slide.png")
        (output / "prompts/02-slide.md").write_text(
            "Repaired page: omit unsupported draft number.\n", encoding="utf-8"
        )
        with (output / "generation-log.md").open("a", encoding="utf-8") as stream:
            stream.write("targeted repair: paper_deck_slide_002 backend=imagegen\n")
        self.llm.needs_repair = False
        self._compose_pdf(output)
        return self._artifact()

    @staticmethod
    def _compose_pdf(output: Path) -> None:
        document = fitz.open()
        for order in range(1, 3):
            image = output / f"images/{order:02d}-slide.png"
            page = document.new_page(width=1600, height=900)
            page.insert_image(page.rect, filename=str(image))
        temporary = output / ".presentation.repair.pdf"
        document.save(temporary)
        document.close()
        temporary.replace(output / "presentation.pdf")
        rendered = output / "rendered"
        rendered.mkdir(exist_ok=True)
        with fitz.open(output / "presentation.pdf") as final_pdf:
            for order, page in enumerate(final_pdf, start=1):
                page.get_pixmap(alpha=False).save(rendered / f"{order:02d}-slide.png")
        presentation = Presentation()
        for order in range(1, 3):
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            picture = slide.shapes.add_picture(
                str(output / f"images/{order:02d}-slide.png"),
                0,
                0,
                width=presentation.slide_width,
                height=presentation.slide_height,
            )
            picture.name = "background:native-raster"
        presentation.save(output / "presentation.pptx")


def _paper_bytes() -> bytes:
    document = fitz.open()
    document.new_page().insert_text(
        (72, 72),
        "Native Runner Paper. The paper proposes a supported method to address the gap.",
    )
    payload = document.tobytes()
    document.close()
    return payload


def test_native_runner_registers_pdf_classroom_interactions_and_live_index(
    tmp_path: Path, monkeypatch
) -> None:
    monkeypatch.setattr(TesseractOCRProvider, "extract_text", lambda _self, _path: "")
    app = create_app(tmp_path)
    services = app.state.services
    llm = RunnerLLM()
    deck_provider = DeckProvider(llm)
    runner = NativePaperDeckWorkflowRunner(
        data_dir=tmp_path,
        materials=services.materials,
        contents=services.contents,
        presentations=services.presentations,
        presentation_repository=services.presentations.repository,
        question_banks=services.question_banks.repository,
        live_questions=services.live_questions,
        llm=llm,
        vision=llm,
        analysis_provider=AnalysisProvider(),
        deck_provider=deck_provider,
    )
    services.paper_workflows.orchestrator = PaperWorkflowOrchestrator(tmp_path, [runner])

    with TestClient(app) as client:
        material = client.post(
            "/api/v1/materials",
            files={"file": ("paper.pdf", _paper_bytes(), "application/pdf")},
        ).json()
        created = client.post(
            "/api/v1/paper-workflows",
            json={
                "material_id": material["id"],
                "strategy": "native_paper_deck",
                "duration_minutes": 15,
                "audience": "graduate students",
            },
        ).json()
        completed = client.post(f"/api/v1/paper-workflows/{created['id']}/run").json()
        assert completed["status"] == "succeeded", completed.get("error")
        course = client.post(
            f"/api/v1/paper-workflows/{created['id']}/create-paper-deck-course"
        ).json()

    plan = services.presentations.get_plan(course["presentation_plan_id"])
    assert plan.mode == "paper_deck"
    assert plan.source_material_id == course["derived_material_id"]
    assert len(plan.slides) == 2
    assert services.question_banks.repository.list_for_plan(plan.id)
    assert services.live_questions.has_index(plan.id)
    assert deck_provider.repair_calls == 1
    assert (
        tmp_path
        / "runtime/paper_workflows"
        / created["id"]
        / "checkpoints/classroom.json"
    ).is_file()

    calls_after_first_run = (llm.vision_calls, llm.mapping_calls, llm.narration_calls)
    workflow = services.paper_workflows
    workspace = tmp_path / "runtime/paper_workflows" / created["id"]
    request = workflow.repository.get_request(created["id"])
    checkpoint = workflow.checkpoints.load(created["id"])
    assert request is not None and checkpoint is not None
    source_bundle = PaperSourceBundle.model_validate_json(
        (workspace / "source/paper_source.json").read_text(encoding="utf-8")
    )
    runner.run(
        PaperProviderContext(
            job_id=created["id"],
            workspace=workspace,
            request=request,
            source_bundle=source_bundle,
            checkpoint=checkpoint,
            report_progress=lambda _progress, _stage: None,
            is_pause_requested=lambda: False,
            persist_checkpoint=lambda _checkpoint: None,
        )
    )

    assert (llm.vision_calls, llm.mapping_calls, llm.narration_calls) == calls_after_first_run
    assert deck_provider.repair_calls == 1
