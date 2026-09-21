import hashlib
import json
import os
import zipfile
from pathlib import Path
from unittest.mock import Mock

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from metaclass.modules.presentation.codex_provider import (
    CodexGenerationError,
    CodexPPTProvider,
)
from metaclass.modules.presentation.providers import (
    FallbackPPTProvider,
    PresentonPPTProvider,
    validate_deck_against_plan,
)
from metaclass.modules.presentation.schemas import (
    PPTArtifact,
    PresentationPlan,
    SlideElement,
    SlideElementStyle,
    SlidePlan,
)
from metaclass.modules.presentation.skill_adapter import PPTSkillAdapter
from metaclass.modules.presentation.themes import PRESENTATION_THEMES, get_presentation_theme


def make_plan() -> PresentationPlan:
    return PresentationPlan(
        id="plan_codex",
        content_id="content_codex",
        title="空间分析",
        slides=[
            SlidePlan(
                id="slide_001",
                order=1,
                source_section_ids=["section_001"],
                title="空间自相关保留地理位置关系",
                key_points=[
                    "传统相关分析适用于非空间观测。",
                    "空间自相关同时考虑观测值与地理位置。",
                ],
                speaker_script="解释两类相关分析之间的区别。",
                suggested_visual="使用左右对照和空间点阵。",
                visual_payload=["空间点阵关系示意"],
            )
        ],
    )


def write_deck(path: Path, values: list[str]) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    for index, value in enumerate(values):
        textbox = slide.shapes.add_textbox(
            Inches(0.8), Inches(0.5 + index), Inches(11.5), Inches(0.7)
        )
        textbox.text_frame.text = value
    deck.save(path)


def test_skill_font_role_is_compiled_into_the_pptx_typeface(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    plan = make_plan()
    element = SlideElement(
        type="text",
        x=0.1,
        y=0.1,
        w=0.8,
        h=0.2,
        text=plan.slides[0].title,
        style=SlideElementStyle(
            font_size=30,
            font_role="serif",
            bold=True,
            color="17324D",
            line_width=0,
        ),
    )
    rendered_plan = plan.model_copy(
        update={"slides": [plan.slides[0].model_copy(update={"elements": [element]})]}
    )
    monkeypatch.setattr(
        "metaclass.modules.presentation.skill_adapter.platform.system",
        lambda: "Windows",
    )
    destination = tmp_path / "skill-font-role.pptx"

    PPTSkillAdapter().render_declarative_pptx(rendered_plan, destination)

    with zipfile.ZipFile(destination) as archive:
        slide_xml = archive.read("ppt/slides/slide1.xml").decode("utf-8")
    assert 'typeface="SimSun"' in slide_xml


def test_exact_plan_copy_preserves_an_embedded_line_break(tmp_path: Path) -> None:
    plan = make_plan()
    point = "first exact line\nsecond exact line"
    slide = plan.slides[0].model_copy(
        update={
            "key_points": [point],
            "elements": [
                SlideElement(
                    type="text",
                    x=0.08,
                    y=0.08,
                    w=0.84,
                    h=0.18,
                    text=plan.slides[0].title,
                    style=SlideElementStyle(font_size=30, bold=True),
                ),
                SlideElement(
                    type="text",
                    x=0.08,
                    y=0.35,
                    w=0.84,
                    h=0.3,
                    text=point,
                    style=SlideElementStyle(font_size=20),
                ),
            ],
        }
    )
    rendered_plan = plan.model_copy(update={"slides": [slide]})
    destination = tmp_path / "exact-line-break.pptx"

    PPTSkillAdapter().render_declarative_pptx(rendered_plan, destination)

    validation = validate_deck_against_plan(
        rendered_plan,
        destination,
        provider_name="Codex",
    )
    assert validation["status"] == "passed"


def test_final_contract_allows_only_materialized_visual_placeholder_text(
    tmp_path: Path,
) -> None:
    plan = make_plan()
    label = CodexPPTProvider._visual_placeholder_label(
        plan.slides[0],
        "visual_payload.0",
    )
    placeholder = SlideElement(
        type="text",
        contract_role="visual_placeholder",
        x=0.1,
        y=0.35,
        w=0.4,
        h=0.25,
        text=label,
        style=SlideElementStyle(
            font_size=16,
            color="17324D",
            fill="FFFFFF",
            line_color="6F8299",
            line_width=1.5,
            align="center",
            valign="middle",
        ),
    )
    materialized_plan = plan.model_copy(
        update={"slides": [plan.slides[0].model_copy(update={"elements": [placeholder]})]}
    )
    destination = tmp_path / "authorized-placeholder.pptx"
    write_deck(
        destination,
        [plan.slides[0].title, *plan.slides[0].key_points, label],
    )

    validation = validate_deck_against_plan(
        plan,
        destination,
        provider_name="Codex",
        materialized_plan=materialized_plan,
    )

    assert validation["status"] == "passed"
    assert validation["checked_auxiliary_text_items"] == 1
    with pytest.raises(RuntimeError, match="extra or rewritten"):
        validate_deck_against_plan(
            plan,
            destination,
            provider_name="Codex",
        )


def test_paper_deck_rejects_fused_body_visual_islands() -> None:
    plan = make_plan()
    body_elements = [
        SlideElement(
            type="text",
            contract_role="plan_copy",
            x=0.1,
            y=0.3,
            w=0.36,
            h=0.2,
            text=plan.slides[0].key_points[0],
        ),
        SlideElement(
            type="text",
            contract_role="plan_copy",
            x=0.45,
            y=0.31,
            w=0.36,
            h=0.2,
            text=plan.slides[0].key_points[1],
        ),
    ]

    with pytest.raises(ValueError, match="fused or too tightly clustered"):
        CodexPPTProvider._validate_distributed_module_spacing(
            body_elements,
            slide=plan.slides[0],
        )


def make_valid_independent_point_objects(plan: PresentationPlan) -> list[SlideElement]:
    elements: list[SlideElement] = []
    for index, (frame_x, text_x) in enumerate(((0.08, 0.1), (0.56, 0.58))):
        content_ref = f"key_points.{index}"
        elements.extend(
            [
                SlideElement(
                    type="shape",
                    contract_role="visual_module",
                    object_id=f"point-{index}-frame",
                    semantic_ref=content_ref,
                    x=frame_x,
                    y=0.28,
                    w=0.36,
                    h=0.26,
                    z=3,
                    shape="rounded_rectangle",
                ),
                SlideElement(
                    type="text",
                    contract_role="plan_copy",
                    object_id=f"copy-key-points-{index}",
                    semantic_ref=content_ref,
                    x=text_x,
                    y=0.31,
                    w=0.32,
                    h=0.2,
                    z=20 + index,
                    text=plan.slides[0].key_points[index],
                ),
            ]
        )
    return elements


def test_paper_deck_requires_an_independent_editorial_anchor_for_every_plan_point() -> None:
    plan = make_plan()
    elements = make_valid_independent_point_objects(plan)
    elements = [element for element in elements if element.object_id != "point-1-frame"]

    with pytest.raises(ValueError, match="no independent editorial anchor for key_points.1"):
        CodexPPTProvider._validate_distributed_module_spacing(
            elements,
            slide=plan.slides[0],
        )


def test_paper_deck_accepts_an_adjacent_rule_instead_of_a_rounded_point_card() -> None:
    plan = make_plan()
    elements = [
        element
        for element in make_valid_independent_point_objects(plan)
        if element.object_id != "point-0-frame"
    ]
    elements.append(
        SlideElement(
            type="line",
            contract_role="visual_module",
            object_id="point-0-editorial-rule",
            semantic_ref="key_points.0",
            x=0.075,
            y=0.31,
            w=0,
            h=0.2,
            z=3,
        )
    )

    CodexPPTProvider._validate_distributed_module_spacing(
        elements,
        slide=plan.slides[0],
    )


def test_paper_deck_accepts_a_skill_side_band_with_editorial_gutter() -> None:
    plan = make_plan()
    elements = [
        element
        for element in make_valid_independent_point_objects(plan)
        if element.object_id != "point-0-frame"
    ]
    elements.append(
        SlideElement(
            type="shape",
            contract_role="visual_module",
            object_id="point-0-side-band",
            semantic_ref="key_points.0",
            x=0.06,
            y=0.31,
            w=0.008,
            h=0.2,
            z=3,
            shape="rectangle",
        )
    )

    # The 0.032 normalized gutter mirrors the failed slide_016 response. It is
    # deliberate Paper Deck whitespace, not a detached or missing anchor.
    CodexPPTProvider._validate_distributed_module_spacing(
        elements,
        slide=plan.slides[0],
    )


def test_paper_deck_rejects_a_shape_that_only_partly_covers_plan_text() -> None:
    plan = make_plan()
    elements = make_valid_independent_point_objects(plan)
    elements.append(
        SlideElement(
            type="shape",
            contract_role="visual_module",
            object_id="partial-accent",
            semantic_ref="decoration",
            x=0.39,
            y=0.34,
            w=0.08,
            h=0.12,
            z=4,
            shape="oval",
        )
    )

    with pytest.raises(ValueError, match="partially overlaps or occludes Plan text"):
        CodexPPTProvider._validate_distributed_module_spacing(
            elements,
            slide=plan.slides[0],
        )


def test_paper_deck_rejects_a_connector_crossing_plan_text() -> None:
    plan = make_plan()
    elements = make_valid_independent_point_objects(plan)
    elements.append(
        SlideElement(
            type="line",
            contract_role="visual_module",
            object_id="crossing-rule",
            semantic_ref="decoration",
            x=0.04,
            y=0.4,
            w=0.44,
            h=0,
            z=4,
        )
    )

    with pytest.raises(ValueError, match="connector crosses a Plan text writing area"):
        CodexPPTProvider._validate_distributed_module_spacing(
            elements,
            slide=plan.slides[0],
        )


def test_paper_deck_local_repair_moves_a_crossing_connector_to_a_safe_anchor() -> None:
    plan = make_plan()
    elements = make_valid_independent_point_objects(plan)
    elements.append(
        SlideElement(
            type="line",
            contract_role="visual_module",
            object_id="crossing-rule",
            semantic_ref="key_points.0",
            x=0.04,
            y=0.4,
            w=0.44,
            h=0,
            z=4,
        )
    )
    raw_elements = [element.model_dump(mode="json") for element in elements]

    assert CodexPPTProvider._repair_paper_deck_connector(
        raw_elements=raw_elements,
        parsed=elements,
        object_id="crossing-rule",
    )

    repaired = [SlideElement.model_validate(element) for element in raw_elements]
    CodexPPTProvider._validate_distributed_module_spacing(
        repaired,
        slide=plan.slides[0],
    )


def test_paper_deck_local_repair_separates_semantic_groups_as_whole_objects() -> None:
    plan = make_plan()
    elements = make_valid_independent_point_objects(plan)
    elements[2] = elements[2].model_copy(update={"x": 0.4})
    elements[3] = elements[3].model_copy(update={"x": 0.42})
    raw_elements = [element.model_dump(mode="json") for element in elements]

    assert CodexPPTProvider._shift_paper_deck_semantic_group(
        raw_elements=raw_elements,
        parsed=elements,
        moving=elements[3],
        fixed=elements[1],
        required_gap=CodexPPTProvider.PAPER_DECK_MIN_MODULE_GAP,
    )

    repaired = [SlideElement.model_validate(element) for element in raw_elements]
    assert repaired[2].x - elements[2].x == pytest.approx(
        repaired[3].x - elements[3].x
    )
    CodexPPTProvider._validate_distributed_module_spacing(
        repaired,
        slide=plan.slides[0],
    )


def test_declarative_preview_renders_chevron_as_a_chevron() -> None:
    plan = make_plan()
    chevron = SlideElement(
        type="shape",
        x=0.2,
        y=0.3,
        w=0.3,
        h=0.2,
        shape="chevron",
        style=SlideElementStyle(fill="FF0000", line_width=0),
    )
    slide = plan.slides[0].model_copy(update={"background": "FFFFFF", "elements": [chevron]})

    preview = PPTSkillAdapter._render_scene_preview(slide)
    x0 = round(chevron.x * preview.width)
    middle_x = round((chevron.x + chevron.w / 2) * preview.width)
    middle_y = round((chevron.y + chevron.h / 2) * preview.height)

    assert preview.getpixel((x0 + 4, middle_y)) == (255, 255, 255)
    assert preview.getpixel((middle_x, middle_y)) == (255, 0, 0)


def make_valid_content_visual_column(
    plan: PresentationPlan,
) -> tuple[SlidePlan, list[SlideElement]]:
    slide = plan.slides[0].model_copy(update={"order": 2})
    elements = [
        SlideElement(
            type="text",
            contract_role="plan_copy",
            object_id="copy-title",
            semantic_ref="title",
            x=0.08,
            y=0.06,
            w=0.84,
            h=0.16,
            z=20,
            text=slide.title,
        )
    ]
    for index, y in enumerate((0.28, 0.58)):
        content_ref = f"key_points.{index}"
        elements.extend(
            [
                SlideElement(
                    type="shape",
                    contract_role="visual_module",
                    object_id=f"point-{index}-frame",
                    semantic_ref=content_ref,
                    x=0.08,
                    y=y,
                    w=0.44,
                    h=0.22,
                    z=3,
                    shape="rounded_rectangle",
                ),
                SlideElement(
                    type="text",
                    contract_role="plan_copy",
                    object_id=f"copy-key-points-{index}",
                    semantic_ref=content_ref,
                    x=0.1,
                    y=y + 0.03,
                    w=0.4,
                    h=0.16,
                    z=21 + index,
                    text=slide.key_points[index],
                ),
            ]
        )
    elements.append(
        SlideElement(
            type="text",
            contract_role="visual_placeholder",
            object_id="placeholder-1",
            semantic_ref="suggested_visual",
            x=0.58,
            y=0.3,
            w=0.34,
            h=0.52,
            z=45,
            text="（此处建议插入：相关图片）",
            style=SlideElementStyle(
                color="17324D",
                fill="FFFFFF",
                line_color="6F8299",
                line_width=1.5,
            ),
        )
    )
    return slide, elements


def test_paper_deck_accepts_a_dominant_right_visual_column() -> None:
    slide, elements = make_valid_content_visual_column(make_plan())

    CodexPPTProvider._validate_distributed_module_spacing(elements, slide=slide)


def test_paper_deck_rejects_a_small_visual_placeholder() -> None:
    slide, elements = make_valid_content_visual_column(make_plan())
    elements[-1] = elements[-1].model_copy(update={"x": 0.66, "w": 0.26, "h": 0.3})

    with pytest.raises(ValueError, match="visual region is too small"):
        CodexPPTProvider._validate_distributed_module_spacing(elements, slide=slide)


def test_paper_deck_accepts_a_centered_visual_placeholder() -> None:
    slide, elements = make_valid_content_visual_column(make_plan())
    elements[1] = elements[1].model_copy(update={"x": 0.08, "y": 0.22, "w": 0.84, "h": 0.1})
    elements[2] = elements[2].model_copy(update={"x": 0.1, "y": 0.24, "w": 0.8, "h": 0.06})
    elements[3] = elements[3].model_copy(update={"x": 0.08, "y": 0.76, "w": 0.84, "h": 0.12})
    elements[4] = elements[4].model_copy(update={"x": 0.1, "y": 0.78, "w": 0.8, "h": 0.08})
    elements[-1] = elements[-1].model_copy(update={"x": 0.33, "y": 0.37, "w": 0.34, "h": 0.32})

    CodexPPTProvider._validate_distributed_module_spacing(elements, slide=slide)


def test_paper_deck_rejects_asset_and_placeholder_on_the_same_page() -> None:
    slide, elements = make_valid_content_visual_column(make_plan())
    elements.append(
        SlideElement(
            type="image",
            contract_role="visual_asset",
            object_id="local-illustration",
            semantic_ref="visual_payload.0",
            x=0.58,
            y=0.3,
            w=0.34,
            h=0.52,
            z=8,
            image_path="generated_visuals/illustration.png",
        )
    )

    with pytest.raises(ValueError, match="cannot combine local illustrations and a placeholder"):
        CodexPPTProvider._validate_distributed_module_spacing(elements, slide=slide)


def test_paper_deck_visual_region_requires_independent_object_clearance() -> None:
    slide, elements = make_valid_content_visual_column(make_plan())
    elements[-1] = elements[-1].model_copy(update={"x": 0.53, "w": 0.37})

    with pytest.raises(ValueError, match="less than 0.015 clearance"):
        CodexPPTProvider._validate_distributed_module_spacing(elements, slide=slide)

    elements[-1] = elements[-1].model_copy(update={"x": 0.54, "w": 0.36})
    CodexPPTProvider._validate_distributed_module_spacing(elements, slide=slide)


def test_paper_deck_visual_column_allows_a_slim_anchor_at_module_clearance() -> None:
    plan = make_plan()
    slide = plan.slides[0].model_copy(update={"order": 2})
    elements = [
        SlideElement(
            type="image",
            contract_role="visual_asset",
            object_id="scientific-illustration",
            semantic_ref="suggested_visual",
            x=0.06,
            y=0.23,
            w=0.4225,
            h=0.64,
            z=1,
            image_path="generated_visuals/illustration.png",
        ),
        SlideElement(
            type="text",
            contract_role="plan_copy",
            object_id="copy-title",
            semantic_ref="title",
            x=0.52,
            y=0.06,
            w=0.41,
            h=0.14,
            z=20,
            text=slide.title,
        ),
    ]
    for index, y in enumerate((0.28, 0.58)):
        content_ref = f"key_points.{index}"
        elements.extend(
            [
                SlideElement(
                    type="shape",
                    contract_role="visual_module",
                    object_id=f"point-{index}-band",
                    semantic_ref=content_ref,
                    x=0.51,
                    y=y,
                    w=0.012,
                    h=0.18,
                    z=3,
                    shape="rectangle",
                ),
                SlideElement(
                    type="text",
                    contract_role="plan_copy",
                    object_id=f"copy-key-points-{index}",
                    semantic_ref=content_ref,
                    x=0.54,
                    y=y,
                    w=0.39,
                    h=0.18,
                    z=21 + index,
                    text=slide.key_points[index],
                ),
            ]
        )

    # Mirrors slide_015: the picture-to-band gap is 0.0275, while the picture-to-copy
    # gap remains 0.0575. The slim band may bridge the fields without overlapping them.
    CodexPPTProvider._validate_distributed_module_spacing(elements, slide=slide)


def make_design_style(
    *,
    color: str = "17324D",
    fill: str = "DCEBFA",
    bold: bool = False,
    size: int = 20,
) -> dict:
    return {
        "font_size": size,
        "bold": bold,
        "color": color,
        "fill": fill,
        "line_color": "2E75B6",
        "line_width": 1,
        "align": "left",
        "valign": "middle",
        "opacity": 100,
    }


def make_structured_result(
    plan: PresentationPlan,
    *,
    image_paths: list[str] | None = None,
) -> dict:
    elements = [
        {
            "type": "text",
            "x": 0.08,
            "y": 0.05,
            "w": 0.84,
            "h": 0.16,
            "z": 10,
            "text": plan.slides[0].title,
            "image_path": "",
            "shape": "rectangle",
            "style": make_design_style(
                color="FFFFFF",
                fill="12365A",
                bold=True,
                size=32,
            ),
        },
        {
            "type": "text",
            "x": 0.08,
            "y": 0.28,
            "w": 0.38,
            "h": 0.24,
            "z": 11,
            "text": plan.slides[0].key_points[0],
            "image_path": "",
            "shape": "rectangle",
            "style": make_design_style(),
        },
        {
            "type": "text",
            "x": 0.08,
            "y": 0.57,
            "w": 0.38,
            "h": 0.24,
            "z": 12,
            "text": plan.slides[0].key_points[1],
            "image_path": "",
            "shape": "rectangle",
            "style": make_design_style(),
        },
    ]
    if image_paths:
        for index, image_path in enumerate(image_paths):
            elements.append(
                {
                    "type": "image",
                    "x": 0.54,
                    "y": 0.28 + index * 0.29,
                    "w": 0.36,
                    "h": 0.24,
                    "z": 5,
                    "text": "",
                    "image_path": image_path,
                    "shape": "rectangle",
                    "style": make_design_style(),
                }
            )
    else:
        elements.append(
            {
                "type": "shape",
                "x": 0.54,
                "y": 0.28,
                "w": 0.36,
                "h": 0.48,
                "z": 5,
                "text": "",
                "image_path": "",
                "shape": "rounded_rectangle",
                "style": make_design_style(),
            }
        )
    return {
        "status": "completed",
        "summary": "structured design",
        "slides": [
            {
                "slide_id": plan.slides[0].id,
                "background": "12365A",
                "elements": elements,
            }
        ],
    }


def make_paper_deck_result(
    plan: PresentationPlan,
    *,
    image_path: str | None = "generated_visuals/illustration.png",
    style_signature: str = "journal-minimal:academic_blue:v3",
    text_blocks: list[dict] | None = None,
    visual_placeholders: list[dict] | None = None,
) -> dict:
    """Return independent Skill-directed objects plus exact-copy block geometry."""

    is_cover = plan.slides[0].order == 1
    text_color = "FFFFFF" if is_cover else "17324D"

    if text_blocks is None:
        if is_cover:
            point_blocks = [
                {
                    "content_ref": "key_points.0",
                    "x": 0.54,
                    "y": 0.31,
                    "w": 0.37,
                    "h": 0.2,
                    "font_size": 20,
                    "min_font_size": 16,
                    "font_role": "serif",
                    "bold": False,
                    "color": text_color,
                    "align": "left",
                    "valign": "middle",
                    "max_lines": 3,
                },
                {
                    "content_ref": "key_points.1",
                    "x": 0.09,
                    "y": 0.63,
                    "w": 0.44,
                    "h": 0.2,
                    "font_size": 19,
                    "min_font_size": 15,
                    "font_role": "sans",
                    "bold": True,
                    "color": text_color,
                    "align": "right",
                    "valign": "bottom",
                    "max_lines": 3,
                },
            ]
        else:
            point_blocks = [
                {
                    "content_ref": "key_points.0",
                    "x": 0.56,
                    "y": 0.31,
                    "w": 0.35,
                    "h": 0.2,
                    "font_size": 20,
                    "min_font_size": 16,
                    "font_role": "serif",
                    "bold": False,
                    "color": text_color,
                    "align": "left",
                    "valign": "middle",
                    "max_lines": 3,
                },
                {
                    "content_ref": "key_points.1",
                    "x": 0.56,
                    "y": 0.61,
                    "w": 0.35,
                    "h": 0.18,
                    "font_size": 19,
                    "min_font_size": 15,
                    "font_role": "sans",
                    "bold": True,
                    "color": text_color,
                    "align": "left",
                    "valign": "middle",
                    "max_lines": 3,
                },
            ]
        if len(plan.slides[0].key_points) > len(point_blocks):
            point_blocks.extend(
                {
                    "content_ref": f"key_points.{index}",
                    "x": 0.09 + 0.27 * ((index - 2) % 3),
                    "y": 0.38 + 0.2 * ((index - 2) // 3),
                    "w": 0.24,
                    "h": 0.16,
                    "font_size": 18,
                    "min_font_size": 14,
                    "font_role": "sans",
                    "bold": False,
                    "color": text_color,
                    "align": "left",
                    "valign": "top",
                    "max_lines": 4,
                }
                for index in range(2, len(plan.slides[0].key_points))
            )
        text_blocks = [
            {
                "content_ref": "title",
                "x": 0.09,
                "y": 0.06,
                "w": 0.82,
                "h": 0.18,
                "font_size": 32,
                "min_font_size": 24,
                "font_role": "display",
                "bold": True,
                "color": text_color,
                "align": "center",
                "valign": "middle",
                "max_lines": 2,
            },
            *point_blocks[: len(plan.slides[0].key_points)],
        ]
    if visual_placeholders is None:
        visual_placeholders = []
    if visual_placeholders:
        image_path = None

    first_frame = (0.52, 0.28, 0.41, 0.26) if is_cover else (0.54, 0.28, 0.39, 0.26)
    second_frame = (0.07, 0.6, 0.48, 0.25) if is_cover else (0.54, 0.58, 0.39, 0.25)
    modules = [
        {
            "object_id": "title-rule",
            "content_ref": "title",
            "type": "line",
            "shape": "rectangle",
            "x": 0.09,
            "y": 0.25,
            "w": 0.82,
            "h": 0,
            "z": 2,
            "style_token": "accent_rule",
        },
        {
            "object_id": "point-0-frame",
            "content_ref": "key_points.0",
            "type": "shape",
            "shape": "rounded_rectangle",
            "x": first_frame[0],
            "y": first_frame[1],
            "w": first_frame[2],
            "h": first_frame[3],
            "z": 3,
            "style_token": "outline_panel" if is_cover else "soft_panel",
        },
        {
            "object_id": "point-1-frame",
            "content_ref": "key_points.1",
            "type": "shape",
            "shape": "rounded_rectangle",
            "x": second_frame[0],
            "y": second_frame[1],
            "w": second_frame[2],
            "h": second_frame[3],
            "z": 3,
            "style_token": "outline_panel" if is_cover else "soft_panel",
        },
        {
            "object_id": "accent-node",
            "content_ref": "decoration",
            "type": "shape",
            "shape": "oval",
            "x": 0.465,
            "y": 0.87,
            "w": 0.03,
            "h": 0.04,
            "z": 4,
            "style_token": "accent_node",
        },
    ]
    for index in range(2, len(plan.slides[0].key_points)):
        modules.append(
            {
                "object_id": f"point-{index}-rule",
                "content_ref": f"key_points.{index}",
                "type": "line",
                "shape": "rectangle",
                "x": 0.09 + 0.27 * ((index - 2) % 3),
                "y": 0.56 + 0.2 * ((index - 2) // 3),
                "w": 0.2,
                "h": 0,
                "z": 3,
                "style_token": "muted_rule",
            }
        )

    visual_geometry = (0.08, 0.3, 0.36, 0.24) if is_cover else (0.08, 0.3, 0.38, 0.53)
    visual_assets = (
        [
            {
                "object_id": "local-illustration",
                "content_ref": "suggested_visual",
                "x": visual_geometry[0],
                "y": visual_geometry[1],
                "w": visual_geometry[2],
                "h": visual_geometry[3],
                "z": 8,
                "image_path": image_path,
                "image_fit": "contain",
            }
        ]
        if image_path
        else []
    )

    return {
        "status": "completed",
        "summary": "generated editable Paper Deck objects and exact-copy block manifest",
        "slides": [
            {
                "slide_id": plan.slides[0].id,
                "background": "12365A" if is_cover else "FFFFFF",
                "manifest_version": 3,
                "layout_mode": "editable-layered-objects",
                "style_signature": style_signature,
                "raster_audit": "inspected-local-assets-text-free",
                "modules": modules,
                "visual_assets": visual_assets,
                "text_blocks": text_blocks,
                "visual_placeholders": visual_placeholders,
            }
        ],
    }


def write_test_image(
    path: Path,
    *,
    image_format: str = "PNG",
    size: tuple[int, int] = (64, 40),
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGBA", size, (40, 110, 170, 255)).save(path, format=image_format)


def test_paper_deck_crop_transform_keeps_text_attached_to_skill_blocks(
    tmp_path: Path,
) -> None:
    source = tmp_path / "three-by-two.png"
    write_test_image(source, size=(1500, 1000))
    blocks = [
        {
            "content_ref": "title",
            "x": 0.2,
            "y": 0.2,
            "w": 0.5,
            "h": 0.2,
        }
    ]

    transformed = CodexPPTProvider._transform_paper_deck_text_blocks_for_crop(
        blocks,
        source_path=source,
        target_size=(1920, 1080),
    )

    assert transformed[0]["x"] == pytest.approx(0.2)
    assert transformed[0]["w"] == pytest.approx(0.5)
    assert transformed[0]["y"] == pytest.approx(0.1444444444)
    assert transformed[0]["h"] == pytest.approx(0.2370370370)
    assert blocks[0]["y"] == 0.2


def test_paper_deck_near_sixteen_by_nine_crop_keeps_boundary_text_safe(
    tmp_path: Path,
) -> None:
    source = tmp_path / "near-sixteen-by-nine.png"
    write_test_image(source, size=(1672, 941))
    transformed = CodexPPTProvider._transform_paper_deck_text_blocks_for_crop(
        [
            {
                "content_ref": "title",
                "x": 0.05,
                "y": 0.04,
                "w": 0.4,
                "h": 0.09,
            }
        ],
        source_path=source,
        target_size=CodexPPTProvider.PAPER_CRAFT_BACKGROUND_SIZE,
    )[0]
    element = SlideElement(
        type="text",
        x=transformed["x"],
        y=transformed["y"],
        w=transformed["w"],
        h=transformed["h"],
        text="DBSCAN algorithm workflow",
        style=SlideElementStyle(font_size=30, line_width=0),
    )

    assert transformed["y"] == pytest.approx(0.03975544816586921)
    assert CodexPPTProvider._is_paper_deck_text_within_safe_canvas(element)


def test_paper_deck_safe_tolerance_does_not_hide_real_crop_overflow(
    tmp_path: Path,
) -> None:
    source = tmp_path / "three-by-two-boundary.png"
    write_test_image(source, size=(1500, 1000))
    transformed = CodexPPTProvider._transform_paper_deck_text_blocks_for_crop(
        [
            {
                "content_ref": "title",
                "x": 0.05,
                "y": 0.08,
                "w": 0.4,
                "h": 0.09,
            }
        ],
        source_path=source,
        target_size=CodexPPTProvider.PAPER_CRAFT_BACKGROUND_SIZE,
    )[0]
    element = SlideElement(
        type="text",
        x=transformed["x"],
        y=transformed["y"],
        w=transformed["w"],
        h=transformed["h"],
        text="DBSCAN algorithm workflow",
        style=SlideElementStyle(font_size=30, line_width=0),
    )

    assert not CodexPPTProvider._is_paper_deck_text_within_safe_canvas(element)


@pytest.mark.parametrize(
    ("background_color", "expected_pixel"),
    [
        ("FFFFFF", (255, 255, 255)),
        ("12365A", (18, 54, 90)),
    ],
)
def test_paper_deck_flattens_alpha_onto_the_trusted_theme_background(
    tmp_path: Path,
    background_color: str,
    expected_pixel: tuple[int, int, int],
) -> None:
    workspace = tmp_path / "workspace"
    staged = workspace / ".codex_image_outputs" / "01.png"
    staged.parent.mkdir(parents=True)
    Image.new("RGBA", (960, 540), (0, 0, 0, 0)).save(staged)

    persisted = CodexPPTProvider._persist_generated_image(
        "generated_visuals/background.png",
        workspace=workspace,
        asset_output_dir=tmp_path / "assets",
        slide_id="slide_001",
        image_number=1,
        fallback_generated_path=staged,
        target_size=CodexPPTProvider.PAPER_CRAFT_BACKGROUND_SIZE,
        prefer_fallback=True,
        require_slide_background=True,
        background_color=background_color,
    )

    with Image.open(persisted) as image:
        assert image.mode == "RGB"
        assert image.size == CodexPPTProvider.PAPER_CRAFT_BACKGROUND_SIZE
        assert image.getpixel((0, 0)) == expected_pixel


def test_paper_deck_contrast_uses_the_visible_composited_background(
    tmp_path: Path,
) -> None:
    workspace = tmp_path / "workspace"
    staged = workspace / ".codex_image_outputs" / "01.png"
    staged.parent.mkdir(parents=True)
    Image.new("RGBA", (1672, 941), (0, 0, 0, 0)).save(staged)
    persisted = CodexPPTProvider._persist_generated_image(
        "generated_visuals/background.png",
        workspace=workspace,
        asset_output_dir=tmp_path / "assets",
        slide_id="slide_010",
        image_number=1,
        fallback_generated_path=staged,
        target_size=CodexPPTProvider.PAPER_CRAFT_BACKGROUND_SIZE,
        prefer_fallback=True,
        require_slide_background=True,
        background_color="FFFFFF",
    )
    title = SlideElement(
        type="text",
        x=0.075,
        y=0.055,
        w=0.44,
        h=0.125,
        text="DBSCAN algorithm workflow",
        style=SlideElementStyle(
            font_size=28,
            bold=True,
            color="17324D",
            line_width=0,
            align="left",
            valign="middle",
        ),
    )

    ratio = CodexPPTProvider._paper_deck_glyph_contrast(persisted, title)

    assert ratio == pytest.approx(
        CodexPPTProvider._contrast_ratio("17324D", "FFFFFF"),
        rel=0.01,
    )
    assert ratio > 13


def write_test_skill_tree(root: Path) -> None:
    for skill_name in CodexPPTProvider.PAPER_CRAFT_SKILLS:
        skill_dir = root / skill_name
        (skill_dir / "references").mkdir(parents=True)
        (skill_dir / "scripts").mkdir()
        (skill_dir / "SKILL.md").write_text(
            f"---\nname: {skill_name}\n---\n# {skill_name}\n",
            encoding="utf-8",
        )
        (skill_dir / "references" / "guidance.md").write_text(
            "approved guidance",
            encoding="utf-8",
        )
        (skill_dir / "references" / "notes.txt").write_text(
            "approved notes",
            encoding="utf-8",
        )
        (skill_dir / "references" / "ignored.png").write_bytes(b"not-guidance")
        (skill_dir / "scripts" / "run.py").write_text(
            "raise RuntimeError('must not be staged')",
            encoding="utf-8",
        )


def calculate_test_skill_digests(root: Path) -> dict[str, str]:
    result = {}
    for skill_name in CodexPPTProvider.PAPER_CRAFT_SKILLS:
        skill_dir = root / skill_name
        sources = [skill_dir / "SKILL.md"]
        sources.extend(
            path
            for path in sorted((skill_dir / "references").rglob("*"))
            if path.is_file() and path.suffix.lower() in {".md", ".txt"}
        )
        digest = hashlib.sha256()
        for source in sources:
            relative_name = source.relative_to(skill_dir).as_posix()
            digest.update(relative_name.encode("utf-8"))
            digest.update(b"\x00")
            digest.update(CodexPPTProvider._canonical_skill_guidance_bytes(source.read_bytes()))
        result[skill_name] = digest.hexdigest()
    return result


def test_vendored_paper_craft_guidance_matches_provider_pins() -> None:
    repository_root = Path(__file__).resolve().parents[3]

    assert calculate_test_skill_digests(repository_root / ".agents" / "skills") == (
        CodexPPTProvider.PAPER_CRAFT_EXPECTED_SHA256
    )


def test_paper_craft_skill_staging_copies_only_audited_guidance(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source_root = tmp_path / "source-skills"
    write_test_skill_tree(source_root)
    monkeypatch.setattr(
        CodexPPTProvider,
        "PAPER_CRAFT_EXPECTED_SHA256",
        calculate_test_skill_digests(source_root),
    )
    workspace = tmp_path / "workspace"
    provider = CodexPPTProvider(
        adapter=Mock(),
        paper_craft_skills_dir=source_root,
    )

    manifest = provider._stage_paper_craft_skills(workspace)
    embedded_guidance = provider._embedded_paper_craft_guidance(workspace, manifest)

    assert [entry["name"] for entry in manifest] == list(CodexPPTProvider.PAPER_CRAFT_SKILLS)
    assert "BEGIN paper-deck/SKILL.md" in embedded_guidance
    assert "BEGIN paper-comic/SKILL.md" in embedded_guidance
    assert "approved guidance" in embedded_guidance
    assert "must not be staged" not in embedded_guidance
    for entry in manifest:
        assert set(entry["files"]) == {
            "SKILL.md",
            "references/guidance.md",
            "references/notes.txt",
        }
        staged = workspace / ".agents" / "skills" / entry["name"]
        assert (staged / "SKILL.md").is_file()
        assert (staged / "references" / "guidance.md").is_file()
        assert not (staged / "references" / "ignored.png").exists()
        assert not (staged / "scripts").exists()
        assert len(entry["sha256"]) == 64


def test_paper_craft_skill_staging_rejects_a_pinned_digest_mismatch(
    tmp_path: Path,
) -> None:
    source_root = tmp_path / "tampered-skills"
    write_test_skill_tree(source_root)
    provider = CodexPPTProvider(
        adapter=Mock(),
        paper_craft_skills_dir=source_root,
    )
    workspace = tmp_path / "workspace"

    with pytest.raises(ValueError, match="pinned skill digest mismatch for paper-deck"):
        provider._stage_paper_craft_skills(workspace)

    assert not (workspace / ".agents" / "skills").exists()


def test_paper_craft_skill_digest_is_stable_across_checkout_whitespace(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source_root = tmp_path / "source-skills"
    write_test_skill_tree(source_root)
    expected = calculate_test_skill_digests(source_root)
    for source in source_root.rglob("*"):
        if source.is_file() and source.suffix.lower() in {".md", ".txt"}:
            text = source.read_text(encoding="utf-8").replace("\r\n", "\n")
            text = "\n".join(f"{line}  " if line else line for line in text.split("\n"))
            source.write_bytes(text.replace("\n", "\r\n").encode("utf-8"))
    monkeypatch.setattr(CodexPPTProvider, "PAPER_CRAFT_EXPECTED_SHA256", expected)
    provider = CodexPPTProvider(adapter=Mock(), paper_craft_skills_dir=source_root)
    workspace = tmp_path / "workspace"

    manifest = provider._stage_paper_craft_skills(workspace)

    assert {entry["name"]: entry["sha256"] for entry in manifest} == expected
    staged_text = (workspace / ".agents" / "skills" / "paper-deck" / "SKILL.md").read_bytes()
    assert b"\r\n" not in staged_text
    assert b"  \n" not in staged_text


def test_vendored_paper_craft_guidance_matches_the_pinned_digests(
    tmp_path: Path,
) -> None:
    repository_root = Path(__file__).resolve().parents[3]
    provider = CodexPPTProvider(
        adapter=Mock(),
        paper_craft_skills_dir=repository_root / ".agents" / "skills",
    )

    manifest = provider._stage_paper_craft_skills(tmp_path / "workspace")

    assert {entry["name"]: entry["sha256"] for entry in manifest} == (
        CodexPPTProvider.PAPER_CRAFT_EXPECTED_SHA256
    )


def test_paper_craft_prompt_and_schema_lock_the_hybrid_contract(tmp_path: Path) -> None:
    plan = make_plan()
    theme = get_presentation_theme()
    schema_path = tmp_path / "schema.json"
    style_signature = "journal-minimal:academic_blue:v3"

    CodexPPTProvider._write_paper_deck_output_schema(
        schema_path,
        text_block_count=1 + len(plan.slides[0].key_points),
        style_signature=style_signature,
        visual_placeholder_refs=("suggested_visual", "visual_payload.0"),
        min_visual_assets=1,
        max_visual_assets=2,
        min_visual_placeholders=0,
        max_visual_placeholders=0,
        require_dominant_visual_column=False,
    )
    schema = json.loads(schema_path.read_text(encoding="utf-8"))
    slide_schema = schema["properties"]["slides"]["items"]
    modules_schema = slide_schema["properties"]["modules"]
    module_schema = modules_schema["items"]
    assets_schema = slide_schema["properties"]["visual_assets"]
    asset_schema = assets_schema["items"]
    text_blocks_schema = slide_schema["properties"]["text_blocks"]
    text_block_schema = text_blocks_schema["items"]
    placeholders_schema = slide_schema["properties"]["visual_placeholders"]
    placeholder_schema = placeholders_schema["items"]
    prompt = CodexPPTProvider._paper_craft_prompt(
        plan,
        "immutable-hash",
        theme,
        max_images=2,
        skill_guidance="PINNED_SKILL_GUIDANCE",
    )

    assert "elements" not in slide_schema["properties"]
    assert modules_schema["minItems"] == 1
    assert modules_schema["maxItems"] == CodexPPTProvider.PAPER_DECK_MAX_VISUAL_MODULES
    shape_module_schema, line_module_schema = module_schema["anyOf"]
    assert shape_module_schema["properties"]["type"]["enum"] == ["shape"]
    assert line_module_schema["properties"]["type"]["enum"] == ["line"]
    assert set(shape_module_schema["properties"]["style_token"]["enum"]).isdisjoint(
        line_module_schema["properties"]["style_token"]["enum"]
    )
    for variant in (shape_module_schema, line_module_schema):
        assert "object_id" in variant["required"]
        assert "content_ref" in variant["required"]
        assert "style_token" in variant["required"]
        assert "text" not in variant["properties"]
    assert assets_schema["minItems"] == 1
    assert assets_schema["maxItems"] == 2
    assert asset_schema["properties"]["w"] == {
        "type": "number",
        "minimum": 0.16,
        "maximum": CodexPPTProvider.PAPER_DECK_VISUAL_REGION_MAX_WIDTH,
    }
    assert asset_schema["properties"]["h"] == {
        "type": "number",
        "minimum": 0.12,
        "maximum": CodexPPTProvider.PAPER_DECK_VISUAL_REGION_MAX_HEIGHT,
    }
    assert asset_schema["properties"]["content_ref"]["enum"] == [
        "suggested_visual",
        "visual_payload.0",
    ]
    assert asset_schema["properties"]["image_fit"]["enum"] == ["contain", "cover"]
    assert "text_anchor" not in slide_schema["properties"]
    assert slide_schema["properties"]["manifest_version"] == {
        "type": "integer",
        "minimum": 3,
        "maximum": 3,
    }
    assert slide_schema["properties"]["layout_mode"]["enum"] == ["editable-layered-objects"]
    assert slide_schema["properties"]["style_signature"]["enum"] == [style_signature]
    assert slide_schema["properties"]["raster_audit"]["enum"] == [
        "inspected-local-assets-text-free"
    ]
    assert text_blocks_schema["minItems"] == 3
    assert text_blocks_schema["maxItems"] == 3
    assert set(text_block_schema["required"]) == {
        "content_ref",
        "x",
        "y",
        "w",
        "h",
        "font_size",
        "min_font_size",
        "font_role",
        "bold",
        "color",
        "align",
        "valign",
        "max_lines",
    }
    assert text_block_schema["properties"]["content_ref"]["pattern"] == (
        "^(title|key_points\\.[0-9]+)$"
    )
    assert text_block_schema["properties"]["font_role"]["enum"] == [
        "sans",
        "serif",
        "handwritten",
        "display",
        "mono",
    ]
    assert placeholders_schema["minItems"] == 0
    assert placeholders_schema["maxItems"] == 0
    assert placeholder_schema["properties"]["w"]["minimum"] == (
        CodexPPTProvider.PAPER_DECK_VISUAL_REGION_MIN_WIDTH
    )
    assert placeholder_schema["properties"]["h"]["minimum"] == (
        CodexPPTProvider.PAPER_DECK_VISUAL_REGION_MIN_HEIGHT
    )
    assert placeholder_schema["properties"]["content_ref"]["enum"] == [
        "suggested_visual",
        "visual_payload.0",
    ]
    assert "text" not in placeholder_schema["properties"]
    assert "$paper-deck" in prompt
    assert "$paper-comic" in prompt
    assert "PINNED_SKILL_GUIDANCE" in prompt
    assert "do not try to read the staged skill files with shell commands" in prompt
    assert "do not pause for confirmation" in prompt
    assert "creates slide 1 of 1" in prompt
    assert "Create 1-2 images in this session" in prompt
    assert "session hard limit is 2" in prompt
    assert "Never generate an entire 16:9 slide image" in prompt
    assert "backend" in prompt.lower()
    assert "exact presentationplan" in prompt.lower()
    assert "generated_visuals/<exact-generated-filename>" in prompt
    assert "reference only the final accepted image" in prompt
    assert "must be entirely text-free" in prompt
    assert "transparent background" in prompt
    assert "inspect the local asset" in prompt
    assert "inspected-local-assets-text-free" in prompt
    assert "no visible copy in modules or visual_assets" in prompt.lower()
    assert "transparent" in prompt.lower()
    assert "separate modules entry" in prompt.lower()
    assert "text_blocks" in prompt
    assert "visual_placeholders" in prompt
    assert "do not hallucinate it" in prompt
    assert "independently movable PowerPoint objects" in prompt
    assert "mega-card" in prompt
    assert "Do not put every point in" in prompt
    assert "independent editorial anchor" in prompt
    assert "No connector, divider, underline, node, or partial shape" in " ".join(prompt.split())
    assert style_signature in prompt
    assert "visual_assets and visual_placeholders are mutually exclusive" in prompt
    assert "PAPER_DECK_COMPOSITION_DIRECTIVE" in prompt
    assert "layout_id=hero_minimal" in prompt
    assert "not required to be a vertical left/right column" in prompt
    assert "at least 0.015" in prompt
    assert "at least 0.015" in prompt
    assert "will not choose a layout, move a block, or rewrite content" in prompt
    assert "deterministically maps module style" in prompt
    assert "cross-platform safe" in prompt
    assert "calculate x+w and y+h" in prompt
    assert "Normal text requires at least 4.5:1" in prompt
    assert "requires at least 3.0:1" in prompt
    assert "Speaker scripts are intentionally omitted" in prompt
    assert plan.slides[0].title in prompt


def test_paper_deck_visual_requirement_forces_local_art_on_mechanism_pages() -> None:
    slide = (
        make_plan()
        .slides[0]
        .model_copy(
            update={
                "order": 2,
                "title": "DBSCAN算法机制与案例",
                "suggested_visual": "无文字的算法流程和局部放大科研插图",
                "visual_payload": ["密度邻域机制", "案例聚类结果局部放大"],
            }
        )
    )

    requirement = CodexPPTProvider._paper_deck_visual_requirement(slide)

    assert requirement["mode"] == "required-local-illustration"
    assert requirement["min_assets"] == 1
    assert requirement["max_assets"] == 2
    assert requirement["max_placeholders"] == 0


def test_paper_deck_visual_requirement_uses_placeholder_only_for_real_source_asset() -> None:
    slide = (
        make_plan()
        .slides[0]
        .model_copy(
            update={
                "order": 2,
                "title": "软件操作案例",
                "suggested_visual": "插入 ArcGIS 软件界面截图",
                "visual_payload": ["参数设置面板"],
            }
        )
    )

    requirement = CodexPPTProvider._paper_deck_visual_requirement(slide)

    assert requirement["mode"] == "source-placeholder"
    assert requirement["max_assets"] == 0
    assert requirement["min_placeholders"] == 1
    assert requirement["max_placeholders"] == 1


def test_paper_deck_composition_plan_varies_page_silhouettes_by_role() -> None:
    base = make_plan().slides[0]
    roles = ("cover", "method", "comparison", "case", "summary")
    slides = [
        base.model_copy(
            update={
                "id": f"slide_{index + 1:03d}",
                "order": index + 1,
                "slide_role": role,
                "title": f"{role} research page",
                "suggested_visual": f"{role} scientific visual",
            }
        )
        for index, role in enumerate(roles)
    ]
    plan = make_plan().model_copy(update={"slides": slides})

    compositions = CodexPPTProvider._paper_deck_composition_plan(plan)
    layout_ids = [item["layout_id"] for item in compositions]
    families = [item["family"] for item in compositions]

    assert layout_ids[0] == "hero_minimal"
    assert len(set(families)) >= 4
    assert all(left != right for left, right in zip(layout_ids, layout_ids[1:]))
    assert {"process", "comparison", "evidence", "summary"}.issubset(set(families))


@pytest.mark.parametrize("theme", PRESENTATION_THEMES, ids=lambda theme: theme.id)
def test_paper_deck_keeps_journal_minimal_for_every_color_theme(theme) -> None:
    plan = make_plan()

    prompt = CodexPPTProvider._paper_craft_prompt(
        plan,
        "immutable-hash",
        theme,
        max_images=1,
    )
    context = json.loads(CodexPPTProvider._paper_deck_deck_context(plan, theme))

    assert CodexPPTProvider._paper_deck_style_preset(theme) == "journal-minimal"
    assert context["style_preset"] == "journal-minimal"
    assert context["color_theme_id"] == theme.id
    assert context["color_theme"]["colors"] == theme.prompt_payload()["colors"]
    assert "regardless of the selected color theme" in prompt
    assert "palette tokens only" in prompt
    assert CodexPPTProvider.PAPER_DECK_STYLE_PROMPT in prompt
    assert f"journal-minimal:{theme.id}:v3" in prompt


def test_paper_deck_allows_extra_lines_when_exact_copy_physically_fits(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    exact_point = "DBSCAN在地理处理中的应用，如ArcGIS Pro中的参数设置面板和流程图。"
    slide = make_plan().slides[0].model_copy(update={"key_points": [exact_point]})
    blocks = [
        {
            "content_ref": "title",
            "x": 0.055,
            "y": 0.045,
            "w": 0.78,
            "h": 0.075,
            "font_size": 34,
            "min_font_size": 28,
            "font_role": "sans",
            "bold": True,
            "color": "17324D",
            "align": "left",
            "valign": "middle",
            "max_lines": 1,
        },
        {
            "content_ref": "key_points.0",
            "x": 0.51,
            "y": 0.20,
            "w": 0.390,
            "h": 0.595,
            "font_size": 17,
            "min_font_size": 14,
            "font_role": "sans",
            "bold": False,
            "color": "17324D",
            "align": "left",
            "valign": "middle",
            "max_lines": 10,
        },
    ]

    def measured_fit(text, width, font_size, *, font_role, bold):
        if text == exact_point:
            return 11, 0.361
        return 1, 0.05

    monkeypatch.setattr(
        CodexPPTProvider,
        "_paper_deck_overlay_metrics",
        staticmethod(measured_fit),
    )

    elements = CodexPPTProvider._build_skill_text_block_overlay(
        slide=slide,
        raw_text_blocks=blocks,
    )

    assert [element["text"] for element in elements] == [slide.title, exact_point]
    assert elements[1]["style"]["font_size"] == 17
    assert elements[1]["x"] == blocks[1]["x"]
    assert elements[1]["y"] == blocks[1]["y"]
    assert elements[1]["w"] == blocks[1]["w"]
    assert elements[1]["h"] == blocks[1]["h"]


def test_paper_deck_still_rejects_copy_that_physically_exceeds_the_block(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    slide = make_plan().slides[0].model_copy(update={"key_points": ["无法容纳的长文本"]})
    blocks = [
        {
            "content_ref": content_ref,
            "x": 0.05,
            "y": 0.05 + index * 0.2,
            "w": 0.4,
            "h": 0.1,
            "font_size": 28 if index == 0 else 17,
            "min_font_size": 24 if index == 0 else 14,
            "font_role": "sans",
            "bold": index == 0,
            "color": "17324D",
            "align": "left",
            "valign": "middle",
            "max_lines": 3,
        }
        for index, content_ref in enumerate(("title", "key_points.0"))
    ]

    monkeypatch.setattr(
        CodexPPTProvider,
        "_paper_deck_overlay_metrics",
        staticmethod(lambda *args, **kwargs: (4, 0.2)),
    )

    with pytest.raises(ValueError, match="requires 4 lines and height 0.200"):
        CodexPPTProvider._build_skill_text_block_overlay(
            slide=slide,
            raw_text_blocks=blocks,
        )


def test_layered_mode_rejects_a_full_bleed_picture(tmp_path: Path) -> None:
    plan = make_plan()
    workspace = tmp_path / "workspace"
    write_test_image(workspace / "generated_visuals" / "background.png")
    result = make_structured_result(
        plan,
        image_paths=["generated_visuals/background.png"],
    )
    background = result["slides"][0]["elements"][-1]
    background.update({"x": 0, "y": 0, "w": 1, "h": 1, "z": 0})
    for element in result["slides"][0]["elements"]:
        if element["type"] == "text":
            element["contract_role"] = "plan_copy"
            element["style"].update({"fill": None, "line_color": None, "line_width": 0})

    with pytest.raises(ValueError, match="forbids a full-slide merged image"):
        CodexPPTProvider._design_plan_from_result(
            plan,
            result,
            get_presentation_theme(),
            workspace=workspace,
            asset_output_dir=tmp_path / "assets",
            allow_images=True,
            max_images=1,
            require_editable_layers_per_slide=True,
        )


def test_paper_deck_generates_independently_movable_objects_per_slide(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    first_plan = make_plan()
    second_slide = first_plan.slides[0].model_copy(
        update={
            "id": "slide_002",
            "order": 2,
            "title": "局部关系需要逐位置解释",
            "key_points": [
                "局部参数刻画空间差异。",
                "结果必须回到具体位置解释。",
            ],
            "speaker_script": "解释局部参数和空间位置之间的关系。",
            "suggested_visual": "插入 ArcGIS 软件界面截图。",
            "visual_payload": ["参数设置面板"],
        }
    )
    third_slide = second_slide.model_copy(
        update={
            "id": "slide_003",
            "order": 3,
            "title": "GWR content anchor continuation",
            "suggested_visual": "无文字的局部回归机制与空间放大示意",
            "visual_payload": ["局部回归机制", "空间放大示意"],
        }
    )
    plan = first_plan.model_copy(
        update={"slides": [first_plan.slides[0], second_slide, third_slide]}
    )
    source_root = tmp_path / "source-skills"
    write_test_skill_tree(source_root)
    monkeypatch.setattr(
        CodexPPTProvider,
        "PAPER_CRAFT_EXPECTED_SHA256",
        calculate_test_skill_digests(source_root),
    )
    real_adapter = PPTSkillAdapter()
    adapter = Mock(spec=PPTSkillAdapter)
    adapter.render_declarative_pptx.side_effect = real_adapter.render_declarative_pptx
    adapter.prepare_external_pptx.return_value = Mock(id="artifact_codex")
    provider = CodexPPTProvider(
        adapter=adapter,
        repair_attempts=1,
        paper_craft_enabled=True,
        paper_craft_max_images=3,
        paper_craft_concurrency=2,
        paper_craft_skills_dir=source_root,
    )
    generated_slide_ids: list[str] = []
    returned_text_blocks: dict[str, list[dict]] = {}
    prompts_by_slide: dict[str, str] = {}

    def fake_execute(
        *,
        workspace,
        output_dir,
        prompt,
        attempt,
        sandbox_mode="read-only",
        run_mode="structured-vector",
        ignore_user_config=True,
    ):
        single_plan = PresentationPlan.model_validate_json(
            (workspace / "presentation_plan.json").read_text(encoding="utf-8")
        )
        generated_slide_ids.append(single_plan.slides[0].id)
        prompts_by_slide[single_plan.slides[0].id] = prompt
        if single_plan.slides[0].id == "slide_002":
            result = make_paper_deck_result(
                single_plan,
                visual_placeholders=[
                    {
                        "content_ref": "visual_payload.0",
                        "x": 0.09,
                        "y": 0.31,
                        "w": 0.38,
                        "h": 0.52,
                    }
                ],
            )
        else:
            write_test_image(
                workspace / ".codex_image_outputs" / "01.png",
                size=(960, 540),
            )
            result = make_paper_deck_result(single_plan)
            if single_plan.slides[0].id == "slide_003":
                write_test_image(
                    workspace / ".codex_image_outputs" / "02.png",
                    size=(960, 540),
                )
                first_asset = result["slides"][0]["visual_assets"][0]
                first_asset.update({"x": 0.08, "y": 0.3, "w": 0.18, "h": 0.53})
                result["slides"][0]["visual_assets"].append(
                    {
                        "object_id": "local-illustration-2",
                        "content_ref": "visual_payload.0",
                        "x": 0.28,
                        "y": 0.3,
                        "w": 0.18,
                        "h": 0.53,
                        "z": 9,
                        "image_path": "generated_visuals/illustration-2.png",
                        "image_fit": "contain",
                    }
                )
        returned_text_blocks[single_plan.slides[0].id] = result["slides"][0]["text_blocks"]
        result["_metaclass_staged_image_count"] = len(result["slides"][0]["visual_assets"])
        return result

    monkeypatch.setattr(provider, "_execute_codex", fake_execute)

    provider.prepare_request(
        plan=plan,
        job_id="job_per_slide_backgrounds",
        output_dir=tmp_path / "output",
    )

    assert sorted(generated_slide_ids) == ["slide_001", "slide_002", "slide_003"]
    rendered_plan = adapter.render_declarative_pptx.call_args.args[0]
    assert len(rendered_plan.slides) == 3
    for expected, rendered in zip(plan.slides, rendered_plan.slides, strict=True):
        visual_modules = [
            element for element in rendered.elements if element.contract_role == "visual_module"
        ]
        assert len(visual_modules) >= 4
        assert all(element.type in {"shape", "line"} for element in visual_modules)
        assert len({element.object_id for element in visual_modules}) == len(visual_modules)
        plan_text_elements = [
            element
            for element in rendered.elements
            if element.type == "text" and element.contract_role == "plan_copy"
        ]
        assert [element.text for element in plan_text_elements] == [
            expected.title,
            *expected.key_points,
        ]
        placeholders = [
            element
            for element in rendered.elements
            if element.type == "text" and element.contract_role == "visual_placeholder"
        ]
        if expected.id == "slide_002":
            assert [element.text for element in placeholders] == [
                "（此处建议插入：参数设置面板相关图片）"
            ]
            assert placeholders[0].style.fill == "FFFFFF"
            assert placeholders[0].style.line_width == 1.5
            assert placeholders[0].w >= CodexPPTProvider.PAPER_DECK_VISUAL_COLUMN_MIN_WIDTH
            assert placeholders[0].h >= CodexPPTProvider.PAPER_DECK_VISUAL_COLUMN_MIN_HEIGHT
        else:
            assert placeholders == []
        local_images = [
            element for element in rendered.elements if element.contract_role == "visual_asset"
        ]
        expected_image_count = {"slide_001": 1, "slide_002": 0, "slide_003": 2}
        assert len(local_images) == expected_image_count[expected.id]
        if expected.order > 1 and local_images:
            left = min(element.x for element in local_images)
            top = min(element.y for element in local_images)
            right = max(element.x + element.w for element in local_images)
            bottom = max(element.y + element.h for element in local_images)
            assert right - left >= CodexPPTProvider.PAPER_DECK_VISUAL_COLUMN_MIN_WIDTH
            assert bottom - top >= CodexPPTProvider.PAPER_DECK_VISUAL_COLUMN_MIN_HEIGHT
        assert not any(
            CodexPPTProvider._is_full_bleed_background_image(element)
            for element in rendered.elements
        )
        text_elements = plan_text_elements
        manifest = returned_text_blocks[expected.id]
        assert [(element.x, element.y, element.w, element.h) for element in text_elements] == [
            (block["x"], block["y"], block["w"], block["h"]) for block in manifest
        ]
        assert [element.style.color for element in text_elements] == [
            block["color"] for block in manifest
        ]
        assert [element.style.font_role for element in text_elements] == [
            block["font_role"] for block in manifest
        ]
        assert [element.style.font_size for element in text_elements] == [
            block["font_size"] for block in manifest
        ]
        assert [element.style.bold for element in text_elements] == [
            block["bold"] for block in manifest
        ]
        assert [element.style.align for element in text_elements] == [
            block["align"] for block in manifest
        ]
        assert [element.style.valign for element in text_elements] == [
            block["valign"] for block in manifest
        ]
        if expected.order == 1:
            assert text_elements[1].x > text_elements[2].x
        else:
            assert text_elements[1].x == text_elements[2].x
        assert text_elements[1].y < text_elements[2].y
        for text_element in text_elements:
            assert text_element.style.fill is None
            assert text_element.style.line_color is None
            assert text_element.style.line_width == 0
        for local_image in local_images:
            with Image.open(local_image.image_path or "") as image:
                assert image.size == (960, 540)
                assert image.mode == "RGBA"
    metadata = adapter.prepare_external_pptx.call_args.kwargs["provider_metadata"]
    assert metadata["generation_mode"] == "paper-craft-layered-safe-render"
    assert metadata["paper_craft"]["generated_image_count"] == 3
    assert metadata["paper_craft"]["required_image_count"] == 2
    assert metadata["paper_craft"]["maximum_generated_image_count"] == 4
    assert [item["slide_id"] for item in metadata["responses"]] == [
        "slide_001",
        "slide_002",
        "slide_003",
    ]
    assert "cover_layout_manifest=" in prompts_by_slide["slide_003"]
    assert "content_layout_manifest=" in prompts_by_slide["slide_003"]
    assert "title font_role and bold value" in prompts_by_slide["slide_003"]
    assert len(list((tmp_path / "output" / "paper_deck_prompts").glob("*.md"))) == 3

    deck_path = tmp_path / "output" / "deck.pptx"
    deck = Presentation(deck_path)
    first_slide = deck.slides[0]
    assert not any(shape.shape_type == MSO_SHAPE_TYPE.GROUP for shape in first_slide.shapes)
    module_a = next(
        shape
        for shape in first_slide.shapes
        if shape.name == "MetaClass visual_module point-0-frame"
    )
    module_b = next(
        shape
        for shape in first_slide.shapes
        if shape.name == "MetaClass visual_module point-1-frame"
    )
    original_a_left = module_a.left
    original_b_position = (module_b.left, module_b.top)
    placeholder = next(
        shape
        for shape in deck.slides[1].shapes
        if shape.name == "MetaClass visual_placeholder placeholder-1"
    )
    placeholder_sibling = next(
        shape
        for shape in deck.slides[1].shapes
        if shape.name == "MetaClass visual_module point-0-frame"
    )
    local_illustration = next(
        shape
        for shape in deck.slides[2].shapes
        if shape.name == "MetaClass visual_asset local-illustration"
    )
    second_local_illustration = next(
        shape
        for shape in deck.slides[2].shapes
        if shape.name == "MetaClass visual_asset local-illustration-2"
    )
    illustration_sibling = next(
        shape
        for shape in deck.slides[2].shapes
        if shape.name == "MetaClass visual_module point-0-frame"
    )
    original_placeholder_left = placeholder.left
    original_placeholder_sibling = (
        placeholder_sibling.left,
        placeholder_sibling.top,
    )
    original_illustration_top = local_illustration.top
    original_illustration_sibling = (
        illustration_sibling.left,
        illustration_sibling.top,
    )
    original_second_illustration = (
        second_local_illustration.left,
        second_local_illustration.top,
    )
    module_a.left += Inches(0.25)
    placeholder.left += Inches(0.15)
    local_illustration.top += Inches(0.15)
    moved_path = tmp_path / "moved-one-module.pptx"
    deck.save(moved_path)

    reopened = Presentation(moved_path)
    reopened_slide = reopened.slides[0]
    moved_a = next(
        shape
        for shape in reopened_slide.shapes
        if shape.name == "MetaClass visual_module point-0-frame"
    )
    untouched_b = next(
        shape
        for shape in reopened_slide.shapes
        if shape.name == "MetaClass visual_module point-1-frame"
    )
    moved_placeholder = next(
        shape
        for shape in reopened.slides[1].shapes
        if shape.name == "MetaClass visual_placeholder placeholder-1"
    )
    untouched_placeholder_sibling = next(
        shape
        for shape in reopened.slides[1].shapes
        if shape.name == "MetaClass visual_module point-0-frame"
    )
    moved_illustration = next(
        shape
        for shape in reopened.slides[2].shapes
        if shape.name == "MetaClass visual_asset local-illustration"
    )
    untouched_illustration_sibling = next(
        shape
        for shape in reopened.slides[2].shapes
        if shape.name == "MetaClass visual_module point-0-frame"
    )
    untouched_second_illustration = next(
        shape
        for shape in reopened.slides[2].shapes
        if shape.name == "MetaClass visual_asset local-illustration-2"
    )
    assert moved_a.left > original_a_left
    assert (untouched_b.left, untouched_b.top) == original_b_position
    assert moved_placeholder.left > original_placeholder_left
    assert (
        untouched_placeholder_sibling.left,
        untouched_placeholder_sibling.top,
    ) == original_placeholder_sibling
    assert moved_illustration.top > original_illustration_top
    assert (
        untouched_illustration_sibling.left,
        untouched_illustration_sibling.top,
    ) == original_illustration_sibling
    assert (
        untouched_second_illustration.left,
        untouched_second_illustration.top,
    ) == original_second_illustration


def test_paper_deck_rejects_a_workspace_image_without_imagegen_provenance(
    tmp_path: Path,
) -> None:
    plan = make_plan()
    workspace = tmp_path / "workspace"
    result = make_paper_deck_result(plan)
    provider = CodexPPTProvider(adapter=Mock(), paper_craft_enabled=True)

    with pytest.raises(ValueError, match="image-generation provenance failed"):
        provider._prepare_paper_deck_slide_result(
            result=result,
            slide=plan.slides[0],
            slide_index=0,
            theme=get_presentation_theme(),
            slide_workspace=workspace,
            generated_visuals=tmp_path / "generated-visuals",
        )


def test_paper_deck_retries_only_the_failed_slide_in_a_fresh_workspace(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    plan = make_plan()
    source_root = tmp_path / "source-skills"
    write_test_skill_tree(source_root)
    monkeypatch.setattr(
        CodexPPTProvider,
        "PAPER_CRAFT_EXPECTED_SHA256",
        calculate_test_skill_digests(source_root),
    )
    real_adapter = PPTSkillAdapter()
    adapter = Mock(spec=PPTSkillAdapter)
    adapter.render_declarative_pptx.side_effect = real_adapter.render_declarative_pptx
    adapter.prepare_external_pptx.return_value = Mock(id="artifact_codex")
    provider = CodexPPTProvider(
        adapter=adapter,
        repair_attempts=1,
        paper_craft_enabled=True,
        paper_craft_max_images=2,
        paper_craft_concurrency=1,
        paper_craft_skills_dir=source_root,
    )
    workspaces: list[str] = []
    prompts: list[str] = []

    def fake_execute(
        *,
        workspace,
        output_dir,
        prompt,
        attempt,
        sandbox_mode="read-only",
        run_mode="structured-vector",
        ignore_user_config=True,
    ):
        workspaces.append(workspace.name)
        prompts.append(prompt)
        single_plan = PresentationPlan.model_validate_json(
            (workspace / "presentation_plan.json").read_text(encoding="utf-8")
        )
        write_test_image(
            workspace / ".codex_image_outputs" / "01.png",
            size=(960, 540),
        )
        result = make_paper_deck_result(single_plan)
        result["_metaclass_staged_image_count"] = 1
        if len(workspaces) == 1:
            result["slides"][0]["modules"][1]["object_id"] = "title-rule"
        return result

    monkeypatch.setattr(provider, "_execute_codex", fake_execute)

    provider.prepare_request(
        plan=plan,
        job_id="job_retry_single_slide",
        output_dir=tmp_path / "output",
    )

    assert workspaces == ["attempt-01", "attempt-02"]
    metadata = adapter.prepare_external_pptx.call_args.kwargs["provider_metadata"]
    assert metadata["attempt_count"] == 2
    assert metadata["paper_craft_attempt_count"] == 2
    assert "missing or duplicated" in prompts[1]
    assert len(list((tmp_path / "output" / "paper_deck_prompts").glob("*.md"))) == 2


@pytest.mark.parametrize(
    ("failure_kind", "expected_retry_reason"),
    [
        ("wrong-content-ref", "content_ref mismatch"),
        ("exact-copy-does-not-fit", "cannot fit exact Plan copy"),
        ("low-contrast", "contrast is too low"),
    ],
)
def test_paper_deck_retries_invalid_skill_text_block_manifests(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    failure_kind: str,
    expected_retry_reason: str,
) -> None:
    plan = make_plan()
    source_root = tmp_path / "source-skills"
    write_test_skill_tree(source_root)
    monkeypatch.setattr(
        CodexPPTProvider,
        "PAPER_CRAFT_EXPECTED_SHA256",
        calculate_test_skill_digests(source_root),
    )
    real_adapter = PPTSkillAdapter()
    adapter = Mock(spec=PPTSkillAdapter)
    adapter.render_declarative_pptx.side_effect = real_adapter.render_declarative_pptx
    adapter.prepare_external_pptx.return_value = Mock(id="artifact_codex")
    provider = CodexPPTProvider(
        adapter=adapter,
        repair_attempts=1,
        paper_craft_enabled=True,
        paper_craft_max_images=2,
        paper_craft_concurrency=1,
        paper_craft_skills_dir=source_root,
    )
    workspaces: list[str] = []
    prompts: list[str] = []

    def fake_execute(
        *,
        workspace,
        output_dir,
        prompt,
        attempt,
        sandbox_mode="read-only",
        run_mode="structured-vector",
        ignore_user_config=True,
    ):
        workspaces.append(workspace.name)
        prompts.append(prompt)
        single_plan = PresentationPlan.model_validate_json(
            (workspace / "presentation_plan.json").read_text(encoding="utf-8")
        )
        write_test_image(
            workspace / ".codex_image_outputs" / "01.png",
            size=(960, 540),
        )
        result = make_paper_deck_result(single_plan)
        result["_metaclass_staged_image_count"] = 1
        if len(workspaces) == 1:
            blocks = result["slides"][0]["text_blocks"]
            if failure_kind == "wrong-content-ref":
                blocks[1]["content_ref"] = "key_points.99"
            elif failure_kind == "exact-copy-does-not-fit":
                blocks[0].update(
                    {
                        "w": 0.05,
                        "h": 0.05,
                        "font_size": 32,
                        "min_font_size": 24,
                        "max_lines": 1,
                    }
                )
            elif failure_kind == "low-contrast":
                blocks[1]["color"] = "286EAA"
            else:  # pragma: no cover - guarded by parametrization
                raise AssertionError(f"unknown failure kind: {failure_kind}")
        return result

    monkeypatch.setattr(provider, "_execute_codex", fake_execute)

    provider.prepare_request(
        plan=plan,
        job_id=f"job_retry_{failure_kind}",
        output_dir=tmp_path / "output",
    )

    if failure_kind == "low-contrast":
        # Deterministic color repair no longer spends a second Codex call.
        assert workspaces == ["attempt-01"]
    else:
        assert workspaces == ["attempt-01", "attempt-02"]
        assert expected_retry_reason in prompts[1]
    rendered_plan = adapter.render_declarative_pptx.call_args.args[0]
    assert [
        element.text for element in rendered_plan.slides[0].elements if element.type == "text"
    ] == [plan.slides[0].title, *plan.slides[0].key_points]
    metadata = adapter.prepare_external_pptx.call_args.kwargs["provider_metadata"]
    expected_attempt_count = 1 if failure_kind == "low-contrast" else 2
    assert metadata["attempt_count"] == expected_attempt_count
    assert metadata["paper_craft_attempt_count"] == expected_attempt_count


def test_paper_deck_retry_prompt_accumulates_failures_and_reaudits_all_gates(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    plan = make_plan()
    source_root = tmp_path / "source-skills"
    write_test_skill_tree(source_root)
    monkeypatch.setattr(
        CodexPPTProvider,
        "PAPER_CRAFT_EXPECTED_SHA256",
        calculate_test_skill_digests(source_root),
    )
    real_adapter = PPTSkillAdapter()
    adapter = Mock(spec=PPTSkillAdapter)
    adapter.render_declarative_pptx.side_effect = real_adapter.render_declarative_pptx
    adapter.prepare_external_pptx.return_value = Mock(id="artifact_codex")
    provider = CodexPPTProvider(
        adapter=adapter,
        repair_attempts=2,
        paper_craft_enabled=True,
        paper_craft_max_images=2,
        paper_craft_concurrency=1,
        paper_craft_skills_dir=source_root,
    )
    prompts: list[str] = []

    def fake_execute(
        *,
        workspace,
        output_dir,
        prompt,
        attempt,
        sandbox_mode="read-only",
        run_mode="structured-vector",
        ignore_user_config=True,
    ):
        prompts.append(prompt)
        single_plan = PresentationPlan.model_validate_json(
            (workspace / "presentation_plan.json").read_text(encoding="utf-8")
        )
        write_test_image(workspace / ".codex_image_outputs" / "01.png", size=(960, 540))
        result = make_paper_deck_result(single_plan)
        result["_metaclass_staged_image_count"] = 1
        blocks = result["slides"][0]["text_blocks"]
        if len(prompts) == 1:
            blocks[0].update({"w": 0.05, "h": 0.05})
        elif len(prompts) == 2:
            blocks[1]["color"] = "286EAA"
        return result

    monkeypatch.setattr(provider, "_execute_codex", fake_execute)

    provider.prepare_request(
        plan=plan,
        job_id="job_cumulative_repair_history",
        output_dir=tmp_path / "output",
    )

    # The second result only has a contrast defect, which is repaired locally
    # instead of triggering a third paid Codex request.
    assert len(prompts) == 2
    assert "cannot fit exact Plan copy" in prompts[1]
    assert "fixing one issue cannot create another" in prompts[1]
    assert "calculate x+w and y+h" in prompts[1]
    assert "never return a full-slide raster or grouped mega-panel" in prompts[1]


def test_paper_deck_publishes_completed_prefix_before_usage_limit(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    first_plan = make_plan()
    second_slide = first_plan.slides[0].model_copy(
        update={
            "id": "slide_002",
            "order": 2,
            "title": "空间权重决定邻域影响",
        }
    )
    plan = first_plan.model_copy(update={"slides": [first_plan.slides[0], second_slide]})
    source_root = tmp_path / "source-skills"
    write_test_skill_tree(source_root)
    monkeypatch.setattr(
        CodexPPTProvider,
        "PAPER_CRAFT_EXPECTED_SHA256",
        calculate_test_skill_digests(source_root),
    )
    provider = CodexPPTProvider(
        adapter=PPTSkillAdapter(),
        repair_attempts=1,
        paper_craft_enabled=True,
        paper_craft_max_images=2,
        paper_craft_concurrency=1,
        paper_craft_skills_dir=source_root,
    )
    calls = 0
    progress: list[tuple[PPTArtifact, int, int]] = []

    def fake_execute(
        *,
        workspace,
        output_dir,
        prompt,
        attempt,
        sandbox_mode="read-only",
        run_mode="structured-vector",
        ignore_user_config=True,
    ):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise CodexGenerationError("You've hit your usage limit")
        single_plan = PresentationPlan.model_validate_json(
            (workspace / "presentation_plan.json").read_text(encoding="utf-8")
        )
        write_test_image(
            workspace / ".codex_image_outputs" / "01.png",
            size=(960, 540),
        )
        result = make_paper_deck_result(single_plan)
        result["_metaclass_staged_image_count"] = 1
        return result

    monkeypatch.setattr(provider, "_execute_codex", fake_execute)
    output_dir = tmp_path / "output"
    output_dir.mkdir()
    legacy_shared_path = output_dir / "deck.pptx"
    legacy_shared_path.write_bytes(b"an older PPTX may be open in PowerPoint")

    with pytest.raises(CodexGenerationError, match="usage limit"):
        provider.prepare_request(
            plan=plan,
            job_id="job_partial_usage_limit",
            output_dir=output_dir,
            progress_callback=lambda artifact, completed, total: progress.append(
                (artifact, completed, total)
            ),
        )

    assert [(completed, total) for _, completed, total in progress] == [(1, 2)]
    partial_path = Path(progress[0][0].pptx_path)
    assert partial_path.is_file()
    assert partial_path.name.startswith("deck.partial.001.")
    assert legacy_shared_path.read_bytes() == b"an older PPTX may be open in PowerPoint"
    assert len(Presentation(partial_path).slides) == 1
    assert (output_dir / "partial_status.json").is_file()
    checkpoint = json.loads((output_dir / "paper_deck_checkpoint.json").read_text(encoding="utf-8"))
    assert checkpoint["completed_slides"] == 1
    assert checkpoint["prepared_slides"][0] is not None
    assert checkpoint["prepared_slides"][1] is None

    # A later resume restores slide 1 from disk and only calls Codex for slide 2.
    progress.clear()
    artifact = provider.prepare_request(
        plan=plan,
        job_id="job_partial_usage_limit",
        output_dir=output_dir,
        progress_callback=lambda current, completed, total: progress.append(
            (current, completed, total)
        ),
    )

    assert calls == 3
    assert [(completed, total) for _, completed, total in progress] == [(1, 2), (2, 2)]
    assert len(Presentation(artifact.pptx_path).slides) == 2
    assert Path(artifact.pptx_path).name.startswith("deck.complete.")
    assert legacy_shared_path.read_bytes() == b"an older PPTX may be open in PowerPoint"
    completed_checkpoint = json.loads(
        (output_dir / "paper_deck_checkpoint.json").read_text(encoding="utf-8")
    )
    assert completed_checkpoint["completed_slides"] == 2


def test_paper_craft_image_is_normalized_and_persisted(tmp_path: Path) -> None:
    plan = make_plan()
    workspace = tmp_path / "workspace"
    source = workspace / "generated_visuals" / "figure.png"
    write_test_image(source, size=(160, 40))
    output_assets = tmp_path / "output" / "codex_assets"
    raw_element = make_structured_result(
        plan,
        image_paths=["generated_visuals/figure.png"],
    )["slides"][0]["elements"][-1]

    element = CodexPPTProvider._parse_design_element(
        raw_element,
        workspace=workspace,
        asset_output_dir=output_assets,
        slide_id="slide:unsafe/id",
        image_number=1,
        allow_images=True,
    )

    persisted = Path(element.image_path or "")
    assert element.type == "image"
    assert persisted == (output_assets / "slide-unsafe-id-01.png").resolve()
    assert persisted.is_file()
    with Image.open(persisted) as image:
        assert image.format == "PNG"
        assert image.mode == "RGBA"
        assert image.size == (160, 40)

    rendered_plan = plan.model_copy(
        update={
            "slides": [
                plan.slides[0].model_copy(
                    update={"elements": [element]},
                )
            ]
        }
    )
    destination = tmp_path / "cropped-image.pptx"
    PPTSkillAdapter().render_declarative_pptx(rendered_plan, destination)
    picture = Presentation(destination).slides[0].shapes[0]
    assert picture.crop_left > 0
    assert picture.crop_right > 0


@pytest.mark.parametrize(
    ("size", "message"),
    [
        ((320, 180), "resolution is too low"),
        ((800, 1000), "landscape image close to 16:9"),
    ],
)
def test_paper_deck_rejects_low_quality_background_sources(
    tmp_path: Path,
    size: tuple[int, int],
    message: str,
) -> None:
    workspace = tmp_path / "workspace"
    staged = workspace / ".codex_image_outputs" / "01.png"
    write_test_image(staged, size=size)

    with pytest.raises(ValueError, match=message):
        CodexPPTProvider._persist_generated_image(
            "generated_visuals/background.png",
            workspace=workspace,
            asset_output_dir=tmp_path / "assets",
            slide_id="slide_001",
            image_number=1,
            fallback_generated_path=staged,
            target_size=CodexPPTProvider.PAPER_CRAFT_BACKGROUND_SIZE,
            prefer_fallback=True,
            require_slide_background=True,
        )


def test_codex_session_image_is_staged_and_recovers_a_corrupt_workspace_copy(
    tmp_path: Path,
) -> None:
    session_id = "019f9e90-1234-7abc-8def-123456789abc"
    codex_home = tmp_path / "codex-home"
    session_dir = codex_home / "generated_images" / session_id
    write_test_image(session_dir / "generated.png", size=(96, 64))
    workspace = tmp_path / "workspace"
    corrupt_copy = workspace / "generated_visuals" / "figure.png"
    corrupt_copy.parent.mkdir(parents=True)
    corrupt_copy.write_bytes(b"incomplete png copied by the model")
    provider = CodexPPTProvider(adapter=Mock(), paper_craft_max_images=2)

    staged_count, error = provider._stage_codex_generated_images(
        workspace=workspace,
        environment={"CODEX_HOME": str(codex_home)},
        stderr=f"Codex event stream\nsession id: {session_id}\n",
    )

    staged = workspace / ".codex_image_outputs" / "01.png"
    assert staged_count == 1
    assert error is None
    assert staged.is_file()
    with Image.open(staged) as image:
        assert image.format == "PNG"
        assert image.size == (96, 64)

    raw_element = make_structured_result(
        make_plan(),
        image_paths=["generated_visuals/figure.png"],
    )["slides"][0]["elements"][-1]
    element = CodexPPTProvider._parse_design_element(
        raw_element,
        workspace=workspace,
        asset_output_dir=tmp_path / "assets",
        slide_id="slide_001",
        image_number=1,
        allow_images=True,
    )

    persisted = Path(element.image_path or "")
    assert persisted.is_file()
    with Image.open(persisted) as image:
        assert image.size == (96, 64)


def test_codex_image_staging_matches_the_final_generated_uuid_among_drafts(
    tmp_path: Path,
) -> None:
    session_id = "019f9e90-1234-7abc-8def-123456789abc"
    codex_home = tmp_path / "codex-home"
    session_dir = codex_home / "generated_images" / session_id
    write_test_image(session_dir / "exec-rejected-draft.png", size=(80, 45))
    write_test_image(session_dir / "exec-final-audit.png", size=(160, 90))
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    provider = CodexPPTProvider(adapter=Mock(), paper_craft_max_images=2)

    staged_count, error = provider._stage_codex_generated_images(
        workspace=workspace,
        environment={"CODEX_HOME": str(codex_home)},
        stderr=f"session id: {session_id}",
        reported_image_paths=("generated_visuals/exec-final-audit.png",),
    )

    assert staged_count == 1
    assert error is None
    with Image.open(workspace / ".codex_image_outputs" / "01.png") as image:
        assert image.size == (160, 90)


def test_codex_image_staging_uses_newest_trusted_source_for_a_logical_name(
    tmp_path: Path,
) -> None:
    session_id = "019f9e90-1234-7abc-8def-123456789abc"
    codex_home = tmp_path / "codex-home"
    session_dir = codex_home / "generated_images" / session_id
    rejected = session_dir / "exec-rejected-draft.png"
    final = session_dir / "exec-final-audit.png"
    write_test_image(rejected, size=(80, 45))
    write_test_image(final, size=(160, 90))
    os.utime(rejected, ns=(1_000_000_000, 1_000_000_000))
    os.utime(final, ns=(2_000_000_000, 2_000_000_000))
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    provider = CodexPPTProvider(adapter=Mock(), paper_craft_max_images=2)

    staged_count, error = provider._stage_codex_generated_images(
        workspace=workspace,
        environment={"CODEX_HOME": str(codex_home)},
        stderr=f"session id: {session_id}",
        reported_image_paths=("generated_visuals/academic-resource-path.png",),
    )

    assert staged_count == 1
    assert error is None
    with Image.open(workspace / ".codex_image_outputs" / "01.png") as image:
        assert image.size == (160, 90)


def test_image_link_checks_work_without_path_is_junction(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    plan = make_plan()
    workspace = tmp_path / "workspace"
    write_test_image(workspace / "generated_visuals" / "figure.png")
    raw_element = make_structured_result(
        plan,
        image_paths=["generated_visuals/figure.png"],
    )["slides"][0]["elements"][-1]
    monkeypatch.delattr(Path, "is_junction", raising=False)

    element = CodexPPTProvider._parse_design_element(
        raw_element,
        workspace=workspace,
        asset_output_dir=tmp_path / "assets",
        allow_images=True,
    )

    assert Path(element.image_path or "").is_file()


def test_codex_image_staging_uses_the_first_trusted_session_id(
    tmp_path: Path,
) -> None:
    trusted_session = "019f9e90-1111-7abc-8def-123456789abc"
    echoed_session = "019f9e90-9999-7abc-8def-123456789abc"
    codex_home = tmp_path / "codex-home"
    generated_root = codex_home / "generated_images"
    write_test_image(
        generated_root / trusted_session / "trusted.png",
        size=(120, 70),
    )
    write_test_image(
        generated_root / echoed_session / "echoed.png",
        size=(31, 29),
    )
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    provider = CodexPPTProvider(adapter=Mock())
    stderr = (
        f"OpenAI Codex\nsession id: {trusted_session}\n"
        "user\nmodel output echoed untrusted prompt text:\n"
        f"session id: {echoed_session}\n"
    )

    staged_count, error = provider._stage_codex_generated_images(
        workspace=workspace,
        environment={"CODEX_HOME": str(codex_home)},
        stderr=stderr,
    )

    assert staged_count == 1
    assert error is None
    with Image.open(workspace / ".codex_image_outputs" / "01.png") as image:
        assert image.size == (120, 70)


def test_codex_image_staging_rejects_a_precreated_recovery_directory(
    tmp_path: Path,
) -> None:
    session_id = "019f9e90-1234-7abc-8def-123456789abc"
    codex_home = tmp_path / "codex-home"
    write_test_image(
        codex_home / "generated_images" / session_id / "generated.png",
    )
    workspace = tmp_path / "workspace"
    staged_root = workspace / ".codex_image_outputs"
    staged_root.mkdir(parents=True)
    sentinel = staged_root / "untrusted.txt"
    sentinel.write_text("pre-created by the model", encoding="utf-8")
    provider = CodexPPTProvider(adapter=Mock())

    staged_count, error = provider._stage_codex_generated_images(
        workspace=workspace,
        environment={"CODEX_HOME": str(codex_home)},
        stderr=f"session id: {session_id}",
    )

    assert staged_count == 0
    assert error
    assert sentinel.read_text(encoding="utf-8") == "pre-created by the model"
    assert not (staged_root / "01.png").exists()


def test_valid_returned_image_takes_priority_over_ordinal_recovery(
    tmp_path: Path,
) -> None:
    plan = make_plan()
    workspace = tmp_path / "workspace"
    write_test_image(
        workspace / "generated_visuals" / "returned.png",
        size=(140, 60),
    )
    write_test_image(
        workspace / ".codex_image_outputs" / "01.png",
        size=(35, 90),
    )
    raw_element = make_structured_result(
        plan,
        image_paths=["generated_visuals/returned.png"],
    )["slides"][0]["elements"][-1]

    element = CodexPPTProvider._parse_design_element(
        raw_element,
        workspace=workspace,
        asset_output_dir=tmp_path / "assets",
        image_number=1,
        allow_images=True,
    )

    with Image.open(Path(element.image_path or "")) as image:
        assert image.size == (140, 60)


def test_codex_image_staging_rejects_an_unknown_session(tmp_path: Path) -> None:
    codex_home = tmp_path / "codex-home"
    (codex_home / "generated_images").mkdir(parents=True)
    provider = CodexPPTProvider(adapter=Mock())

    staged_count, error = provider._stage_codex_generated_images(
        workspace=tmp_path / "workspace",
        environment={"CODEX_HOME": str(codex_home)},
        stderr="session id: 019f9e90-ffff-7abc-8def-123456789abc",
    )

    assert staged_count == 0
    assert error == "Codex generated-image session directory is unavailable"


def test_codex_image_staging_rejects_a_linked_session_directory(tmp_path: Path) -> None:
    session_id = "019f9e90-1234-7abc-8def-123456789abc"
    generated_root = tmp_path / "codex-home" / "generated_images"
    real_session = generated_root / "real-session"
    write_test_image(real_session / "generated.png")
    linked_session = generated_root / session_id
    try:
        linked_session.symlink_to(real_session, target_is_directory=True)
    except (NotImplementedError, OSError) as exc:
        pytest.skip(f"directory symlinks are unavailable: {exc}")
    provider = CodexPPTProvider(adapter=Mock())

    staged_count, error = provider._stage_codex_generated_images(
        workspace=tmp_path / "workspace",
        environment={"CODEX_HOME": str(tmp_path / "codex-home")},
        stderr=f"session id: {session_id}",
    )

    assert staged_count == 0
    assert error == "Codex generated-image session directory is unavailable"


@pytest.mark.parametrize(
    ("raw_path", "message"),
    [
        ("../escape.png", "stay inside generated_visuals"),
        ("generated_visuals/../escape.png", "stay inside generated_visuals"),
        ("C:\\temp\\figure.png", "relative to generated_visuals"),
        ("/tmp/figure.png", "relative to generated_visuals"),
        ("https://example.com/figure.png", "stay inside generated_visuals"),
        ("other/figure.png", "stay inside generated_visuals"),
        ("generated_visuals/figure.svg", "PNG, JPEG, or WebP"),
    ],
)
def test_paper_craft_rejects_unsafe_image_paths(
    tmp_path: Path,
    raw_path: str,
    message: str,
) -> None:
    raw_element = make_structured_result(
        make_plan(),
        image_paths=[raw_path],
    )["slides"][0]["elements"][-1]

    with pytest.raises(ValueError, match=message):
        CodexPPTProvider._parse_design_element(
            raw_element,
            workspace=tmp_path / "workspace",
            asset_output_dir=tmp_path / "assets",
            allow_images=True,
        )


def test_paper_craft_rejects_missing_and_corrupt_images(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    result = make_structured_result(
        make_plan(),
        image_paths=["generated_visuals/figure.png"],
    )
    raw_element = result["slides"][0]["elements"][-1]

    with pytest.raises(ValueError, match="generated image is missing"):
        CodexPPTProvider._parse_design_element(
            raw_element,
            workspace=workspace,
            asset_output_dir=tmp_path / "assets",
            allow_images=True,
        )

    corrupt = workspace / "generated_visuals" / "figure.png"
    corrupt.parent.mkdir(parents=True)
    corrupt.write_bytes(b"this is not a raster image")
    with pytest.raises(ValueError, match="generated image is unreadable"):
        CodexPPTProvider._parse_design_element(
            raw_element,
            workspace=workspace,
            asset_output_dir=tmp_path / "assets",
            allow_images=True,
        )


def test_paper_craft_rejects_a_linked_generated_visuals_root(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    outside = tmp_path / "outside"
    outside.mkdir()
    visual_root = workspace / "generated_visuals"
    try:
        visual_root.symlink_to(outside, target_is_directory=True)
    except (NotImplementedError, OSError) as exc:
        pytest.skip(f"directory symlinks are unavailable: {exc}")
    raw_element = make_structured_result(
        make_plan(),
        image_paths=["generated_visuals/figure.png"],
    )["slides"][0]["elements"][-1]

    with pytest.raises(ValueError, match="directory cannot be a link or junction"):
        CodexPPTProvider._parse_design_element(
            raw_element,
            workspace=workspace,
            asset_output_dir=tmp_path / "assets",
            allow_images=True,
        )


def test_paper_craft_rejects_more_images_than_the_configured_budget(
    tmp_path: Path,
) -> None:
    plan = make_plan()
    workspace = tmp_path / "workspace"
    for name in ("one.png", "two.png"):
        write_test_image(workspace / "generated_visuals" / name)
    result = make_structured_result(
        plan,
        image_paths=[
            "generated_visuals/one.png",
            "generated_visuals/two.png",
        ],
    )

    with pytest.raises(ValueError, match="image budget"):
        CodexPPTProvider._design_plan_from_result(
            plan,
            result,
            get_presentation_theme(),
            workspace=workspace,
            asset_output_dir=tmp_path / "assets",
            allow_images=True,
            max_images=1,
        )


@pytest.mark.parametrize("paper_result_kind", ["blocked", "zero-image"])
def test_paper_craft_failure_refuses_silent_vector_degradation(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    paper_result_kind: str,
) -> None:
    plan = make_plan()
    source_root = tmp_path / "source-skills"
    write_test_skill_tree(source_root)
    monkeypatch.setattr(
        CodexPPTProvider,
        "PAPER_CRAFT_EXPECTED_SHA256",
        calculate_test_skill_digests(source_root),
    )
    adapter = Mock()
    adapter.prepare_external_pptx.return_value = Mock(id="artifact_codex")
    provider = CodexPPTProvider(
        adapter=adapter,
        repair_attempts=0,
        paper_craft_enabled=True,
        paper_craft_max_images=2,
        paper_craft_skills_dir=source_root,
    )
    calls: list[tuple[int, str, str, bool]] = []

    def fake_execute(
        *,
        workspace,
        output_dir,
        prompt,
        attempt,
        sandbox_mode="read-only",
        run_mode="structured-vector",
        ignore_user_config=True,
    ):
        calls.append((attempt, sandbox_mode, run_mode, ignore_user_config))
        assert "approved guidance" in prompt
        assert "VERIFIED_PAPER_CRAFT_GUIDANCE_BEGIN" in prompt
        if run_mode == "paper-deck-slide-hybrid":
            if paper_result_kind == "blocked":
                return {
                    "status": "blocked",
                    "summary": "built-in image generation is unavailable",
                    "slides": [],
                }
            return make_structured_result(plan)
        raise AssertionError("strict Paper Deck mode must not invoke vector fallback")

    monkeypatch.setattr(provider, "_execute_codex", fake_execute)

    with pytest.raises(CodexGenerationError, match="refusing silent vector degradation"):
        provider.prepare_request(
            plan=plan,
            job_id=f"job_paper_{paper_result_kind}",
            output_dir=tmp_path / "output",
        )

    assert calls == [
        (1, "workspace-write", "paper-deck-slide-hybrid", False),
    ]
    adapter.prepare_external_pptx.assert_not_called()


def test_corrupt_job_image_does_not_disable_paper_craft_process_wide(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    plan = make_plan()
    source_root = tmp_path / "source-skills"
    write_test_skill_tree(source_root)
    monkeypatch.setattr(
        CodexPPTProvider,
        "PAPER_CRAFT_EXPECTED_SHA256",
        calculate_test_skill_digests(source_root),
    )
    adapter = Mock()
    adapter.prepare_external_pptx.return_value = Mock(id="artifact_codex")
    provider = CodexPPTProvider(
        adapter=adapter,
        repair_attempts=0,
        paper_craft_enabled=True,
        paper_craft_skills_dir=source_root,
    )

    def fake_execute(
        *,
        workspace,
        output_dir,
        prompt,
        attempt,
        sandbox_mode="read-only",
        run_mode="structured-vector",
        ignore_user_config=True,
    ):
        if run_mode == "paper-deck-slide-hybrid":
            result = make_paper_deck_result(
                plan,
                image_path="generated_visuals/missing.png",
            )
            result["_metaclass_staged_image_count"] = 1
            return result
        raise AssertionError("strict Paper Deck mode must not invoke vector fallback")

    monkeypatch.setattr(provider, "_execute_codex", fake_execute)

    with pytest.raises(CodexGenerationError, match="trusted staged visual asset is missing"):
        provider.prepare_request(
            plan=plan,
            job_id="job_missing_image",
            output_dir=tmp_path / "output",
        )

    assert provider._paper_craft_capability is None
    adapter.prepare_external_pptx.assert_not_called()


def test_codex_request_does_not_leak_the_configured_skills_directory(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    plan = make_plan()
    skills_dir = tmp_path / "private-skills-location"
    adapter = Mock()
    adapter.prepare_external_pptx.return_value = Mock(id="artifact_codex")
    provider = CodexPPTProvider(
        adapter=adapter,
        repair_attempts=0,
        paper_craft_enabled=False,
        paper_craft_max_images=0,
        paper_craft_skills_dir=skills_dir,
    )

    def fake_execute(*, workspace, output_dir, prompt, attempt):
        write_deck(
            workspace / "deck.pptx",
            [plan.slides[0].title, *plan.slides[0].key_points],
        )
        return {
            "status": "completed",
            "deck_path": "deck.pptx",
            "summary": "safe vector fallback",
        }

    monkeypatch.setattr(provider, "_execute_codex", fake_execute)
    output_dir = tmp_path / "output"

    provider.prepare_request(
        plan=plan,
        job_id="job_no_skills_path_leak",
        output_dir=output_dir,
    )

    request_text = (output_dir / "codex_request.json").read_text(encoding="utf-8")
    request = json.loads(request_text)
    assert "skills_dir" not in request["paper_craft"]
    assert skills_dir.name not in request_text


def test_paper_craft_capability_failure_is_retried_by_the_next_request(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    plan = make_plan()
    source_root = tmp_path / "source-skills"
    write_test_skill_tree(source_root)
    monkeypatch.setattr(
        CodexPPTProvider,
        "PAPER_CRAFT_EXPECTED_SHA256",
        calculate_test_skill_digests(source_root),
    )
    adapter = Mock()
    adapter.prepare_external_pptx.return_value = Mock(id="artifact_codex")
    provider = CodexPPTProvider(
        adapter=adapter,
        repair_attempts=0,
        paper_craft_enabled=True,
        paper_craft_max_images=2,
        paper_craft_skills_dir=source_root,
    )
    calls: list[tuple[str, bool]] = []

    def fake_execute(
        *,
        workspace,
        output_dir,
        prompt,
        attempt,
        sandbox_mode="read-only",
        run_mode="structured-vector",
        ignore_user_config=True,
    ):
        calls.append((run_mode, ignore_user_config))
        if run_mode == "paper-deck-slide-hybrid":
            assert (workspace / "generated_visuals").is_dir()
            return {
                "status": "blocked",
                "summary": "built-in image generation is unavailable",
                "slides": [],
            }
        raise AssertionError("strict Paper Deck mode must not invoke vector fallback")

    monkeypatch.setattr(provider, "_execute_codex", fake_execute)

    for index in (1, 2):
        with pytest.raises(CodexGenerationError, match="image generation is unavailable"):
            provider.prepare_request(
                plan=plan,
                job_id=f"job_capability_cache_{index}",
                output_dir=tmp_path / f"output-{index}",
            )

    assert calls == [
        ("paper-deck-slide-hybrid", False),
        ("paper-deck-slide-hybrid", False),
    ]
    assert provider._paper_craft_capability is None
    adapter.prepare_external_pptx.assert_not_called()


def test_execute_codex_applies_requested_sandbox_and_retains_safety_flags(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    codex_bin = tmp_path / "codex.exe"
    codex_bin.write_bytes(b"placeholder")
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    process = Mock(returncode=0)
    process.communicate.return_value = ("", "")
    popen = Mock(return_value=process)
    monkeypatch.setattr(
        "metaclass.modules.presentation.codex_provider.subprocess.Popen",
        popen,
    )
    provider = CodexPPTProvider(
        adapter=Mock(),
        api_key="test-key",
        codex_bin=str(codex_bin),
    )

    provider._execute_codex(
        workspace=workspace,
        output_dir=tmp_path,
        prompt="generate",
        attempt=1,
        sandbox_mode="workspace-write",
        run_mode="paper-craft-hybrid",
        ignore_user_config=False,
    )

    provider._execute_codex(
        workspace=workspace,
        output_dir=tmp_path,
        prompt="fallback",
        attempt=2,
        run_mode="paper-craft-vector-fallback",
    )

    hybrid_command = popen.call_args_list[0].args[0]
    assert hybrid_command[hybrid_command.index("--sandbox") + 1] == "workspace-write"
    assert "--ephemeral" in hybrid_command
    assert "--ignore-user-config" not in hybrid_command
    assert "--ignore-rules" in hybrid_command
    assert 'web_search="disabled"' in hybrid_command
    assert 'model_reasoning_effort="low"' in hybrid_command
    if os.name == "nt":
        assert 'windows.sandbox="unelevated"' in hybrid_command
    assert hybrid_command[hybrid_command.index("-C") + 1] == str(workspace)
    vector_command = popen.call_args_list[1].args[0]
    assert vector_command[vector_command.index("--sandbox") + 1] == "read-only"
    assert "--ignore-user-config" in vector_command
    assert 'model_reasoning_effort="low"' not in vector_command
    run_log = json.loads((tmp_path / "codex_attempt_1.json").read_text(encoding="utf-8"))
    assert run_log["run_mode"] == "paper-craft-hybrid"
    assert run_log["sandbox_mode"] == "workspace-write"
    assert run_log["ignore_user_config"] is False
    vector_log = json.loads((tmp_path / "codex_attempt_2.json").read_text(encoding="utf-8"))
    assert vector_log["run_mode"] == "paper-craft-vector-fallback"
    assert vector_log["sandbox_mode"] == "read-only"
    assert vector_log["ignore_user_config"] is True


def test_codex_provider_accepts_only_an_exact_plan_deck(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    plan = make_plan()
    adapter = Mock()
    adapter.prepare_external_pptx.return_value = Mock(id="artifact_codex")
    provider = CodexPPTProvider(adapter=adapter, repair_attempts=1)

    def fake_execute(*, workspace, output_dir, prompt, attempt):
        write_deck(
            workspace / "deck.pptx",
            [plan.slides[0].title, *plan.slides[0].key_points],
        )
        return {
            "status": "completed",
            "deck_path": "deck.pptx",
            "summary": "generated",
        }

    monkeypatch.setattr(provider, "_execute_codex", fake_execute)

    artifact = provider.prepare_request(
        plan=plan,
        job_id="job_codex",
        output_dir=tmp_path,
    )

    assert artifact.id == "artifact_codex"
    assert (tmp_path / "deck.pptx").exists()
    call = adapter.prepare_external_pptx.call_args.kwargs
    assert call["provider_name"] == "codex"
    assert call["provider_metadata"]["validation"]["status"] == "passed"
    assert call["provider_metadata"]["attempt_count"] == 1
    assert call["provider_metadata"]["content_contract"].startswith("exact")


def test_codex_structured_design_is_safely_compiled_to_pptx(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    plan = make_plan()
    adapter = PPTSkillAdapter()
    monkeypatch.setattr(adapter, "_render_slide_images", lambda *args, **kwargs: [])
    package_spy = Mock(wraps=adapter.prepare_external_pptx)
    monkeypatch.setattr(adapter, "prepare_external_pptx", package_spy)
    provider = CodexPPTProvider(adapter=adapter, repair_attempts=1)
    attempts = []

    def style(*, color: str, fill: str, bold: bool = False, size: int = 20):
        return {
            "font_size": size,
            "bold": bold,
            "color": color,
            "fill": fill,
            "line_color": "2E75B6",
            "line_width": 1,
            "align": "left",
            "valign": "middle",
            "opacity": 100,
        }

    def fake_execute(*, workspace, output_dir, prompt, attempt):
        attempts.append(attempt)
        return {
            "status": "completed",
            "summary": "structured design",
            "slides": [
                {
                    "slide_id": plan.slides[0].id,
                    "background": "12365A",
                    "elements": [
                        {
                            "type": "text",
                            "x": 0.1,
                            "y": 0.1,
                            "w": 0.8,
                            "h": 0.03 if attempt == 1 else 0.16,
                            "z": 10,
                            "text": plan.slides[0].title,
                            "shape": "rectangle",
                            "style": style(color="FFFFFF", fill="12365A", bold=True, size=32),
                        },
                        {
                            "type": "text",
                            "x": 0.1,
                            "y": 0.36,
                            "w": 0.36,
                            "h": 0.24,
                            "z": 11,
                            "text": plan.slides[0].key_points[0],
                            "shape": "rectangle",
                            "style": style(color="17324D", fill="DCEBFA"),
                        },
                        {
                            "type": "text",
                            "x": 0.54,
                            "y": 0.36,
                            "w": 0.36,
                            "h": 0.24,
                            "z": 12,
                            "text": plan.slides[0].key_points[1],
                            "shape": "rectangle",
                            "style": style(color="17324D", fill="DCEBFA"),
                        },
                        {
                            "type": "shape",
                            "x": 0.08,
                            "y": 0.33,
                            "w": 0.4,
                            "h": 0.3,
                            "z": 1,
                            "text": "",
                            "shape": "rounded_rectangle",
                            "style": style(color="17324D", fill="DCEBFA"),
                        },
                        {
                            "type": "shape",
                            "x": 0.52,
                            "y": 0.33,
                            "w": 0.4,
                            "h": 0.3,
                            "z": 1,
                            "text": "",
                            "shape": "rounded_rectangle",
                            "style": style(color="17324D", fill="DCEBFA"),
                        },
                    ],
                }
            ],
        }

    monkeypatch.setattr(provider, "_execute_codex", fake_execute)

    artifact = provider.prepare_request(
        plan=plan,
        job_id="job_codex_structured",
        output_dir=tmp_path,
    )

    assert Path(artifact.pptx_path).exists()
    request = json.loads(Path(artifact.skill_request_path).read_text(encoding="utf-8"))
    assert request["provider"] == "codex"
    assert request["provider_metadata"]["generation_mode"] == ("structured-design-safe-render")
    assert attempts == [1, 2]
    assert request["provider_metadata"]["repair_attempts_used"] == 1
    preview_plan = package_spy.call_args.kwargs["preview_plan"]
    assert preview_plan is not None
    assert preview_plan.slides[0].speaker_script == plan.slides[0].speaker_script
    assert [
        element.text for element in preview_plan.slides[0].elements if element.type == "text"
    ] == [plan.slides[0].title, *plan.slides[0].key_points]

    deck = Presentation(artifact.pptx_path)
    text_shapes = {
        shape.text_frame.text: shape for shape in deck.slides[0].shapes if shape.has_text_frame
    }
    point_box = text_shapes[plan.slides[0].key_points[0]]
    assert str(point_box.fill.fore_color.rgb) == "DCEBFA"
    assert str(point_box.line.color.rgb) == "2E75B6"


def test_scene_renderer_applies_text_fill_line_and_opacity(tmp_path: Path) -> None:
    plan = make_plan()
    slide = plan.slides[0].model_copy(
        update={
            "elements": [
                SlideElement(
                    type="text",
                    x=0.1,
                    y=0.1,
                    w=0.8,
                    h=0.2,
                    text=plan.slides[0].title,
                    style=SlideElementStyle(
                        font_size=32,
                        bold=True,
                        color="FFFFFF",
                        fill="12365A",
                        line_color="2E75B6",
                        line_width=2,
                        opacity=40,
                    ),
                )
            ]
        }
    )
    rendered_plan = plan.model_copy(update={"slides": [slide]})
    destination = tmp_path / "scene-style.pptx"

    PPTSkillAdapter().render_declarative_pptx(rendered_plan, destination)

    deck = Presentation(destination)
    textbox = deck.slides[0].shapes[0]
    assert str(textbox.fill.fore_color.rgb) == "12365A"
    assert str(textbox.line.color.rgb) == "2E75B6"
    assert textbox._element.xml.count('a:alpha val="40000"') >= 3


def test_codex_rejects_low_contrast_text_before_render() -> None:
    plan = make_plan()

    def style(color: str, fill: str, *, size: int = 20, bold: bool = False):
        return {
            "font_size": size,
            "bold": bold,
            "color": color,
            "fill": fill,
            "line_color": fill,
            "line_width": 0,
            "align": "left",
            "valign": "middle",
            "opacity": 100,
        }

    result = {
        "status": "completed",
        "summary": "low contrast",
        "slides": [
            {
                "slide_id": plan.slides[0].id,
                "background": "12365A",
                "elements": [
                    {
                        "type": "text",
                        "x": 0.1,
                        "y": 0.08,
                        "w": 0.8,
                        "h": 0.14,
                        "z": 10,
                        "text": plan.slides[0].title,
                        "shape": "rectangle",
                        "style": style("12365A", "12365A", size=32, bold=True),
                    },
                    {
                        "type": "text",
                        "x": 0.1,
                        "y": 0.3,
                        "w": 0.8,
                        "h": 0.18,
                        "z": 11,
                        "text": plan.slides[0].key_points[0],
                        "shape": "rectangle",
                        "style": style("FFFFFF", "12365A"),
                    },
                    {
                        "type": "text",
                        "x": 0.1,
                        "y": 0.58,
                        "w": 0.8,
                        "h": 0.18,
                        "z": 12,
                        "text": plan.slides[0].key_points[1],
                        "shape": "rectangle",
                        "style": style("FFFFFF", "12365A"),
                    },
                    {
                        "type": "shape",
                        "x": 0.07,
                        "y": 0.28,
                        "w": 0.86,
                        "h": 0.5,
                        "z": 1,
                        "text": "",
                        "shape": "rounded_rectangle",
                        "style": style("FFFFFF", "12365A"),
                    },
                ],
            }
        ],
    }

    with pytest.raises(ValueError, match="contrast"):
        CodexPPTProvider._design_plan_from_result(
            plan,
            result,
            get_presentation_theme(),
        )


def test_codex_design_payload_includes_visual_plan_but_omits_script() -> None:
    plan = make_plan()

    payload = json.loads(CodexPPTProvider._design_payload_json(plan))

    slide = payload["slides"][0]
    assert slide["layout_id"] == plan.slides[0].layout_id
    assert slide["visual_payload"] == plan.slides[0].visual_payload
    assert slide["suggested_visual"] == plan.slides[0].suggested_visual
    assert "speaker_script" not in slide


def test_codex_provider_repairs_a_contract_violation_once(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    plan = make_plan()
    adapter = Mock()
    adapter.prepare_external_pptx.return_value = Mock(id="artifact_codex")
    provider = CodexPPTProvider(adapter=adapter, repair_attempts=1)
    attempts = []

    def fake_execute(*, workspace, output_dir, prompt, attempt):
        attempts.append(attempt)
        title = "被改写的标题" if attempt == 1 else plan.slides[0].title
        write_deck(workspace / "deck.pptx", [title, *plan.slides[0].key_points])
        return {
            "status": "completed",
            "deck_path": "deck.pptx",
            "summary": f"attempt {attempt}",
        }

    monkeypatch.setattr(provider, "_execute_codex", fake_execute)

    provider.prepare_request(
        plan=plan,
        job_id="job_codex_repair",
        output_dir=tmp_path,
    )

    metadata = adapter.prepare_external_pptx.call_args.kwargs["provider_metadata"]
    assert attempts == [1, 2]
    assert metadata["repair_attempts_used"] == 1
    assert "content mismatch" in metadata["validation_errors"][0]


def test_codex_provider_rejects_punctuation_or_case_changes(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    plan = make_plan()
    provider = CodexPPTProvider(adapter=Mock(), repair_attempts=0)

    def fake_execute(*, workspace, output_dir, prompt, attempt):
        changed_points = [plan.slides[0].key_points[0].rstrip("。")]
        changed_points.extend(plan.slides[0].key_points[1:])
        write_deck(workspace / "deck.pptx", [plan.slides[0].title, *changed_points])
        return {
            "status": "completed",
            "deck_path": "deck.pptx",
            "summary": "changed punctuation",
        }

    monkeypatch.setattr(provider, "_execute_codex", fake_execute)

    with pytest.raises(CodexGenerationError, match="immutable PresentationPlan contract"):
        provider.prepare_request(
            plan=plan,
            job_id="job_codex_invalid",
            output_dir=tmp_path,
        )


def test_codex_authentication_failure_is_reported_before_generation(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        "metaclass.modules.presentation.codex_provider.subprocess.run",
        lambda *args, **kwargs: Mock(returncode=1),
    )

    with pytest.raises(CodexGenerationError, match="CODEX_API_KEY"):
        CodexPPTProvider._ensure_authenticated("codex")


def test_codex_falls_back_to_presenton_and_records_the_reason(tmp_path: Path) -> None:
    plan = make_plan()
    primary = Mock()
    primary.prepare_request.side_effect = CodexGenerationError("Codex timed out")
    request_path = tmp_path / "skill_request.json"
    request_path.write_text(
        json.dumps({"provider": "presenton"}),
        encoding="utf-8",
    )
    fallback = Mock()
    fallback.prepare_request.return_value = PPTArtifact(
        id="artifact_presenton",
        job_id="job_fallback",
        presentation_plan_id=plan.id,
        pptx_path=str(tmp_path / "deck.pptx"),
        skill_request_path=str(request_path),
    )
    provider = FallbackPPTProvider(
        primary=primary,
        fallback=fallback,
        primary_name="codex",
        fallback_name="presenton",
        fallback_exceptions=(CodexGenerationError,),
    )

    artifact = provider.prepare_request(
        plan=plan,
        job_id="job_fallback",
        output_dir=tmp_path,
    )

    assert artifact.id == "artifact_presenton"
    payload = json.loads(request_path.read_text(encoding="utf-8"))
    assert payload["fallback"] == {
        "from": "codex",
        "to": "presenton",
        "reason": "Codex timed out",
    }


def test_external_provider_preserves_speaker_scripts_by_slide_id(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    plan = make_plan()
    deck_path = tmp_path / "deck.pptx"
    write_deck(deck_path, [plan.slides[0].title, *plan.slides[0].key_points])
    adapter = PPTSkillAdapter()
    monkeypatch.setattr(adapter, "_render_slide_images", lambda *args, **kwargs: [])

    adapter.prepare_external_pptx(
        plan=plan,
        job_id="job_scripts",
        output_dir=tmp_path,
        pptx_path=deck_path,
        provider_name="codex",
        provider_metadata={"validation": {"status": "passed"}},
    )

    scripts = json.loads((tmp_path / "speaker_scripts.json").read_text(encoding="utf-8"))
    assert scripts["presentation_plan_id"] == plan.id
    assert scripts["slides"] == [
        {
            "slide_id": plan.slides[0].id,
            "order": plan.slides[0].order,
            "title": plan.slides[0].title,
            "speaker_script": plan.slides[0].speaker_script,
        }
    ]


def test_codex_runtime_is_primary_and_presenton_is_the_only_fallback(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from metaclass.core.application import build_services
    from metaclass.core.config import settings

    monkeypatch.setattr(settings, "ppt_provider", "codex")
    monkeypatch.setattr(settings, "presenton_api_key", "test-key")
    services = build_services(tmp_path, force_fake_llm=False)
    try:
        provider = services.presentations.ppt_adapter
        assert isinstance(provider, FallbackPPTProvider)
        assert isinstance(provider.primary, CodexPPTProvider)
        assert isinstance(provider.fallback, PresentonPPTProvider)
        assert provider.fallback_name == "presenton"
    finally:
        services.database.dispose()
